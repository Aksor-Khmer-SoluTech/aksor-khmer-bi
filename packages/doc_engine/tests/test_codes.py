"""QR codes and barcodes drawn from a value at render time (codes.py) -- in Word, Excel and HTML templates, by template
function or by a spec in the data. Every test reads the codes back out of the output with a real decoder (zxing-cpp),
so "it rendered a picture" is never mistaken for "it scans"."""
import io
import os
import re
import shutil
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image

from doc_engine import codes, render
from doc_engine.config import SOFFICE_BIN
from doc_engine.segmentation import segment_generic

zxingcpp = pytest.importorskip("zxingcpp")

INVOICE = "EACL2007091100012163"
KHMER = "វិក្កយបត្រ លេខ ០០១"  # unspaced Khmer: segmentation will put zero-width spaces into it

needs_soffice = pytest.mark.skipif(not (shutil.which(SOFFICE_BIN) or os.path.exists(SOFFICE_BIN)), reason="LibreOffice isn't installed")


def read(png: bytes) -> set[tuple[str, str]]:
    return {(r.format.name, r.text) for r in zxingcpp.read_barcodes(Image.open(io.BytesIO(png)))}


# --- the codes themselves ----------------------------------------------------------------------------------------------


@pytest.mark.parametrize("kind, value, expected", [
    ("code128", INVOICE, ("Code128", INVOICE)),
    ("Code 128", "ab-12", ("Code128", "ab-12")),
    ("code39", "inv-2026-001", ("Code39", "INV-2026-001")),  # capitals only; no check character added
    ("EAN-13", "590123412345", ("EAN13", "5901234123457")),  # the check digit is worked out
    ("ean13", "5901234123457", ("EAN13", "5901234123457")),
])
def test_each_barcode_scans_back(kind, value, expected):
    assert read(codes.barcode_png(value, kind)) == {expected}


def test_a_qr_code_holds_khmer_without_segmentation_marks():
    assert read(codes.qr_png(KHMER, "H")) == {("QRCode", KHMER)}
    segmented = segment_generic(KHMER)
    assert "​" in segmented  # what reaches a template function...
    assert read(codes.qr_png(segmented)) == {("QRCode", KHMER)}  # ...and what the code holds


@pytest.mark.parametrize("call, message", [
    (lambda: codes.barcode_png("59012341234X", "ean13"), "12 digits"),
    (lambda: codes.barcode_png("5901234123458", "ean13"), "check digit should be 7"),
    (lambda: codes.barcode_png("ក", "code39"), "Code 39 can't hold"),
    (lambda: codes.barcode_png("ក", "code128"), "plain ASCII only"),
    (lambda: codes.barcode_png("1", "upc"), "Unknown barcode type"),
    (lambda: codes.qr_png(""), "needs a value"),
    (lambda: codes.qr_png("x", "Z"), "L, M, Q or H"),
])
def test_a_value_that_cant_be_drawn_says_why(call, message):
    with pytest.raises(codes.CodeError, match=message):
        call()


def test_a_barcode_svg_scales_to_its_css_size():
    # A viewBox, so a browser scales it to `width: 40mm` instead of showing it at its natural size.
    svg = codes.barcode_svg("A-1", "code128").decode()
    assert re.search(r'<svg\b[^>]*\bviewBox="0 0 [\d.]+ [\d.]+"', svg)


def test_only_a_real_spec_is_taken_for_a_code():
    assert codes.is_code_spec({"qr": "x"}) and codes.is_code_spec({"qr": 12, "width_mm": 20, "level": "H"})
    assert codes.is_code_spec({"barcode": "ean-13", "value": "590123412345", "text": True})
    # Ordinary rows that happen to have a "qr" or "barcode" column stay data.
    assert not codes.is_code_spec({"qr": "https://example.com", "name": "Shop"})
    assert not codes.is_code_spec({"barcode": "8850123456789", "value": 10})
    assert not codes.is_code_spec({"barcode": "code128"})


def test_segmentation_leaves_a_spec_alone():
    spec = {"qr": KHMER, "width_mm": 20}
    assert segment_generic({"code": spec})["code"] == spec


# --- Word ----------------------------------------------------------------------------------------------------------


def _docx(tmp_path, *paragraphs, header=None, table_rows=None):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, text in enumerate(row):
                table.cell(r, c).text = text
    if header:
        doc.sections[0].header.paragraphs[0].text = header
    path = tmp_path / "t.docx"
    doc.save(path)
    return path


def _docx_pictures(docx: bytes) -> list[set]:
    with zipfile.ZipFile(io.BytesIO(docx)) as z:
        return [read(z.read(n)) for n in sorted(z.namelist()) if n.startswith("word/media/")]


def test_word_functions_body_table_and_header(tmp_path):
    path = _docx(
        tmp_path,
        "{{ qrcode(no, width_mm=30) }}",
        "{{ barcode(sku, 'ean13', width_mm=40, text=true) }}",
        header="{{ barcode(no) }}",
        table_rows=[["{%tr for i in items %}"], ["{{ barcode(i.code, 'code39') }}"], ["{%tr endfor %}"]],
    )
    out = render({"no": INVOICE, "sku": "590123412345", "items": [{"code": "A-1"}, {"code": "B-2"}]}, "docx", template_path=path)
    found = set().union(*_docx_pictures(out))
    assert found == {("QRCode", INVOICE), ("EAN13", "5901234123457"), ("Code128", INVOICE), ("Code39", "A-1"), ("Code39", "B-2")}
    doc = Document(io.BytesIO(out))
    assert len(doc.inline_shapes) == 4  # body pictures (the header's lives in its own part)


def test_word_spec_in_the_data_and_a_field_named_barcode_still_wins(tmp_path):
    path = _docx(tmp_path, "{{ label }}", "Item barcode: {{ barcode }}")
    out = render({"label": {"barcode": "code128", "value": "PKG-77", "width_mm": 45}, "barcode": "8850123456789"}, "docx", template_path=path)
    assert set().union(*_docx_pictures(out)) == {("Code128", "PKG-77")}
    assert "Item barcode: 8850123456789" in [p.text for p in Document(io.BytesIO(out)).paragraphs]


@pytest.mark.parametrize("value", ["5901234123457", 590123412345])
def test_a_label_draws_its_own_barcode_field(tmp_path, value):
    path = _docx(tmp_path, "{{ barcode(barcode, 'ean13') }}", "EAN {{ barcode }}")
    out = render({"barcode": value}, "docx", template_path=path)
    assert set().union(*_docx_pictures(out)) == {("EAN13", "5901234123457")}
    assert f"EAN {value}" in [p.text for p in Document(io.BytesIO(out)).paragraphs]  # still prints as itself


def test_a_field_named_qrcode_that_is_a_list_stays_data(tmp_path):
    path = _docx(tmp_path, "{% for q in qrcode %}{{ q }} {% endfor %}")
    out = render({"qrcode": ["a", "b"]}, "docx", template_path=path)
    assert Document(io.BytesIO(out)).paragraphs[0].text == "a b "


@needs_soffice
def test_word_codes_scan_from_the_printed_page(tmp_path):
    path = _docx(tmp_path, "{{ qrcode(no, width_mm=35) }}", "{{ barcode(no, width_mm=90) }}")
    page = render({"no": INVOICE}, "png", template_path=path)
    assert {("QRCode", INVOICE), ("Code128", INVOICE)} <= read(page)


@pytest.mark.parametrize("text", [False, True])
def test_a_barcode_height_sets_its_proportions(tmp_path, text):
    path = _docx(tmp_path, "{{ barcode(no, width_mm=50, height_mm=10.5, text=t) }}")
    out = render({"no": INVOICE, "t": text}, "docx", template_path=path)
    shape = Document(io.BytesIO(out)).inline_shapes[0]
    assert abs(shape.width.mm - 50) < 0.1 and abs(shape.height.mm - 10.5) < 0.3
    assert set().union(*_docx_pictures(out)) == {("Code128", INVOICE)}
    with pytest.raises(codes.CodeError, match="height_mm"):
        codes.barcode_png(INVOICE, height_mm=-1)


def test_a_bad_value_fails_the_render_with_its_reason(tmp_path):
    path = _docx(tmp_path, "{{ barcode(sku, 'ean13') }}")
    with pytest.raises(codes.CodeError, match="check digit"):
        render({"sku": "5901234123458"}, "docx", template_path=path)


# --- Excel -------------------------------------------------------------------------------------------------------------


def _xlsx(tmp_path, cells: dict[str, str]):
    wb = Workbook()
    ws = wb.active
    for ref, text in cells.items():
        ws[ref] = text
    path = tmp_path / "t.xlsx"
    wb.save(path)
    return path


def test_excel_functions_and_specs_become_pictures_at_their_cells(tmp_path):
    path = _xlsx(tmp_path, {
        "A1": "Invoice", "B1": "{{ qrcode(no, width_mm=30) }}",
        "A2": "{{ label }}",
        "A3": "SKU {{ barcode(sku, 'ean-13') }}",
    })
    out = render({"no": KHMER, "label": {"qr": INVOICE}, "sku": "590123412345"}, "xlsx", template_path=path)
    ws = load_workbook(io.BytesIO(out)).active
    pictures = {p.anchor._from.row + 1: (p.anchor._from.col + 1, read(p._data())) for p in ws._images}
    assert pictures == {
        1: (2, {("QRCode", KHMER)}),
        2: (1, {("QRCode", INVOICE)}),
        3: (1, {("EAN13", "5901234123457")}),
    }
    assert ws["B1"].value is None and ws["A2"].value is None and ws["A3"].value == "SKU"  # the marker is gone, the text stays
    assert ws["A1"].value == "Invoice"
    qr = next(p for p in ws._images if p.anchor._from.row == 0)
    assert qr.anchor.ext.width == pixels_to_EMU(round(30 * 96 / 25.4))  # placed 30 mm wide


def test_excel_codes_in_a_repeated_row(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["B2"], ws["A3"] = "{% for i in items %}", "{{ i.name }}", "{{ barcode(i.code, 'code128', width_mm=40) }}", "{% endfor %}"
    path = tmp_path / "rows.xlsx"
    wb.save(path)
    out = render({"items": [{"name": "ប្រេង", "code": "F-1"}, {"name": "សាំង", "code": "F-2"}]}, "xlsx", template_path=path)
    sheet = load_workbook(io.BytesIO(out)).active
    got = sorted((p.anchor._from.row, read(p._data()).pop()[1]) for p in sheet._images)
    assert [code for _, code in got] == ["F-1", "F-2"] and got[0][0] < got[1][0]  # one per row, in order


# --- HTML --------------------------------------------------------------------------------------------------------------


def test_html_functions_draw_sharp_svg_that_scans(tmp_path):
    path = tmp_path / "t.html"
    path.write_text(
        '<html><body><p>{{ title }}</p>'
        '<img src="{{ qrcode(no, level=\'Q\') }}" style="width:40mm">'
        '<div style="height:10mm"></div>'
        '<img src="{{ barcode(sku, \'ean13\', text=true) }}" style="width:70mm">'
        "</body></html>",
        encoding="utf-8",
    )
    page = render({"title": KHMER, "no": KHMER, "sku": "590123412345"}, "png", "weasyprint", template_path=path)
    assert read(page) == {("QRCode", KHMER), ("EAN13", "5901234123457")}
    assert codes.svg_data_uri(codes.qr_svg("x")).startswith("data:image/svg+xml;base64,")
