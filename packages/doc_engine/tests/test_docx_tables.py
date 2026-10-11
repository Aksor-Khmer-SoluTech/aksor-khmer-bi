"""Excel (tables): a rendered .docx report becomes an .xlsx workbook (engines/docx_tables.py) -- a sheet per table,
or per page when the template has page breaks; the text around the tables comes along; numbers stay numbers and keep
their look."""
import io

from docx import Document
from docx.enum.section import WD_SECTION
from openpyxl import load_workbook

from doc_engine.engines.docx_tables import docx_tables_to_xlsx

ZWSP = "​"


def _docx(build) -> bytes:
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            table.rows[r].cells[c].text = text
    return table


def _book(build):
    return load_workbook(io.BytesIO(docx_tables_to_xlsx(_docx(build))))


def _rows(ws) -> list[list]:
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_each_table_gets_a_sheet_with_the_lines_around_it():
    def build(doc):
        doc.add_heading("Monthly invoices", 1)
        doc.add_paragraph("Company: Aksor Co")
        doc.add_heading("Invoices", 2)
        _table(doc, [["Invoice", "Total"], ["INV-1", "1,240.00"]])
        doc.add_paragraph("Customers: who bought")  # a ':' can't be in a sheet name
        _table(doc, [["Name"], ["Sokha"]])
        _table(doc, [["Untitled"], ["x"]])
        doc.add_paragraph("Printed by Aksor")

    wb = _book(build)
    assert wb.sheetnames == ["Invoices", "Customers  who bought", "Table 3"]  # the nearest heading names the sheet
    assert _rows(wb["Invoices"]) == [
        ["Monthly invoices", None],
        ["Company: Aksor Co", None],
        ["Invoices", None],
        [None, None],
        ["Invoice", "Total"],
        ["INV-1", 1240],
    ]
    assert wb["Invoices"]["A1"].font.bold and not wb["Invoices"]["A2"].font.bold  # headings stand out
    assert _rows(wb["Table 3"]) == [["Untitled"], ["x"], [None], ["Printed by Aksor"]]  # lines after go below


def test_numbers_are_numbers_and_keep_their_look():
    def build(doc):
        _table(
            doc,
            [
                ["a", "b", "c", "d", "e", "f", "g", "h", "i"],
                ["1240", "1,240.00", "-3.5", "12.5%", "(1,240.00)", "00123", "012 345 678", "2026-09-02", "1,234,567"],
            ],
        )

    ws = _book(build).active
    values = [c.value for c in ws[2]]
    formats = [c.number_format for c in ws[2]]
    assert values == [1240, 1240, -3.5, 0.125, -1240, "00123", "012 345 678", "2026-09-02", 1234567]
    assert formats[:5] == ["General", "#,##0.00", "0.0", "0.0%", "#,##0.00;(#,##0.00)"]
    assert formats[8] == "#,##0"


def test_a_paged_template_gives_a_sheet_per_page():
    def build(doc):
        doc.add_heading("Summary", 1)
        _table(doc, [["Branch", "Total"], ["PP", "10"]])
        doc.add_page_break()
        doc.add_heading("Detail", 1)
        doc.add_paragraph("Every invoice")
        _table(doc, [["Invoice"], ["INV-1"]])
        _table(doc, [["Note"], ["n"]])  # a second table on the same page stays on its sheet
        doc.add_section(WD_SECTION.NEW_PAGE)
        doc.add_paragraph("Signatures")

    wb = _book(build)
    assert wb.sheetnames == ["Summary", "Detail", "Signatures"]
    assert _rows(wb["Detail"]) == [["Detail"], ["Every invoice"], [None], ["Invoice"], ["INV-1"], [None], ["Note"], ["n"]]
    assert wb["Detail"].auto_filter.ref == "A4:A5"  # one filter per sheet: the first table's


def test_khmer_keeps_its_text_without_the_layout_breaks():
    def build(doc):
        doc.add_paragraph(f"ហាង{ZWSP}លក់")
        _table(doc, [["ឈ្មោះ"], [f"ហាង{ZWSP}លក់ទំនិញ{ZWSP}ស្រីពៅ"]])

    wb = _book(build)
    assert wb.sheetnames == ["ហាងលក់"]
    assert _rows(wb.active)[-1] == ["ហាងលក់ទំនិញស្រីពៅ"]


def test_a_merged_cell_is_written_once():
    def build(doc):
        table = _table(doc, [["Group", ""], ["a", "b"]])
        table.rows[0].cells[0].merge(table.rows[0].cells[1])
        table.rows[0].cells[0].text = "Group"

    rows = _rows(_book(build).active)
    assert rows[0] == ["Group", None] and rows[1] == ["a", "b"]


def test_the_header_row_is_bold_frozen_and_filterable():
    def build(doc):
        doc.add_paragraph("Staff")
        _table(doc, [["Name", "Total"], ["a", "1"], ["b", "2"]])

    ws = _book(build).active
    assert ws["A3"].font.bold and not ws["A4"].font.bold
    assert ws.freeze_panes == "A4" and ws.auto_filter.ref == "A3:B5"


def test_a_document_without_tables_gives_its_lines():
    def build(doc):
        doc.add_heading("Certificate", 1)
        doc.add_paragraph("Awarded to Sokha")

    wb = _book(build)
    assert wb.sheetnames == ["Report"] and _rows(wb.active) == [["Certificate"], ["Awarded to Sokha"]]


def test_duplicate_names_get_numbered():
    def build(doc):
        for _ in range(2):
            doc.add_paragraph("Totals")
            _table(doc, [["x"], ["1"]])

    assert _book(build).sheetnames == ["Totals", "Totals (2)"]


# --- the table's look ----------------------------------------------------------------------------------------------


def _shade(cell, hex_colour):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(shd)


def test_fonts_shading_and_alignment_come_along():
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    def build(doc):
        table = _table(doc, [["Name", "Total"], ["Sokha", "1,240.00"]])
        header = table.rows[0].cells[0]
        _shade(header, "0F3D6E")
        run = header.paragraphs[0].runs[0]
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(12)
        amount = table.rows[1].cells[1]
        amount.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        amount.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        name = table.rows[1].cells[0].paragraphs[0].runs[0]
        name.italic = True

    ws = _book(build).active
    assert ws["A1"].fill.fgColor.rgb == "FF0F3D6E"
    assert ws["A1"].font.color.rgb == "FFFFFFFF" and ws["A1"].font.size == 12 and ws["A1"].font.bold
    assert ws["A2"].font.italic and not ws["A2"].font.bold
    assert ws["B2"].alignment.horizontal == "right" and ws["B2"].alignment.vertical == "center"
    assert ws["B2"].value == 1240 and ws["B2"].number_format == "#,##0.00"  # the look doesn't cost the number


def test_merged_cells_are_real_merges_across_and_down():
    def build(doc):
        table = _table(doc, [["Group", "", "x"], ["a", "b", "y"], ["c", "d", ""]])
        table.rows[0].cells[0].merge(table.rows[0].cells[1])
        table.rows[0].cells[0].text = "Group"
        table.rows[1].cells[2].merge(table.rows[2].cells[2])
        table.rows[1].cells[2].text = "y"

    ws = _book(build).active
    assert sorted(str(r) for r in ws.merged_cells.ranges) == ["A1:B1", "C2:C3"]
    assert ws["A1"].value == "Group" and ws["C2"].value == "y"
    assert _rows(ws)[2][:2] == ["c", "d"]


def test_borders_only_when_the_word_table_has_them():
    def build(doc):
        grid = _table(doc, [["Bordered"], ["1"]])
        grid.style = "Table Grid"
        doc.add_paragraph("Plain")
        _table(doc, [["Plain"], ["2"]])

    wb = _book(build)
    assert wb["Table 1"]["A2"].border.left.style == "thin"
    assert wb["Plain"]["A3"].border.left.style is None


def test_a_line_keeps_its_own_look():
    from docx.shared import RGBColor

    def build(doc):
        run = doc.add_paragraph().add_run("Confidential")
        run.bold = True
        run.font.color.rgb = RGBColor(0xA0, 0x3E, 0x26)
        _table(doc, [["x"], ["1"]])

    cell = _book(build).active["A1"]
    assert cell.value == "Confidential" and cell.font.bold and cell.font.color.rgb == "FFA03E26"


def test_a_title_row_across_the_table_is_not_taken_for_the_headers():
    def build(doc):
        table = _table(doc, [["September 2026", "", ""], ["Invoice", "Customer", "Total"], ["INV-1", "Sokha", "10"]])
        title = table.rows[0].cells[0].merge(table.rows[0].cells[2])
        title.text = "September 2026"

    ws = _book(build).active
    assert ws.freeze_panes == "A3" and ws.auto_filter.ref == "A2:C3"
    assert ws["A2"].font.bold and ws["B2"].font.bold


def test_a_built_in_table_style_colours_the_header_and_bands():
    def build(doc):
        table = _table(doc, [["Name", "Total"], ["a", "1"], ["b", "2"], ["c", "3"]])
        table.style = "Medium Shading 1 Accent 1"

    ws = _book(build).active
    assert ws["A1"].fill.fgColor.rgb == "FF4F81BD"  # the style's header row
    assert ws["A2"].fill.fgColor.rgb == "FFD3DFEE"  # band 1
    assert ws["A3"].fill.fill_type is None  # band 2: unshaded in this style
    assert ws["A4"].fill.fgColor.rgb == "FFD3DFEE"


def test_a_colour_from_the_theme_palette_is_resolved():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def build(doc):
        table = _table(doc, [["Name"], ["Sokha"]])
        run = table.rows[1].cells[0].paragraphs[0].runs[0]
        colour = OxmlElement("w:color")
        colour.set(qn("w:val"), "auto")
        colour.set(qn("w:themeColor"), "accent2")  # Word's default theme: accent2 = C0504D
        run._r.get_or_add_rPr().append(colour)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "auto")
        shd.set(qn("w:themeFill"), "accent1")  # 4F81BD
        shd.set(qn("w:themeFillTint"), "33")  # 20% of it, the rest white
        table.rows[1].cells[0]._tc.get_or_add_tcPr().append(shd)

    cell = _book(build).active["A2"]
    assert cell.font.color.rgb == "FFC0504D"
    assert cell.fill.fgColor.rgb == "FFDCE6F2"


def test_mixed_formatting_in_one_cell_becomes_rich_text():
    def build(doc):
        table = _table(doc, [["Note"], [""]])
        paragraph = table.rows[1].cells[0].paragraphs[0]
        paragraph.add_run("Total: ").bold = True
        paragraph.add_run("1,240 (unaudited)")

    value = load_workbook(io.BytesIO(docx_tables_to_xlsx(_docx(build))), rich_text=True).active["A2"].value
    assert str(value) == "Total: 1,240 (unaudited)"
    blocks = list(value)
    assert blocks[0].font.b is True and blocks[0].text == "Total: "
    assert isinstance(blocks[1], str) or not blocks[1].font.b  # the rest stays plain


def test_column_widths_follow_the_word_table():
    from docx.shared import Inches

    def build(doc):
        table = _table(doc, [["Narrow", "Wide"], ["a", "b"]])
        table.autofit = False
        for row in table.rows:
            row.cells[0].width = Inches(0.8)
            row.cells[1].width = Inches(3.0)
        table.columns[0].width = Inches(0.8)
        table.columns[1].width = Inches(3.0)

    ws = _book(build).active
    assert ws.column_dimensions["B"].width > 2.5 * ws.column_dimensions["A"].width
