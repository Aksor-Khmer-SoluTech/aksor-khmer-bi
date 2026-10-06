"""docx templates that real-world tools produce, and values real-world data contains."""
import io
import re
import zipfile
from collections import Counter

import pytest
from docx import Document

from doc_engine.engines import libreoffice_engine as lo

STANDARD = "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties"
WPS = "http://schemas.openxmlformats.org/officedocument/2006/relationships/metadata/core-properties"


def _template(tmp_path, *paragraphs, wps=False):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    if wps:  # what WPS Office writes: the core-properties relationship under the wrong type URL
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                body = src.read(item.filename)
                if item.filename == "_rels/.rels":
                    assert STANDARD.encode() in body
                    body = body.replace(STANDARD.encode(), WPS.encode())
                dst.writestr(item, body)
        data = out.getvalue()
    path = tmp_path / "template.docx"
    path.write_bytes(data)
    return path


def _names(docx_bytes):
    return [i.filename for i in zipfile.ZipFile(io.BytesIO(docx_bytes)).infolist()]


def test_wps_core_properties_type_no_longer_duplicates_the_core_part(tmp_path):
    out = lo.render_docx({"n": "x"}, _template(tmp_path, "{{ n }}", wps=True))
    names = _names(out)
    assert [n for n, c in Counter(names).items() if c > 1] == []
    assert [p.text for p in Document(io.BytesIO(out)).paragraphs][:1] == ["x"]


def test_a_wps_template_converts_to_pdf(tmp_path):
    pdf = lo.render_pdf({"n": "x"}, _template(tmp_path, "{{ n }}", wps=True))
    assert pdf.startswith(b"%PDF")


def test_a_standard_template_is_left_byte_for_byte(tmp_path):
    path = _template(tmp_path, "{{ n }}")
    assert lo._normalized_template(path).getvalue() == path.read_bytes()


def test_duplicate_entries_are_dropped_keeping_the_first():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        import warnings

        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            z.writestr("a.xml", "first")
            z.writestr("a.xml", "second")
            z.writestr("b.xml", "b")
    fixed = lo._without_duplicate_entries(buf.getvalue())
    z = zipfile.ZipFile(io.BytesIO(fixed))
    assert z.namelist() == ["a.xml", "b.xml"] and z.read("a.xml") == b"first"


@pytest.mark.parametrize("value", ["R&D Dept", "a<b>c", "5 > 3 & 2 < 4", "x'y \"z\""])
def test_special_characters_in_data_survive(tmp_path, value):
    out = lo.render_docx({"v": value}, _template(tmp_path, "Name: {{ v }}"))
    assert [p.text for p in Document(io.BytesIO(out)).paragraphs][0] == f"Name: {value}"
