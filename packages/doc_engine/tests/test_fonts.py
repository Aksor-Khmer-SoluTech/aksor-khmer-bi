"""Fonts added at runtime (doc_engine.fonts): found by their own family name, and handed to each rendering path.

The font used here is built on the spot -- a tiny TrueType font with one visible glyph, named "AksorTestSans" -- so
the tests need no font file and can tell, from the font names inside the output PDF, whether it was really used."""
import io
import os
import shutil
import subprocess
import sys

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from fontTools.fontBuilder import FontBuilder
from fontTools.pens.ttGlyphPen import TTGlyphPen

from doc_engine import fonts
from doc_engine.config import SOFFICE_BIN
from doc_engine.engines import libreoffice_engine, weasyprint_engine

FAMILY = "AksorTestSans"


def build_font(family: str = FAMILY) -> bytes:
    pen = TTGlyphPen(None)
    pen.moveTo((100, 0))
    pen.lineTo((100, 700))
    pen.lineTo((600, 700))
    pen.lineTo((600, 0))
    pen.closePath()
    square = pen.glyph()
    empty = TTGlyphPen(None).glyph()
    builder = FontBuilder(1000, isTTF=True)
    builder.setupGlyphOrder([".notdef", "A", "space"])
    builder.setupCharacterMap({ord("A"): "A", ord(" "): "space"})
    builder.setupGlyf({".notdef": square, "A": square, "space": empty})
    builder.setupHorizontalMetrics({".notdef": (700, 100), "A": (700, 100), "space": (300, 0)})
    builder.setupHorizontalHeader(ascent=800, descent=-200)
    builder.setupNameTable({"familyName": family, "styleName": "Regular", "uniqueFontIdentifier": family, "fullName": family, "psName": family + "-Regular", "version": "Version 1.000"})
    builder.setupOS2(sTypoAscender=800, sTypoDescender=-200, usWinAscent=800, usWinDescent=200)
    builder.setupPost()
    out = io.BytesIO()
    builder.save(out)
    return out.getvalue()


@pytest.fixture
def font_dir(tmp_path, monkeypatch):
    folder = tmp_path / "fonts"
    folder.mkdir()
    (folder / "one.ttf").write_bytes(build_font())
    monkeypatch.setenv(fonts.ENV_VAR, str(folder))
    return folder


def test_no_folder_means_no_fonts(monkeypatch):
    monkeypatch.delenv(fonts.ENV_VAR, raising=False)
    assert fonts.installed() == []
    assert fonts.font_face_css("<p>anything</p>") == ""


def test_a_font_is_known_by_the_family_inside_the_file_not_its_filename(font_dir):
    assert [family for _path, family in fonts.installed()] == [FAMILY]


def test_files_that_are_not_fonts_are_ignored(font_dir):
    (font_dir / "broken.ttf").write_bytes(b"not a font at all")
    (font_dir / "notes.txt").write_text("hello")
    assert [family for _path, family in fonts.installed()] == [FAMILY]


def test_the_fonts_are_linked_into_a_libreoffice_profile(font_dir, tmp_path):
    profile = tmp_path / "profile"
    assert fonts.install_into_profile(profile) == 1
    assert (profile / "user" / "fonts" / "one.ttf").read_bytes() == (font_dir / "one.ttf").read_bytes()
    assert fonts.install_into_profile(tmp_path / "other") == 1  # a second conversion gets its own profile


def test_font_face_css_covers_only_the_families_the_html_names(font_dir):
    assert fonts.font_face_css("<p style='font-family: Arial'>x</p>") == ""
    css = fonts.font_face_css(f"<p style='font-family: {FAMILY.lower()}'>x</p>")
    assert f'font-family:"{FAMILY}"' in css and "src:url(data:font/ttf;base64," in css


@pytest.mark.parametrize("name, ok", [("Noto Sans Khmer", True), ("", False), ('a"b', False), ("x" * 200, False), (7, False), (None, False)])
def test_a_family_named_in_a_chart_spec_is_checked(name, ok):
    assert (fonts.safe_family(name) is not None) is ok


def _fonts_in(pdf: bytes, tmp_path) -> str:
    """The font names `pdffonts` (poppler) reads out of the PDF -- WeasyPrint packs its font dictionaries into
    compressed streams, so grepping the bytes would find nothing."""
    path = tmp_path / "out.pdf"
    path.write_bytes(pdf)
    return subprocess.run(["pdffonts", str(path)], capture_output=True, text=True, check=True).stdout


needs_pdffonts = pytest.mark.skipif(shutil.which("pdffonts") is None, reason="poppler's pdffonts isn't installed")


@needs_pdffonts
def test_weasyprint_renders_html_with_an_added_font(font_dir, tmp_path):
    template = tmp_path / "t.html"
    template.write_text(f"<html><body><p style='font-family: \"{FAMILY}\"; font-size: 30px'>AAA</p></body></html>")
    assert FAMILY in _fonts_in(weasyprint_engine.render_pdf({}, template), tmp_path)


@needs_pdffonts
def test_weasyprint_without_the_font_does_not_use_it(tmp_path, monkeypatch):
    monkeypatch.delenv(fonts.ENV_VAR, raising=False)
    template = tmp_path / "t.html"
    template.write_text(f"<html><body><p style='font-family: \"{FAMILY}\"'>AAA</p></body></html>")
    assert FAMILY not in _fonts_in(weasyprint_engine.render_pdf({}, template), tmp_path)


@pytest.mark.skipif(not (shutil.which(SOFFICE_BIN) or os.path.exists(SOFFICE_BIN)), reason="LibreOffice isn't installed")
@pytest.mark.skipif(sys.platform == "darwin", reason="LibreOffice for macOS ignores <profile>/user/fonts (install fonts in ~/Library/Fonts there)")
def test_libreoffice_renders_a_docx_with_an_added_font(font_dir, tmp_path):
    doc = Document()
    run = doc.add_paragraph().add_run("AAA")
    run.font.size = Pt(24)
    fonts_el = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts_el.set(qn(attr), FAMILY)
    template = tmp_path / "t.docx"
    doc.save(str(template))
    pdf = libreoffice_engine.render_pdf({}, template)
    assert FAMILY.encode() in pdf
