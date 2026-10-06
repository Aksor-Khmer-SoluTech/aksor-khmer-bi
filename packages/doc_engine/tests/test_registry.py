import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from doc_engine import render


def test_unknown_format_raises(tmp_path):
    doc = Document()
    doc.add_paragraph("Hello {{ customer_name }}")
    path = tmp_path / "custom.docx"
    doc.save(str(path))
    with pytest.raises(ValueError):
        render({"customer_name": "x"}, "bmp", template_path=path)  # type: ignore[arg-type]


@pytest.fixture
def custom_template(tmp_path):
    """A minimal docx with its own Jinja2 placeholder, standing in for a
    user-uploaded report template.
    """
    doc = Document()
    doc.add_paragraph("Hello {{ customer_name }}, total due: {{ amount }}")
    path = tmp_path / "custom.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def html_template(tmp_path):
    """A minimal HTML template, standing in for a user-uploaded
    template_ext="html" report."""
    path = tmp_path / "custom.html"
    path.write_text("<html><body>Hello {{ customer_name }}</body></html>", encoding="utf-8")
    return path


def test_custom_template_path_docx(custom_template):
    result = render({"customer_name": "សុខ សុភា", "amount": 100}, "docx", template_path=custom_template)
    assert result[:2] == b"PK"


def test_render_extra_terms_reaches_segmentation(custom_template):
    # A made-up word ICU would otherwise split -- proves render()'s
    # extra_terms kwarg actually reaches segment_generic/process_text,
    # not just that the parameter is accepted.
    made_up_word = "សាកល្បងដប់ពីរ"
    without = render({"customer_name": made_up_word, "amount": 1}, "docx", template_path=custom_template)
    with_terms = render(
        {"customer_name": made_up_word, "amount": 1}, "docx", template_path=custom_template, extra_terms=[made_up_word]
    )
    assert without != with_terms


def test_render_exclude_terms_reaches_segmentation(custom_template):
    builtin_term = "ឥណ្ឌូនេស៊ី"  # Indonesia -- a real built-in protected term
    without_exclude = render({"customer_name": builtin_term, "amount": 1}, "docx", template_path=custom_template)
    with_exclude = render(
        {"customer_name": builtin_term, "amount": 1},
        "docx",
        template_path=custom_template,
        exclude_terms=[builtin_term],
    )
    assert without_exclude != with_exclude


def test_custom_template_path_pdf(custom_template):
    result = render({"customer_name": "សុខ សុភា", "amount": 100}, "pdf", template_path=custom_template)
    assert result.startswith(b"%PDF")


def test_custom_template_path_png(custom_template):
    result = render({"customer_name": "សុខ សុភា", "amount": 100}, "png", template_path=custom_template)
    assert result.startswith(b"\x89PNG")


def test_html_template_pdf_weasyprint_backend(html_template):
    result = render({"customer_name": "សុខ សុភា"}, "pdf", backend="weasyprint", template_path=html_template)
    assert result.startswith(b"%PDF")


def test_html_template_png_weasyprint_backend(html_template):
    result = render({"customer_name": "សុខ សុភា"}, "png", backend="weasyprint", template_path=html_template)
    assert result.startswith(b"\x89PNG")


def test_html_template_rejects_libreoffice_backend(html_template):
    with pytest.raises(ValueError):
        render({"customer_name": "x"}, "pdf", backend="libreoffice", template_path=html_template)


@pytest.fixture
def custom_xlsx_template(tmp_path):
    """A minimal xlsx with a `{{ field }}` placeholder plus a hidden
    for/endfor row pair around a body row, for the line-item loop — see
    excel_engine.render_xlsx_template's docstring for why the loop needs
    three rows and why the marker rows are hidden.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "Customer: {{ customer_name }}"
    ws["A2"] = "{% for item in items %}"
    ws.row_dimensions[2].hidden = True
    ws["A3"] = "{{ item.label }}"
    ws["B3"] = "{{ item.amount }}"
    ws["A4"] = "{% endfor %}"
    ws.row_dimensions[4].hidden = True
    path = tmp_path / "custom.xlsx"
    wb.save(str(path))
    return path


def test_custom_template_path_xlsx(custom_xlsx_template):
    result = render(
        {"customer_name": "សុខ សុភា", "items": [{"label": "A", "amount": 1}, {"label": "B", "amount": 2}]},
        "xlsx",
        template_path=custom_xlsx_template,
    )
    assert result[:2] == b"PK"

    from io import BytesIO

    wb = load_workbook(BytesIO(result))
    ws = wb.active
    rows = [tuple(r) for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
    # customer_name went through segment_generic (ZWSP inserted), same as
    # any other string in a custom-template context.
    customer_cell = rows[0][0]
    assert customer_cell.replace("​", "") == "Customer: សុខ សុភា"
    assert ("A", 1) in rows
    assert ("B", 2) in rows


def test_custom_template_path_rejects_weasyprint_backend(custom_template):
    with pytest.raises(ValueError):
        render(
            {"customer_name": "x", "amount": 1},
            "pdf",
            backend="weasyprint",
            template_path=custom_template,
        )


@pytest.fixture
def chart_template(tmp_path):
    """A docx template with a plain {{ field }} placeholder where a
    chart's InlineImage is expected to land — same as any other docxtpl
    placeholder from the template author's point of view.
    """
    doc = Document()
    doc.add_paragraph("Report for {{ customer_name }}")
    doc.add_paragraph("{{ sales_chart }}")
    path = tmp_path / "chart_template.docx"
    doc.save(str(path))
    return path


def test_custom_template_with_embedded_chart(chart_template):
    data = {
        "customer_name": "សុខ សុភា",
        "sales_chart": {
            "chart": "bar",
            "title": "ការលក់ប្រចាំខែ",
            "labels": ["មករា", "កុម្ភៈ", "មីនា"],
            "series": [{"name": "ការលក់", "values": [100, 150, 120]}],
        },
    }
    result = render(data, "docx", template_path=chart_template)
    assert result[:2] == b"PK"

    # A real embedded image increases the docx's zip size meaningfully
    # over the same template with no chart — confirms an image actually
    # landed in the document, not just that rendering didn't crash.
    baseline = render(
        {"customer_name": "សុខ សុភា", "sales_chart": ""}, "docx", template_path=chart_template
    )
    assert len(result) > len(baseline) + 5000

    import zipfile
    from io import BytesIO

    with zipfile.ZipFile(BytesIO(result)) as zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
        assert media, "expected an embedded image under word/media/"


@pytest.fixture
def image_template(tmp_path):
    """Same recipe as chart_template above -- an image lands in a plain
    {{ field }} placeholder exactly like a chart or any other value.
    """
    doc = Document()
    doc.add_paragraph("Report for {{ customer_name }}")
    doc.add_paragraph("{{ employee_photo }}")
    path = tmp_path / "image_template.docx"
    doc.save(str(path))
    return path


def _make_png_bytes() -> bytes:
    from io import BytesIO as _BytesIO

    from PIL import Image

    buf = _BytesIO()
    Image.new("RGB", (40, 20), (0, 128, 255)).save(buf, format="PNG")
    return buf.getvalue()


def test_custom_template_with_embedded_image(image_template):
    """Mirrors test_custom_template_with_embedded_chart above, but for an
    already-resolved image reference (the shape api/app/context_media.py
    hands to doc_engine -- see doc_engine.images' module docstring for
    why doc_engine never sees a bare image_id).
    """
    data = {
        "customer_name": "សុខ សុភា",
        "employee_photo": {"image_bytes": _make_png_bytes(), "width_mm": 30},
    }
    result = render(data, "docx", template_path=image_template)
    assert result[:2] == b"PK"

    baseline = render({"customer_name": "សុខ សុភា", "employee_photo": ""}, "docx", template_path=image_template)
    assert len(result) > len(baseline) + 500

    import zipfile
    from io import BytesIO

    with zipfile.ZipFile(BytesIO(result)) as zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
        assert media, "expected an embedded image under word/media/"


def test_image_spec_type_discriminator_survives_segmentation():
    from doc_engine.segmentation import segment_generic

    spec = {"image_bytes": b"\x89PNG\r\n\x1a\n", "width_mm": 40}
    assert segment_generic({"photo": spec}) == {"photo": spec}


def test_chart_spec_type_discriminator_survives_segmentation():
    # The regression this guards against: segment_generic ICU-segmenting
    # the literal string "bar" would silently break is_chart_spec's
    # `value.get("chart") in CHART_TYPES` check.
    from doc_engine.segmentation import segment_generic

    spec = {"chart": "bar", "title": "ចំណងជើង", "labels": ["ក"], "series": [{"values": [1]}]}
    result = segment_generic({"sales_chart": spec})
    assert result["sales_chart"] == spec  # byte-for-byte untouched
