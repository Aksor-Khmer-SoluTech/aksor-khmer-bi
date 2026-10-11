"""Excel (tables): a rendered .docx report as an .xlsx workbook -- its data, readable as the report.

There is no Word-to-spreadsheet conversion to lean on: LibreOffice has no export filter between the two, and the only
library that does it directly (Aspose.Words) is commercial. So this reads the document (python-docx) and writes the
workbook (openpyxl) itself, the way Oracle BI Publisher turns an RTF layout into Excel:

- **Sheets.** If the template has page breaks (an inserted page break, "page break before", or a section break), each
  page becomes a sheet -- BI Publisher's rule. Without any, each table gets its own sheet. A sheet is named after the
  heading (or line) above its table, else "Table 1" / "Page 1".
- **Text comes along.** The lines around a table -- a title, "Company: Aksor Co" -- sit above it on its sheet (headings
  bold), and lines after the last table below it, so the workbook reads like the report. A document with no table at
  all becomes one sheet of its lines.
- **Numbers are numbers, and look the same.** "1,240.00" is stored as 1240 and shown as 1,240.00 (#,##0.00); "12.5%"
  as 0.125 shown 12.5%; "(1,240.00)" as -1240. Values that only look numeric stay text: IDs with leading zeros
  ("00123"), phone numbers, dates.
- **The table's look comes along:** bold, italic, font size and colour; cell shading; horizontal and vertical
  alignment; merged cells as real Excel merges; thin borders when the Word table has them. The first table's header
  row on a sheet is also frozen and filterable.
- The invisible word breaks Khmer text carries for page layout are removed: in a cell they only get in the way of
  searching, comparing and copying.

- Also carried: Word's built-in table styles (the header row, total row and banded rows they colour), colours picked
  from the document's theme palette, mixed formatting inside a cell (as Excel rich text), and the table's own column
  widths.

Not copied: images and page positions. For a spreadsheet designed cell by cell, use an .xlsx template.
"""
from __future__ import annotations

import io
import re
from contextvars import ContextVar
from dataclasses import dataclass, field

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_ZWSP = "​"
_NUMBER = re.compile(r"^(?P<sign>[-+]?)(?P<int>\d{1,3}(?:,\d{3})+|\d+)(?:\.(?P<dec>\d+))?(?P<pct>%?)$")
_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")


def _value(text: str) -> tuple[object, str | None]:
    """A cell's text as Excel should hold it -- (value, number format or None) -- a number when it is one, else text."""
    clean = text.replace(_ZWSP, "").strip()
    negative = clean.startswith("(") and clean.endswith(")")  # accounting style: (1,240.00)
    body = clean[1:-1].strip() if negative else clean
    m = _NUMBER.match(body)
    if not m or (len(m["int"]) > 1 and m["int"].startswith("0")):  # 00123 is an ID, not 123
        return clean, None
    number = float(f"{m['int'].replace(',', '')}.{m['dec'] or 0}")
    if m["sign"] == "-" or negative:
        number = -number
    decimals = len(m["dec"] or "")
    if m["pct"]:
        return number / 100, "0" + ("." + "0" * decimals if decimals else "") + "%"
    grouped = "," in m["int"]
    if not decimals and abs(number) < 2**53:
        number = int(number)
    if not grouped and not decimals and not negative:
        return number, None  # plain 1240: Excel's General shows it as written
    fmt = ("#,##0" if grouped else "0") + ("." + "0" * decimals if decimals else "")
    return number, f"{fmt};({fmt})" if negative else fmt


# --- reading the document in order -------------------------------------------------------------------------------


@dataclass
class _Line:
    text: str
    heading: bool
    look: dict = field(default_factory=dict)  # Font arguments taken from the paragraph (_font_look)


_H_ALIGN = {"LEFT": "left", "CENTER": "center", "RIGHT": "right", "JUSTIFY": "justify", "DISTRIBUTE": "distributed"}
_V_ALIGN = {"top": "top", "center": "center", "bottom": "bottom", "both": "center"}
_THIN = Side(style="thin", color="FF9AA3AE")
_BORDERED = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


# --- colours: explicit ones, and the document theme's ---------------------------------------------------------------

# The document's theme colours (dk1, lt1, accent1...), read once per conversion -- a colour picked from Word's theme
# palette is stored as a name ("accent1", lightened by a tint), not a value.
_THEME: ContextVar[dict] = ContextVar("docx_theme", default={})
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_THEME_ALIASES = {
    "text1": "dk1", "dark1": "dk1", "tx1": "dk1", "background1": "lt1", "light1": "lt1", "bg1": "lt1",
    "text2": "dk2", "dark2": "dk2", "tx2": "dk2", "background2": "lt2", "light2": "lt2", "bg2": "lt2",
    "hyperlink": "hlink", "followedhyperlink": "folHlink",
}


def _read_theme(document) -> dict:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    try:
        root = etree.fromstring(document.part.part_related_by(RT.THEME).blob)
    except (KeyError, ValueError, etree.XMLSyntaxError):
        return {}
    scheme = root.find(f".//{_A}clrScheme")
    colours: dict = {}
    for child in scheme if scheme is not None else []:
        srgb, system = child.find(f"{_A}srgbClr"), child.find(f"{_A}sysClr")
        value = srgb.get("val") if srgb is not None else system.get("lastClr") if system is not None else None
        if value:
            colours[child.tag.replace(_A, "")] = value.upper()
    return colours


def _theme_colour(name: str | None, tint: str | None = None, shade: str | None = None) -> str | None:
    """A theme colour by Word's name ("accent1", "text1", "ACCENT_1"), lightened by a tint or darkened by a shade
    (hex 00-FF fractions, as Word stores them), as ARGB -- None when the theme doesn't have it."""
    if not name:
        return None
    key = name.replace("_", "").lower()
    key = _THEME_ALIASES.get(key, key)
    hex6 = _THEME.get().get(key)
    if not hex6:
        return None
    channels = [int(hex6[i : i + 2], 16) for i in (0, 2, 4)]
    if tint:
        f = int(tint, 16) / 255
        channels = [round(c * f + 255 * (1 - f)) for c in channels]
    if shade:
        f = int(shade, 16) / 255
        channels = [round(c * f) for c in channels]
    return "FF" + "".join(f"{c:02X}" for c in channels)


def _run_colour(run) -> str | None:
    colour = run.font.color
    if colour is None or colour.type is None:
        return None
    if colour.rgb is not None:
        return "FF" + str(colour.rgb)
    element = run._r.rPr.find(qn("w:color")) if run._r.rPr is not None else None
    if element is None:
        return None
    return _theme_colour(element.get(qn("w:themeColor")), element.get(qn("w:themeTint")), element.get(qn("w:themeShade")))


def _run_look(run, paragraph) -> dict:
    """openpyxl Font arguments for one run: bold, italic, size and colour, each falling back to the run's and then the
    paragraph's style."""
    styles = [s.font for s in (run.style, paragraph.style) if s is not None]

    def pick(name):
        value = getattr(run.font, name)
        for font in styles:
            if value is not None:
                break
            value = getattr(font, name)
        return value

    look: dict = {}
    if pick("bold"):
        look["bold"] = True
    if pick("italic"):
        look["italic"] = True
    size = pick("size")
    if size is not None:
        look["size"] = round(size.pt, 1)
    colour = _run_colour(run)
    if colour:
        look["color"] = colour
    return look


def _font_look(paragraphs) -> dict:
    """The look of the first run that has text. Empty when nothing is set."""
    for paragraph in paragraphs:
        for run in paragraph.runs:
            if run.text.strip():
                return _run_look(run, paragraph)
    return {}


def _rich_text(paragraphs, base: dict) -> CellRichText | None:
    """The cell as Excel rich text when its words don't all look the same ("Total: 1,240" with only "Total:" bold),
    else None -- a uniform cell just takes one Font."""
    pieces: list[tuple[str, dict]] = []
    for index, paragraph in enumerate(paragraphs):
        if index and pieces:
            pieces.append(("\n", {}))
        for run in paragraph.runs:
            text = run.text.replace(_ZWSP, "")
            if text:
                pieces.append((text, {**base, **_run_look(run, paragraph)}))
    looks = {tuple(sorted(look.items())) for text, look in pieces if text.strip()}
    if len(looks) < 2:
        return None
    blocks = []
    for text, look in pieces:
        if not look:
            blocks.append(text)
            continue
        blocks.append(
            TextBlock(InlineFont(b=look.get("bold"), i=look.get("italic"), sz=look.get("size"), color=look.get("color")), text)
        )
    return CellRichText(*blocks)


def _alignment(paragraphs, tc) -> Alignment | None:
    horizontal = None
    for paragraph in paragraphs:
        if paragraph.text.strip() and paragraph.alignment is not None:
            horizontal = _H_ALIGN.get(paragraph.alignment.name)
            break
    vertical = None
    tc_pr = tc.tcPr
    if tc_pr is not None:
        v_align = tc_pr.find(qn("w:vAlign"))
        if v_align is not None:
            vertical = _V_ALIGN.get(v_align.get(qn("w:val")) or "")
    if horizontal is None and vertical is None:
        return None
    return Alignment(horizontal=horizontal, vertical=vertical, wrap_text=True)


def _shading_colour(shading) -> str | None:
    """ARGB from a w:shd: its explicit fill, else the theme colour it names (themeFill + tint/shade)."""
    if shading is None:
        return None
    colour = shading.get(qn("w:fill"))
    if colour and colour.lower() != "auto" and re.fullmatch(r"[0-9A-Fa-f]{6}", colour):
        return "FF" + colour.upper()
    return _theme_colour(shading.get(qn("w:themeFill")), shading.get(qn("w:themeFillTint")), shading.get(qn("w:themeFillShade")))


def _fill(tc) -> PatternFill | None:
    tc_pr = tc.tcPr
    colour = _shading_colour(tc_pr.find(qn("w:shd")) if tc_pr is not None else None)
    return PatternFill(fill_type="solid", fgColor=colour) if colour else None


# --- Word's built-in table styles: header row, total row, banded rows -------------------------------------------


@dataclass
class _StylePart:
    fill: str | None = None
    look: dict = field(default_factory=dict)


def _style_part(style_pr) -> _StylePart:
    part = _StylePart()
    tc_pr = style_pr.find(qn("w:tcPr"))
    if tc_pr is not None:
        part.fill = _shading_colour(tc_pr.find(qn("w:shd")))
    r_pr = style_pr.find(qn("w:rPr"))
    if r_pr is not None:
        bold = r_pr.find(qn("w:b"))
        if bold is not None and bold.get(qn("w:val")) not in ("0", "false"):
            part.look["bold"] = True
        colour = r_pr.find(qn("w:color"))
        if colour is not None:
            value = colour.get(qn("w:val"))
            if value and value.lower() != "auto" and re.fullmatch(r"[0-9A-Fa-f]{6}", value):
                part.look["color"] = "FF" + value.upper()
            else:
                themed = _theme_colour(colour.get(qn("w:themeColor")), colour.get(qn("w:themeTint")), colour.get(qn("w:themeShade")))
                if themed:
                    part.look["color"] = themed
    return part


def _table_style(table: Table) -> dict[str, _StylePart]:
    """What the table's Word style draws per region -- "firstRow", "lastRow", "band1Horz" -- following the style's
    basedOn chain, limited to the regions the table turns on (its w:tblLook)."""
    parts: dict[str, _StylePart] = {}
    style = table.style
    while style is not None:
        for style_pr in style.element.findall(qn("w:tblStylePr")):
            kind = style_pr.get(qn("w:type"))
            if kind and kind not in parts:
                parts[kind] = _style_part(style_pr)
        style = style.base_style
    tbl_pr = table._tbl.tblPr
    look = tbl_pr.find(qn("w:tblLook")) if tbl_pr is not None else None
    on = {"firstRow": True, "lastRow": False, "band1Horz": True}
    if look is not None:
        bits = int(look.get(qn("w:val")) or "0", 16) if look.get(qn("w:val")) else None
        def flag(attr: str, bit: int, default: bool) -> bool:
            value = look.get(qn(f"w:{attr}"))
            if value is not None:
                return value in ("1", "true")
            return bool(bits & bit) if bits is not None else default
        on = {"firstRow": flag("firstRow", 0x0020, True), "lastRow": flag("lastRow", 0x0040, False),
              "band1Horz": not flag("noHBand", 0x0200, False)}
    return {kind: part for kind, part in parts.items() if on.get(kind, False)}


def _grid_widths(table: Table) -> list[float]:
    """The Word table's column widths, in Excel's units (about one character of the default font each)."""
    grid = table._tbl.tblGrid
    widths = []
    for col in grid.findall(qn("w:gridCol")) if grid is not None else []:
        twips = col.get(qn("w:w"))
        if twips and twips.isdigit():
            widths.append(max(6.0, min(round(int(twips) / 105, 1), 80.0)))
    return widths


def _has_borders(table: Table) -> bool:
    """Whether the Word table draws cell borders: set on the table itself, or by its style (Table Grid and the like)."""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders")) if tbl_pr is not None else None
    if borders is not None:
        return any((edge.get(qn("w:val")) or "none") not in ("none", "nil") for edge in borders)
    style = table.style
    while style is not None:
        el = style.element.find(qn("w:tblPr"))
        found = el.find(qn("w:tblBorders")) if el is not None else None
        if found is not None:
            return any((edge.get(qn("w:val")) or "none") not in ("none", "nil") for edge in found)
        style = style.base_style
    return False


@dataclass
class _Page:
    items: list = field(default_factory=list)  # _Line and Table, in document order

    def has_content(self) -> bool:
        return any(isinstance(i, Table) or i.text for i in self.items)


def _is_heading(paragraph: Paragraph) -> bool:
    name = (paragraph.style.name if paragraph.style is not None else "") or ""
    return name.startswith("Heading") or name == "Title"


def _pages(document) -> list[_Page]:
    """The body split at the author's own page breaks: an inserted page break, "page break before", a section break.
    (Word's remembered page positions, lastRenderedPageBreak, are left alone: they go stale as soon as data flows in.)"""
    pages = [_Page()]
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "tbl":
            pages[-1].items.append(Table(child, document))
            continue
        if tag != "p":
            continue
        paragraph = Paragraph(child, document)
        ppr = child.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:pageBreakBefore")) is not None and pages[-1].has_content():
            pages.append(_Page())
        text = paragraph.text.replace(_ZWSP, "").strip()
        if text:
            pages[-1].items.append(_Line(text, _is_heading(paragraph), _font_look([paragraph])))
        page_break = any(br.get(qn("w:type")) == "page" for br in child.iter(qn("w:br")))
        section_break = ppr is not None and ppr.find(qn("w:sectPr")) is not None
        if (page_break or section_break) and pages[-1].has_content():
            pages.append(_Page())
    return [p for p in pages if p.has_content()]


# --- writing ------------------------------------------------------------------------------------------------------


def _sheet_name(wanted: str, taken: set[str]) -> str:
    base = _SHEET_BAD.sub(" ", wanted).strip()[:31] or "Sheet"
    name, n = base, 2
    while name.lower() in taken:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    taken.add(name.lower())
    return name


class _SheetWriter:
    def __init__(self, ws):
        self.ws = ws
        self.row = 1
        self.widths: dict[int, int] = {}  # from the text, when Word gives no column widths
        self.grid_widths: dict[int, float] = {}  # the Word tables' own column widths
        self.first_table = True

    def _gap(self) -> None:
        if self.row > 1:
            self.row += 1

    def line(self, line: _Line) -> None:
        value, fmt = _value(line.text)
        cell = self.ws.cell(row=self.row, column=1, value=value)
        if fmt:
            cell.number_format = fmt
        look = dict(line.look)
        if line.heading:
            look["bold"] = True
            look.setdefault("size", 13)
        if look:
            cell.font = Font(**look)
        self.row += 1

    def table(self, table: Table) -> None:
        if not table.rows:
            return
        self._gap()
        top = self.row
        columns = 0
        bordered = _has_borders(table)
        # A first row that is one cell across the whole table is a title ("September 2026"), not the column headers:
        # the headers -- bold, frozen, filtered -- are then the row under it.
        first_cells = table.rows[0].cells
        titled = len(table.rows) > 2 and len(first_cells) > 1 and all(c._tc is first_cells[0]._tc for c in first_cells)
        header = top + 1 if titled else top
        styled = _table_style(table)
        last = top + len(table.rows) - 1
        for c, width in enumerate(_grid_widths(table), start=1):
            self.grid_widths[c] = max(self.grid_widths.get(c, 0), width)
        # python-docx repeats a merged cell's object at every grid position it covers (across and down): collect
        # those positions to merge them for real in Excel, writing the value and its look once, at the top left.
        # Keyed by id() *and* holding the element itself: while it is referenced, no other cell can reuse its id
        # (lxml hands out short-lived wrappers, whose ids are recycled the moment one is freed).
        spans: dict[int, tuple[object, list[tuple[int, int]]]] = {}
        for row in table.rows:
            for c, cell in enumerate(row.cells, start=1):
                columns = max(columns, c)
                positions = spans.setdefault(id(cell._tc), (cell._tc, []))[1]
                positions.append((self.row, c))
                target = self.ws.cell(row=self.row, column=c)
                if bordered:
                    target.border = _BORDERED
                if len(positions) > 1:
                    continue  # covered by a merge: borders only
                value, fmt = _value(cell.text)
                target.value = value if value != "" else None
                if fmt:
                    target.number_format = fmt
                # What the table style draws for this row (header, total, banding), under the cell's own formatting.
                # A region the style doesn't turn on falls through to the next: a last row with no total-row style
                # is still a banded row.
                if self.row == top and "firstRow" in styled:
                    region = styled["firstRow"]
                elif self.row == last and "lastRow" in styled:
                    region = styled["lastRow"]
                elif self.row != top and (self.row - top) % 2 == 1:
                    region = styled.get("band1Horz")
                else:
                    region = None
                look = {**(region.look if region else {}), **_font_look(cell.paragraphs)}
                if self.row == header:
                    look["bold"] = True
                rich = _rich_text(cell.paragraphs, {"bold": True} if self.row == header else {}) if not fmt else None
                if rich is not None:
                    target.value = rich
                elif look:
                    target.font = Font(**look)
                fill = _fill(cell._tc)
                if fill is None and region is not None and region.fill:
                    fill = PatternFill(fill_type="solid", fgColor=region.fill)
                if fill is not None:
                    target.fill = fill
                alignment = _alignment(cell.paragraphs, cell._tc)
                if alignment is not None:
                    target.alignment = alignment
                shown = cell.text.replace(_ZWSP, "").strip()
                self.widths[c] = max(self.widths.get(c, 8), min(len(shown) + 2, 60))
            self.row += 1
        for _, positions in spans.values():
            if len(positions) < 2:
                continue
            rows = [r for r, _ in positions]
            cols = [c for _, c in positions]
            first_row, last_row, first_col, last_col = min(rows), max(rows), min(cols), max(cols)
            if len(set(positions)) == (last_row - first_row + 1) * (last_col - first_col + 1):  # a clean rectangle
                self.ws.merge_cells(start_row=first_row, start_column=first_col, end_row=last_row, end_column=last_col)
        if self.first_table and self.row - header > 1:
            # Excel allows one filter per sheet: the first table gets it, with its header row kept in view.
            self.ws.freeze_panes = f"A{header + 1}"
            self.ws.auto_filter.ref = f"A{header}:{get_column_letter(columns)}{self.row - 1}"
        self.first_table = False

    def items(self, items) -> None:
        after_table = False
        for item in items:
            if isinstance(item, Table):
                self.table(item)
                after_table = True
            else:
                if after_table:
                    self._gap()
                    after_table = False
                self.line(item)

    def finish(self) -> None:
        for c in set(self.widths) | set(self.grid_widths):
            self.ws.column_dimensions[get_column_letter(c)].width = self.grid_widths.get(c) or self.widths[c]


def _title(items, fallback: str, *, nearest: bool) -> str:
    """A sheet's name from its lines: a heading over any other line; the one nearest the table (`nearest`, a table's
    own sheet) or the first (a page's sheet, whose top heading names the page)."""
    lines = [i for i in items if isinstance(i, _Line) and len(i.text) <= 60]
    pick = [i for i in lines if i.heading] or lines
    if not pick:
        return fallback
    return (pick[-1] if nearest else pick[0]).text


def docx_tables_to_xlsx(docx_bytes: bytes) -> bytes:
    """A rendered .docx as an .xlsx workbook (see the module docstring)."""
    document = Document(io.BytesIO(docx_bytes))
    theme = _THEME.set(_read_theme(document))
    try:
        return _convert(document)
    finally:
        _THEME.reset(theme)


def _convert(document) -> bytes:
    pages = _pages(document)
    workbook = Workbook()
    workbook.remove(workbook.active)
    taken: set[str] = set()

    def sheet(name: str, items) -> None:
        writer = _SheetWriter(workbook.create_sheet(_sheet_name(name, taken)))
        writer.items(items)
        writer.finish()

    if len(pages) > 1:
        # The template is paged: one sheet per page, like BI Publisher.
        for number, page in enumerate(pages, start=1):
            sheet(_title(page.items, f"Page {number}", nearest=False), page.items)
    else:
        items = pages[0].items if pages else []
        tables = [i for i in items if isinstance(i, Table)]
        if not tables:
            sheet("Report", items or [_Line("This report has no content.", False)])
        else:
            # One sheet per table, carrying the lines above it; the lines after the last table go below it.
            group: list = []
            groups: list[list] = []
            for item in items:
                group.append(item)
                if isinstance(item, Table):
                    groups.append(group)
                    group = []
            groups[-1].extend(group)
            for number, items_ in enumerate(groups, start=1):
                above = [i for i in items_[: next(k for k, i in enumerate(items_) if isinstance(i, Table))]]
                sheet(_title(above, f"Table {number}", nearest=True), items_)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
