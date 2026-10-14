"""An Analyze result as a file: PDF (the HTML engine, Khmer segmented like every report), Word (built here from the
same segmented text, the chart as a picture) or Excel (numbers kept as numbers, in each measure's format).

`table()` lays a pivot out once -- header rows, body rows, the totals -- and every format reads it, so the three
can't disagree about what the export holds."""
from __future__ import annotations

import base64
import io
from datetime import datetime, timezone
from html import escape
from pathlib import Path

_TEMPLATES = Path(__file__).with_name("export_templates")
SERIES = ["#1f5fc4", "#e07a1f", "#18794a", "#9b3fb4", "#c23b4a", "#0f8a8a"]
FORMATS = ("pdf", "docx", "xlsx")
CONTENT_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _name(entry: dict) -> str:
    return entry.get("label_km") or entry.get("label") or entry.get("name") or ""


def key_text(values: list) -> str:
    return " · ".join("(blank)" if v is None else str(v) for v in values) or "All"


def table(result: dict) -> dict:
    """{"headers": [[{text, span, cls}]], "body": [{cls, cells: [{text, value, cls}]}], "value_keys": [...]}: the
    pivot as a plain grid. Columns: one per row field, then each column key x each value, then the row totals (when
    there are columns)."""
    labels = result["labels"]
    values = labels["values"]
    dims = labels["rows"] or [{"label": "", "label_km": None}]
    columns = result["column_keys"] if labels["columns"] else []
    first = [{"text": _name(d), "span": 1, "cls": "dim"} for d in dims]
    if columns:
        top = [{"text": " · ".join(_name(c) for c in labels["columns"]), "span": len(dims), "cls": "dim"}]
        top += [{"text": key_text(ck), "span": len(values), "cls": ""} for ck in columns]
        top += [{"text": "សរុប · Total", "span": len(values), "cls": ""}]
        second = first + [{"text": _name(v), "span": 1, "cls": ""} for _ in range(len(columns) + 1) for v in values]
        headers = [top, second]
    else:
        headers = [first + [{"text": _name(v), "span": 1, "cls": ""} for v in values]]

    def cells(entries: list | None) -> list[dict]:
        if entries is None:
            return [{"text": "—", "value": None, "cls": "empty"} for _ in values]
        return [{"text": e["t"] if e["v"] is not None else "—", "value": e["v"], "cls": "" if e["v"] is not None else "empty"} for e in entries]

    body = []
    for i, rk in enumerate(result["row_keys"]):
        row = [{"text": "(blank)" if v is None else str(v), "value": None, "cls": "dim"} for v in (rk or [""])]
        if columns:
            for cell in result["cells"][i]:
                row += cells(cell)
            row += cells(result["row_totals"][i])
        else:
            row += cells(result["cells"][i][0] if result["cells"][i] else None)
        body.append({"cls": "", "cells": row})
    total = [{"text": "សរុប · Total", "value": None, "cls": "dim"}] + [{"text": "", "value": None, "cls": "dim"} for _ in dims[1:]]
    if columns:
        for ct in result["column_totals"]:
            total += cells(ct)
    total += cells(result["total"])
    body.append({"cls": "total", "cells": total})
    return {"headers": headers, "body": body, "value_keys": [v["key"] for v in values]}


# --- the chart, as SVG ----------------------------------------------------------------------------------------------


def nice_top(value: float) -> float:
    """The axis' top: the first of 1, 1.2, 1.5, 2, 2.5, 3, 4, 5, 6, 8, 10 x a power of ten that reaches `value` --
    so 106M tops at 120M, not 200M."""
    magnitude = 10 ** (len(str(int(max(value, 1)))) - 1)
    return next(step * magnitude for step in (1, 1.2, 1.5, 2, 2.5, 3, 4, 5, 6, 8, 10) if step * magnitude >= value)


def chart(result: dict, kind: str) -> dict | None:
    """A bar or line chart of the first value -- row keys along the axis, one series per column key -- as
    {uri, categories, series, pad_left, pad_right}. Only the marks, the grid and the axis numbers are drawn in SVG:
    WeasyPrint draws SVG text without Khmer shaping (stacked consonants come apart), so the category names and the
    legend are HTML text under the picture, lined up with it by the same margins."""
    if kind not in ("bar", "line") or not result["row_keys"]:
        return None
    labels = result["labels"]
    has_columns = bool(labels["columns"])
    series = [key_text(ck) for ck in result["column_keys"]] if has_columns else [_name(labels["values"][0])]
    cats = [key_text(rk) for rk in result["row_keys"]]
    data = [
        [((result["cells"][ri][si] or [{"v": None}])[0]["v"] if has_columns else (result["cells"][ri][0] or [{"v": None}])[0]["v"]) or 0
         for ri in range(len(cats))]
        for si in range(len(series))
    ]
    top = nice_top(max([v for row in data for v in row] + [1]))
    W, H, L, R, T, B = 760, 220, 92, 16, 14, 6
    iw, ih = W - L - R, H - T - B

    def y(v: float) -> float:
        return T + ih - (v / top) * ih

    def short(v: float) -> str:
        for div, unit in ((1e9, "B"), (1e6, "M"), (1e3, "K")):
            if top >= div * 10:
                return f"{v / div:,.1f}".rstrip("0").rstrip(".") + unit
        return f"{v:,.0f}"

    font = 'font-family="DejaVu Sans, sans-serif" font-size="11" fill="#56607a"'
    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}">']
    for f in (0, 0.25, 0.5, 0.75, 1):
        parts.append(f'<line x1="{L}" x2="{W - R}" y1="{y(f * top):.1f}" y2="{y(f * top):.1f}" stroke="#dfe3ec"/>')
        parts.append(f'<text x="{L - 8}" y="{y(f * top) + 4:.1f}" text-anchor="end" {font}>{escape(short(f * top))}</text>')
    band = iw / len(cats)
    if kind == "bar":
        width = min(46.0, band * 0.78 / len(series))
        for ci in range(len(cats)):
            for si in range(len(series)):
                x = L + ci * band + (band - width * len(series)) / 2 + si * width
                parts.append(f'<rect x="{x:.1f}" y="{y(data[si][ci]):.1f}" width="{max(width - 3, 1):.1f}" height="{T + ih - y(data[si][ci]):.1f}" rx="2" fill="{SERIES[si % len(SERIES)]}"/>')
    else:
        for si in range(len(series)):
            points = " ".join(f"{L + ci * band + band / 2:.1f},{y(data[si][ci]):.1f}" for ci in range(len(cats)))
            parts.append(f'<polyline points="{points}" fill="none" stroke="{SERIES[si % len(SERIES)]}" stroke-width="2.5"/>')
            for ci in range(len(cats)):
                parts.append(f'<circle cx="{L + ci * band + band / 2:.1f}" cy="{y(data[si][ci]):.1f}" r="3.5" fill="{SERIES[si % len(SERIES)]}"/>')
    parts.append("</svg>")
    svg = "".join(parts)
    return {
        "uri": f"data:image/svg+xml;base64,{base64.b64encode(svg.encode()).decode()}",
        "categories": cats,
        "series": [{"name": name, "color": SERIES[si % len(SERIES)]} for si, name in enumerate(series)],
        "pad_left": round(L / W * 100, 2),
        "pad_right": round(R / W * 100, 2),
    }


# --- the files --------------------------------------------------------------------------------------------------------


def _context(result: dict, meta: dict, picture: dict | None) -> dict:
    grid = table(result)
    return {
        "title": meta["title"],
        "subtitle": f"{meta['dataset_title']} · {result['row_count']:,} rows",
        "filters": meta["filters"],
        # The chart's names go through segmentation with everything else; its picture (base64) has nothing to segment.
        # ("graph", not "chart": the engine takes a dict with a "chart" key for one of its own chart specs.)
        "graph": picture,
        "headers": grid["headers"],
        "body": [{"cls": r["cls"], "cells": [{"text": c["text"], "cls": c["cls"]} for c in r["cells"]]} for r in grid["body"]],
        "footer": f"Aksor Khmer · {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')} UTC · {meta['who']}",
        "orientation": "landscape" if len(grid["headers"][-1]) > 7 else "portrait",
    }


def to_pdf(result: dict, meta: dict, kind: str) -> bytes:
    from doc_engine import render

    return render(_context(result, meta, chart(result, kind)), "pdf", "weasyprint", template_path=_TEMPLATES / "analysis.html")


def to_docx(result: dict, meta: dict, kind: str) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt
    from doc_engine import render
    from doc_engine.segmentation import segment_generic

    context = segment_generic(_context(result, meta, None))
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Khmer OS Siemreap"
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set(qn("w:cs"), "Khmer OS Siemreap")
    if context["orientation"] == "landscape":
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    heading = doc.add_paragraph()
    run = heading.add_run(context["title"])
    run.font.size = Pt(15)
    run.font.name = "Khmer OS Muol Light"
    run._element.rPr.rFonts.set(qn("w:cs"), "Khmer OS Muol Light")
    doc.add_paragraph(context["subtitle"])
    if context["filters"]:
        doc.add_paragraph("    ".join(f"{f['label']}: {f['value']}" for f in context["filters"]))
    picture = chart(result, kind)
    if picture:
        # The chart page -- picture, then its names as shaped HTML text -- rasterized for Word.
        png = render({"graph": picture}, "png", "weasyprint", template_path=_TEMPLATES / "chart.html")
        doc.add_picture(io.BytesIO(png), width=Mm(165))
    headers, body = context["headers"], context["body"]
    width = sum(c["span"] for c in headers[-1])
    grid = doc.add_table(rows=len(headers) + len(body), cols=width)
    grid.style = "Table Grid"
    for r, header in enumerate(headers):
        col = 0
        for cell in header:
            target = grid.cell(r, col)
            if cell["span"] > 1:
                target = target.merge(grid.cell(r, col + cell["span"] - 1))
            target.text = cell["text"]
            for p in target.paragraphs:
                for run in p.runs:
                    run.bold = True
            col += cell["span"]
    for r, row in enumerate(body, start=len(headers)):
        for c, cell in enumerate(row["cells"]):
            grid.cell(r, c).text = cell["text"]
            if row["cls"] == "total":
                for run in grid.cell(r, c).paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph(context["footer"])
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _number_format(measure_format: str | None) -> str:
    fmt = measure_format or ""
    if "%" in fmt:
        return "0.0%" if "0.0" in fmt else "0%"
    base = "#,##0.00" if ".00" in fmt else "#,##0"
    if "៛" in fmt:
        return f'{base} "៛"'
    if fmt.startswith("$"):
        return f'"$"{base}'
    return base


def to_xlsx(result: dict, meta: dict, measure_formats: dict[str, str | None]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    grid = table(result)
    wb = Workbook()
    ws = wb.active
    ws.title = "Analysis"
    ws.append([meta["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{meta['dataset_title']} · {result['row_count']:,} rows"])
    for f in meta["filters"]:
        ws.append([f["label"], f["value"]])
    ws.append([])
    head_fill = PatternFill("solid", fgColor="E7ECF5")
    for header in grid["headers"]:
        row_index = ws.max_row + 1
        col = 1
        for cell in header:
            ws.cell(row=row_index, column=col, value=cell["text"]).font = Font(bold=True)
            ws.cell(row=row_index, column=col).fill = head_fill
            if cell["span"] > 1:
                ws.merge_cells(start_row=row_index, start_column=col, end_row=row_index, end_column=col + cell["span"] - 1)
                ws.cell(row=row_index, column=col).alignment = Alignment(horizontal="center")
            col += cell["span"]
    first_data = ws.max_row + 1
    keys = grid["value_keys"]
    dims = len(result["labels"]["rows"]) or 1
    for row in grid["body"]:
        values = []
        for i, cell in enumerate(row["cells"]):
            values.append(cell["value"] if i >= dims and cell["value"] is not None else (None if i >= dims else cell["text"]))
        ws.append(values)
        excel_row = ws.max_row
        for i in range(dims, len(row["cells"])):
            key = keys[(i - dims) % len(keys)]
            ws.cell(row=excel_row, column=i + 1).number_format = "#,##0" if key == "__rows" else _number_format(measure_formats.get(key))
        if row["cls"] == "total":
            for c in range(1, len(row["cells"]) + 1):
                ws.cell(row=excel_row, column=c).font = Font(bold=True)
    ws.freeze_panes = ws.cell(row=first_data, column=dims + 1)
    for column in ws.columns:
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = min(40, max(10, *(len(str(c.value or "")) + 2 for c in column)))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
