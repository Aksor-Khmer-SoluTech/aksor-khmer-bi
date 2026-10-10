"""The New report wizard's working parts: start from the data, not the template.

Someone picks a data source and runs it once (routers/report_wizard.py). From what it returned, this module:

- lists its **fields** -- every path a template can use, with its type and an example value (`data_fields`);
- writes a **starter template** in .docx, .xlsx or .html with each field already placed: a line per single value, a
  table per list, a loop row already wired up (`starter_template`). It is the draft report's first version, so the
  report can be previewed before anyone has designed anything;
- **checks** a template someone uploads against those fields (`check_template`): which placeholders match, which
  don't exist in the data (with the nearest real name, for a typo), and which fields the template never uses.

The check follows loops -- in `{%tr for inv in invoices %}...{{ inv.total }}`, `inv.total` is the field
`invoices[].total` -- which the plain "undeclared variables" scan the Placeholders tab uses (template_fields.py)
can't do. It is still a best-effort reading of the template, not a proof: a field reached through something the
walker doesn't follow is simply not judged.
"""
from __future__ import annotations

import difflib
import html
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from jinja2 import Environment, TemplateSyntaxError, nodes

# How much of the run's data a draft keeps as its sample: enough rows to design and preview against, without
# storing a whole table in the report's row.
SAMPLE_LIST_ITEMS = 20
SAMPLE_STRING_CHARS = 2000
# Columns a starter table gets at most -- wider tables are rarely what anyone wants on a page.
STARTER_MAX_COLUMNS = 8
STARTER_FONT = "Khmer OS Siemreap"

_IDENT = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


# --- the data's fields --------------------------------------------------------------------------------------------


def trim_sample(value: Any) -> Any:
    """A copy of `value` small enough to keep: lists cut to SAMPLE_LIST_ITEMS, long strings shortened."""
    if isinstance(value, dict):
        return {k: trim_sample(v) for k, v in value.items()}
    if isinstance(value, list):
        return [trim_sample(v) for v in value[:SAMPLE_LIST_ITEMS]]
    if isinstance(value, str) and len(value) > SAMPLE_STRING_CHARS:
        return value[:SAMPLE_STRING_CHARS]
    return value


def _kind(value: Any) -> str:
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, list):
        return "list"
    if isinstance(value, dict):
        return "object"
    if value is None:
        return "empty"
    if isinstance(value, str) and re.match(r"^\d{4}-\d{2}-\d{2}([T ]\d{2}:\d{2}.*)?$", value):
        return "date"
    return "text"


def _example(value: Any) -> str | None:
    if isinstance(value, (dict, list)) or value is None:
        return None
    text = str(value)
    return text if len(text) <= 60 else text[:57] + "…"


def data_fields(data: dict) -> list[dict]:
    """Every path a template can use in `data`, in the data's own order: {path, kind, example, count, depth}.
    A list of objects contributes its own path (kind "list", count = items) and then its items' fields as
    `list[].field`; the items' keys are the union over the sampled rows."""
    out: list[dict] = []

    def walk(value: Any, path: str, depth: int) -> None:
        kind = _kind(value)
        entry = {"path": path, "kind": kind, "example": _example(value), "count": None, "depth": depth}
        if kind == "list":
            entry["count"] = len(value)
        out.append(entry)
        if kind == "object":
            for key, child in value.items():
                walk(child, f"{path}.{key}", depth + 1)
        elif kind == "list":
            merged: dict[str, Any] = {}
            scalars = []
            for item in value[:SAMPLE_LIST_ITEMS]:
                if isinstance(item, dict):
                    for key, child in item.items():
                        if key not in merged or merged[key] is None:
                            merged[key] = child
                else:
                    scalars.append(item)
            for key, child in merged.items():
                walk(child, f"{path}[].{key}", depth + 1)
            if scalars and not merged:
                out.append({"path": f"{path}[]", "kind": _kind(scalars[0]), "example": _example(scalars[0]), "count": None, "depth": depth + 1})

    for key, value in data.items():
        walk(value, key, 0)
    return out


# --- starter templates --------------------------------------------------------------------------------------------


def _label(key: str) -> str:
    text = key.replace("_", " ").replace("-", " ").strip()
    return (text[:1].upper() + text[1:]) if text else key


def _literal(text: str) -> str:
    """Text placed in a template as-is must not be read as Jinja itself."""
    return text.replace("{", "(").replace("}", ")")


def _loop_name(list_name: str, taken: set[str]) -> str:
    base = list_name.split(".")[-1]
    if base.endswith("ies") and len(base) > 4:
        name = base[:-3] + "y"
    elif base.endswith("s") and len(base) > 3 and not base.endswith("ss"):
        name = base[:-1]
    else:
        name = "item"
    if not _IDENT.match(name) or name in taken or name == base:
        name = "row" if "row" not in taken else "item"
    return name


def _ref(base: str, key: str) -> str:
    return f"{base}.{key}" if _IDENT.match(key) else f'{base}["{key}"]'


@dataclass
class _Table:
    title: str
    expr: str  # the list, as written in a template: invoices / summary.rows
    var: str  # the loop variable: invoice
    columns: list[tuple[str, str, bool]] = field(default_factory=list)  # (label, cell expression, numeric)


@dataclass
class _Plan:
    singles: list[tuple[str, str]] = field(default_factory=list)  # (label, expression)
    tables: list[_Table] = field(default_factory=list)


def _plan(data: dict, parameters: list[str]) -> _Plan:
    """What a starter shows: the filters and single values first, then a table for each list of rows."""
    plan = _Plan()
    taken = set(data) | set(parameters)
    for name in parameters:
        if _IDENT.match(name):
            plan.singles.append((_label(name), name))

    def visit(value: Any, expr: str, label: str, depth: int) -> None:
        if isinstance(value, dict):
            if depth >= 2:
                return
            for key, child in value.items():
                visit(child, _ref(expr, key), f"{label} {_label(key).lower()}" if label else _label(key), depth + 1)
        elif isinstance(value, list):
            rows = [item for item in value[:SAMPLE_LIST_ITEMS] if isinstance(item, dict)]
            var = _loop_name(expr, taken)
            taken.add(var)
            table = _Table(title=label, expr=expr, var=var)
            if rows:
                keys: list[str] = []
                for row in rows:
                    for key, child in row.items():
                        if key not in keys and not isinstance(child, (dict, list)):
                            keys.append(key)
                for key in keys[:STARTER_MAX_COLUMNS]:
                    numeric = all(isinstance(r.get(key), (int, float)) and not isinstance(r.get(key), bool) for r in rows if key in r)
                    table.columns.append((_label(key), _ref(var, key), numeric))
            elif value:
                table.columns.append((label, var, all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in value)))
            if table.columns:
                plan.tables.append(table)
        else:
            plan.singles.append((label, expr))

    for key, value in data.items():
        if key in parameters or not _IDENT.match(key):
            continue  # a filter is already listed; a top-level key that isn't a name can't be reached from a template
        visit(value, key, _label(key), 0)
    return plan


def _total(table: _Table, column_expr: str) -> str:
    attribute = column_expr.split(".", 1)[1] if "." in column_expr else None
    return f'{table.expr} | sum(attribute="{attribute}")' if attribute else f"{table.expr} | sum"


def starter_template(name: str, data: dict, fmt: str, parameters: list[str] | None = None) -> bytes:
    """A first template for `data` in `fmt` ("docx", "xlsx" or "html"), every field already placed. Plain on
    purpose: it is there to be restyled in Word, Excel or an editor, not to be the finished look."""
    plan = _plan(data, parameters or [])
    title = _literal(name.strip() or "Report")
    if fmt == "docx":
        return _starter_docx(title, plan)
    if fmt == "xlsx":
        return _starter_xlsx(title, plan)
    if fmt == "html":
        return _starter_html(title, plan).encode("utf-8")
    raise ValueError(f"Unsupported starter format {fmt!r}")


def _set_font(style, name: str) -> None:
    from docx.oxml.ns import qn

    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(theme_attr), None)


def _starter_docx(title: str, plan: _Plan) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Table Grid"):
        try:
            _set_font(doc.styles[style_name], STARTER_FONT)
        except KeyError:
            pass
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading(title, level=1)
    for label, expr in plan.singles:
        p = doc.add_paragraph()
        p.add_run(f"{_literal(label)}: ").bold = True
        p.add_run("{{ " + expr + " }}")
    for table in plan.tables:
        doc.add_heading(_literal(table.title), level=2)
        has_total = any(numeric for _, _, numeric in table.columns)
        grid = doc.add_table(rows=4 + (1 if has_total else 0), cols=len(table.columns))
        grid.style = "Table Grid"
        for i, (label, _, numeric) in enumerate(table.columns):
            cell = grid.rows[0].cells[i]
            cell.paragraphs[0].add_run(_literal(label)).bold = True
            if numeric:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        # docxtpl's row loop: a row holding only the opening tag, the row that repeats, a row holding only the
        # closing tag (docs/create-a-template.md, 5.1). The tag rows are removed when the document is rendered.
        grid.rows[1].cells[0].paragraphs[0].add_run("{%tr for " + table.var + " in " + table.expr + " %}")
        for i, (_, expr, numeric) in enumerate(table.columns):
            para = grid.rows[2].cells[i].paragraphs[0]
            para.add_run("{{ " + expr + " }}")
            if numeric:
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        grid.rows[3].cells[0].paragraphs[0].add_run("{%tr endfor %}")
        if has_total:
            total_row = grid.rows[4]
            total_row.cells[0].paragraphs[0].add_run("Total").bold = True
            for i, (_, expr, numeric) in enumerate(table.columns):
                if numeric and i > 0:
                    para = total_row.cells[i].paragraphs[0]
                    para.add_run("{{ " + _total(table, expr) + " }}").bold = True
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _starter_xlsx(title: str, plan: _Plan) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.cell(row=1, column=1, value=title).font = Font(name=STARTER_FONT, size=14, bold=True)
    row = 3
    for label, expr in plan.singles:
        ws.cell(row=row, column=1, value=_literal(label)).font = Font(name=STARTER_FONT, bold=True)
        ws.cell(row=row, column=2, value="{{ " + expr + " }}").font = Font(name=STARTER_FONT)
        row += 1
    header_fill = PatternFill("solid", fgColor="EEF1F5")
    widths: dict[int, int] = {1: 18, 2: 18}
    for table in plan.tables:
        row += 1
        ws.cell(row=row, column=1, value=_literal(table.title)).font = Font(name=STARTER_FONT, size=12, bold=True)
        row += 1
        for i, (label, _, numeric) in enumerate(table.columns, start=1):
            cell = ws.cell(row=row, column=i, value=_literal(label))
            cell.font = Font(name=STARTER_FONT, bold=True)
            cell.fill = header_fill
            if numeric:
                cell.alignment = Alignment(horizontal="right")
            widths[i] = max(widths.get(i, 10), min(len(label) + 4, 40))
        row += 1
        # xltpl's row loop: the opening and closing tags each on a row of their own. Those two rows stay in the
        # output as empty rows, so they are hidden here (docs/create-a-template.md, 6.1).
        ws.cell(row=row, column=1, value="{% for " + table.var + " in " + table.expr + " %}")
        ws.row_dimensions[row].hidden = True
        row += 1
        for i, (_, expr, numeric) in enumerate(table.columns, start=1):
            # A cell holding only the placeholder keeps a number a number (Excel can sum and format it).
            cell = ws.cell(row=row, column=i, value="{{ " + expr + " }}")
            cell.font = Font(name=STARTER_FONT)
            if numeric:
                cell.number_format = "#,##0.00"
        row += 1
        ws.cell(row=row, column=1, value="{% endfor %}")
        ws.row_dimensions[row].hidden = True
        row += 1
        if any(numeric for _, _, numeric in table.columns):
            ws.cell(row=row, column=1, value="Total").font = Font(name=STARTER_FONT, bold=True)
            for i, (_, expr, numeric) in enumerate(table.columns, start=1):
                if numeric and i > 1:
                    cell = ws.cell(row=row, column=i, value="{{ " + _total(table, expr) + " }}")
                    cell.font = Font(name=STARTER_FONT, bold=True)
                    cell.number_format = "#,##0.00"
            row += 1
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _starter_html(title: str, plan: _Plan) -> str:
    esc = lambda text: html.escape(_literal(text))  # noqa: E731
    lines = [
        "<!doctype html>",
        '<html lang="km">',
        "<head>",
        '<meta charset="utf-8">',
        f"<title>{esc(title)}</title>",
        "<style>",
        f'  body {{ font-family: "{STARTER_FONT}", sans-serif; font-size: 11pt; color: #1f2328; margin: 24px; }}',
        "  h1 { font-size: 18pt; margin: 0 0 12px; }",
        "  h2 { font-size: 13pt; margin: 20px 0 8px; }",
        "  table { border-collapse: collapse; width: 100%; }",
        "  th, td { border: 1px solid #c2c7d1; padding: 6px 8px; text-align: left; }",
        "  th { background: #eef1f5; }",
        "  .num { text-align: right; }",
        "</style>",
        "</head>",
        "<body>",
        f"<h1>{esc(title)}</h1>",
    ]
    for label, expr in plan.singles:
        lines.append(f"<p><strong>{esc(label)}:</strong> {{{{ {expr} }}}}</p>")
    for table in plan.tables:
        lines.append(f"<h2>{esc(table.title)}</h2>")
        lines.append("<table>")
        num = ' class="num"'
        head = "".join(f"<th{num if numeric else ''}>{esc(label)}</th>" for label, _, numeric in table.columns)
        lines.append(f"  <thead><tr>{head}</tr></thead>")
        lines.append("  <tbody>")
        lines.append(f"  {{% for {table.var} in {table.expr} %}}")
        cells = "".join(f"<td{num if numeric else ''}>{{{{ {expr} }}}}</td>" for _, expr, numeric in table.columns)
        lines.append(f"    <tr>{cells}</tr>")
        lines.append("  {% endfor %}")
        lines.append("  </tbody>")
        if any(numeric for _, _, numeric in table.columns):
            total_cells = []
            for i, (_, expr, numeric) in enumerate(table.columns):
                if i == 0:
                    total_cells.append("<th>Total</th>")
                elif numeric:
                    total_cells.append(f'<th class="num">{{{{ {_total(table, expr)} }}}}</th>')
                else:
                    total_cells.append("<th></th>")
            lines.append(f"  <tfoot><tr>{''.join(total_cells)}</tr></tfoot>")
        lines.append("</table>")
    lines += ["</body>", "</html>", ""]
    return "\n".join(lines)


# --- checking a template against the data ------------------------------------------------------------------------


def template_source(path: Path, template_ext: str) -> str:
    """The Jinja source of a template file -- for a .docx, its body, headers and footers with docxtpl's `{%tr %}`
    style tags turned into plain Jinja, the way docxtpl itself reads them; for an .xlsx, its cells' text in order."""
    if template_ext == "html":
        return path.read_text(encoding="utf-8")
    if template_ext == "docx":
        from docx import Document
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml import parse_xml
        from docxtpl import DocxTemplate

        tpl = DocxTemplate(str(path))
        document = Document(str(path))
        source = tpl.patch_xml(tpl.xml_to_string(document._element.body))
        for rel in document.part.rels.values():
            if rel.reltype in (RT.HEADER, RT.FOOTER) and rel.target_part.blob:
                source += tpl.patch_xml(tpl.xml_to_string(parse_xml(rel.target_part.blob)))
        return source
    if template_ext == "xlsx":
        from openpyxl import load_workbook

        texts = []
        workbook = load_workbook(io.BytesIO(path.read_bytes()))
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and ("{{" in cell.value or "{%" in cell.value):
                        texts.append(cell.value)
        return "\n".join(texts)
    raise ValueError(f"Unsupported template type {template_ext!r}")


_LOCAL = object()  # a name the template defines itself ({% set %}, a macro, a loop over something unknown)
_COLLECTION_FILTERS = {"sum", "map", "sort", "groupby", "selectattr", "rejectattr", "unique", "min", "max"}
# Filters that hand back the same list's items: a loop over `invoices | sort(attribute="date")` is still a loop
# over the invoices.
_SAME_ITEMS_FILTERS = {"sort", "reverse", "selectattr", "rejectattr", "unique", "list"}


@dataclass
class _Ref:
    path: str  # the data path: invoices[].total_usd
    text: str  # as written: inv.total_usd


def _declared_names(ast: nodes.Template) -> set[str]:
    names: set[str] = set()
    for node in ast.find_all((nodes.Assign, nodes.AssignBlock, nodes.Macro, nodes.CallBlock, nodes.With)):
        if isinstance(node, (nodes.Assign, nodes.AssignBlock)):
            names |= {n.name for n in node.target.find_all(nodes.Name)} | ({node.target.name} if isinstance(node.target, nodes.Name) else set())
        elif isinstance(node, nodes.Macro):
            names.add(node.name)
            names |= {a.name for a in node.args}
        elif isinstance(node, nodes.CallBlock):
            names |= {a.name for a in node.args}
        elif isinstance(node, nodes.With):
            names |= {t.name for t in node.targets if isinstance(t, nodes.Name)}
    return names


def _chain(node: nodes.Node, scope: dict) -> tuple[str, str] | None:
    """(data path, text as written) for a Name / attribute / item chain, or None when it isn't one we can follow."""
    if isinstance(node, nodes.Filter) and node.name in _SAME_ITEMS_FILTERS and node.node is not None:
        return _chain(node.node, scope)
    if isinstance(node, nodes.Name):
        bound = scope.get(node.name)
        if bound is _LOCAL:
            return None
        return (bound if bound is not None else node.name), node.name
    if isinstance(node, nodes.Getattr):
        base = _chain(node.node, scope)
        return (f"{base[0]}.{node.attr}", f"{base[1]}.{node.attr}") if base else None
    if isinstance(node, nodes.Getitem):
        base = _chain(node.node, scope)
        if base is None:
            return None
        if isinstance(node.arg, nodes.Const) and isinstance(node.arg.value, str):
            return f"{base[0]}.{node.arg.value}", f'{base[1]}["{node.arg.value}"]'
        if isinstance(node.arg, nodes.Const) and isinstance(node.arg.value, int):
            return f"{base[0]}[]", f"{base[1]}[{node.arg.value}]"
        return base  # a slice (date[:4]) or a computed index: the value itself is what's used
    return None


def referenced_fields(source: str) -> tuple[list[_Ref], set[str]]:
    """The data fields a template uses: (value references, lists it loops over). Raises TemplateSyntaxError."""
    ast = Environment(extensions=["jinja2.ext.do", "jinja2.ext.loopcontrols"]).parse(source)
    locals_ = _declared_names(ast)
    refs: list[_Ref] = []
    loops: set[str] = set()

    def record(found: tuple[str, str] | None) -> None:
        if found is not None:
            refs.append(_Ref(path=found[0], text=found[1]))

    def visit(node: nodes.Node, scope: dict) -> None:
        if isinstance(node, nodes.For):
            visit(node.iter, scope)
            iterated = _chain(node.iter, scope)
            if iterated is not None:
                loops.add(iterated[0])
            inner = dict(scope)
            inner["loop"] = _LOCAL
            targets = [node.target] if isinstance(node.target, nodes.Name) else list(node.target.find_all(nodes.Name))
            for target in targets:
                inner[target.name] = f"{iterated[0]}[]" if (iterated and isinstance(node.target, nodes.Name)) else _LOCAL
            if node.test is not None:
                visit(node.test, inner)
            for child in node.body:
                visit(child, inner)
            for child in node.else_:
                visit(child, scope)
            return
        if isinstance(node, (nodes.Name, nodes.Getattr, nodes.Getitem)):
            found = _chain(node, scope)
            if found is not None:
                record(found)
                if isinstance(node, nodes.Getitem) and not isinstance(node.arg, nodes.Const):
                    visit(node.arg, scope)
                return
        if isinstance(node, nodes.Call):
            callee = node.node
            if isinstance(callee, nodes.Getattr):
                visit(callee.node, scope)  # x.replace(...): x is the field, replace is a method
            elif not isinstance(callee, nodes.Name):
                visit(callee, scope)  # resource("logo"), range(3): functions, not fields
            for child in [*node.args, *(kw.value for kw in node.kwargs), node.dyn_args, node.dyn_kwargs]:
                if child is not None:
                    visit(child, scope)
            return
        if isinstance(node, nodes.Filter) and node.name in _COLLECTION_FILTERS and node.node is not None:
            base = _chain(node.node, scope)
            for kw in node.kwargs:
                if kw.key == "attribute" and isinstance(kw.value, nodes.Const) and isinstance(kw.value.value, str) and base:
                    refs.append(_Ref(path=f"{base[0]}[].{kw.value.value}", text=f'{base[1]} | {node.name}(attribute="{kw.value.value}")'))
        for child in node.iter_child_nodes():
            visit(child, scope)

    visit(ast, {name: _LOCAL for name in locals_})
    return refs, loops


def _suggest(ref: _Ref, fields: dict) -> str | None:
    """The nearest real name for a placeholder that names no field -- at the first part of it that's wrong, so
    both `compnay.phone` and `invoice.customer` get an answer. None when nothing is close."""
    path_parts = ref.path.split(".")
    text_parts = ref.text.split(".")
    if "[" in ref.text or len(path_parts) != len(text_parts):
        return None
    for i in range(len(path_parts)):
        prefix = ".".join(path_parts[: i + 1])
        if prefix in fields or prefix.removesuffix("[]") in fields:
            continue
        parent = ".".join(path_parts[:i])
        siblings = [f.rpartition(".")[2] for f in fields if f.rpartition(".")[0] == parent] if i else [f for f in fields if "." not in f]
        siblings = [name.removesuffix("[]") for name in siblings]
        close = difflib.get_close_matches(path_parts[i].removesuffix("[]"), siblings, n=1, cutoff=0.6)
        if not close:
            return None
        return ".".join(text_parts[:i] + [close[0]] + text_parts[i + 1 :])
    return None


def check_template(path: Path, template_ext: str, data: dict, parameters: list[str] | None = None) -> dict:
    """Compare a template with the data it will get: {checked, reason, matched, unknown, unused}.

    - matched: placeholders that name a real field (or a filter);
    - unknown: placeholders naming no field in the data -- each with the nearest real name, when there is one,
      since that's usually a typo -- because those print nothing;
    - unused: the data's single values the template never shows. Informational: leaving data out is often
      deliberate."""
    try:
        source = template_source(path, template_ext)
        refs, loops = referenced_fields(source)
    except TemplateSyntaxError as exc:
        return {"checked": False, "reason": f"The template has a syntax error on line {exc.lineno}: {exc.message}", "matched": [], "unknown": [], "unused": []}
    except Exception:  # noqa: BLE001 -- an unreadable file is reported, not raised: the upload itself was accepted
        return {"checked": False, "reason": "Aksor couldn't read the placeholders in this file", "matched": [], "unknown": [], "unused": []}

    fields = {f["path"]: f for f in data_fields(data)}
    params = set(parameters or [])

    def known(p: str) -> bool:
        root = p.split(".", 1)[0].split("[", 1)[0]
        return p in fields or root in params

    matched: list[str] = []
    unknown: list[dict] = []
    seen: set[str] = set()
    for ref in refs:
        if ref.text in seen:
            continue
        seen.add(ref.text)
        if known(ref.path):
            matched.append(ref.text)
            continue
        unknown.append({"placeholder": ref.text, "suggestion": _suggest(ref, fields)})

    used_paths = {ref.path for ref in refs if known(ref.path)}

    def used(p: str) -> bool:
        # Printing a whole object shows its fields; looping over (or counting) a list doesn't show its items' fields.
        return any(u == p or (p.startswith(u + ".") and fields.get(u, {}).get("kind") == "object") for u in used_paths)

    unused = [
        p for p, f in fields.items()
        if f["kind"] not in ("list", "object") and not used(p) and p.split(".", 1)[0] not in params
    ]
    return {"checked": True, "reason": None, "matched": matched, "unknown": unknown, "unused": unused}
