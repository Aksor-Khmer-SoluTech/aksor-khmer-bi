"""Datasets (app/datasets.py, routers/datasets.py): one governed source + filters + fields + measures that reports
build on -- checked end to end: a report built on a dataset runs with the dataset's data and measures, follows the
dataset when it changes, keeps the filter grants it had, and a report made before datasets can be turned into one."""
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import datasets
from app.main import app
from app.rbac import ROOT_ORG_ID

client = TestClient(app)

BRANCHES = [{"value": "BR01", "label": "ភ្នំពេញ"}, {"value": "BR02", "label": "សៀមរាប"}]
SALES = {
    "rows": [
        {"branch": "BR01", "product": "Diesel", "amount": 40000, "discount": 1000, "order_id": "SO-1"},
        {"branch": "BR01", "product": "Petrol", "amount": "25,000", "discount": 0, "order_id": "SO-1"},
        {"branch": "BR02", "product": "Diesel", "amount": 35000, "discount": 500, "order_id": "SO-2"},
        {"branch": "BR02", "product": "Service", "amount": None, "discount": 0, "order_id": "SO-3"},
    ],
    "title": "របាយការណ៍លក់",
}
MEASURES = [
    {"key": "gross", "label": "Gross", "kind": "sum", "field": "amount", "format": "#,##0 ៛"},
    {"key": "discounts", "label": "Discounts", "kind": "sum", "field": "discount"},
    {"key": "orders", "label": "Orders", "label_km": "ការបញ្ជាទិញ", "kind": "count_distinct", "field": "order_id"},
    {"key": "lines", "label": "Lines", "kind": "count"},
    {"key": "diesel", "label": "Diesel", "kind": "sum", "field": "amount", "where": {"field": "product", "value": "Diesel"}},
    {"key": "diesel_share", "label": "Diesel share", "kind": "ratio", "numerator": "diesel", "denominator": "gross", "format": "0.0%"},
    {"key": "biggest", "label": "Biggest sale", "kind": "max", "field": "amount", "format": "$#,##0.00"},
    {"key": "average", "label": "Average sale", "kind": "avg", "field": "amount"},
]


def _dataset(auth_headers, name="branch-sales", static=None, measures=None, parameters=None, expect=200, **extra):
    body = {
        "name": name, "title": "Branch sales", "description": "Every sale line",
        "parameters": parameters if parameters is not None else [{"name": "p_branch", "label": "Branch", "options": BRANCHES}],
        "data_source": {"type": "static", "static_data": static or SALES},
        "fields": [{"source": "rows", "name": "amount", "type": "money", "label": "Amount", "label_km": "ចំនួនទឹកប្រាក់", "format": "#,##0 ៛"}],
        "measures": MEASURES if measures is None else measures,
        **extra,
    }
    resp = client.post("/api/v1/datasets", json=body, headers=auth_headers)
    assert resp.status_code == expect, resp.text
    return resp.json()


def _report(auth_headers, text="{{ title }} · {{ p_branch }} · {{ measures.gross }} · {{ measures.orders }} · {{ rows | length }}") -> str:
    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    resp = client.post("/api/v1/reports", files={"file": ("t.docx", buf.getvalue(), "application/octet-stream")},
                       data={"name": "Sales on a dataset"}, headers=auth_headers)
    assert resp.status_code == 200, resp.text
    return resp.json()["report_id"]


def _text(content: bytes) -> str:
    return "\n".join(p.text for p in Document(BytesIO(content)).paragraphs).replace("​", "")


def _run(headers, report_id, branch="BR01"):
    return client.post(f"/api/v1/reports/{report_id}/run", json={"parameters": {"p_branch": branch}, "format": "docx"}, headers=headers)


# --- measures ---------------------------------------------------------------------------------------------------------


def test_measures_are_computed_over_the_rows():
    measures = datasets.validate_measures(MEASURES)
    texts, values = datasets.compute_measures(measures, SALES)
    assert values["gross"] == 100000  # "25,000" read as a number, None skipped
    assert texts["gross"] == "100,000 ៛"
    assert values["orders"] == 3 and values["lines"] == 4
    assert values["diesel"] == 75000 and texts["diesel_share"] == "75.0%"
    assert texts["biggest"] == "$40,000.00"
    assert values["average"] == pytest.approx(100000 / 3) and texts["average"] == "33,333.33"
    assert texts["discounts"] == "1,500"


STATS = {"rows": [{"v": n} for n in (4, 1, 3, 3, 9, None, "x", "2")]}  # numbers 1 2 3 3 4 9; None and "x" skipped


@pytest.mark.parametrize("kind, extra, expected", [
    ("median", {}, 3.0),
    ("percentile", {"percentile": 50}, 3.0),
    ("percentile", {"percentile": 0}, 1.0),
    ("percentile", {"percentile": 100}, 9.0),
    ("percentile", {"percentile": 90}, 6.5),  # Excel PERCENTILE.INC: 4 + (9 - 4) * 0.5
    ("percentile", {"percentile": 25}, 2.25),
    ("stdev", {}, pytest.approx(2.80475786, rel=1e-6)),  # of a sample, as Excel's STDEV.S
    ("variance", {}, pytest.approx(7.86666667, rel=1e-6)),
    ("mode", {}, 3.0),
])
def test_statistical_measures(kind, extra, expected):
    measures = datasets.validate_measures([{"key": "m", "label": "M", "kind": kind, "field": "v", **extra}])
    _, values = datasets.compute_measures(measures, STATS)
    assert values["m"] == expected


def test_statistical_measures_on_too_few_numbers():
    measures = datasets.validate_measures([
        {"key": kind, "label": kind, "kind": kind, "field": "v", **({"percentile": 75} if kind == "percentile" else {})}
        for kind in ("median", "percentile", "stdev", "variance", "mode")
    ])
    _, empty = datasets.compute_measures(measures, {"rows": [{"v": None}]})
    assert all(v is None for v in empty.values())
    _, one = datasets.compute_measures(measures, {"rows": [{"v": 7}]})
    assert one == {"median": 7.0, "percentile": 7.0, "stdev": None, "variance": None, "mode": 7.0}


def test_mode_on_a_tie_is_the_smallest():
    measures = datasets.validate_measures([{"key": "m", "label": "M", "kind": "mode", "field": "v"}])
    _, values = datasets.compute_measures(measures, {"rows": [{"v": n} for n in (5, 2, 5, 2, 8)]})
    assert values["m"] == 2.0


def test_a_ratio_can_use_a_statistical_measure():
    measures = datasets.validate_measures([
        {"key": "med", "label": "Median", "kind": "median", "field": "v"},
        {"key": "avg", "label": "Average", "kind": "avg", "field": "v"},
        {"key": "skew", "label": "Median / average", "kind": "ratio", "numerator": "med", "denominator": "avg"},
    ])
    _, values = datasets.compute_measures(measures, STATS)
    assert values["skew"] == pytest.approx(3 / (22 / 6))


def test_a_measure_saved_before_percentiles_still_computes():
    # Stored before 5a: no "percentile" key at all.
    stored = {"key": "m", "label": "M", "kind": "median", "source": None, "field": "v", "where": None,
              "numerator": None, "denominator": None, "format": None}
    _, values = datasets.compute_measures([stored], STATS)
    assert values["m"] == 3.0


@pytest.mark.parametrize("value, pattern, text", [
    (1234.5, None, "1,234.50"), (1234, None, "1,234"), (1234.567, "#,##0.00", "1,234.57"), (0.256, "0%", "26%"),
    (12, "0.0", "12.0"), (1500, "#,##0 ៛", "1,500 ៛"), (None, "#,##0", "—"),
])
def test_number_formats(value, pattern, text):
    assert datasets.format_number(value, pattern) == text


@pytest.mark.parametrize("measures, fragment", [
    ([{"key": "Net", "label": "x", "kind": "sum", "field": "a"}], "lowercase letters"),
    ([{"key": "a", "label": "x", "kind": "geomean", "field": "a"}], "a measure is one of"),
    ([{"key": "a", "label": "x", "kind": "median"}], "choose the field"),
    ([{"key": "a", "label": "x", "kind": "percentile", "field": "a"}], "which percentile"),
    ([{"key": "a", "label": "x", "kind": "percentile", "field": "a", "percentile": 101}], "which percentile"),
    ([{"key": "a", "label": "x", "kind": "percentile", "field": "a", "percentile": True}], "which percentile"),
    ([{"key": "a", "label": "x", "kind": "sum"}], "choose the field"),
    ([{"key": "a", "label": "x", "kind": "ratio", "numerator": "a", "denominator": "b"}], "isn't another measure"),
    ([{"key": "a", "label": "x", "kind": "sum", "field": "f"}, {"key": "a", "label": "y", "kind": "count"}], "Two measures"),
])
def test_bad_measures_are_refused(measures, fragment):
    with pytest.raises(datasets.DataConfigError, match=fragment):
        datasets.validate_measures(measures)


def test_detected_fields_keep_their_labels():
    known = [{"source": "rows", "name": "amount", "type": "money", "label": "Amount", "label_km": "ចំនួនទឹកប្រាក់", "format": "#,##0 ៛", "hidden": False}]
    fields = datasets.detect_fields({**SALES, "day": "2026-10-13"}, known)
    by_name = {(f["source"], f["name"]): f for f in fields}
    assert by_name[("rows", "amount")]["label_km"] == "ចំនួនទឹកប្រាក់"  # kept
    assert by_name[("rows", "discount")]["type"] == "number"
    assert by_name[("rows", "product")]["type"] == "text"
    assert by_name[(None, "day")]["type"] == "date"


# --- the API ----------------------------------------------------------------------------------------------------------


def test_create_list_update_and_the_rules(auth_headers):
    created = _dataset(auth_headers)
    assert created["name"] == "branch-sales" and len(created["measures"]) == len(MEASURES) and created["reports"] == []
    assert _dataset(auth_headers, expect=409)["detail"].startswith("A dataset named")
    assert "lowercase letters" in _dataset(auth_headers, name="Branch Sales", expect=400)["detail"]
    nested = _dataset(auth_headers, name="nested", expect=422 if False else 400, data_source={"type": "dataset", "dataset": "branch-sales"})
    assert "not from another dataset" in nested["detail"]

    listed = client.get("/api/v1/datasets", headers=auth_headers).json()
    assert [(d["name"], d["measure_count"], d["source_type"]) for d in listed] == [("branch-sales", len(MEASURES), "static")]

    body = {k: created[k] for k in ("title", "description", "parameters", "data_source", "fields")}
    updated = client.put(f"/api/v1/datasets/{created['id']}", headers=auth_headers, json={**body, "title": "Sales", "measures": MEASURES[:2]})
    assert updated.status_code == 200 and updated.json()["title"] == "Sales" and len(updated.json()["measures"]) == 2


def test_preview_gives_data_fields_and_measures(auth_headers):
    resp = client.post("/api/v1/datasets/preview", headers=auth_headers, json={
        "data_source": {"type": "static", "static_data": SALES}, "measures": MEASURES,
        "fields": [{"source": "rows", "name": "amount", "type": "money", "label_km": "ចំនួនទឹកប្រាក់"}],
    })
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["measures"]["gross"] == {"text": "100,000 ៛", "value": 100000}
    fields = {(f["source"], f["name"]): f for f in body["fields"]}
    assert fields[("rows", "amount")]["label_km"] == "ចំនួនទឹកប្រាក់" and fields[("rows", "order_id")]["type"] == "text"
    assert body["data"]["title"] == "របាយការណ៍លក់"


def test_statistical_measures_save_and_preview(auth_headers):
    stats = [
        {"key": "typical", "label": "Typical sale", "kind": "median", "field": "amount", "format": "#,##0 ៛"},
        {"key": "p90", "label": "90th percentile", "kind": "percentile", "field": "amount", "percentile": 90},
        {"key": "spread", "label": "Spread", "kind": "stdev", "field": "amount", "format": "#,##0"},
    ]
    resp = client.post("/api/v1/datasets/preview", headers=auth_headers, json={"data_source": {"type": "static", "static_data": SALES}, "measures": stats})
    assert resp.status_code == 200, resp.text
    measures = resp.json()["measures"]
    assert measures["typical"] == {"text": "35,000 ៛", "value": 35000}
    assert measures["p90"]["value"] == pytest.approx(39000)  # 25,000 35,000 40,000: 35,000 + 5,000 * 0.8
    assert measures["spread"]["text"] == "7,638"
    saved = _dataset(auth_headers, name="sales-stats", measures=stats)
    assert {m["key"]: m["percentile"] for m in saved["measures"]} == {"typical": None, "p90": 90, "spread": None}
    bad = [{**stats[1], "percentile": 120}]
    _dataset(auth_headers, name="sales-stats-bad", measures=bad, expect=422)


def test_a_report_built_on_a_dataset_runs_with_its_data_and_follows_it(auth_headers):
    dataset = _dataset(auth_headers)
    report_id = _report(auth_headers)
    saved = client.put(f"/api/v1/reports/{report_id}/data-config", headers=auth_headers,
                       json={"parameters": [{"name": "ignored", "label": "x"}], "data_source": {"type": "dataset", "dataset": "branch-sales"}})
    assert saved.status_code == 200, saved.text
    assert saved.json()["parameters"] == [] and saved.json()["data_source"]["dataset"] == "branch-sales"

    # The run form shows the dataset's filters; the run gets its data and its measures.
    form = client.get(f"/api/v1/reports/{report_id}/run-form", headers=auth_headers).json()
    assert [p["name"] for p in form["parameters"]] == ["p_branch"] and form["has_data_source"] is True
    run = _run(auth_headers, report_id, "BR02")
    assert run.status_code == 200, run.text
    assert _text(run.content) == "របាយការណ៍លក់ · BR02 · 100,000 ៛ · 3 · 4"

    # Change the dataset: the report follows on its next run.
    body = {k: dataset[k] for k in ("title", "description", "parameters", "fields", "measures")}
    smaller = {**SALES, "rows": SALES["rows"][:1]}
    client.put(f"/api/v1/datasets/{dataset['id']}", headers=auth_headers, json={**body, "data_source": {"type": "static", "static_data": smaller}})
    assert _text(_run(auth_headers, report_id).content) == "របាយការណ៍លក់ · BR01 · 40,000 ៛ · 1 · 1"

    # It can't be deleted while a report is built on it; the dataset lists the report.
    assert client.get(f"/api/v1/datasets/{dataset['id']}", headers=auth_headers).json()["reports"][0]["report_id"] == report_id
    gone = client.delete(f"/api/v1/datasets/{dataset['id']}", headers=auth_headers)
    assert gone.status_code == 409 and "built on it" in gone.json()["detail"]


def test_a_report_cant_name_a_dataset_that_isnt_there(auth_headers):
    report_id = _report(auth_headers)
    resp = client.put(f"/api/v1/reports/{report_id}/data-config", headers=auth_headers,
                      json={"parameters": [], "data_source": {"type": "dataset", "dataset": "nope"}})
    assert resp.status_code == 400 and "no dataset named 'nope'" in resp.json()["detail"]


def test_filter_limits_still_hold_on_a_report_built_on_a_dataset(auth_headers, make_local_user):
    _dataset(auth_headers)
    report_id = _report(auth_headers)
    client.put(f"/api/v1/reports/{report_id}/data-config", headers=auth_headers,
               json={"parameters": [], "data_source": {"type": "dataset", "dataset": "branch-sales"}})
    user_id, headers = make_local_user("pp-staff", "pw-staff-1", ())
    grant = client.post("/api/v1/grants/reports", headers=auth_headers, json={
        "subject_type": "user", "subject_id": user_id, "report_id": report_id, "permission_level": "render",
        "parameter_limits": {"p_branch": ["BR01"]},
    })
    assert grant.status_code == 200, grant.text
    assert _run(headers, report_id, "BR01").status_code == 200
    assert _run(headers, report_id, "BR02").status_code == 403


def test_a_report_made_before_datasets_becomes_one(auth_headers):
    report_id = _report(auth_headers, "{{ title }} · {{ p_branch }} · {{ rows | length }}")
    client.put(f"/api/v1/reports/{report_id}/data-config", headers=auth_headers, json={
        "parameters": [{"name": "p_branch", "label": "Branch", "options": BRANCHES}],
        "data_source": {"type": "static", "static_data": SALES},
    })
    resp = client.post("/api/v1/datasets/from-report", headers=auth_headers, json={"report_id": report_id, "name": "sales", "title": "Sales"})
    assert resp.status_code == 200, resp.text
    assert [p["name"] for p in resp.json()["parameters"]] == ["p_branch"] and resp.json()["reports"][0]["report_id"] == report_id
    config = client.get(f"/api/v1/reports/{report_id}/data-config", headers=auth_headers).json()
    assert config["data_source"]["type"] == "dataset" and config["parameters"] == []
    assert _text(_run(auth_headers, report_id).content) == "របាយការណ៍លក់ · BR01 · 4"  # runs exactly as before


def test_who_may_see_and_manage_datasets(auth_headers, make_local_user):
    _dataset(auth_headers)
    _, author = make_local_user("author", "pw-author-1", ("ROLE_REPORT_ADMIN",))
    _, viewer = make_local_user("viewer", "pw-viewer-1", ("ROLE_REPORT_VIEWER",))
    assert client.get("/api/v1/datasets", headers=author).status_code == 200
    assert client.get("/api/v1/datasets", headers=viewer).status_code == 403
    assert _dataset(author, name="author-sales", expect=200)["name"] == "author-sales"  # report admins manage datasets too


def test_a_connection_a_dataset_reads_through_cant_be_deleted(auth_headers):
    conn = client.post("/api/v1/connections", headers=auth_headers, json={"name": "erp-api", "config": {"base_url": "https://erp.example.test"}}).json()
    _dataset(auth_headers, name="erp-sales", parameters=[], data_source={"type": "rest", "connection": "erp-api", "url": "/sales"})
    resp = client.delete(f"/api/v1/connections/{conn['id']}", headers=auth_headers)
    assert resp.status_code == 409 and "Branch sales" in resp.json()["detail"]
