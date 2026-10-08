"""Resources > Fonts: add a font to the whole server, check it, see what uses it, remove it.

The fonts are built on the spot (a tiny TrueType font with one visible glyph) so no font file is needed. What each
rendering path does with an added font is tested in packages/doc_engine/tests/test_fonts.py; here: the API around it."""
import io
import os
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from fontTools.fontBuilder import FontBuilder
from fontTools.pens.ttGlyphPen import TTGlyphPen

from app import font_catalog, font_store
from app.main import app

client = TestClient(app)


def build_font(family="Aksor Test Sans", style="Regular", version="Version 1.000", khmer=False) -> bytes:
    pen = TTGlyphPen(None)
    pen.moveTo((100, 0))
    pen.lineTo((100, 700))
    pen.lineTo((600, 700))
    pen.lineTo((600, 0))
    pen.closePath()
    square = pen.glyph()
    empty = TTGlyphPen(None).glyph()
    cmap = {ord("A"): "A", ord(" "): "space"}
    order = [".notdef", "A", "space"]
    glyphs = {".notdef": square, "A": square, "space": empty}
    if khmer:  # one Khmer consonant, so coverage is above zero
        order.append("ka")
        cmap[0x1780] = "ka"
        glyphs["ka"] = square
    builder = FontBuilder(1000, isTTF=True)
    builder.setupGlyphOrder(order)
    builder.setupCharacterMap(cmap)
    builder.setupGlyf(glyphs)
    builder.setupHorizontalMetrics({name: (700, 100) for name in order})
    builder.setupHorizontalHeader(ascent=800, descent=-200)
    builder.setupNameTable({
        "familyName": family, "styleName": style, "uniqueFontIdentifier": f"{family} {style} {version}",
        "fullName": f"{family} {style}", "psName": f"{family.replace(' ', '')}-{style}", "version": version,
        "copyright": "Copyright 2026 Test", "licenseDescription": "Free to use in tests", "licenseInfoURL": "https://example.test/license",
    })
    builder.setupOS2(sTypoAscender=800, sTypoDescender=-200, usWinAscent=800, usWinDescent=200)
    builder.setupPost()
    out = io.BytesIO()
    builder.save(out)
    return out.getvalue()


@pytest.fixture(autouse=True)
def isolated_fonts(tmp_path, monkeypatch):
    folder = tmp_path / "font_resources"
    monkeypatch.setattr(font_store, "FONT_DIR", folder)
    monkeypatch.setenv("DOC_ENGINE_FONT_DIRS", str(folder))
    monkeypatch.setattr(
        font_catalog, "system_families", lambda: {"arial": "Arial", "khmer os siemreap": "Khmer OS Siemreap", "liberation serif": "Liberation Serif"}
    )
    font_catalog.forget()
    yield folder
    font_catalog.forget()


def upload(headers, data: bytes, name="font.ttf", note=""):
    return client.post("/api/v1/fonts", files={"file": (name, data, "font/ttf")}, data={"note": note}, headers=headers)


def _docx_naming(*families: str) -> bytes:
    doc = Document()
    for family in families:
        run = doc.add_paragraph().add_run("text")
        fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(attr), family)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def register(headers, name, families):
    resp = client.post(
        "/api/v1/reports", files={"file": ("t.docx", _docx_naming(*families), "application/octet-stream")}, data={"name": name}, headers=headers
    )
    assert resp.status_code == 200, resp.text
    return resp.json()["report_id"]


# --- adding a font ------------------------------------------------------------------------------------------------


def test_a_font_is_added_read_and_stored_by_id(auth_headers, isolated_fonts):
    resp = upload(auth_headers, build_font(khmer=True), name="whatever-name.ttf", note="from the brand team")
    assert resp.status_code == 200, resp.text
    font = resp.json()
    assert (font["family"], font["subfamily"], font["version"]) == ("Aksor Test Sans", "Regular", "Version 1.000")
    assert font["glyph_count"] == 4 and font["khmer_coverage"] > 0 and font["license"] == "Free to use in tests"
    assert font["license_url"] == "https://example.test/license" and font["note"] == "from the brand team"
    # stored under its id, never the uploaded name
    assert (isolated_fonts / f"{font['id']}.ttf").is_file() and not (isolated_fonts / "whatever-name.ttf").exists()
    listed = client.get("/api/v1/fonts", headers=auth_headers).json()
    assert [f["id"] for f in listed] == [font["id"]]


def test_a_font_with_no_khmer_says_so(auth_headers):
    warnings = upload(auth_headers, build_font()).json()["warnings"]
    assert any("no Khmer glyphs" in w for w in warnings)


def test_the_font_file_can_be_fetched_for_previewing(auth_headers):
    font = upload(auth_headers, build_font()).json()
    resp = client.get(f"/api/v1/fonts/{font['id']}/file", headers=auth_headers)
    assert resp.status_code == 200 and resp.headers["content-type"] == "font/ttf" and resp.content[:4] == b"\x00\x01\x00\x00"


@pytest.mark.parametrize(
    "data, message",
    [
        (b"", "empty"),
        (b"this is not a font", "isn't a TrueType"),
        (b"ttcf" + b"\0" * 40, "collection"),
        (b"wOFF" + b"\0" * 40, "web font"),
        (b"\x00\x01\x00\x00" + b"x" * 60, "damaged"),
    ],
)
def test_a_file_that_is_not_a_usable_font_is_refused(auth_headers, isolated_fonts, data, message):
    resp = upload(auth_headers, data)
    assert resp.status_code == 400 and message in resp.json()["detail"]
    assert not isolated_fonts.exists() or list(isolated_fonts.iterdir()) == []


def test_an_oversized_file_is_refused(auth_headers, monkeypatch):
    monkeypatch.setattr(font_store, "MAX_BYTES", 100)
    assert upload(auth_headers, build_font()).status_code == 400


def test_the_same_file_twice_and_two_files_claiming_one_font_are_refused(auth_headers):
    first = build_font()
    assert upload(auth_headers, first).status_code == 200
    again = upload(auth_headers, first)
    assert again.status_code == 409 and "already been added" in again.json()["detail"]
    other_version = upload(auth_headers, build_font(version="Version 2.000"))
    assert other_version.status_code == 409 and "already there" in other_version.json()["detail"]
    # a different style of the same family is fine
    assert upload(auth_headers, build_font(style="Bold")).status_code == 200


def test_a_family_the_server_already_has_gets_a_warning(auth_headers):
    resp = upload(auth_headers, build_font(family="Khmer OS Siemreap", version="Version 1.00"))
    assert resp.status_code == 200
    assert resp.json()["shadows_installed"] is True and any("already has a font named" in w for w in resp.json()["warnings"])


# --- who may --------------------------------------------------------------------------------------------------------


def test_only_font_managers_add_or_remove_but_report_managers_can_look(auth_headers, make_local_user):
    font = upload(auth_headers, build_font()).json()
    _, manager = make_local_user("report_admin", "pw12345", role_names=("ROLE_REPORT_ADMIN",))
    assert client.get("/api/v1/fonts", headers=manager).status_code == 200
    assert client.get(f"/api/v1/fonts/{font['id']}", headers=manager).status_code == 200
    assert client.get(f"/api/v1/fonts/{font['id']}/file", headers=manager).status_code == 200
    assert upload(manager, build_font(style="Italic")).status_code == 403
    assert client.delete(f"/api/v1/fonts/{font['id']}", headers=manager).status_code == 403

    _, nobody = make_local_user("just_a_viewer", "pw12345", role_names=("ROLE_USER",))
    assert client.get("/api/v1/fonts", headers=nobody).status_code == 403
    assert client.get(f"/api/v1/fonts/{font['id']}/file", headers=nobody).status_code == 403


# --- what uses a font, and removing it -----------------------------------------------------------------------------


def test_a_font_a_template_names_cannot_be_removed_without_force(auth_headers, isolated_fonts):
    font = upload(auth_headers, build_font()).json()
    register(auth_headers, "Brand letter", ["Aksor Test Sans"])
    detail = client.get(f"/api/v1/fonts/{font['id']}", headers=auth_headers).json()
    assert detail["used_by"] == ["Brand letter"]

    blocked = client.delete(f"/api/v1/fonts/{font['id']}", headers=auth_headers)
    assert blocked.status_code == 409 and "Brand letter" in blocked.json()["detail"]
    assert (isolated_fonts / f"{font['id']}.ttf").exists()

    assert client.delete(f"/api/v1/fonts/{font['id']}?force=true", headers=auth_headers).status_code == 204
    assert not (isolated_fonts / f"{font['id']}.ttf").exists()
    assert client.get("/api/v1/fonts", headers=auth_headers).json() == []


def test_an_unused_font_is_removed_with_its_file(auth_headers, isolated_fonts):
    font = upload(auth_headers, build_font()).json()
    assert client.delete(f"/api/v1/fonts/{font['id']}", headers=auth_headers).status_code == 204
    assert list(isolated_fonts.iterdir()) == []
    assert client.get(f"/api/v1/fonts/{font['id']}", headers=auth_headers).status_code == 404


def test_coverage_shows_which_characters_the_font_has(auth_headers):
    font = upload(auth_headers, build_font(khmer=True)).json()
    blocks = {b["name"]: b for b in client.get(f"/api/v1/fonts/{font['id']}/coverage", headers=auth_headers).json()}
    assert blocks["Khmer"]["present"] == [0x1780] and 0x1781 in blocks["Khmer"]["missing"]
    assert 0x17DE not in blocks["Khmer"]["missing"]  # unassigned code points are in neither list
    assert blocks["Basic Latin"]["present"] == [0x20, 0x41]
    assert client.get("/api/v1/fonts/nope/coverage", headers=auth_headers).status_code == 404


def test_the_installed_list_has_the_system_fonts_and_the_added_ones(auth_headers):
    upload(auth_headers, build_font())
    listed = {(f["family"], f["source"]) for f in client.get("/api/v1/fonts/installed", headers=auth_headers).json()}
    assert ("Aksor Test Sans", "uploaded") in listed and ("Arial", "system") in listed


def test_changes_are_in_the_audit_trail(auth_headers):
    font = upload(auth_headers, build_font()).json()
    client.delete(f"/api/v1/fonts/{font['id']}", headers=auth_headers)
    page = client.get("/api/v1/audit?action=font", headers=auth_headers).json()
    actions = [e["action"] for e in page["items"]]
    assert "font.upload" in actions and "font.delete" in actions


# --- a template's fonts ---------------------------------------------------------------------------------------------


def test_a_template_lists_the_fonts_it_names_with_their_status(auth_headers):
    upload(auth_headers, build_font())
    report = register(auth_headers, "Mixed", ["Aksor Test Sans", "Arial", "Times New Roman", "Khmer UI", "Totally Unknown Font"])
    by_name = {f["name"]: f for f in client.get(f"/api/v1/reports/{report}/fonts").json()}
    assert by_name["Aksor Test Sans"]["status"] == "uploaded"
    assert by_name["Arial"]["status"] == "installed"
    # not installed, but a metric-compatible stand-in is: the layout holds
    assert by_name["Times New Roman"] == {"name": "Times New Roman", "status": "substituted", "resolved_to": "Liberation Serif"}
    # not installed and no stand-in: the renderer picks a default, so the layout can change
    assert by_name["Khmer UI"] == {"name": "Khmer UI", "status": "missing", "resolved_to": None}
    assert by_name["Totally Unknown Font"]["status"] == "missing"


def test_the_render_code_is_told_where_the_fonts_are(isolated_fonts):
    assert os.environ["DOC_ENGINE_FONT_DIRS"] == str(isolated_fonts)
