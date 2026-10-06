import zipfile
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import app

client = TestClient(app)


def _make_docx_bytes(paragraph_text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(paragraph_text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_xlsx_bytes() -> bytes:
    """A minimal xlsx template with a `{{ field }}` placeholder and a
    hidden for/endfor row pair around a body row — see
    excel_engine.render_xlsx_template's docstring for the recipe.
    """
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Customer: {{ customer_name }}"
    ws["A2"] = "{% for item in items %}"
    ws.row_dimensions[2].hidden = True
    ws["A3"] = "{{ item.label }}"
    ws["B3"] = "{{ item.amount }}"
    ws["A4"] = "{% endfor %}"
    ws.row_dimensions[4].hidden = True
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _register_report(auth_headers, paragraph_text="Hello {{ customer_name }}", name="Test Report"):
    resp = client.post(
        "/api/v1/reports",
        files={"file": ("template.docx", _make_docx_bytes(paragraph_text), "application/octet-stream")},
        data={"name": name, "description": "A test report"},
        headers=auth_headers,
    )
    return resp


def _register_xlsx_report(auth_headers, name="Test Xlsx Report"):
    resp = client.post(
        "/api/v1/reports",
        files={"file": ("template.xlsx", _make_xlsx_bytes(), "application/octet-stream")},
        data={"name": name, "description": "A test xlsx report"},
        headers=auth_headers,
    )
    return resp


def test_create_report(auth_headers):
    resp = _register_report(auth_headers)
    assert resp.status_code == 200
    body = resp.json()
    assert body["name"] == "Test Report"
    assert body["description"] == "A test report"
    assert "report_id" in body and body["report_id"]
    assert body["version"] == 1
    assert body["created_at"] == body["updated_at"]


def test_create_report_requires_auth():
    resp = client.post(
        "/api/v1/reports",
        files={"file": ("template.docx", _make_docx_bytes("hi"), "application/octet-stream")},
        data={"name": "No auth"},
    )
    assert resp.status_code == 401


def test_create_report_rejects_non_docx(auth_headers):
    resp = client.post(
        "/api/v1/reports",
        files={"file": ("template.txt", b"not a docx", "text/plain")},
        data={"name": "Bad"},
        headers=auth_headers,
    )
    assert resp.status_code == 400


def test_create_report_rejects_invalid_docx_content(auth_headers):
    resp = client.post(
        "/api/v1/reports",
        files={"file": ("template.docx", b"not actually a zip", "application/octet-stream")},
        data={"name": "Bad"},
        headers=auth_headers,
    )
    assert resp.status_code == 400


def test_list_reports_includes_created_report(auth_headers):
    created = _register_report(auth_headers).json()
    resp = client.get("/api/v1/reports")  # list stays public
    assert resp.status_code == 200
    ids = [r["report_id"] for r in resp.json()]
    assert created["report_id"] in ids


def test_get_report_metadata(auth_headers):
    created = _register_report(auth_headers, name="Invoice").json()
    resp = client.get(f"/api/v1/reports/{created['report_id']}")  # get stays public
    assert resp.status_code == 200
    assert resp.json()["name"] == "Invoice"


def test_get_report_not_found():
    resp = client.get("/api/v1/reports/does-not-exist")
    assert resp.status_code == 404


def test_render_report_docx(auth_headers):
    created = _register_report(auth_headers, "Invoice for {{ customer_name }}: {{ amount }}").json()
    resp = client.post(
        f"/api/v1/reports/{created['report_id']}/render?format=docx",
        json={"customer_name": "សុខ សុភា", "amount": "100"},
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert resp.content[:2] == b"PK"


def test_render_report_pdf(auth_headers):
    created = _register_report(auth_headers, "Invoice for {{ customer_name }}: {{ amount }}").json()
    resp = client.post(
        f"/api/v1/reports/{created['report_id']}/render?format=pdf",
        json={"customer_name": "សុខ សុភា", "amount": "100"},
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content.startswith(b"%PDF")


def test_render_docx_report_rejects_xlsx_format(auth_headers):
    created = _register_report(auth_headers).json()
    resp = client.post(
        f"/api/v1/reports/{created['report_id']}/render?format=xlsx",
        json={"customer_name": "x"},
    )
    assert resp.status_code == 400  # this report is a .docx template


def test_create_xlsx_report(auth_headers):
    resp = _register_xlsx_report(auth_headers)
    assert resp.status_code == 200
    assert resp.json()["name"] == "Test Xlsx Report"


def test_render_xlsx_report(auth_headers):
    created = _register_xlsx_report(auth_headers).json()
    resp = client.post(
        f"/api/v1/reports/{created['report_id']}/render?format=xlsx",
        json={
            "customer_name": "សុខ សុភា",
            "items": [{"label": "Widget A", "amount": 10}, {"label": "Widget B", "amount": 20}],
        },
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")

    wb = load_workbook(BytesIO(resp.content))
    ws = wb.active
    # Every string (including plain Latin text) goes through
    # segment_generic for custom templates, so ZWSP may appear at word
    # boundaries — strip it before comparing, same as elsewhere.
    rows = [
        tuple(c.replace("​", "") if isinstance(c, str) else c for c in r)
        for r in ws.iter_rows(values_only=True)
        if any(c is not None for c in r)
    ]
    assert ("Widget A", 10) in rows
    assert ("Widget B", 20) in rows


def test_render_xlsx_report_rejects_pdf_format(auth_headers):
    created = _register_xlsx_report(auth_headers).json()
    resp = client.post(
        f"/api/v1/reports/{created['report_id']}/render?format=pdf",
        json={"customer_name": "x", "items": []},
    )
    assert resp.status_code == 400  # this report is a .xlsx template


def test_render_report_not_found():
    resp = client.post(
        "/api/v1/reports/does-not-exist/render?format=pdf",
        json={"customer_name": "x"},
    )
    assert resp.status_code == 404


def test_delete_report(auth_headers):
    created = _register_report(auth_headers).json()
    resp = client.delete(f"/api/v1/reports/{created['report_id']}", headers=auth_headers)
    assert resp.status_code == 204
    assert client.get(f"/api/v1/reports/{created['report_id']}").status_code == 404


def test_delete_report_not_found(auth_headers):
    resp = client.delete("/api/v1/reports/does-not-exist", headers=auth_headers)
    assert resp.status_code == 404


def test_delete_report_requires_auth(auth_headers):
    created = _register_report(auth_headers).json()
    resp = client.delete(f"/api/v1/reports/{created['report_id']}")
    assert resp.status_code == 401
    # never actually deleted
    assert client.get(f"/api/v1/reports/{created['report_id']}").status_code == 200


def test_update_report_metadata(auth_headers):
    created = _register_report(auth_headers, name="Old Name").json()
    resp = client.patch(
        f"/api/v1/reports/{created['report_id']}",
        json={"name": "New Name", "description": "New description"},
        headers=auth_headers,
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["name"] == "New Name"
    assert body["description"] == "New description"
    assert body["version"] == 1  # metadata edit doesn't bump version
    assert body["updated_at"] >= body["created_at"]


def test_is_public_defaults_to_false_on_create(auth_headers):
    created = _register_report(auth_headers, name="Fresh").json()
    assert created["is_public"] is False


def test_update_report_can_toggle_is_public(auth_headers):
    created = _register_report(auth_headers, name="Togglable").json()
    report_id = created["report_id"]

    resp = client.patch(f"/api/v1/reports/{report_id}", json={"is_public": True}, headers=auth_headers)
    assert resp.status_code == 200
    assert resp.json()["is_public"] is True

    resp = client.patch(f"/api/v1/reports/{report_id}", json={"is_public": False}, headers=auth_headers)
    assert resp.status_code == 200
    assert resp.json()["is_public"] is False


def test_update_report_metadata_requires_auth(auth_headers):
    created = _register_report(auth_headers).json()
    resp = client.patch(f"/api/v1/reports/{created['report_id']}", json={"name": "x"})
    assert resp.status_code == 401


def test_update_report_metadata_not_found(auth_headers):
    resp = client.patch("/api/v1/reports/does-not-exist", json={"name": "x"}, headers=auth_headers)
    assert resp.status_code == 404


def test_replace_report_file_bumps_version(auth_headers):
    created = _register_report(auth_headers, "Original {{ x }}").json()
    resp = client.put(
        f"/api/v1/reports/{created['report_id']}/file",
        files={"file": ("new.docx", _make_docx_bytes("Replaced {{ x }}"), "application/octet-stream")},
        headers=auth_headers,
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["version"] == 2
    assert body["report_id"] == created["report_id"]

    render_resp = client.post(
        f"/api/v1/reports/{created['report_id']}/render?format=docx",
        json={"x": "hello"},
    )
    assert render_resp.status_code == 200


def test_replace_report_file_requires_auth(auth_headers):
    created = _register_report(auth_headers).json()
    resp = client.put(
        f"/api/v1/reports/{created['report_id']}/file",
        files={"file": ("new.docx", _make_docx_bytes("x"), "application/octet-stream")},
    )
    assert resp.status_code == 401


def test_replace_report_file_rejects_invalid_content(auth_headers):
    created = _register_report(auth_headers).json()
    resp = client.put(
        f"/api/v1/reports/{created['report_id']}/file",
        files={"file": ("new.docx", b"not a zip", "application/octet-stream")},
        headers=auth_headers,
    )
    assert resp.status_code == 400


# --- sample_context ---------------------------------------------------


def test_create_report_defaults_sample_context_to_none(auth_headers):
    created = _register_report(auth_headers).json()
    assert created["sample_context"] is None


def test_patch_report_saves_sample_context(auth_headers):
    created = _register_report(auth_headers).json()
    resp = client.patch(
        f"/api/v1/reports/{created['report_id']}",
        json={"sample_context": {"customer_name": "សុខ សុភា", "amount": "100"}},
        headers=auth_headers,
    )
    assert resp.status_code == 200
    assert resp.json()["sample_context"] == {"customer_name": "សុខ សុភា", "amount": "100"}

    # persists across a fresh GET, not just the PATCH response
    fetched = client.get(f"/api/v1/reports/{created['report_id']}")
    assert fetched.json()["sample_context"] == {"customer_name": "សុខ សុភា", "amount": "100"}


def test_patch_report_without_sample_context_leaves_it_unchanged(auth_headers):
    created = _register_report(auth_headers).json()
    client.patch(
        f"/api/v1/reports/{created['report_id']}",
        json={"sample_context": {"a": 1}},
        headers=auth_headers,
    )
    resp = client.patch(
        f"/api/v1/reports/{created['report_id']}",
        json={"name": "Renamed only"},
        headers=auth_headers,
    )
    assert resp.status_code == 200
    assert resp.json()["sample_context"] == {"a": 1}


# --- schema -------------------------------------------------------------


def test_get_docx_report_schema(auth_headers):
    created = _register_report(auth_headers, "Invoice for {{ customer_name }}: {{ amount }}").json()
    resp = client.get(f"/api/v1/reports/{created['report_id']}/schema")
    assert resp.status_code == 200
    body = resp.json()
    assert body["engine"] == "docxtpl"
    assert set(body["fields"]) >= {"customer_name", "amount"}


def test_get_xlsx_report_schema(auth_headers):
    created = _register_xlsx_report(auth_headers).json()
    resp = client.get(f"/api/v1/reports/{created['report_id']}/schema")
    assert resp.status_code == 200
    body = resp.json()
    assert body["engine"] == "regex-scan"
    assert set(body["fields"]) >= {"customer_name", "item", "items"}
    assert "for" not in body["fields"] and "endfor" not in body["fields"]


def test_get_report_schema_not_found():
    resp = client.get("/api/v1/reports/does-not-exist/schema")
    assert resp.status_code == 404


# --- batch render ---------------------------------------------------------


def test_render_report_batch_docx(auth_headers):
    created = _register_report(auth_headers, "Invoice for {{ customer_name }}: {{ amount }}").json()
    contexts = [
        {"customer_name": "សុខ សុភា", "amount": "100"},
        {"customer_name": "ចាន់ណារិទ្ធ", "amount": "200"},
        {"customer_name": "Third Customer", "amount": "300"},
    ]
    resp = client.post(
        f"/api/v1/reports/{created['report_id']}/render/batch?format=docx",
        json=contexts,
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/zip"

    zf = zipfile.ZipFile(BytesIO(resp.content))
    names = sorted(zf.namelist())
    assert names == [
        f"{created['report_id']}-0000.docx",
        f"{created['report_id']}-0001.docx",
        f"{created['report_id']}-0002.docx",
    ]
    for name in names:
        assert zf.read(name)[:2] == b"PK"


def test_render_report_batch_empty_rejected(auth_headers):
    created = _register_report(auth_headers).json()
    resp = client.post(f"/api/v1/reports/{created['report_id']}/render/batch?format=docx", json=[])
    assert resp.status_code == 400


def test_render_report_batch_over_cap_rejected(auth_headers, monkeypatch):
    monkeypatch.setenv("MAX_BATCH_DOCX_XLSX", "2")
    created = _register_report(auth_headers).json()
    resp = client.post(
        f"/api/v1/reports/{created['report_id']}/render/batch?format=docx",
        json=[{"customer_name": "a"}, {"customer_name": "b"}, {"customer_name": "c"}],
    )
    assert resp.status_code == 400


def test_render_report_batch_rejects_wrong_format(auth_headers):
    created = _register_report(auth_headers).json()
    resp = client.post(
        f"/api/v1/reports/{created['report_id']}/render/batch?format=xlsx",
        json=[{"customer_name": "a"}],
    )
    assert resp.status_code == 400  # this report is a .docx template


def test_render_report_batch_not_found():
    resp = client.post(
        "/api/v1/reports/does-not-exist/render/batch?format=pdf",
        json=[{"x": "1"}],
    )
    assert resp.status_code == 404
