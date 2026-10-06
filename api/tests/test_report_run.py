"""Filter parameters, grant limits, and the authenticated run path.

The property everything here protects: a user may only run a report with
parameter values their grants allow, and the report's data source is only
ever called *after* that check passes -- so a forbidden value never causes
a fetch at all, and the browser never supplies the data.
"""
from dataclasses import dataclass, field
from io import BytesIO

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import db, report_data, report_store
from app.main import app
from app.rbac import ROOT_ORG_ID, create_organization

client = TestClient(app)

BRANCHES = [
    {"value": "BR01", "label": "Phnom Penh"},
    {"value": "BR02", "label": "Siem Reap"},
    {"value": "BR03", "label": "Battambang"},
]
P_BRANCH = {"name": "p_branch", "label": "Branch", "options": BRANCHES}
DATA_SOURCE = {"url": "https://erp.example.test/api/sales?branch={{ p_branch }}", "method": "GET"}


# --- helpers -------------------------------------------------------------


def _docx_bytes(text: str = "Branch {{ p_branch }}: {{ branch_name }} total {{ total }}") -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _docx_text(content: bytes) -> str:
    return "\n".join(p.text for p in Document(BytesIO(content)).paragraphs).replace("​", "")


def _register(auth_headers, text: str | None = None, name: str = "Sales") -> str:
    resp = client.post(
        "/api/v1/reports",
        files={"file": ("t.docx", _docx_bytes() if text is None else _docx_bytes(text), "application/octet-stream")},
        data={"name": name},
        headers=auth_headers,
    )
    assert resp.status_code == 200, resp.text
    return resp.json()["report_id"]


def _configure(auth_headers, report_id, parameters=None, data_source=None):
    resp = client.put(
        f"/api/v1/reports/{report_id}/data-config",
        json={"parameters": [P_BRANCH] if parameters is None else parameters, "data_source": data_source},
        headers=auth_headers,
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


def _grant(auth_headers, subject_type, subject_id, report_id, level="render", limits=None, expect=200):
    body = {"subject_type": subject_type, "subject_id": subject_id, "report_id": report_id, "permission_level": level}
    if limits is not None:
        body["parameter_limits"] = limits
    resp = client.post("/api/v1/grants/reports", json=body, headers=auth_headers)
    assert resp.status_code == expect, resp.text
    return resp


def _run(headers, report_id, parameters, fmt="docx"):
    return client.post(f"/api/v1/reports/{report_id}/run", json={"parameters": parameters, "format": fmt}, headers=headers)


def _form(headers, report_id):
    return client.get(f"/api/v1/reports/{report_id}/run-form", headers=headers)


@dataclass
class FakeDataSource:
    """Stands in for the report's REST data source: records every request
    the server makes and answers with `respond(request)`."""

    calls: list = field(default_factory=list)
    respond: object = None

    def __post_init__(self):
        def default(request: httpx.Request) -> httpx.Response:
            branch = request.url.params.get("branch", "?")
            return httpx.Response(200, json={"branch_name": f"Branch {branch} Office", "total": "1200"})

        self.respond = default


@pytest.fixture
def fake_source(monkeypatch) -> FakeDataSource:
    fake = FakeDataSource()

    def factory():
        def handler(request):
            fake.calls.append(request)
            return fake.respond(request)

        return httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=False)

    monkeypatch.setattr(report_data, "_make_client", factory)
    return fake


@pytest.fixture
def sales_report(auth_headers) -> str:
    report_id = _register(auth_headers)
    _configure(auth_headers, report_id, data_source=DATA_SOURCE)
    return report_id


# --- validate_data_config ------------------------------------------------


def _validate(parameters=None, data_source=None):
    return report_data.validate_data_config(parameters if parameters is not None else [P_BRANCH], data_source)


@pytest.mark.parametrize("name", ["", "1abc", "p-branch", "p branch", "p.branch"])
def test_parameter_names_must_be_identifiers(name):
    with pytest.raises(report_data.DataConfigError, match="invalid"):
        _validate([{"name": name}])


def test_duplicate_parameter_names_rejected():
    with pytest.raises(report_data.DataConfigError, match="more than once"):
        _validate([{"name": "a"}, {"name": "a"}])


def test_empty_option_list_rejected():
    with pytest.raises(report_data.DataConfigError, match="empty option list"):
        _validate([{"name": "p_branch", "options": []}])


def test_duplicate_option_values_rejected():
    with pytest.raises(report_data.DataConfigError, match="more than once"):
        _validate([{"name": "p", "options": [{"value": "A"}, {"value": " A "}]}])


def test_parameters_are_normalized():
    cleaned, _ = _validate([{"name": " p ", "label": "  ", "options": [{"value": " A ", "label": " Alpha "}]}])
    assert cleaned == [
        {
            "name": "p",
            "label": None,
            "type": "text",
            "required": True,
            "default_value": None,
            "options": [{"value": "A", "label": "Alpha"}],
            "options_source": None,
        }
    ]


def test_parameter_type_defaults_to_text():
    cleaned, _ = _validate([{"name": "q"}])
    assert cleaned[0]["type"] == "text"


@pytest.mark.parametrize("type_", ["date", "datetime", "time", "number"])
def test_parameter_type_is_accepted(type_):
    cleaned, _ = _validate([{"name": "q", "type": type_}])
    assert cleaned[0]["type"] == type_


def test_parameter_type_rejects_unknown_values():
    with pytest.raises(report_data.DataConfigError, match="invalid type"):
        _validate([{"name": "q", "type": "datepicker"}])


def test_parameter_required_defaults_to_true():
    cleaned, _ = _validate([{"name": "q"}])
    assert cleaned[0]["required"] is True


def test_free_text_parameter_can_be_made_optional():
    cleaned, _ = _validate([{"name": "q", "required": False}])
    assert cleaned[0]["required"] is False


def test_choice_list_parameter_cannot_be_made_optional():
    with pytest.raises(report_data.DataConfigError, match="can't be optional"):
        _validate([{"name": "p_branch", "required": False, "options": [{"value": "A"}]}])


def test_options_source_parameter_cannot_be_made_optional():
    with pytest.raises(report_data.DataConfigError, match="can't be optional"):
        _validate([{"name": "p_branch", "required": False, "options_source": OPTIONS_SOURCE}])


@pytest.mark.parametrize(
    "url",
    ["ftp://x.test/a", "{{ p_branch }}://x.test/a", "https://{{ p_branch }}.test/a", "/relative", "", "javascript:alert(1)"],
)
def test_data_source_url_needs_literal_scheme_and_host(url):
    with pytest.raises(report_data.DataConfigError):
        _validate(data_source={"url": url})


def test_data_source_url_may_template_after_the_host():
    _, source = _validate(data_source={"url": "https://x.test/{{ p_branch }}/sales?b={{ p_branch }}"})
    assert source["method"] == "GET"


def test_data_source_unknown_variable_is_flagged():
    with pytest.raises(report_data.DataConfigError, match=r"\{\{ p_bracnh \}\}"):
        _validate(data_source={"url": "https://x.test/?b={{ p_bracnh }}"})


def test_data_source_template_syntax_error_is_readable():
    with pytest.raises(report_data.DataConfigError, match="syntax error"):
        _validate(data_source={"url": "https://x.test/?b={{ p_branch "})


def test_body_only_allowed_with_post():
    with pytest.raises(report_data.DataConfigError, match="only sent with POST"):
        _validate(data_source={"url": "https://x.test/", "method": "GET", "body_template": {"b": "{{ p_branch }}"}})
    _, source = _validate(data_source={"url": "https://x.test/", "method": "POST", "body_template": {"b": "{{ p_branch }}"}})
    assert source["body_template"] == {"b": "{{ p_branch }}"}


def test_body_variables_are_checked_too():
    with pytest.raises(report_data.DataConfigError, match=r"\{\{ nope \}\}"):
        _validate(data_source={"url": "https://x.test/", "method": "POST", "body_template": {"a": {"b": "{{ nope }}"}}})


def test_auth_and_header_rules():
    with pytest.raises(report_data.DataConfigError, match="environment variable"):
        _validate(data_source={"url": "https://x.test/", "auth": {"type": "bearer", "token_env": "not valid"}})
    with pytest.raises(report_data.DataConfigError, match="username"):
        _validate(data_source={"url": "https://x.test/", "auth": {"type": "basic", "password_env": "PW"}})
    with pytest.raises(report_data.DataConfigError, match="set automatically"):
        _validate(data_source={"url": "https://x.test/", "headers": {"Host": "evil.test"}})
    with pytest.raises(report_data.DataConfigError, match="Authorization header"):
        _validate(
            data_source={
                "url": "https://x.test/",
                "headers": {"authorization": "x"},
                "auth": {"type": "bearer", "token_env": "TOKEN"},
            }
        )


# --- options_source (a parameter's own options, fetched live) -----------

OPTIONS_SOURCE = {"url": "https://branches.example.test/list", "value_field": "code", "label_field": "name"}
P_BRANCH_DYNAMIC = {"name": "p_branch", "label": "Branch", "options_source": OPTIONS_SOURCE}


def test_options_source_and_static_options_are_mutually_exclusive():
    with pytest.raises(report_data.DataConfigError, match="choose one"):
        _validate([{"name": "p_branch", "options": [{"value": "A"}], "options_source": OPTIONS_SOURCE}])


def test_options_source_needs_a_value_field():
    with pytest.raises(report_data.DataConfigError, match="needs a value"):
        _validate([{"name": "p_branch", "options_source": {"url": "https://x.test/"}}])


def test_options_source_reuses_data_source_url_and_header_rules():
    # Same validators as data_source (_validate_url / _validate_headers_and_auth) --
    # spot-check one of each rather than re-testing every case.
    with pytest.raises(report_data.DataConfigError, match="literal http"):
        _validate([{"name": "p_branch", "options_source": {**OPTIONS_SOURCE, "url": "ftp://x.test/"}}])
    with pytest.raises(report_data.DataConfigError, match="set automatically"):
        _validate([{"name": "p_branch", "options_source": {**OPTIONS_SOURCE, "headers": {"Host": "evil.test"}}}])
    with pytest.raises(report_data.DataConfigError, match="only sent with POST"):
        _validate([{"name": "p_branch", "options_source": {**OPTIONS_SOURCE, "body": {"active": True}}}])


def test_options_source_round_trips_through_data_config(auth_headers):
    report_id = _register(auth_headers)
    saved = _configure(auth_headers, report_id, parameters=[P_BRANCH_DYNAMIC])
    assert saved["parameters"][0]["options"] is None
    assert saved["parameters"][0]["options_source"]["value_field"] == "code"


# --- data-config API -----------------------------------------------------


def test_data_config_round_trip_and_never_public(auth_headers):
    report_id = _register(auth_headers)
    saved = _configure(auth_headers, report_id, data_source={**DATA_SOURCE, "headers": {"X-Api-Key": "k"}})
    assert saved["parameters"][0]["name"] == "p_branch"
    assert client.get(f"/api/v1/reports/{report_id}/data-config", headers=auth_headers).json() == saved

    # Neither the public single-report route nor the public list carries it.
    for body in (client.get(f"/api/v1/reports/{report_id}").json(), client.get("/api/v1/reports").json()[0]):
        assert "parameters" not in body and "data_source" not in body


def test_data_config_can_be_cleared(auth_headers):
    report_id = _register(auth_headers)
    _configure(auth_headers, report_id, data_source=DATA_SOURCE)
    cleared = _configure(auth_headers, report_id, parameters=[], data_source=None)
    assert cleared == {"parameters": [], "data_source": None}


def test_data_config_requires_manage(auth_headers, make_local_user):
    report_id = _register(auth_headers)
    user_id, headers = make_local_user("renderer", "pw12345", role_names=())
    _grant(auth_headers, "user", user_id, report_id, "render")
    assert client.get(f"/api/v1/reports/{report_id}/data-config", headers=headers).status_code == 403
    assert client.put(f"/api/v1/reports/{report_id}/data-config", json={"parameters": []}, headers=headers).status_code == 403
    assert client.get(f"/api/v1/reports/{report_id}/data-config").status_code == 401


def test_data_config_invalid_returns_a_readable_400(auth_headers):
    report_id = _register(auth_headers)
    resp = client.put(
        f"/api/v1/reports/{report_id}/data-config",
        json={"parameters": [{"name": "1bad"}], "data_source": None},
        headers=auth_headers,
    )
    assert resp.status_code == 400
    assert isinstance(resp.json()["detail"], str) and "1bad" in resp.json()["detail"]


# --- grants with parameter limits ---------------------------------------


def test_grant_limits_are_stored_sorted_and_returned(auth_headers, make_local_user):
    report_id = _register(auth_headers)
    _configure(auth_headers, report_id)
    user_id, _ = make_local_user("limited", "pw12345", role_names=())
    resp = _grant(auth_headers, "user", user_id, report_id, "render", {"p_branch": ["BR02", "BR01", "BR02"]})
    assert resp.json()["parameter_limits"] == {"p_branch": ["BR01", "BR02"]}
    listed = client.get("/api/v1/grants/reports", params={"report_id": report_id}, headers=auth_headers).json()
    assert listed[0]["parameter_limits"] == {"p_branch": ["BR01", "BR02"]}


@pytest.mark.parametrize(
    "limits,level,message",
    [
        ({"p_nope": ["BR01"]}, "render", "isn't a filter parameter"),
        ({"p_branch": ["BR99"]}, "render", "Not an option"),
        ({"p_branch": []}, "render", "empty"),
        ({"p_branch": ["BR01"]}, "view", "only apply to render or manage"),
    ],
)
def test_grant_limit_validation(auth_headers, make_local_user, limits, level, message):
    report_id = _register(auth_headers)
    _configure(auth_headers, report_id)
    user_id, _ = make_local_user("bad_limits", "pw12345", role_names=())
    resp = _grant(auth_headers, "user", user_id, report_id, level, limits, expect=400)
    assert message in resp.json()["detail"]


def test_cannot_limit_a_free_text_parameter(auth_headers, make_local_user):
    report_id = _register(auth_headers)
    _configure(auth_headers, report_id, parameters=[{"name": "q"}])
    user_id, _ = make_local_user("free_text", "pw12345", role_names=())
    resp = _grant(auth_headers, "user", user_id, report_id, "render", {"q": ["x"]}, expect=400)
    assert "no option list" in resp.json()["detail"]


def test_can_limit_a_dynamic_options_parameter_without_a_live_fetch(auth_headers, make_local_user):
    # No fake_source fixture here on purpose -- creating the grant must
    # not need to reach the options_source at all (see
    # _validate_parameter_limits's docstring).
    report_id = _register(auth_headers)
    _configure(auth_headers, report_id, parameters=[P_BRANCH_DYNAMIC])
    user_id, _ = make_local_user("dyn_limited", "pw12345", role_names=())
    resp = _grant(auth_headers, "user", user_id, report_id, "render", {"p_branch": ["PP01"]})
    assert resp.json()["parameter_limits"] == {"p_branch": ["PP01"]}


# --- run-form ------------------------------------------------------------


def test_run_form_narrows_options_to_the_grant(auth_headers, make_local_user):
    report_id = _register(auth_headers)
    _configure(auth_headers, report_id, parameters=[P_BRANCH, {"name": "note", "label": "Note"}])
    user_id, headers = make_local_user("branch_user", "pw12345", role_names=())
    _grant(auth_headers, "user", user_id, report_id, "render", {"p_branch": ["BR01"]})

    form = _form(headers, report_id).json()
    assert form["access_level"] == "render"
    assert form["formats"] == ["pdf", "png", "docx"]
    by_name = {p["name"]: p for p in form["parameters"]}
    assert by_name["p_branch"]["options"] == [{"value": "BR01", "label": "Phnom Penh"}]
    assert by_name["note"]["options"] is None  # free text
    # The other branches must not appear anywhere in the response.
    assert "BR02" not in _form(headers, report_id).text and "Siem Reap" not in _form(headers, report_id).text


def test_run_form_exposes_parameter_type(auth_headers):
    report_id = _register(auth_headers)
    _configure(
        auth_headers,
        report_id,
        parameters=[{"name": "fromDate", "label": "From date", "type": "date"}, {"name": "q", "label": "Search"}],
    )
    by_name = {p["name"]: p for p in _form(auth_headers, report_id).json()["parameters"]}
    assert by_name["fromDate"]["type"] == "date"
    assert by_name["q"]["type"] == "text"  # default, unset


def test_run_form_exposes_required(auth_headers):
    report_id = _register(auth_headers)
    _configure(
        auth_headers,
        report_id,
        parameters=[{"name": "q", "label": "Search", "required": False}, P_BRANCH],
    )
    by_name = {p["name"]: p for p in _form(auth_headers, report_id).json()["parameters"]}
    assert by_name["q"]["required"] is False
    assert by_name["p_branch"]["required"] is True  # choice list -- always required


def test_run_form_is_unrestricted_for_global_render_permission(auth_headers, make_local_user):
    report_id = _register(auth_headers)
    _configure(auth_headers, report_id)
    _, headers = make_local_user("global_viewer", "pw12345", role_names=("ROLE_REPORT_VIEWER",))
    values = [o["value"] for o in _form(headers, report_id).json()["parameters"][0]["options"]]
    assert values == ["BR01", "BR02", "BR03"]


def test_run_form_access_rules(auth_headers, make_local_user):
    report_id = _register(auth_headers)
    _configure(auth_headers, report_id)
    view_id, view_headers = make_local_user("view_only", "pw12345", role_names=())
    _grant(auth_headers, "user", view_id, report_id, "view")
    _, nobody_headers = make_local_user("nobody", "pw12345", role_names=())

    assert _form(view_headers, report_id).status_code == 403  # can view, can't run
    assert _form(nobody_headers, report_id).status_code == 404  # existence isn't theirs to learn
    assert client.get(f"/api/v1/reports/{report_id}/run-form").status_code == 401
    assert _form(auth_headers, "does-not-exist").status_code == 404


def test_reports_in_another_org_are_invisible_to_run(auth_headers, make_local_user):
    with db.SessionLocal() as session:
        create_organization(session, "other-org", "Other Org")
        session.commit()
    foreign = report_store.create_report(name="Foreign", content=_docx_bytes(), template_ext="docx", org_id="other-org")["report_id"]
    _, headers = make_local_user("root_renderer", "pw12345", role_names=("ROLE_REPORT_VIEWER",))
    assert _form(headers, foreign).status_code == 404
    assert _run(headers, foreign, {}).status_code == 404


# --- run-form: options_source resolved live -------------------------------


@pytest.fixture
def dynamic_branch_report(auth_headers) -> str:
    report_id = _register(auth_headers)
    _configure(auth_headers, report_id, parameters=[P_BRANCH_DYNAMIC])
    return report_id


def _respond_branches(request: httpx.Request) -> httpx.Response:
    return httpx.Response(200, json=[{"code": "PP01", "name": "Phnom Penh"}, {"code": "SR01", "name": "Siem Reap"}])


def test_run_form_resolves_options_source_live(auth_headers, dynamic_branch_report, fake_source):
    fake_source.respond = _respond_branches
    form = _form(auth_headers, dynamic_branch_report).json()
    assert form["parameters"][0]["options"] == [{"value": "PP01", "label": "Phnom Penh"}, {"value": "SR01", "label": "Siem Reap"}]
    assert fake_source.calls[0].url.host == "branches.example.test"


def test_run_form_options_source_is_still_narrowed_by_grant(auth_headers, make_local_user, dynamic_branch_report, fake_source):
    fake_source.respond = _respond_branches
    user_id, headers = make_local_user("dyn_branch_user", "pw12345", role_names=())
    _grant(auth_headers, "user", user_id, dynamic_branch_report, "render", {"p_branch": ["PP01"]})
    form = _form(headers, dynamic_branch_report).json()
    assert form["parameters"][0]["options"] == [{"value": "PP01", "label": "Phnom Penh"}]


def test_run_accepts_a_value_from_the_live_options_and_rejects_one_not_in_it(auth_headers, dynamic_branch_report, fake_source):
    fake_source.respond = _respond_branches
    assert _run(auth_headers, dynamic_branch_report, {"p_branch": "PP01"}).status_code == 200
    # Not free text just because the list is fetched, not stored -- a
    # value the live source never advertised is still rejected, closing
    # the gap resolve_parameter_definitions exists to close.
    resp = _run(auth_headers, dynamic_branch_report, {"p_branch": "MADE-UP"})
    assert resp.status_code == 400 and "valid choice" in resp.json()["detail"]


@pytest.mark.parametrize(
    "respond",
    [
        lambda request: httpx.Response(500, text="boom"),
        lambda request: httpx.Response(200, json={"not": "an array"}),
        lambda request: httpx.Response(200, json=[{"no_code_field": "x"}]),
    ],
    ids=["http-500", "not-an-array", "item-missing-value-field"],
)
def test_options_source_failures_are_a_502_on_run_form_and_run(auth_headers, dynamic_branch_report, fake_source, respond):
    fake_source.respond = respond
    form_resp = _form(auth_headers, dynamic_branch_report)
    assert form_resp.status_code == 502 and "Branch" in form_resp.json()["detail"]
    run_resp = _run(auth_headers, dynamic_branch_report, {"p_branch": "PP01"})
    assert run_resp.status_code == 502


def test_options_source_result_is_deduplicated_and_capped(auth_headers, dynamic_branch_report, fake_source, monkeypatch):
    monkeypatch.setattr(report_data, "MAX_OPTIONS", 2)
    fake_source.respond = lambda request: httpx.Response(
        200, json=[{"code": "A"}, {"code": "A"}, {"code": "B"}, {"code": "C"}]
    )
    form = _form(auth_headers, dynamic_branch_report).json()
    assert [o["value"] for o in form["parameters"][0]["options"]] == ["A", "B"]


# --- run: the protection -------------------------------------------------


def test_run_with_an_allowed_value_fetches_and_renders(auth_headers, make_local_user, sales_report, fake_source):
    user_id, headers = make_local_user("br01_user", "pw12345", role_names=())
    _grant(auth_headers, "user", user_id, sales_report, "render", {"p_branch": ["BR01"]})

    resp = _run(headers, sales_report, {"p_branch": "BR01"})
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert len(fake_source.calls) == 1 and fake_source.calls[0].url.params["branch"] == "BR01"
    assert "Branch BR01: Branch BR01 Office total 1200" in _docx_text(resp.content)


def test_run_with_a_forbidden_value_is_403_and_never_calls_the_data_source(auth_headers, make_local_user, sales_report, fake_source):
    user_id, headers = make_local_user("br01_only", "pw12345", role_names=())
    _grant(auth_headers, "user", user_id, sales_report, "render", {"p_branch": ["BR01"]})

    resp = _run(headers, sales_report, {"p_branch": "BR02"})
    assert resp.status_code == 403
    assert fake_source.calls == []


def test_forbidden_and_nonexistent_values_are_indistinguishable(auth_headers, make_local_user, sales_report, fake_source):
    """A restricted user must not be able to probe which branches exist."""
    user_id, headers = make_local_user("prober", "pw12345", role_names=())
    _grant(auth_headers, "user", user_id, sales_report, "render", {"p_branch": ["BR01"]})
    real_but_forbidden = _run(headers, sales_report, {"p_branch": "BR02"})
    made_up = _run(headers, sales_report, {"p_branch": "BR99"})
    assert (real_but_forbidden.status_code, real_but_forbidden.json()) == (made_up.status_code, made_up.json())
    assert fake_source.calls == []


def test_unrestricted_caller_gets_400_for_a_nonexistent_value(auth_headers, sales_report, fake_source):
    resp = _run(auth_headers, sales_report, {"p_branch": "BR99"})
    assert resp.status_code == 400
    assert fake_source.calls == []


def test_superuser_and_global_render_permission_are_unrestricted(auth_headers, make_local_user, sales_report, fake_source):
    assert _run(auth_headers, sales_report, {"p_branch": "BR03"}).status_code == 200
    _, headers = make_local_user("global_renderer", "pw12345", role_names=("ROLE_REPORT_VIEWER",))
    assert _run(headers, sales_report, {"p_branch": "BR02"}).status_code == 200


def test_view_only_cannot_run(auth_headers, make_local_user, sales_report, fake_source):
    user_id, headers = make_local_user("viewer", "pw12345", role_names=())
    _grant(auth_headers, "user", user_id, sales_report, "view")
    assert _run(headers, sales_report, {"p_branch": "BR01"}).status_code == 403
    assert fake_source.calls == []


def test_public_report_view_only_cannot_run(auth_headers, make_local_user, sales_report, fake_source):
    from app import report_store

    report_store.update_report_meta(sales_report, is_public=True)
    _, headers = make_local_user("public_viewer_2", "pw12345", role_names=())
    assert _run(headers, sales_report, {"p_branch": "BR01"}).status_code == 403
    assert fake_source.calls == []


def test_run_requires_authentication(sales_report, fake_source):
    assert client.post(f"/api/v1/reports/{sales_report}/run", json={"parameters": {"p_branch": "BR01"}}).status_code == 401
    assert fake_source.calls == []


def test_grants_combine_as_a_union_and_an_unrestricted_grant_wins(auth_headers, make_local_user, sales_report, fake_source):
    user_id, headers = make_local_user("multi", "pw12345", role_names=())
    role_id = client.post("/api/v1/roles", json={"org_id": ROOT_ORG_ID, "name": "Branch Two"}, headers=auth_headers).json()["id"]
    client.post("/api/v1/grants/roles", json={"user_id": user_id, "role_id": role_id}, headers=auth_headers)
    _grant(auth_headers, "user", user_id, sales_report, "render", {"p_branch": ["BR01"]})
    _grant(auth_headers, "role", role_id, sales_report, "render", {"p_branch": ["BR02"]})

    assert _run(headers, sales_report, {"p_branch": "BR01"}).status_code == 200
    assert _run(headers, sales_report, {"p_branch": "BR02"}).status_code == 200
    assert _run(headers, sales_report, {"p_branch": "BR03"}).status_code == 403

    # A third grant that doesn't limit p_branch at all opens everything.
    _grant(auth_headers, "user", user_id, sales_report, "render")
    assert _run(headers, sales_report, {"p_branch": "BR03"}).status_code == 200


def test_a_view_grant_does_not_widen_a_limited_render_grant(auth_headers, make_local_user, sales_report, fake_source):
    user_id, headers = make_local_user("view_plus_limited", "pw12345", role_names=())
    _grant(auth_headers, "user", user_id, sales_report, "view")
    _grant(auth_headers, "user", user_id, sales_report, "render", {"p_branch": ["BR01"]})
    assert _run(headers, sales_report, {"p_branch": "BR02"}).status_code == 403


def test_expired_grant_no_longer_runs(auth_headers, make_local_user, sales_report, fake_source):
    from datetime import datetime, timedelta, timezone

    user_id, headers = make_local_user("expired", "pw12345", role_names=())
    resp = client.post(
        "/api/v1/grants/reports",
        json={
            "subject_type": "user",
            "subject_id": user_id,
            "report_id": sales_report,
            "permission_level": "render",
            "expires_at": (datetime.now(timezone.utc) - timedelta(days=1)).isoformat(),
        },
        headers=auth_headers,
    )
    assert resp.status_code == 200
    assert _run(headers, sales_report, {"p_branch": "BR01"}).status_code == 404


def test_missing_unknown_and_oversized_parameters_are_400(auth_headers, sales_report, fake_source):
    assert _run(auth_headers, sales_report, {}).status_code == 400
    assert _run(auth_headers, sales_report, {"p_branch": ""}).status_code == 400
    unknown = _run(auth_headers, sales_report, {"p_branch": "BR01", "extra": "x"})
    assert unknown.status_code == 400 and "extra" in unknown.json()["detail"]
    report_id = _register(auth_headers)
    _configure(auth_headers, report_id, parameters=[{"name": "q"}], data_source={"url": "https://x.test/?q={{ q }}"})
    assert _run(auth_headers, report_id, {"q": "x" * 201}).status_code == 400
    assert fake_source.calls == []


@pytest.mark.parametrize(
    "type_,good,bad",
    [
        ("date", "2026-10-15", "15/10/2026"),
        ("datetime", "2026-10-15T09:30", "2026-10-15 09:30"),
        ("time", "09:30", "9:30am"),
        ("number", "42.5", "forty-two"),
    ],
)
def test_run_rejects_a_value_that_does_not_match_its_declared_type(auth_headers, type_, good, bad):
    report_id = _register(auth_headers, text="When: {{ when }}")
    _configure(auth_headers, report_id, parameters=[{"name": "when", "label": "When", "type": type_}])
    assert _run(auth_headers, report_id, {"when": good}, fmt="docx").status_code == 200
    resp = _run(auth_headers, report_id, {"when": bad}, fmt="docx")
    assert resp.status_code == 400 and "valid" in resp.json()["detail"]


def test_optional_parameter_can_be_omitted_but_is_still_validated_if_given(auth_headers):
    report_id = _register(auth_headers, text="Q: {{ q }}")
    _configure(auth_headers, report_id, parameters=[{"name": "q", "label": "Search", "required": False, "type": "number"}])

    omitted = _run(auth_headers, report_id, {}, fmt="docx")
    assert omitted.status_code == 200
    assert "Q: " in _docx_text(omitted.content)

    empty_string = _run(auth_headers, report_id, {"q": ""}, fmt="docx")
    assert empty_string.status_code == 200

    # Given a value, an optional parameter is still held to its declared type.
    bad = _run(auth_headers, report_id, {"q": "not-a-number"}, fmt="docx")
    assert bad.status_code == 400


def test_required_parameter_cannot_be_omitted(auth_headers):
    report_id = _register(auth_headers, text="Q: {{ q }}")
    _configure(auth_headers, report_id, parameters=[{"name": "q", "label": "Search"}])
    resp = _run(auth_headers, report_id, {}, fmt="docx")
    assert resp.status_code == 400 and "Choose a value" in resp.json()["detail"]


def test_format_not_allowed_for_the_template_is_rejected_before_any_fetch(auth_headers, sales_report, fake_source):
    resp = _run(auth_headers, sales_report, {"p_branch": "BR01"}, fmt="xlsx")
    assert resp.status_code == 400
    assert fake_source.calls == []


# --- run: the data source ------------------------------------------------


def test_validated_parameters_override_same_named_keys_from_the_data_source(auth_headers, sales_report, fake_source):
    fake_source.respond = lambda request: httpx.Response(200, json={"p_branch": "BR99", "branch_name": "X", "total": "1"})
    resp = _run(auth_headers, sales_report, {"p_branch": "BR01"})
    assert "Branch BR01:" in _docx_text(resp.content)


def test_report_without_a_data_source_renders_from_parameters_alone(auth_headers, fake_source):
    report_id = _register(auth_headers, text="Chosen branch: {{ p_branch }}")
    _configure(auth_headers, report_id, data_source=None)
    resp = _run(auth_headers, report_id, {"p_branch": "BR02"})
    assert resp.status_code == 200
    assert "Chosen branch: BR02" in _docx_text(resp.content)
    assert fake_source.calls == []


def test_free_text_values_are_percent_encoded_in_the_url(auth_headers, fake_source):
    report_id = _register(auth_headers, text="{{ q }}")
    _configure(auth_headers, report_id, parameters=[{"name": "q"}], data_source={"url": "https://x.test/search?q={{ q }}"})
    fake_source.respond = lambda request: httpx.Response(200, json={})
    assert _run(auth_headers, report_id, {"q": "a&admin=1 b"}).status_code == 200
    request = fake_source.calls[0]
    assert request.url.query == b"q=a%26admin%3D1%20b"
    assert "admin" not in request.url.params


def test_post_body_template_is_rendered_into_json(auth_headers, fake_source):
    report_id = _register(auth_headers, text="{{ n }}")
    _configure(
        auth_headers,
        report_id,
        parameters=[P_BRANCH],
        data_source={"url": "https://x.test/q", "method": "POST", "body_template": {"filter": {"branch": "{{ p_branch }}"}}},
    )
    fake_source.respond = lambda request: httpx.Response(200, json={"n": "ok"})
    assert _run(auth_headers, report_id, {"p_branch": "BR03"}).status_code == 200
    request = fake_source.calls[0]
    assert request.method == "POST"
    import json

    assert json.loads(request.content) == {"filter": {"branch": "BR03"}}


def test_bearer_auth_reads_the_token_from_the_environment(auth_headers, fake_source, monkeypatch):
    monkeypatch.setenv("ERP_TOKEN", "s3cret-token")
    report_id = _register(auth_headers)
    _configure(auth_headers, report_id, data_source={**DATA_SOURCE, "auth": {"type": "bearer", "token_env": "ERP_TOKEN"}})
    assert _run(auth_headers, report_id, {"p_branch": "BR01"}).status_code == 200
    assert fake_source.calls[0].headers["authorization"] == "Bearer s3cret-token"


def test_missing_credentials_env_var_is_a_502_that_does_not_leak_the_name(auth_headers, fake_source, monkeypatch):
    monkeypatch.delenv("ERP_TOKEN", raising=False)
    report_id = _register(auth_headers)
    _configure(auth_headers, report_id, data_source={**DATA_SOURCE, "auth": {"type": "bearer", "token_env": "ERP_TOKEN"}})
    resp = _run(auth_headers, report_id, {"p_branch": "BR01"})
    assert resp.status_code == 502
    assert "ERP_TOKEN" not in resp.text
    assert fake_source.calls == []


@pytest.mark.parametrize(
    "respond",
    [
        lambda request: httpx.Response(500, text="boom: internal stack trace"),
        lambda request: httpx.Response(302, headers={"location": "http://169.254.169.254/"}),
        lambda request: httpx.Response(200, text="not json"),
        lambda request: httpx.Response(200, json=[1, 2, 3]),
    ],
    ids=["http-500", "redirect-not-followed", "not-json", "not-an-object"],
)
def test_data_source_failures_are_a_502_with_no_upstream_details(auth_headers, sales_report, fake_source, respond):
    fake_source.respond = respond
    resp = _run(auth_headers, sales_report, {"p_branch": "BR01"})
    assert resp.status_code == 502
    for leaked in ("stack trace", "erp.example.test", "169.254"):
        assert leaked not in resp.text
    assert len(fake_source.calls) == 1  # a redirect is not followed


def test_unreachable_data_source_is_a_502(auth_headers, sales_report, fake_source):
    def refuse(request):
        raise httpx.ConnectError("connection refused", request=request)

    fake_source.respond = refuse
    resp = _run(auth_headers, sales_report, {"p_branch": "BR01"})
    assert resp.status_code == 502 and "Couldn't reach" in resp.json()["detail"]


def test_oversized_data_source_response_is_rejected(auth_headers, sales_report, fake_source, monkeypatch):
    monkeypatch.setattr(report_data, "MAX_RESPONSE_BYTES", 64)
    fake_source.respond = lambda request: httpx.Response(200, json={"rows": ["x" * 200]})
    resp = _run(auth_headers, sales_report, {"p_branch": "BR01"})
    assert resp.status_code == 502 and "too large" in resp.json()["detail"]
