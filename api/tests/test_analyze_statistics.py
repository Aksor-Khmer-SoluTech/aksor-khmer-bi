"""Analyze statistics (Phase 5b, app/analyze.py): a value shown as a percentage, a running total or a change; number
fields grouped into ranges (a histogram); a trend down the rows -- line, moving average, forecast; outliers marked by
z-score or IQR -- and a data alert that emails when a group stands out. Checked against numbers worked out by hand,
with Khmer keys and labels, in the pivot, the exports, dashboards and alerts."""
from datetime import datetime, timezone
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import analyze, data_alerts, db, job_delivery
from app.celery_app import celery_app
from app.main import app
from app.rbac import ROOT_ORG_ID
from tests.test_analyze import BB, PP, SR, _dataset, _pivot
from tests.test_job_delivery import FakeSmtp, _mail_server

client = TestClient(app)

# Sales by month with March missing: Jan 100, Feb 120, Apr 160 -- a straight line of +20 a month, if March is
# counted as a month (it is: the axis is real months, not rows).
MONTHLY = {"rows": [
    {"month": "2026-01-15", "store": "A", "sales": 60}, {"month": "2026-01-20", "store": "B", "sales": 40},
    {"month": "2026-02-10", "store": "A", "sales": 70}, {"month": "2026-02-11", "store": "B", "sales": 50},
    {"month": "2026-04-03", "store": "A", "sales": 90}, {"month": "2026-04-09", "store": "B", "sales": 70},
]}
# Seven stores, one far ahead: 50 against about 10.
STORES = {"rows": [{"store": s, "sales": v} for s, v in
                   [("ភ្នំពេញ", 50), ("សៀមរាប", 10), ("បាត់ដំបង", 11), ("កំពត", 9), ("តាកែវ", 10), ("កែប", 12), ("ព្រៃវែង", 10)]]}


@pytest.fixture(autouse=True)
def fresh(tmp_path, monkeypatch):
    analyze.clear_cache()
    monkeypatch.setattr(job_delivery, "OUTPUT_DIR", tmp_path / "job_outputs")
    monkeypatch.setattr(celery_app, "send_task", lambda name, args=None, **kw: type("R", (), {"id": "stub"})())
    yield
    analyze.clear_cache()


@pytest.fixture
def smtp():
    server = FakeSmtp()
    yield server
    server.close()


def _static(headers, name, data, fields, measures):
    resp = client.post("/api/v1/datasets", headers=headers, json={
        "name": name, "title": name.title(), "org_id": ROOT_ORG_ID,
        "data_source": {"type": "static", "static_data": data}, "fields": fields, "measures": measures,
    })
    assert resp.status_code == 200, resp.text


def _monthly(headers):
    _static(headers, "monthly", MONTHLY, [
        {"source": "rows", "name": "month", "type": "date", "label": "Month"},
        {"source": "rows", "name": "store", "type": "text", "label": "Store"},
        {"source": "rows", "name": "sales", "type": "number", "label": "Sales"},
    ], [{"key": "sales", "label": "Sales", "kind": "sum", "source": "rows", "field": "sales", "format": "#,##0"}])


def _stores(headers):
    _static(headers, "stores", STORES, [
        {"source": "rows", "name": "store", "type": "text", "label": "Store", "label_km": "ហាង"},
        {"source": "rows", "name": "sales", "type": "number", "label": "Sales"},
        {"source": "rows", "name": "secret", "type": "text", "label": "Secret", "hidden": True},
    ], [{"key": "sales", "label": "Sales", "label_km": "ការលក់", "kind": "sum", "source": "rows", "field": "sales", "format": "#,##0 ៛"}])


def _own(headers, dataset, layout):
    """A pivot of one of this file's datasets (they have no filters)."""
    resp = client.post("/api/v1/analyze/pivot", headers=headers, json={"dataset": dataset, "values": {}, "layout": layout})
    assert resp.status_code == 200, resp.text
    return resp.json()


def _v(result, row, column=0, value=0):
    cell = result["cells"][row][column]
    return None if cell is None else cell[value]["v"]


def _close(a, b):
    return a is not None and abs(a - b) < 1e-9


# --- show value as ----------------------------------------------------------------------------------------------------


def test_percent_of_total(auth_headers):
    _dataset(auth_headers)
    result = _pivot(auth_headers, {"rows": ["branch"], "values": ["net", "orders"], "show_as": {"net": {"mode": "percent_of_total"}}})
    by = {k[0]: c[0] for k, c in zip(result["row_keys"], result["cells"])}
    assert _close(by[PP][0]["v"], 125000 / 185000) and by[PP][0]["t"] == "67.6%"
    assert by[BB][0]["t"] == "10.8%" and result["total"][0] == {"v": 1.0, "t": "100.0%"}
    assert by[PP][1]["v"] == 11  # the other value is left as it was
    label = result["labels"]["values"][0]
    assert label["label"] == "Net sales (% of total)" and label["label_km"] == "ការលក់សុទ្ធ (% នៃសរុប)" and label["show_as"] == "percent_of_total"


def test_percent_of_row_and_of_column(auth_headers):
    _dataset(auth_headers)
    layout = {"rows": ["branch"], "columns": ["sold_on"], "values": ["net"], "buckets": {"sold_on": "month"}}
    row = _pivot(auth_headers, {**layout, "show_as": {"net": {"mode": "percent_of_row"}}})
    pp = row["row_keys"].index([PP])
    assert _close(_v(row, pp, 0), 65000 / 125000) and row["row_totals"][pp][0]["v"] == 1.0
    assert _close(row["column_totals"][0][0]["v"], 95000 / 185000) and row["total"][0]["v"] == 1.0
    col = _pivot(auth_headers, {**layout, "show_as": {"net": {"mode": "percent_of_column"}}})
    assert _close(_v(col, pp, 0), 65000 / 95000) and col["column_totals"][0][0]["v"] == 1.0
    assert _close(col["row_totals"][pp][0]["v"], 125000 / 185000)
    assert _v(col, col["row_keys"].index([BB]), 0) is None  # nothing sold there that month stays empty


def test_running_total_across_the_months(auth_headers):
    _dataset(auth_headers)
    result = _pivot(auth_headers, {"rows": ["branch"], "columns": ["sold_on"], "values": ["net"], "buckets": {"sold_on": "month"},
                                   "show_as": {"net": {"mode": "running_total", "along": "columns"}}})
    pp, bb = result["row_keys"].index([PP]), result["row_keys"].index([BB])
    assert [_v(result, pp, c) for c in (0, 1)] == [65000, 125000]
    assert [_v(result, bb, c) for c in (0, 1)] == [None, 20000]
    assert [t[0]["v"] for t in result["column_totals"]] == [95000, 185000]
    assert result["row_totals"][pp][0]["v"] is None and result["total"][0]["t"] == "—"  # nothing to run along


def test_a_running_total_starts_again_for_each_outer_value(auth_headers):
    _dataset(auth_headers)
    result = _pivot(auth_headers, {"rows": ["branch", "sold_on"], "values": ["net"], "buckets": {"sold_on": "month"},
                                   "show_as": {"net": {"mode": "running_total"}}})
    got = {tuple(k): c[0][0]["v"] for k, c in zip(result["row_keys"], result["cells"])}
    assert got[(PP, "2026-08")] == 65000 and got[(PP, "2026-09")] == 125000
    assert got[(SR, "2026-08")] == 30000 and got[(SR, "2026-09")] == 40000  # Siem Reap starts from its own first month


def test_change_from_the_previous_month(auth_headers):
    _dataset(auth_headers)
    layout = {"rows": ["sold_on"], "values": ["net"], "buckets": {"sold_on": "month"}}
    diff = _pivot(auth_headers, {**layout, "show_as": {"net": {"mode": "difference"}}})
    assert [_v(diff, 0), _v(diff, 1)] == [None, -5000] and diff["cells"][1][0][0]["t"] == "-5,000 ៛"
    pct = _pivot(auth_headers, {**layout, "show_as": {"net": {"mode": "percent_difference"}}})
    assert _close(_v(pct, 1), -5000 / 95000) and pct["cells"][1][0][0]["t"] == "-5.3%"


@pytest.mark.parametrize("show_as, fragment", [
    ({"orders": {"mode": "percent_of_total"}}, "isn't one of the values shown"),
    ({"net": {"mode": "running_total", "along": "columns"}}, "put one in columns first"),
    ({"net": {"mode": "share"}}, "Input should be"),
])
def test_show_as_is_checked(auth_headers, show_as, fragment):
    _dataset(auth_headers)
    resp = client.post("/api/v1/analyze/pivot", headers=auth_headers, json={
        "dataset": "branch-sales", "values": {"p_branch": "BR01"}, "layout": {"rows": ["branch"], "values": ["net"], "show_as": show_as},
    })
    assert resp.status_code in (400, 422) and fragment in resp.text, resp.text


# --- number ranges ----------------------------------------------------------------------------------------------------


def test_amounts_in_ranges(auth_headers):
    _dataset(auth_headers)
    result = _pivot(auth_headers, {"rows": ["amount"], "values": ["__rows"], "ranges": {"amount": 20000}})
    # 10,000 | 20,000 25,000 30,000 | 40,000 | 60,000: each range from its lower edge up to, not including, the next.
    assert result["row_keys"] == [["0–20,000"], ["20,000–40,000"], ["40,000–60,000"], ["60,000–80,000"]]
    assert [_v(result, r) for r in range(4)] == [1, 3, 1, 1]
    assert result["ranges"] == {"amount": {"width": 20000.0, "width_t": "20,000"}}


def test_an_empty_range_is_still_listed(auth_headers):
    _dataset(auth_headers)
    result = _pivot(auth_headers, {"rows": ["amount"], "values": ["__rows"], "ranges": {"amount": 10000}})
    keys = [k[0] for k in result["row_keys"]]
    assert keys[0] == "10,000–20,000" and keys[-1] == "60,000–70,000" and len(keys) == 6
    assert result["cells"][keys.index("50,000–60,000")][0] is None  # nobody sold that much -- a gap, not left out


def test_ranges_chosen_for_you_and_across_the_columns(auth_headers):
    _dataset(auth_headers)
    result = _pivot(auth_headers, {"rows": ["branch"], "columns": ["amount"], "values": ["__rows"], "ranges": {"amount": None}})
    assert result["ranges"]["amount"]["width"] == 5000  # 10,000 to 60,000: about ten ranges, of a round width
    assert result["column_keys"][0] == ["10,000–15,000"] and len(result["column_keys"]) == 11


@pytest.mark.parametrize("ranges, fragment", [
    ({"branch": 100}, "isn't a number"),
    ({"amount": 1}, "Make them wider"),
    ({"amount": -5}, "above 0"),
])
def test_ranges_are_checked(auth_headers, ranges, fragment):
    _dataset(auth_headers)
    resp = client.post("/api/v1/analyze/pivot", headers=auth_headers, json={
        "dataset": "branch-sales", "values": {"p_branch": "BR01"}, "layout": {"rows": ["branch", "amount"], "ranges": ranges},
    })
    assert resp.status_code == 400 and fragment in resp.text, resp.text


# --- trend ------------------------------------------------------------------------------------------------------------


def test_a_trend_counts_real_months(auth_headers):
    _monthly(auth_headers)
    result = _own(auth_headers, "monthly", {"rows": ["month"], "values": ["sales"], "buckets": {"month": "month"},
                                   "trend": {"moving_average": 2, "forecast": 2}})
    assert [k[0] for k in result["row_keys"]] == ["2026-01", "2026-02", "2026-04"]
    trend = result["trend"]
    line = trend["series"][0]
    # Jan 100, Feb 120, Apr 160: +20 a month exactly -- March missing is a gap of two months, not one.
    assert _close(line["slope"], 20) and _close(line["r2"], 1.0) and line["slope_t"] == "20" and trend["per"] == "month"
    assert trend["forecast_keys"] == ["2026-05", "2026-06"] and [f["t"] for f in line["forecast"]] == ["180", "200"]
    assert line["moving_average"] == [None, 110, 140]
    assert [round(v, 9) for v in line["fitted"]] == [100, 120, 160]


def test_a_trend_for_each_column(auth_headers):
    _monthly(auth_headers)
    result = _own(auth_headers, "monthly", {"rows": ["month"], "columns": ["store"], "values": ["sales"], "buckets": {"month": "month"},
                                   "trend": {"forecast": 1}})
    a, b = result["trend"]["series"]
    assert _close(a["slope"], 10) and _close(b["slope"], 10)  # A: 60 70 _ 90; B: 40 50 _ 70
    assert [f["v"] for f in a["forecast"]] == pytest.approx([100]) and [f["v"] for f in b["forecast"]] == pytest.approx([80])


def test_a_trend_over_rows_that_arent_dates(auth_headers):
    _dataset(auth_headers)
    result = _pivot(auth_headers, {"rows": ["branch"], "values": ["net"], "trend": {"forecast": 1}})
    assert result["trend"]["per"] is None and result["trend"]["forecast_keys"] == ["+1"]
    ranged = _pivot(auth_headers, {"rows": ["amount"], "values": ["__rows"], "ranges": {"amount": 20000}, "trend": {"forecast": 1}})
    assert ranged["trend"]["per"] == "range" and ranged["trend"]["forecast_keys"] == ["80,000–100,000"]  # the next range


@pytest.mark.parametrize("layout, fragment", [
    ({"rows": ["branch", "product"], "trend": None}, None),
    ({"rows": ["branch", "product"], "trend": {"forecast": 1}}, "exactly one"),
    ({"rows": ["branch"], "trend": {"line": False, "forecast": 2}}, "turn the line on"),
    ({"rows": ["branch"], "trend": {"value": "orders"}}, "isn't one of the values shown"),
])
def test_a_trend_is_checked(auth_headers, layout, fragment):
    _dataset(auth_headers)
    resp = client.post("/api/v1/analyze/pivot", headers=auth_headers, json={
        "dataset": "branch-sales", "values": {"p_branch": "BR01"}, "layout": {"values": ["net"], **layout},
    })
    if fragment is None:
        assert resp.status_code == 200 and "trend" not in resp.json()
    else:
        assert resp.status_code == 400 and fragment in resp.text, resp.text


# --- outliers ---------------------------------------------------------------------------------------------------------


@pytest.mark.parametrize("method", ["zscore", "iqr"])
def test_the_store_far_ahead_is_an_outlier(auth_headers, method):
    _stores(auth_headers)
    result = _own(auth_headers, "stores", {"rows": ["store"], "values": ["sales"], "outliers": {"method": method}})
    marked = {k[0]: c[0][0].get("o") for k, c in zip(result["row_keys"], result["cells"])}
    assert marked["ភ្នំពេញ"] == "high" and sum(1 for o in marked.values() if o) == 1
    found = result["outliers"]["values"][0]
    assert found["count"] == 1 and result["outliers"]["threshold"] == (2.0 if method == "zscore" else 1.5)
    assert "o" not in result["total"][0]  # the totals aren't compared with the cells
    if method == "iqr":  # quartiles 10 and 11.5: fences at 7.75 and 13.75
        assert found["low_t"] == "8 ៛" and found["high_t"] == "14 ៛"


def test_a_stricter_threshold_and_too_few_values(auth_headers):
    _stores(auth_headers)
    strict = _own(auth_headers, "stores", {"rows": ["store"], "values": ["sales"], "outliers": {"threshold": 3}})
    assert strict["outliers"]["values"][0]["count"] == 0  # 50 is 2.45 standard deviations out: not beyond 3
    _dataset(auth_headers)
    few = _pivot(auth_headers, {"rows": ["product"], "values": ["net"], "outliers": {"method": "iqr"}})
    assert few["outliers"]["values"][0]["low"] is None  # two products: nothing to compare


# --- exports and dashboards -------------------------------------------------------------------------------------------


def _export(headers, fmt, layout, dataset="stores", values=None, chart="bar"):
    resp = client.post("/api/v1/analyze/export", headers=headers, json={
        "dataset": dataset, "format": fmt, "chart": chart, "values": values or {}, "layout": layout,
    })
    assert resp.status_code == 200, resp.text
    return resp.content


def test_excel_export_marks_outliers_and_keeps_percentages(auth_headers):
    _stores(auth_headers)
    content = _export(auth_headers, "xlsx", {"rows": ["store"], "values": ["sales", "__rows"], "outliers": {},
                                             "show_as": {"__rows": {"mode": "percent_of_total"}}})
    ws = load_workbook(BytesIO(content)).active
    notes = [r[0].value for r in ws.iter_rows() if isinstance(r[0].value, str) and r[0].value.startswith("Outliers")]
    assert len(notes) == 2 and "ការលក់: 1 (below" in notes[0]
    row = next(r for r in ws.iter_rows() if r[0].value == "ភ្នំពេញ")
    assert row[1].value == 50 and row[1].fill.fgColor.rgb.endswith("FDE2C8")
    assert row[2].number_format == "0.0%" and abs(row[2].value - 1 / 7) < 1e-9  # a share, kept as a number


def test_word_and_pdf_exports_carry_the_trend(auth_headers):
    _monthly(auth_headers)
    layout = {"rows": ["month"], "values": ["sales"], "buckets": {"month": "month"}, "trend": {"moving_average": 2, "forecast": 2}}
    doc = Document(BytesIO(_export(auth_headers, "docx", layout, dataset="monthly", chart="line")))
    text = "\n".join(p.text for p in doc.paragraphs).replace("​", "")
    assert "និន្នាការ" in text and "20 per month, R² 1.00; forecast 2026-05 180, 2026-06 200" in text
    assert len(doc.inline_shapes) == 1
    assert _export(auth_headers, "pdf", layout, dataset="monthly", chart="bar").startswith(b"%PDF")
    diff = {"rows": ["month"], "values": ["sales"], "buckets": {"month": "month"}, "show_as": {"sales": {"mode": "difference"}}}
    assert _export(auth_headers, "pdf", diff, dataset="monthly").startswith(b"%PDF")  # a chart that goes below zero


def test_a_dashboard_tile_keeps_its_statistics(auth_headers, make_local_user):
    _stores(auth_headers)
    _, owner = make_local_user("owner", "pw-owner-1", ("ROLE_REPORT_ADMIN",))
    tile = {"id": "t1", "title": "Stores", "dataset": "stores", "values": {}, "view": "table", "width": 2,
            "layout": {"rows": ["store"], "values": ["sales"], "outliers": {"method": "iqr"}, "show_as": {"sales": {"mode": "percent_of_total"}}}}
    created = client.post("/api/v1/dashboards", headers=owner, json={"title": "Stores"}).json()
    resp = client.put(f"/api/v1/dashboards/{created['id']}", headers=owner, json={"title": "Stores", "tiles": [tile], "filters": {}})
    assert resp.status_code == 200, resp.text
    assert resp.json()["tiles"][0]["layout"]["outliers"] == {"method": "iqr", "threshold": 1.5}  # kept, cleaned
    result = client.post(f"/api/v1/dashboards/{created['id']}/tiles/t1/run", headers=owner, json={}).json()
    assert result["outliers"]["values"][0]["count"] == 1 and result["total"][0]["t"] == "100.0%"


# --- an alert when a group stands out ---------------------------------------------------------------------------------


def _alert(headers, expect=200, **config):
    body = {
        "name": "Unusual stores", "job_type": "data_alert", "trigger_type": "cron", "cron_expression": "0 8 * * *",
        "config": {"dataset": "stores", "values": {}, "measure": "sales", "condition": "outlier", "group_by": "store",
                   "email": {"connection": "company-mail", "to": ["ops@example.com"]}, **config},
    }
    resp = client.post("/api/v1/jobs", params={"org_id": ROOT_ORG_ID}, headers=headers, json=body)
    assert resp.status_code == expect, resp.text
    return resp.json()


def _check(headers, job) -> str:
    run_id = client.post(f"/api/v1/jobs/{job['id']}/run", headers=headers).json()["id"]
    summary = data_alerts.run(job["config"], job["id"], run_id)
    with db.SessionLocal() as session:
        run = session.get(db.JobRun, run_id)
        run.status, run.result_summary, run.finished_at = "succeeded", summary, datetime.now(timezone.utc).isoformat()
        session.commit()
    return summary


def test_an_alert_when_a_store_stands_out(auth_headers, smtp):
    _stores(auth_headers)
    _mail_server(auth_headers, smtp.port)
    job = _alert(auth_headers)
    assert job["config"]["method"] == "zscore" and job["config"]["threshold"] == 2.0
    summary = _check(auth_headers, job)
    assert summary == "Triggered: 1 Store stands out on Sales: ភ្នំពេញ 50 ៛ (high) -- emailed 1 person"
    body = smtp.messages[0][1].get_content()
    assert "ភ្នំពេញ: 50 ៛ (unusually high)" in body and "from the other 6" in body
    assert "already emailed" in _check(auth_headers, job) and len(smtp.messages) == 1


def test_an_outlier_alert_that_finds_none(auth_headers, smtp):
    _stores(auth_headers)
    _mail_server(auth_headers, smtp.port)
    job = _alert(auth_headers, threshold=3)
    assert _check(auth_headers, job) == "Not triggered: no Store stands out on Sales (7 compared, z-score beyond 3)"
    assert smtp.messages == []


@pytest.mark.parametrize("config, fragment", [
    ({"group_by": None}, "Choose the field to compare across"),
    ({"group_by": "secret"}, "isn't a field you can analyze"),
    ({"method": "mad"}, "z-score or by interquartile range"),
    ({"threshold": 0}, "above 0"),
])
def test_an_outlier_alert_is_checked(auth_headers, smtp, config, fragment):
    _stores(auth_headers)
    _mail_server(auth_headers, smtp.port)
    resp = _alert(auth_headers, expect=400, **config)
    assert fragment in resp["detail"]
