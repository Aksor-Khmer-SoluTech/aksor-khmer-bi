"""Analyze (app/analyze.py, routers/analyze.py): pivot a dataset by its fields and measures, export it, keep the view.
Checked: the numbers (a ratio worked out per cell, not summed), Khmer keys and labels, dates grouped by month, who may
analyze what (analyze:use; a restricted dataset only for its grants, within their values), the layout rules, the
cache, the three export formats, and that saved views are their owner's alone."""
import shutil
import subprocess
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import analyze, db, report_data
from app.main import app
from app.rbac import ROOT_ORG_ID

client = TestClient(app)

PP, SR, BB = "ភ្នំពេញ", "សៀមរាប", "បាត់ដំបង"
SALES = {
    "title": "របាយការណ៍លក់",
    "rows": [
        {"branch": PP, "product": "Diesel", "sold_on": "2026-08-03", "amount": 40000, "orders": 4, "cost": 30000},
        {"branch": PP, "product": "Petrol", "sold_on": "2026-08-19", "amount": 25000, "orders": 2, "cost": 20000},
        {"branch": PP, "product": "Diesel", "sold_on": "2026-09-02", "amount": 60000, "orders": 5, "cost": 45000},
        {"branch": SR, "product": "Diesel", "sold_on": "2026-08-07", "amount": 30000, "orders": 3, "cost": 24000},
        {"branch": SR, "product": "Petrol", "sold_on": "2026-09-21", "amount": 10000, "orders": 1, "cost": 9000},
        {"branch": BB, "product": "Diesel", "sold_on": "2026-09-30", "amount": 20000, "orders": 2, "cost": 15000},
    ],
}
FIELDS = [
    {"source": "rows", "name": "branch", "type": "text", "label": "Branch", "label_km": "សាខា"},
    {"source": "rows", "name": "product", "type": "text", "label": "Product"},
    {"source": "rows", "name": "sold_on", "type": "date", "label": "Sold on"},
    {"source": "rows", "name": "amount", "type": "money", "label": "Amount"},
    {"source": "rows", "name": "orders", "type": "number", "label": "Orders"},
    {"source": "rows", "name": "cost", "type": "money", "label": "Cost", "hidden": True},
]
MEASURES = [
    {"key": "net", "label": "Net sales", "label_km": "ការលក់សុទ្ធ", "kind": "sum", "source": "rows", "field": "amount", "format": "#,##0 ៛"},
    {"key": "orders", "label": "Orders", "kind": "sum", "source": "rows", "field": "orders"},
    {"key": "cost", "label": "Cost", "kind": "sum", "source": "rows", "field": "cost"},
    {"key": "margin", "label": "Cost share", "kind": "ratio", "numerator": "cost", "denominator": "net", "format": "0.0%"},
]
BRANCHES = [{"value": "BR01", "label": PP}, {"value": "BR02", "label": SR}, {"value": "BR03", "label": BB}]


@pytest.fixture(autouse=True)
def fresh_cache():
    analyze.clear_cache()
    yield
    analyze.clear_cache()


def _dataset(headers, name="branch-sales", **extra):
    resp = client.post("/api/v1/datasets", headers=headers, json={
        "name": name, "title": "Branch sales", "org_id": ROOT_ORG_ID,
        "parameters": [{"name": "p_branch", "label": "Branch", "options": BRANCHES}],
        "data_source": {"type": "static", "static_data": SALES}, "fields": FIELDS, "measures": MEASURES, **extra,
    })
    assert resp.status_code == 200, resp.text
    return resp.json()


def _pivot(headers, layout, values=None, expect=200, dataset="branch-sales"):
    resp = client.post("/api/v1/analyze/pivot", headers=headers, json={"dataset": dataset, "values": values or {"p_branch": "BR01"}, "layout": layout})
    assert resp.status_code == expect, resp.text
    return resp.json()


def _analyst(make_local_user, username="analyst"):
    """A person who may analyze (analyze:use) but manages no dataset -- so a restricted dataset limits them."""
    user_id, headers = make_local_user(username, f"pw-{username}-1", ("ROLE_REPORT_VIEWER",))
    return user_id, headers


def _grant(admin_headers, user_id, code="analyze:use"):
    resp = client.post("/api/v1/grants/permissions", headers=admin_headers, json={"user_id": user_id, "permission_code": code})
    assert resp.status_code == 200, resp.text


# --- the numbers ------------------------------------------------------------------------------------------------------


def test_a_pivot_by_branch_and_month(auth_headers):
    _dataset(auth_headers)
    result = _pivot(auth_headers, {"rows": ["branch"], "columns": ["sold_on"], "values": ["net", "margin"], "buckets": {"sold_on": "month"}})
    assert result["row_keys"] == [[BB], [PP], [SR]]  # Khmer keys, whole (no segmentation marks)
    assert result["column_keys"] == [["2026-08"], ["2026-09"]]
    cell = lambda branch, month: result["cells"][result["row_keys"].index([branch])][result["column_keys"].index([month])]  # noqa: E731
    assert cell(PP, "2026-08")[0] == {"v": 65000, "t": "65,000 ៛"}
    assert cell(BB, "2026-08") is None  # nothing sold there that month
    # A ratio is worked out over each cell's rows -- cost 50,000 / net 65,000 -- never summed from other ratios.
    assert cell(PP, "2026-08")[1]["t"] == "76.9%"
    assert result["row_totals"][result["row_keys"].index([PP])][0]["v"] == 125000
    assert [t[0]["v"] for t in result["column_totals"]] == [95000, 90000]
    assert result["total"][0] == {"v": 185000, "t": "185,000 ៛"} and result["total"][1]["t"] == "77.3%"
    assert result["labels"]["rows"][0] == {"name": "branch", "label": "Branch", "label_km": "សាខា", "type": "text"}
    assert result["labels"]["values"][0]["label_km"] == "ការលក់សុទ្ធ"


def test_a_pivot_without_columns_and_the_row_count(auth_headers):
    _dataset(auth_headers)
    result = _pivot(auth_headers, {"rows": ["product"], "values": ["__rows", "orders"]})
    assert result["row_keys"] == [["Diesel"], ["Petrol"]]
    assert [c[0] for c in result["cells"]] == [[{"v": 4, "t": "4"}, {"v": 14.0, "t": "14"}], [{"v": 2, "t": "2"}, {"v": 3.0, "t": "3"}]]
    assert result["total"][0]["v"] == 6 and result["row_count"] == 6


@pytest.mark.parametrize("layout, fragment", [
    ({"rows": ["cost"]}, "isn't a field you can analyze"),  # hidden from authors, hidden here
    ({"rows": ["nope"]}, "isn't a field"),
    ({"rows": ["branch"], "columns": ["branch"]}, "rows or in columns, not both"),
    ({"rows": ["branch"], "values": ["profit"]}, "isn't a measure"),
    ({"rows": ["branch"], "buckets": {"branch": "week"}}, "Input should be"),
    ({"rows": ["branch", "product", "sold_on", "amount"]}, "at most 3"),
])
def test_a_layout_is_checked(auth_headers, layout, fragment):
    _dataset(auth_headers)
    resp = client.post("/api/v1/analyze/pivot", headers=auth_headers, json={"dataset": "branch-sales", "values": {"p_branch": "BR01"}, "layout": layout})
    assert resp.status_code in (400, 422) and fragment in resp.text, resp.text


def test_too_many_cells_is_a_clear_error(auth_headers, monkeypatch):
    _dataset(auth_headers)
    monkeypatch.setattr(analyze, "MAX_ROW_KEYS", 2)
    resp = client.post("/api/v1/analyze/pivot", headers=auth_headers, json={"dataset": "branch-sales", "values": {"p_branch": "BR01"}, "layout": {"rows": ["branch"]}})
    assert resp.status_code == 400 and "Group a date by month" in resp.json()["detail"]


def test_the_data_is_fetched_once_per_filter_set(auth_headers, monkeypatch):
    dataset = _dataset(auth_headers)
    calls = []
    real = report_data.fetch_report_data
    monkeypatch.setattr(report_data, "fetch_report_data", lambda *a, **k: calls.append(a[1]) or real(*a, **k))
    _pivot(auth_headers, {"rows": ["branch"]})
    _pivot(auth_headers, {"rows": ["product"], "columns": ["branch"]})
    assert len(calls) == 1  # re-arranging doesn't query again
    _pivot(auth_headers, {"rows": ["branch"]}, values={"p_branch": "BR02"})
    assert len(calls) == 2  # other filter values do
    body = {k: dataset[k] for k in ("title", "description", "parameters", "fields", "measures", "data_source")}
    client.put(f"/api/v1/datasets/{dataset['id']}", headers=auth_headers, json=body)
    _pivot(auth_headers, {"rows": ["branch"]})
    assert len(calls) == 3  # a changed dataset is fetched afresh


# --- who may analyze what --------------------------------------------------------------------------------------------


def test_analyze_needs_its_permission(auth_headers, make_local_user):
    _dataset(auth_headers)
    user_id, viewer = _analyst(make_local_user)
    assert client.get("/api/v1/analyze/datasets", headers=viewer).status_code == 403
    _grant(auth_headers, user_id)
    assert [d["name"] for d in client.get("/api/v1/analyze/datasets", headers=viewer).json()] == ["branch-sales"]
    assert _pivot(viewer, {"rows": ["branch"]})["total"][0]["v"] == 6
    _, report_admin = make_local_user("author", "pw-author-1", ("ROLE_REPORT_ADMIN",))
    assert client.get("/api/v1/analyze/datasets", headers=report_admin).status_code == 200  # given it by default


def test_a_restricted_dataset_only_gives_its_grants_values(auth_headers, make_local_user):
    dataset = _dataset(auth_headers)
    user_id, analyst = _analyst(make_local_user)
    _grant(auth_headers, user_id)
    client.put(f"/api/v1/datasets/{dataset['id']}/access/mode", headers=auth_headers, json={"mode": "restricted"})
    # No grant: the dataset isn't theirs to see at all.
    assert client.get("/api/v1/analyze/datasets", headers=analyst).json() == []
    assert client.get("/api/v1/analyze/datasets/branch-sales", headers=analyst).status_code == 404
    client.post(f"/api/v1/datasets/{dataset['id']}/access", headers=auth_headers, json={
        "subject_type": "user", "subject_id": user_id, "level": "run", "parameter_limits": {"p_branch": ["BR02"]},
    })
    form = client.get("/api/v1/analyze/datasets/branch-sales", headers=analyst).json()
    assert [o["value"] for o in form["parameters"][0]["options"]] == ["BR02"]  # only their branch is offered
    assert [f["name"] for f in form["fields"]] == ["branch", "product", "sold_on", "amount", "orders"]  # hidden left out
    assert _pivot(analyst, {"rows": ["branch"]}, values={"p_branch": "BR01"}, expect=403)
    assert _pivot(analyst, {"rows": ["branch"]}, values={"p_branch": "BR02"})["row_count"] == 6


# --- exports ---------------------------------------------------------------------------------------------------------


def _export(headers, fmt, chart="bar"):
    resp = client.post("/api/v1/analyze/export", headers=headers, json={
        "dataset": "branch-sales", "format": fmt, "chart": chart, "title": "លក់តាមសាខា",
        "values": {"p_branch": "BR01"},
        "layout": {"rows": ["branch"], "columns": ["sold_on"], "values": ["net"], "buckets": {"sold_on": "month"}},
    })
    assert resp.status_code == 200, resp.text
    return resp


def test_export_as_excel_keeps_numbers(auth_headers):
    _dataset(auth_headers)
    resp = _export(auth_headers, "xlsx")
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
    ws = load_workbook(BytesIO(resp.content)).active
    assert ws["A1"].value == "លក់តាមសាខា"
    assert ("Branch", f"{PP} (BR01)") in [(r[0].value, r[1].value) for r in ws.iter_rows()]  # the filter, its label too
    rows = {r[0].value: r for r in ws.iter_rows() if r[0].value in (PP, SR, BB, "សរុប · Total")}
    assert rows[PP][1].value == 65000 and rows[PP][1].number_format == '#,##0 "៛"'  # a number, shown in riel
    assert rows["សរុប · Total"][-1].value == 185000


def test_export_as_word_has_the_khmer_text_and_the_chart(auth_headers):
    _dataset(auth_headers)
    doc = Document(BytesIO(_export(auth_headers, "docx").content))
    text = "\n".join(p.text for p in doc.paragraphs).replace("​", "")
    assert "លក់តាមសាខា" in text and f"{PP} (BR01)" in text
    cells = [c.text.replace("​", "") for c in doc.tables[0]._cells]
    assert PP in cells and "65,000 ៛" in cells and "185,000 ៛" in cells
    assert len(doc.inline_shapes) == 1  # the bar chart, as a picture


def test_export_as_pdf_is_audited(auth_headers):
    _dataset(auth_headers)
    resp = _export(auth_headers, "pdf", chart="line")
    # The chart is a picture, not its SVG source printed as text (segmentation once broke its tags apart).
    if shutil.which("pdftotext"):
        text = subprocess.run(["pdftotext", "-", "-"], input=resp.content, capture_output=True, check=True).stdout.decode()
        assert "Branch" in text and "xmlns" not in text and "<rect" not in text
    assert resp.content.startswith(b"%PDF") and resp.headers["content-disposition"].startswith('attachment; filename="branch-sales-analysis.pdf"')
    with db.SessionLocal() as session:
        event = session.query(db.AuditEvent).filter(db.AuditEvent.action == "analyze.export").one()
    assert event.summary == "Exported an analysis of branch-sales as PDF (6 rows)" and event.details["values"] == {"p_branch": "BR01"}


# --- saved views ------------------------------------------------------------------------------------------------------


def test_saved_views_are_their_owners(auth_headers, make_local_user):
    _dataset(auth_headers)
    owner_id, owner = _analyst(make_local_user, "owner")
    other_id, other = _analyst(make_local_user, "other")
    _grant(auth_headers, owner_id)
    _grant(auth_headers, other_id)
    config = {"values": {"p_branch": "BR01"}, "layout": {"rows": ["branch"], "values": ["net"]}, "view": "bar"}
    created = client.post("/api/v1/analyze/views", headers=owner, json={"title": "Sales by branch", "dataset": "branch-sales", "config": config})
    assert created.status_code == 200, created.text
    view = created.json()
    assert [v["title"] for v in client.get("/api/v1/analyze/views", headers=owner).json()] == ["Sales by branch"]
    assert client.get("/api/v1/analyze/views", headers=other).json() == []
    assert client.get(f"/api/v1/analyze/views/{view['id']}", headers=other).status_code == 404
    assert client.delete(f"/api/v1/analyze/views/{view['id']}", headers=other).status_code == 404
    renamed = client.put(f"/api/v1/analyze/views/{view['id']}", headers=owner, json={"title": "By branch", "dataset": "branch-sales", "config": config})
    assert renamed.json()["title"] == "By branch" and renamed.json()["config"]["view"] == "bar"
    bad = {**config, "layout": {"rows": ["cost"]}}
    assert client.post("/api/v1/analyze/views", headers=owner, json={"title": "x", "dataset": "branch-sales", "config": bad}).status_code == 400
    assert client.delete(f"/api/v1/analyze/views/{view['id']}", headers=owner).status_code == 204
    # The break-glass sign-in has no account to keep a view under.
    resp = client.post("/api/v1/analyze/views", headers=auth_headers, json={"title": "x", "dataset": "branch-sales", "config": config})
    assert resp.status_code == 400
