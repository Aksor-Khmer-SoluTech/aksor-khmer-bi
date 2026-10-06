"""Report codes: a report can be addressed by a human-chosen code as well as its
random id, on every /api/v1/reports/{ref}/... route (app/report_ref.py,
specs/report_codes_design.md).
"""
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select

from app import db, job_executors, report_ref, report_store
from app.main import app

client = TestClient(app)


def _docx(text="Hello {{ name }}") -> bytes:
    buf = BytesIO()
    doc = Document()
    doc.add_paragraph(text)
    doc.save(buf)
    return buf.getvalue()


def _create(headers, name="Revenue Comparison", code=None):
    data = {"name": name}
    if code is not None:
        data["code"] = code
    return client.post(
        "/api/v1/reports", files={"file": ("t.docx", _docx(), "application/octet-stream")}, data=data, headers=headers
    )


def _make(headers, code="revenue-comparison", name="Revenue Comparison") -> dict:
    resp = _create(headers, name=name, code=code)
    assert resp.status_code == 200, resp.text
    return resp.json()


def _patch(headers, ref, **body):
    return client.patch(f"/api/v1/reports/{ref}", json=body, headers=headers)


def _events(action):
    with db.SessionLocal() as session:
        return [e for e in session.execute(select(db.AuditEvent).order_by(db.AuditEvent.created_at)).scalars() if e.action == action]


# --- the basics: one report, two ways in -----------------------------------------


def test_create_with_a_code_normalizes_it_and_both_addresses_reach_the_same_report(auth_headers):
    created = _make(auth_headers, code="  Revenue-Comparison ")
    assert created["code"] == "revenue-comparison"  # trimmed, lowercased

    by_id = client.get(f"/api/v1/reports/{created['report_id']}").json()
    by_code = client.get("/api/v1/reports/revenue-comparison").json()
    assert by_code == by_id
    assert by_code["report_id"] == created["report_id"] and by_code["code"] == "revenue-comparison"
    # URLs are case-insensitive to a human; the code is stored lowercase.
    assert client.get("/api/v1/reports/REVENUE-COMPARISON").json()["report_id"] == created["report_id"]


def test_a_report_without_a_code_is_unchanged(auth_headers):
    created = _make(auth_headers, code=None)
    assert created["code"] is None
    assert client.get(f"/api/v1/reports/{created['report_id']}").status_code == 200


def test_the_public_render_routes_work_by_code_and_keep_their_query_string(auth_headers):
    created = _make(auth_headers)
    by_id = client.post(f"/api/v1/reports/{created['report_id']}/render?format=docx", json={"name": "Bopha"})
    by_code = client.post("/api/v1/reports/revenue-comparison/render?format=docx", json={"name": "Bopha"})
    assert by_id.status_code == by_code.status_code == 200
    assert Document(BytesIO(by_code.content)).paragraphs[0].text == "Hello Bopha"

    batch = client.post("/api/v1/reports/revenue-comparison/render/batch?format=docx", json=[{"name": "A"}, {"name": "B"}])
    assert batch.status_code == 200 and batch.headers["content-type"] == "application/zip"
    assert client.get("/api/v1/reports/revenue-comparison/schema").json()["fields"] == ["name"]
    assert client.get("/api/v1/reports/nope-nope/schema").status_code == 404


def test_the_protected_routes_work_by_code_and_still_need_credentials(auth_headers):
    created = _make(auth_headers)
    rid = created["report_id"]

    assert client.patch("/api/v1/reports/revenue-comparison", json={"description": "x"}).status_code == 401
    assert _patch(auth_headers, "revenue-comparison", description="via the code").status_code == 200
    assert client.get(f"/api/v1/reports/{rid}").json()["description"] == "via the code"

    for path in ("data-config", "protected-terms-config", "changelog", "file"):
        resp = client.get(f"/api/v1/reports/revenue-comparison/{path}", headers=auth_headers)
        assert resp.status_code == 200, f"{path}: {resp.status_code}"
    assert client.get("/api/v1/reports/revenue-comparison/file").status_code == 401

    # ...and the audit trail records the canonical id, never the alias.
    assert [e.entity_id for e in _events("report.update")] == [rid]
    assert [e.entity_id for e in _events("report.template_download")] == [rid]

    assert client.delete("/api/v1/reports/revenue-comparison", headers=auth_headers).status_code == 204
    assert client.get(f"/api/v1/reports/{rid}").status_code == 404


def test_an_unknown_code_is_the_same_404_as_an_unknown_id(auth_headers):
    _make(auth_headers)
    by_code = client.get("/api/v1/reports/no-such-report")
    by_id = client.get("/api/v1/reports/aaaaaaaaaaaa")
    assert by_code.status_code == by_id.status_code == 404
    assert by_code.json() == by_id.json()


# --- the rules ---------------------------------------------------------------------


@pytest.mark.parametrize(
    "bad",
    [
        "ab",  # too short
        "a" * 65,  # too long
        "has space",
        "under_score",
        "dot.ted",
        "-leading",
        "trailing-",
        "dou--ble",
        "ünïcode",
        "1f094af3df23",  # looks like an id
        "accessible", "batch-limits", "parse-template", "analytics",  # static route segments
    ],
)
def test_codes_that_break_the_rules_are_rejected_at_create_and_update(auth_headers, bad):
    resp = _create(auth_headers, code=bad)
    assert resp.status_code == 400, f"{bad!r} -> {resp.status_code}"
    assert report_store.list_reports() == []  # nothing half-created

    ok = _make(auth_headers, code=None)
    assert _patch(auth_headers, ok["report_id"], code=bad).status_code == 400
    assert client.get(f"/api/v1/reports/{ok['report_id']}").json()["code"] is None


@pytest.mark.parametrize("good", ["abc", "a-b-c", "revenue-comparison-2026", "x1y", "a" * 64])
def test_codes_that_follow_the_rules_are_accepted(auth_headers, good):
    assert _make(auth_headers, code=good)["code"] == good


def test_a_reserved_code_is_rejected_and_the_static_routes_still_work(auth_headers, make_local_user):
    _make(auth_headers, code="accessible-report")  # near-miss, allowed
    _, viewer = make_local_user("vera", "pw-vera-12345", ("ROLE_REPORT_VIEWER",))
    assert client.get("/api/v1/reports/accessible", headers=viewer).status_code == 200
    assert client.get("/api/v1/reports/batch-limits").status_code == 200
    assert client.get("/api/v1/reports/analytics/summary", headers=auth_headers).status_code == 200


# --- uniqueness, and who owns a code ----------------------------------------------------


def test_a_code_can_only_belong_to_one_report_case_insensitively(auth_headers):
    _make(auth_headers, code="revenue-comparison")
    assert _create(auth_headers, name="Another", code="revenue-comparison").status_code == 409
    assert _create(auth_headers, name="Another", code="Revenue-Comparison").status_code == 409
    other = _make(auth_headers, code=None, name="Other")
    assert _patch(auth_headers, other["report_id"], code="revenue-comparison").status_code == 409

    # A refused claim leaves nothing behind: no row, no orphaned template file.
    assert len(report_store.list_reports()) == 2
    assert len([p for p in report_store.STORE_DIR.iterdir() if p.is_dir()]) == 2

    # Re-saving a report's own code is fine.
    first = client.get("/api/v1/reports/revenue-comparison").json()
    assert _patch(auth_headers, first["report_id"], code="revenue-comparison").status_code == 200


def test_codes_are_unique_across_organizations(auth_headers, make_local_user):
    client.post("/api/v1/organizations", json={"id": "acme", "name": "Acme"}, headers=auth_headers)
    _, acme = make_local_user("acme-admin", "pw-acme-12345", ("ROLE_ORG_ADMIN",), org_id="acme")
    _make(auth_headers, code="shared-code")
    assert _create(acme, name="Acme's", code="shared-code").status_code == 409


def test_a_freed_code_cannot_be_claimed_by_another_organization(auth_headers, make_local_user):
    """The takeover: tenant A renames or deletes a report, tenant B claims the freed
    code, and A's integrations start rendering B's template. The code stays A's."""
    client.post("/api/v1/organizations", json={"id": "acme", "name": "Acme"}, headers=auth_headers)
    _, acme = make_local_user("acme-admin", "pw-acme-12345", ("ROLE_ORG_ADMIN",), org_id="acme")

    mine = _make(auth_headers, code="partner-notice")
    # ...freed by renaming it,
    assert _patch(auth_headers, mine["report_id"], code="partner-notice-v2").status_code == 200
    assert _create(acme, name="B's", code="partner-notice").status_code == 409
    # ...and by clearing it,
    assert _patch(auth_headers, mine["report_id"], code=None).status_code == 200
    assert _create(acme, name="B's", code="partner-notice").status_code == 409
    # ...and by deleting the report outright.
    assert client.delete(f"/api/v1/reports/{mine['report_id']}", headers=auth_headers).status_code == 204
    assert _create(acme, name="B's", code="partner-notice").status_code == 409
    assert _create(acme, name="B's", code="partner-notice-v2").status_code == 409

    # The organization that first used it can have it back.
    assert _create(auth_headers, name="Reborn", code="partner-notice").status_code == 200
    # An unrelated code is free for anyone.
    assert _create(acme, name="B's", code="acme-only").status_code == 200


def test_clearing_and_changing_a_code(auth_headers):
    created = _make(auth_headers)
    rid = created["report_id"]

    # Not mentioning the code leaves it alone.
    assert _patch(auth_headers, rid, description="unrelated").json()["code"] == "revenue-comparison"

    assert _patch(auth_headers, rid, code="revenue-report").json()["code"] == "revenue-report"
    assert client.get("/api/v1/reports/revenue-comparison").status_code == 404  # the old one stops working
    assert client.get("/api/v1/reports/revenue-report").json()["report_id"] == rid

    assert _patch(auth_headers, rid, code="").json()["code"] is None  # "" clears, like null
    assert client.get("/api/v1/reports/revenue-report").status_code == 404
    assert _patch(auth_headers, rid, code="again-here").status_code == 200
    assert _patch(auth_headers, rid, code=None).json()["code"] is None
    assert client.get(f"/api/v1/reports/{rid}").status_code == 200  # the id never stops working


def test_setting_changing_and_clearing_a_code_is_audited_with_before_and_after(auth_headers):
    rid = _make(auth_headers, code=None)["report_id"]
    _patch(auth_headers, rid, code="first-code")
    _patch(auth_headers, rid, code="second-code")
    _patch(auth_headers, rid, code=None)
    events = _events("report.update")
    assert [e.summary for e in events] == [
        'Set code "first-code"',
        'Changed code "first-code" to "second-code"',
        'Removed code "second-code"',
    ]
    assert events[1].changes == [{"field": "code", "before": "first-code", "after": "second-code"}]
    assert _events("report.create")[0].details["code"] is None


# --- access control still applies through a code ------------------------------------------


def test_a_report_level_grant_applies_through_the_code_and_denials_log_the_id(auth_headers, make_local_user):
    rid = _make(auth_headers)["report_id"]
    user_id, manager = make_local_user("mo", "pw-mo-1234567", ("ROLE_USER",))
    _, stranger = make_local_user("sam", "pw-sam-1234567", ("ROLE_USER",))

    assert _patch(manager, "revenue-comparison", description="x").status_code == 403
    granted = client.post(
        "/api/v1/grants/reports",
        json={"subject_type": "user", "subject_id": user_id, "report_id": rid, "permission_level": "manage"},
        headers=auth_headers,
    )
    assert granted.status_code == 200
    assert _patch(manager, "revenue-comparison", description="now allowed").status_code == 200
    assert _patch(stranger, "revenue-comparison", description="nope").status_code == 403

    with db.SessionLocal() as session:
        resources = set(session.execute(select(db.AccessDeniedEvent.resource)).scalars())
    assert resources == {rid}  # the canonical id -- not the alias -- in the security feed


def test_another_organizations_manager_cannot_reach_a_template_through_its_code(auth_headers, make_local_user):
    _make(auth_headers)
    client.post("/api/v1/organizations", json={"id": "acme", "name": "Acme"}, headers=auth_headers)
    _, acme = make_local_user("acme-admin", "pw-acme-12345", ("ROLE_ORG_ADMIN",), org_id="acme")
    assert client.get("/api/v1/reports/revenue-comparison/file", headers=acme).status_code == 404
    assert client.get("/api/v1/reports/revenue-comparison/changelog", headers=acme).status_code == 404


# --- scheduled jobs -----------------------------------------------------------------------


def test_a_scheduled_job_can_reference_its_report_by_code(auth_headers):
    _make(auth_headers)
    out = job_executors._render_bytes("revenue-comparison", "docx", {"name": "Job"})
    assert Document(BytesIO(out)).paragraphs[0].text == "Hello Job"
    with pytest.raises(job_executors.PermanentJobError):
        job_executors._render_bytes("no-such-report", "docx", {})


# --- the middleware itself ------------------------------------------------------------------


def test_only_a_plausible_unknown_code_costs_a_database_lookup(auth_headers, monkeypatch):
    created = _make(auth_headers)
    calls = []
    real = report_ref._lookup
    monkeypatch.setattr(report_ref, "_lookup", lambda code: (calls.append(code), real(code))[1])

    client.get(f"/api/v1/reports/{created['report_id']}")  # an id: no lookup
    client.get("/api/v1/reports/accessible", headers=auth_headers)  # a static route: none
    client.get("/api/v1/reports/batch-limits")
    client.get("/api/v1/reports/ab")  # too short to be a code
    client.get("/api/v1/reports/Not%20A%20Code")  # can't be a code
    client.get("/api/v1/reports")  # no segment
    assert calls == []

    client.get("/api/v1/reports/revenue-comparison")
    assert calls == ["revenue-comparison"]


def test_nothing_from_the_request_reaches_the_rewritten_path(auth_headers):
    """Only a syntactically valid code that exists is rewritten, and what goes in
    is the id the database returned -- so path tricks in the segment do nothing."""
    created = _make(auth_headers)
    for tricky in ("../reports/" + created["report_id"], "revenue-comparison%2f..%2f..", "revenue-comparison/../x", "%2e%2e"):
        resp = client.get(f"/api/v1/reports/{tricky}")
        assert resp.status_code in (404, 405, 307, 308) or resp.json().get("report_id") in (None, created["report_id"])
    # The deeper routes still see exactly the tail that was sent.
    assert client.get("/api/v1/reports/revenue-comparison/schema").status_code == 200
    assert client.get("/api/v1/reports/revenue-comparison/nonexistent-tail").status_code == 404


def test_rules_are_the_same_in_the_store_as_at_the_router(auth_headers):
    """The store enforces them too, so no other caller can get a bad code in."""
    with pytest.raises(report_ref.InvalidCodeError):
        report_store.create_report("x", _docx(), "docx", code="Bad Code")
    assert report_store.list_reports() == []
    rid = report_store.create_report("x", _docx(), "docx", code="ok-code")["report_id"]
    with pytest.raises(report_ref.InvalidCodeError):
        report_store.update_report_meta(rid, code="1f094af3df23")  # id-shaped
    assert report_store.get_report(rid)["code"] == "ok-code"  # the refused update changed nothing
    assert report_store.resolve_ref("ok-code") == rid == report_store.resolve_ref(rid)
    with pytest.raises(report_store.ReportNotFoundError):
        report_store.resolve_ref("missing")
