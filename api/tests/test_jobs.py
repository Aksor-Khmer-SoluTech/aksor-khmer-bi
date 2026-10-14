import pytest
from fastapi.testclient import TestClient

from app.celery_app import celery_app
from app.main import app
from app.rbac import ROOT_ORG_ID

client = TestClient(app)


@pytest.fixture(autouse=True)
def stub_celery_send_task(monkeypatch):
    """These router tests exercise permission checks, org-scoping, and
    JobRun bookkeeping -- not Celery's ability to actually publish to a
    broker, which would otherwise require a live Redis just to run this
    file. Real end-to-end task execution (a live Redis + a real Celery
    worker actually running app/job_executors.py code, including a
    forced-failure retry case) is verified separately -- see
    docs/deployment.md's job-scheduling section for how, matching the
    same standard applied to Postgres/LDAP elsewhere in this project.
    """
    sent = []

    class _StubResult:
        id = "stub-task-id"

    def _stub_send_task(name, args=None, **kwargs):
        sent.append((name, args))
        return _StubResult()

    monkeypatch.setattr(celery_app, "send_task", _stub_send_task)
    return sent


def _create_job(auth_headers, **overrides) -> dict:
    body = {
        "name": "Nightly sweep",
        "job_type": "file_output",
        "config": {"destination_path": "/tmp/out-{job_id}.txt", "source": "static", "content": "hi"},
        "trigger_type": "cron",
        "cron_expression": "0 2 * * *",
    }
    body.update(overrides)
    resp = client.post("/api/v1/jobs", params={"org_id": ROOT_ORG_ID}, json=body, headers=auth_headers)
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_create_and_get_job(auth_headers):
    created = _create_job(auth_headers)
    resp = client.get(f"/api/v1/jobs/{created['id']}", headers=auth_headers)
    assert resp.status_code == 200
    assert resp.json()["name"] == "Nightly sweep"
    assert resp.json()["is_enabled"] is True


def test_create_job_requires_permission(make_local_user):
    _, headers = make_local_user("nojobperm", "pw12345", role_names=())
    resp = client.post(
        "/api/v1/jobs",
        params={"org_id": ROOT_ORG_ID},
        json={"name": "x", "job_type": "file_output", "config": {}, "trigger_type": "cron", "cron_expression": "0 2 * * *"},
        headers=headers,
    )
    assert resp.status_code == 403


def test_job_operator_can_trigger_but_not_manage(make_local_user, auth_headers):
    created = _create_job(auth_headers)
    _, operator_headers = make_local_user("operator1", "pw12345", role_names=("ROLE_JOB_OPERATOR",))

    run_resp = client.post(f"/api/v1/jobs/{created['id']}/run", headers=operator_headers)
    assert run_resp.status_code == 200

    update_resp = client.patch(f"/api/v1/jobs/{created['id']}", json={"name": "renamed"}, headers=operator_headers)
    assert update_resp.status_code == 403


def test_create_cron_job_validates_expression(auth_headers):
    resp = client.post(
        "/api/v1/jobs",
        params={"org_id": ROOT_ORG_ID},
        json={
            "name": "bad cron",
            "job_type": "file_output",
            "config": {"destination_path": "/tmp/x", "source": "static", "content": "x"},
            "trigger_type": "cron",
            "cron_expression": "not a cron expression",
        },
        headers=auth_headers,
    )
    assert resp.status_code == 400


def test_create_interval_job_requires_positive_seconds(auth_headers):
    resp = client.post(
        "/api/v1/jobs",
        params={"org_id": ROOT_ORG_ID},
        json={
            "name": "bad interval",
            "job_type": "file_output",
            "config": {},
            "trigger_type": "interval",
            "interval_seconds": 0,
        },
        headers=auth_headers,
    )
    assert resp.status_code == 400


def test_create_one_off_job_requires_run_at(auth_headers):
    resp = client.post(
        "/api/v1/jobs",
        params={"org_id": ROOT_ORG_ID},
        json={"name": "bad one-off", "job_type": "file_output", "config": {}, "trigger_type": "one_off"},
        headers=auth_headers,
    )
    assert resp.status_code == 400


def test_update_job_disable(auth_headers):
    created = _create_job(auth_headers)
    resp = client.patch(f"/api/v1/jobs/{created['id']}", json={"is_enabled": False}, headers=auth_headers)
    assert resp.status_code == 200
    assert resp.json()["is_enabled"] is False


def test_delete_job(auth_headers):
    created = _create_job(auth_headers)
    resp = client.delete(f"/api/v1/jobs/{created['id']}", headers=auth_headers)
    assert resp.status_code == 204
    assert client.get(f"/api/v1/jobs/{created['id']}", headers=auth_headers).status_code == 404


def test_run_job_creates_a_queued_run_and_appears_in_history(auth_headers):
    created = _create_job(auth_headers)
    run = client.post(f"/api/v1/jobs/{created['id']}/run", headers=auth_headers)
    assert run.status_code == 200
    assert run.json()["triggered_by"] == "manual"
    assert run.json()["status"] == "queued"

    history = client.get(f"/api/v1/jobs/{created['id']}/runs", headers=auth_headers)
    assert history.status_code == 200
    assert len(history.json()) == 1
    assert history.json()[0]["id"] == run.json()["id"]


def test_run_job_requires_job_trigger_permission(make_local_user, auth_headers):
    created = _create_job(auth_headers)
    _, headers = make_local_user("notrigger", "pw12345", role_names=())
    resp = client.post(f"/api/v1/jobs/{created['id']}/run", headers=headers)
    assert resp.status_code == 403


def test_jobs_are_org_scoped(auth_headers, make_local_user):
    client.post("/api/v1/organizations", json={"id": "jobsorg", "name": "Jobs Org"}, headers=auth_headers)
    root_org_job = _create_job(auth_headers, name="root org job")
    other_org_job = client.post(
        "/api/v1/jobs",
        params={"org_id": "jobsorg"},
        json={
            "name": "org-scoped job",
            "job_type": "file_output",
            "config": {"destination_path": "/tmp/x", "source": "static", "content": "x"},
            "trigger_type": "cron",
            "cron_expression": "0 2 * * *",
        },
        headers=auth_headers,
    ).json()

    _, root_operator_headers = make_local_user(
        "jobsorgoperator", "pw12345", role_names=("ROLE_JOB_OPERATOR",), org_id=ROOT_ORG_ID
    )
    # A ROOT_ORG_ID job operator cannot see a job belonging to jobsorg...
    forbidden = client.get(f"/api/v1/jobs/{other_org_job['id']}", headers=root_operator_headers)
    assert forbidden.status_code == 404
    # ...but can see the root-org one.
    ok = client.get(f"/api/v1/jobs/{root_org_job['id']}", headers=root_operator_headers)
    assert ok.status_code == 200


# --- a job that runs a report: its data, its filters, its schedule ---------------------------------------------------

from io import BytesIO  # noqa: E402

import httpx  # noqa: E402
from docx import Document  # noqa: E402

from app import report_data  # noqa: E402
from app.job_executors import PermanentJobError, RetryableJobError, execute_render_report  # noqa: E402

BRANCHES = [{"value": "BR01", "label": "Phnom Penh"}, {"value": "BR02", "label": "Siem Reap"}]


def _sales_report(auth_headers, parameters=None) -> str:
    doc = Document()
    doc.add_paragraph("Branch {{ p_branch }} on {{ p_day }}: {{ branch_name }} total {{ total }}")
    buf = BytesIO()
    doc.save(buf)
    report_id = client.post(
        "/api/v1/reports", files={"file": ("t.docx", buf.getvalue(), "application/octet-stream")}, data={"name": "Daily sales"},
        headers=auth_headers,
    ).json()["report_id"]
    resp = client.put(
        f"/api/v1/reports/{report_id}/data-config",
        json={
            "parameters": parameters if parameters is not None else [
                {"name": "p_branch", "label": "Branch", "options": BRANCHES},
                {"name": "p_day", "label": "Day", "type": "date", "default_value": "now()"},
            ],
            "data_source": {"url": "https://erp.example.test/api/sales?branch={{ p_branch }}&day={{ p_day }}", "method": "GET"},
        },
        headers=auth_headers,
    )
    assert resp.status_code == 200, resp.text
    return report_id


@pytest.fixture
def erp(monkeypatch) -> list:
    """The report's REST data source: records each request, answers with that branch's figures."""
    calls: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(request)
        return httpx.Response(200, json={"branch_name": f"{request.url.params['branch']} Office", "total": "1200"})

    monkeypatch.setattr(report_data, "_make_client", lambda: httpx.Client(transport=httpx.MockTransport(handler)))
    return calls


def _report_job(auth_headers, config: dict, expect=200):
    resp = client.post(
        "/api/v1/jobs",
        params={"org_id": ROOT_ORG_ID},
        json={"name": "Daily sales", "job_type": "render_report", "config": config, "trigger_type": "cron", "cron_expression": "0 8 * * *"},
        headers=auth_headers,
    )
    assert resp.status_code == expect, resp.text
    return resp.json()


def _docx_text(path) -> str:
    return "\n".join(p.text for p in Document(str(path)).paragraphs).replace("​", "")


def test_a_report_job_fetches_the_reports_own_data_with_its_filters(auth_headers, erp, tmp_path):
    report_id = _sales_report(auth_headers)
    out = tmp_path / "sales.docx"
    job = _report_job(auth_headers, {"report_id": report_id, "format": "docx", "parameters": {"p_branch": "BR02"}, "destination_path": str(out)})

    execute_render_report(job["config"], job["id"])

    # The data source was called with the chosen branch, and the date filter took its now() default -- the day it fired.
    assert len(erp) == 1 and erp[0].url.params["branch"] == "BR02"
    day = erp[0].url.params["day"]
    assert len(day) == 10 and day[4] == "-"
    assert _docx_text(out) == f"Branch BR02 on {day}: BR02 Office total 1200"


def test_a_report_jobs_context_is_laid_over_the_fetched_data(auth_headers, erp, tmp_path):
    report_id = _sales_report(auth_headers)
    out = tmp_path / "sales.docx"
    config = {"report_id": report_id, "format": "docx", "parameters": {"p_branch": "BR01", "p_day": "2026-10-01"},
              "context": {"total": "999"}, "destination_path": str(out)}
    execute_render_report(_report_job(auth_headers, config)["config"], "j1")
    assert _docx_text(out) == "Branch BR01 on 2026-10-01: BR01 Office total 999"


def test_an_older_job_with_only_context_still_renders_just_that(auth_headers, erp, tmp_path):
    report_id = _sales_report(auth_headers)
    out = tmp_path / "sales.docx"
    config = {"report_id": report_id, "format": "docx", "destination_path": str(out),
              "context": {"p_branch": "X", "p_day": "-", "branch_name": "Given", "total": "1"}}
    execute_render_report(config, "j1")
    assert erp == []  # the data source isn't called
    assert _docx_text(out) == "Branch X on -: Given total 1"


def test_a_report_job_whose_data_source_is_down_is_retried(auth_headers, monkeypatch, tmp_path):
    report_id = _sales_report(auth_headers)

    def down(request):
        raise httpx.ConnectError("refused", request=request)

    monkeypatch.setattr(report_data, "_make_client", lambda: httpx.Client(transport=httpx.MockTransport(down)))
    config = {"report_id": report_id, "format": "docx", "parameters": {"p_branch": "BR01"}, "destination_path": str(tmp_path / "x.docx")}
    with pytest.raises(RetryableJobError):
        execute_render_report(config, "j1")


def test_a_report_job_with_a_filter_value_gone_from_the_report_fails_for_good(auth_headers, erp, tmp_path):
    report_id = _sales_report(auth_headers)
    config = {"report_id": report_id, "format": "docx", "parameters": {"p_branch": "BR09"}, "destination_path": str(tmp_path / "x.docx")}
    with pytest.raises(PermanentJobError, match="valid choice"):
        execute_render_report(config, "j1")
    assert erp == []


def test_saving_a_report_job_checks_the_report_format_and_filters(auth_headers):
    report_id = _sales_report(auth_headers)
    base = {"report_id": report_id, "format": "pdf", "parameters": {"p_branch": "BR01"}, "destination_path": "/tmp/s.pdf"}

    assert _report_job(auth_headers, {**base, "report_id": "nope"}, expect=404)["detail"] == "Report not found"
    assert "not 'html'" in _report_job(auth_headers, {**base, "format": "html"}, expect=400)["detail"]
    assert _report_job(auth_headers, {**base, "parameters": {"p_branch": "BR09"}}, expect=400)["detail"] == "That isn't a valid choice for Branch"
    assert _report_job(auth_headers, {**base, "parameters": {"p_nope": "1"}}, expect=400)["detail"] == "Unknown parameter(s): p_nope"
    assert _report_job(auth_headers, {**base, "parameters": {}}, expect=400)["detail"] == "Choose a value for Branch"
    assert "save it to a folder, email it, or both" in _report_job(auth_headers, {**base, "destination_path": ""}, expect=400)["detail"]
    job = _report_job(auth_headers, base)

    # The same checks on an edit of the config -- but not on an edit that leaves it alone.
    bad = client.patch(f"/api/v1/jobs/{job['id']}", json={"config": {**base, "parameters": {"p_branch": "BR09"}}}, headers=auth_headers)
    assert bad.status_code == 400
    assert client.patch(f"/api/v1/jobs/{job['id']}", json={"name": "Renamed"}, headers=auth_headers).status_code == 200


def test_a_report_job_cant_use_a_report_of_another_organization(auth_headers):
    from app import db as app_db
    from app.rbac import create_organization

    report_id = _sales_report(auth_headers)
    with app_db.SessionLocal() as session:
        create_organization(session, org_id="joborg2", name="Job Org 2")
        session.commit()
    resp = client.post(
        "/api/v1/jobs",
        params={"org_id": "joborg2"},
        json={"name": "x", "job_type": "render_report", "trigger_type": "cron", "cron_expression": "0 8 * * *",
              "config": {"report_id": report_id, "format": "pdf", "parameters": {"p_branch": "BR01"}, "destination_path": "/tmp/x.pdf"}},
        headers=auth_headers,
    )
    assert resp.status_code == 400
    assert "another organization" in resp.json()["detail"]


def test_schedule_preview_gives_the_next_runs_in_the_schedulers_time_zone(auth_headers):
    resp = client.post("/api/v1/jobs/schedule-preview", json={"cron_expression": "30 8 * * 1-5", "count": 5}, headers=auth_headers)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["timezone"]
    from datetime import datetime as dt

    runs = [dt.fromisoformat(r) for r in body["next_runs"]]
    assert len(runs) == 5 and runs == sorted(runs) and len(set(runs)) == 5
    assert all((r.hour, r.minute) == (8, 30) and r.weekday() < 5 for r in runs)


def test_schedule_preview_stops_at_the_end_date_and_rejects_nonsense(auth_headers):
    from datetime import datetime as dt, timedelta, timezone as tz

    end = (dt.now(tz.utc) + timedelta(days=2)).isoformat()
    body = client.post("/api/v1/jobs/schedule-preview", json={"cron_expression": "0 0 * * *", "end_date": end, "count": 10}, headers=auth_headers).json()
    assert 1 <= len(body["next_runs"]) <= 2
    bad = client.post("/api/v1/jobs/schedule-preview", json={"cron_expression": "every day"}, headers=auth_headers)
    assert bad.status_code == 400 and "Invalid cron_expression" in bad.json()["detail"]
