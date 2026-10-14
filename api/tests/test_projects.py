"""Projects (app/projects.py, routers/projects.py): a top-level folder marked as a project, exported as one package
and imported elsewhere. "Elsewhere" here is a second organization in the same database -- resources are matched per
organization, so it behaves like another environment. Checked: what a project holds and uses; the package (shared
resources left out on request); a first import (everything created, ids re-linked, the templates render); a second
import (identical files linked, a changed one overwritten or kept beside, templates gaining a version); missing
things warned about; codes that clash; history; who may; a failure half-way undone."""
import io
import json
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from app import db, font_catalog, font_store, image_store, projects, report_store, stylesheet_store
from tests.test_fonts import build_font
from app.main import app
from app.rbac import ROOT_ORG_ID

client = TestClient(app)
TARGET = "acme"  # the "other environment"


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch, auth_headers):
    monkeypatch.setattr(image_store, "STORE_DIR", tmp_path / "images")
    monkeypatch.setattr(stylesheet_store, "STORE_DIR", tmp_path / "styles")
    monkeypatch.setattr(projects, "IMPORT_DIR", tmp_path / "imports")
    monkeypatch.setattr(font_store, "FONT_DIR", tmp_path / "fonts")
    monkeypatch.setenv("DOC_ENGINE_FONT_DIRS", str(tmp_path / "fonts"))
    monkeypatch.setattr(font_catalog, "system_families", lambda: {"arial": "Arial"})
    font_catalog.forget()
    client.post("/api/v1/organizations", json={"id": TARGET, "name": "Acme"}, headers=auth_headers)


def png(color=(200, 30, 30), size=(24, 12)) -> bytes:
    out = io.BytesIO()
    Image.new("RGB", size, color).save(out, format="PNG")
    return out.getvalue()


def docx(*paragraphs) -> bytes:
    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def folder(headers, name, parent=None, org=ROOT_ORG_ID):
    resp = client.post("/api/v1/folders", json={"org_id": org, "name": name, "parent_folder_id": parent}, headers=headers)
    assert resp.status_code == 200, resp.text
    return resp.json()


def image(headers, name, content=None, folder_id=None, org=ROOT_ORG_ID):
    data = {"name": name, "org_id": org, **({"folder_id": folder_id} if folder_id else {})}
    resp = client.post("/api/v1/images", files={"file": (f"{name}.png", content or png(), "image/png")}, data=data, headers=headers)
    assert resp.status_code == 200, resp.text
    return resp.json()


def stylesheet(headers, name, css=b"p { color: #06519e; }", folder_id=None, org=ROOT_ORG_ID):
    data = {"name": name, "org_id": org, **({"folder_id": folder_id} if folder_id else {})}
    resp = client.post("/api/v1/stylesheets", files={"file": (f"{name}.css", css, "text/css")}, data=data, headers=headers)
    assert resp.status_code == 200, resp.text
    return resp.json()


def register(headers, filename, content, name, folder_id, bindings=None, code=None):
    data = {"name": name, "folder_id": folder_id, **({"resource_bindings": json.dumps(bindings)} if bindings else {}), **({"code": code} if code else {})}
    resp = client.post("/api/v1/reports", files={"file": (filename, content, "application/octet-stream")}, data=data, headers=headers)
    assert resp.status_code == 200, resp.text
    return resp.json()


HTML = (
    '<link rel="stylesheet" href="{{ resource(\'theme\') }}">'
    "<h1>{{ title }}</h1>"
    '<img src="{{ resource(\'logo\') }}"><img src="{{ resource(\'seal\') }}">'
)


def build_project(headers):
    """A project "Receipts": an html receipt using its own logo plus a shared stylesheet and seal, and a Word form in
    a subfolder whose data points at a project image."""
    root = folder(headers, "Receipts")
    assert client.patch(f"/api/v1/folders/{root['id']}", json={"is_project": True}, headers=headers).status_code == 200
    forms = folder(headers, "Forms", parent=root["id"])
    logo = image(headers, "logo", png((10, 80, 160)), folder_id=root["id"])
    photo = image(headers, "photo", png((0, 120, 0)), folder_id=forms["id"])
    seal = image(headers, "seal", png((180, 150, 0)))  # shared: at the top level
    theme = stylesheet(headers, "theme")  # shared
    html = register(headers, "receipt.html", HTML.encode(), "E-receipt", root["id"], code="e-receipt",
                    bindings={"theme": {"kind": "stylesheet", "id": theme["id"]}, "logo": {"kind": "image", "id": logo["id"]},
                              "seal": {"kind": "image", "id": seal["id"]}})
    form = register(headers, "form.docx", docx("Name: {{ name }}", "{{ PHOTO }}"), "Request form", forms["id"])
    client.patch(f"/api/v1/reports/{form['report_id']}", json={"sample_context": {"name": "សុខ", "PHOTO": {"image_id": photo["id"], "width_mm": 20}}}, headers=headers)
    return {"root": root, "forms": forms, "logo": logo, "photo": photo, "seal": seal, "theme": theme, "html": html, "form": form}


def export(headers, project_id, **body):
    resp = client.post(f"/api/v1/projects/{project_id}/export", json=body, headers=headers)
    assert resp.status_code == 200, resp.text
    return resp.content


def preview(headers, package, into=None, org=TARGET, expect=200):
    data = {"org_id": org, **({"into": into} if into else {})}
    resp = client.post("/api/v1/projects/import/preview", files={"file": ("p.zip", package, "application/zip")}, data=data, headers=headers)
    assert resp.status_code == expect, resp.text
    return resp.json()


def apply(headers, token, decisions=None, expect=200):
    resp = client.post(f"/api/v1/projects/import/{token}", json={"decisions": decisions or {}}, headers=headers)
    assert resp.status_code == expect, resp.text
    return resp.json()


def by_name(items, name):
    return next(i for i in items if i["name"] == name)


# --- a project ----------------------------------------------------------------------------------------------------


def test_only_a_top_level_folder_is_a_project(auth_headers):
    root = folder(auth_headers, "Top")
    child = folder(auth_headers, "Child", parent=root["id"])
    assert client.patch(f"/api/v1/folders/{child['id']}", json={"is_project": True}, headers=auth_headers).status_code == 400
    assert client.patch(f"/api/v1/folders/{root['id']}", json={"is_project": True}, headers=auth_headers).json()["is_project"] is True
    other = folder(auth_headers, "Other")
    moved = client.patch(f"/api/v1/folders/{root['id']}", json={"parent_folder_id": other["id"]}, headers=auth_headers)
    assert moved.status_code == 400 and "top level" in moved.text
    assert any(f["is_project"] for f in client.get("/api/v1/folders", headers=auth_headers).json() if f["id"] == root["id"])


def test_what_a_project_holds_and_uses(auth_headers):
    p = build_project(auth_headers)
    resp = client.get(f"/api/v1/projects/{p['root']['id']}", headers=auth_headers)
    assert resp.status_code == 200, resp.text
    info = resp.json()
    assert [t["name"] for t in info["templates"]] == ["E-receipt", "Request form"]
    assert by_name(info["templates"], "Request form")["folder"] == "Forms"
    res = {r["name"]: r for r in info["resources"]}
    assert res["logo"]["scope"] == "project" and res["photo"]["scope"] == "project" and res["photo"]["folder"] == "Forms"
    assert res["seal"]["scope"] == "shared" and res["theme"]["scope"] == "shared" and res["seal"]["used_by"] == ["E-receipt"]
    assert res["photo"]["used_by"] == ["Request form"]  # found in the form's data, not just its folder
    assert info["warnings"] == []


# --- export and a first import ------------------------------------------------------------------------------------


def test_a_package_and_a_first_import(auth_headers):
    p = build_project(auth_headers)
    package = export(auth_headers, p["root"]["id"])
    with zipfile.ZipFile(io.BytesIO(package)) as z:
        manifest = json.loads(z.read("manifest.json"))
        assert {r["name"] for r in manifest["resources"] if r["file"]} == {"logo", "photo", "seal", "theme"}
        assert all(r["file"] in z.namelist() for r in manifest["resources"])
    plan = preview(auth_headers, package)
    assert plan["project"] == {"name": "Receipts", "existing_id": None, "exported_at": manifest["exported_at"], "exported_by": "testadmin"}
    assert {r["name"]: r["status"] for r in plan["resources"]} == {"logo": "new", "photo": "new", "seal": "new", "theme": "new"}
    assert by_name(plan["resources"], "seal")["default"] == "create_shared"
    assert [t["status"] for t in plan["templates"]] == ["new", "new"]

    done = apply(auth_headers, plan["token"])
    assert 'Project "Receipts"' in done["created"] and 'Template "E-receipt"' in done["created"]
    with db.SessionLocal() as session:
        project = session.get(db.Folder, done["project"]["id"])
        assert project.org_id == TARGET and project.is_project
        reports = {r.name: r for r in session.query(db.ReportRow).filter(db.ReportRow.org_id == TARGET)}
        images = {i.name: i for i in session.query(db.ImageResource).filter(db.ImageResource.org_id == TARGET)}
    receipt, form = reports["E-receipt"], reports["Request form"]
    assert receipt.code is None or receipt.code == "e-receipt"  # the code is the source report's -- see the clash test
    assert receipt.resource_bindings["logo"]["id"] == images["logo"].id != p["logo"]["id"]  # re-linked, not the old id
    assert images["seal"].folder_id is None  # shared stays shared
    assert form.sample_context["PHOTO"]["image_id"] == images["photo"].id  # a Word template's picture re-linked too
    rendered = client.post(f"/api/v1/reports/{receipt.report_id}/render?format=pdf", json={"title": "វិក្កយបត្រ"}, headers=auth_headers)
    assert rendered.status_code == 200 and rendered.content.startswith(b"%PDF"), rendered.text[:300]
    word = client.post(f"/api/v1/reports/{form.report_id}/render?format=docx", json=form.sample_context, headers=auth_headers)
    assert word.status_code == 200 and len(Document(io.BytesIO(word.content)).inline_shapes) == 1


def test_a_code_taken_here_is_left_off_and_said(auth_headers):
    p = build_project(auth_headers)
    package = export(auth_headers, p["root"]["id"])
    plan = preview(auth_headers, package)
    receipt = by_name(plan["templates"], "E-receipt")
    assert "e-receipt" in receipt["code_note"]  # the code is the source's own report, in another organization
    done = apply(auth_headers, plan["token"])
    assert any("e-receipt" in w for w in done["warnings"])


def test_history_travels(auth_headers):
    p = build_project(auth_headers)
    rid = p["form"]["report_id"]
    client.put(f"/api/v1/reports/{rid}/file", files={"file": ("f.docx", docx("v2 {{ name }}"), "application/octet-stream")},
               data={"note": "Second", "version_label": "1.1"}, headers=auth_headers)
    done = apply(auth_headers, preview(auth_headers, export(auth_headers, p["root"]["id"]))["token"])
    assert 'Template "Request form" (2 versions)' in done["created"]
    with db.SessionLocal() as session:
        new_id = session.query(db.ReportRow).filter(db.ReportRow.org_id == TARGET, db.ReportRow.name == "Request form").one().report_id
    versions = report_store.list_versions(new_id)
    assert [v["version"] for v in versions] == [2, 1] and versions[0]["version_label"] == "1.1"
    without = export(auth_headers, p["root"]["id"], versions=False)
    assert json.loads(zipfile.ZipFile(io.BytesIO(without)).read("manifest.json"))["templates"][1]["history"] == []


# --- importing again ------------------------------------------------------------------------------------------------


def test_importing_again_updates_and_asks_about_changed_files(auth_headers):
    p = build_project(auth_headers)
    first = apply(auth_headers, preview(auth_headers, export(auth_headers, p["root"]["id"]))["token"])
    target_id = first["project"]["id"]
    # Back at the source: the logo and the shared seal change, a template gets a new version.
    image_store.replace_image_content(p["logo"]["id"], png((255, 255, 0)))
    image_store.replace_image_content(p["seal"]["id"], png((0, 0, 0)))
    client.put(f"/api/v1/reports/{p['html']['report_id']}/file", files={"file": ("r.html", (HTML + "<p>v2</p>").encode(), "text/html")},
               headers=auth_headers)
    plan = preview(auth_headers, export(auth_headers, p["root"]["id"]), into=target_id)
    status = {r["name"]: (r["status"], r["default"]) for r in plan["resources"]}
    assert status["photo"] == ("identical", None) and status["theme"] == ("identical", None)  # linked, never copied again
    assert status["logo"] == ("differs", "overwrite")  # the project's own file: newer wins by default
    # A shared file with the same name but other content: keeping the one here is the safe default -- overwriting a
    # shared file would change every other template that uses it.
    assert status["seal"] == ("differs", "use_existing")
    seal_key = by_name(plan["resources"], "seal")["key"]
    receipt = by_name(plan["templates"], "E-receipt")
    assert receipt["status"] == "exists" and receipt["default"] == "update"

    logo_key = by_name(plan["resources"], "logo")["key"]
    done = apply(auth_headers, plan["token"], {logo_key: "keep_both", seal_key: "keep_both"})
    with db.SessionLocal() as session:
        target_logos = session.query(db.ImageResource).filter(db.ImageResource.org_id == TARGET, db.ImageResource.name == "logo").all()
        report = session.query(db.ReportRow).filter(db.ReportRow.org_id == TARGET, db.ReportRow.name == "E-receipt").one()
        reports = session.query(db.ReportRow).filter(db.ReportRow.org_id == TARGET).count()
    assert len(target_logos) == 2  # kept both
    with db.SessionLocal() as session:
        seals = session.query(db.ImageResource).filter(db.ImageResource.org_id == TARGET, db.ImageResource.name == "seal").all()
    assert sorted(s.folder_id is None for s in seals) == [False, True]  # the shared one untouched, the import's own copy in the project
    assert report.version == 2 and reports == 2  # updated, not duplicated
    assert report.resource_bindings["logo"]["id"] in {i.id for i in target_logos}
    assert any("new version" in u for u in done["updated"])


def test_overwrite_changes_the_existing_file_in_place(auth_headers):
    p = build_project(auth_headers)
    target_id = apply(auth_headers, preview(auth_headers, export(auth_headers, p["root"]["id"]))["token"])["project"]["id"]
    with db.SessionLocal() as session:
        old = session.query(db.ImageResource).filter(db.ImageResource.org_id == TARGET, db.ImageResource.name == "logo").one()
    image_store.replace_image_content(p["logo"]["id"], png((1, 2, 3)))
    plan = preview(auth_headers, export(auth_headers, p["root"]["id"]), into=target_id)
    apply(auth_headers, plan["token"])  # defaults: overwrite the project's own changed logo
    assert image_store.fingerprint(old.id) == image_store.fingerprint(p["logo"]["id"])  # same id, new content


def test_term_sets_and_fonts_travel_too(auth_headers):
    p = build_project(auth_headers)
    terms = client.post("/api/v1/protected-term-sets", json={"org_id": ROOT_ORG_ID, "name": "Bank names", "terms": ["ធនាគារ"], "exclude_terms": []},
                        headers=auth_headers).json()
    client.put(f"/api/v1/reports/{p['html']['report_id']}/protected-terms-config", json={"set_ids": [terms["id"]], "terms": [], "exclude_terms": []},
               headers=auth_headers)
    font = client.post("/api/v1/fonts", files={"file": ("brand.ttf", build_font("Aksor Brand"), "font/ttf")}, headers=auth_headers).json()
    stylesheet_store.replace_stylesheet_content(p["theme"]["id"], b'h1 { font-family: "Aksor Brand"; }')  # the theme names it
    info = client.get(f"/api/v1/projects/{p['root']['id']}", headers=auth_headers).json()
    assert [f["family"] for f in info["fonts"]] == ["Aksor Brand"] and info["term_sets"][0]["name"] == "Bank names"
    package = export(auth_headers, p["root"]["id"])
    client.delete(f"/api/v1/fonts/{font['id']}?force=true", headers=auth_headers)  # a server that doesn't have it
    plan = preview(auth_headers, package)
    assert plan["fonts"][0]["status"] == "new" and plan["fonts"][0]["default"] == "install"
    assert plan["term_sets"][0]["status"] == "new"
    done = apply(auth_headers, plan["token"])
    assert "Font Aksor Brand Regular" in done["created"] and 'Protected terms "Bank names"' in done["created"]
    with db.SessionLocal() as session:
        report = session.query(db.ReportRow).filter(db.ReportRow.org_id == TARGET, db.ReportRow.name == "E-receipt").one()
        new_set = session.query(db.ProtectedTermSet).filter(db.ProtectedTermSet.org_id == TARGET).one()
        assert session.query(db.FontResource).filter(db.FontResource.family == "Aksor Brand").count() == 1
    assert report.protected_terms_config["set_ids"] == [new_set.id]
    # Exported without fonts this time: none listed; the term set is already here, so it's linked.
    again = preview(auth_headers, export(auth_headers, p["root"]["id"], fonts=False), into=done["project"]["id"])
    assert again["fonts"] == [] and again["term_sets"][0]["status"] == "identical"


# --- leaving shared things out -------------------------------------------------------------------------------------


def test_a_shared_resource_left_out_is_found_or_warned_about(auth_headers):
    p = build_project(auth_headers)
    package = export(auth_headers, p["root"]["id"], leave_out=[f"image:{p['seal']['id']}", f"stylesheet:{p['theme']['id']}"])
    manifest = json.loads(zipfile.ZipFile(io.BytesIO(package)).read("manifest.json"))
    assert {r["name"] for r in manifest["resources"] if not r["file"]} == {"seal", "theme"}
    # The target has the same theme (identical content) but no seal.
    stylesheet(auth_headers, "site-theme", org=TARGET)
    plan = preview(auth_headers, package)
    status = {r["name"]: r["status"] for r in plan["resources"]}
    assert status["theme"] == "identical" and status["seal"] == "missing"
    assert any('"seal"' in w and "E-receipt" in w for w in plan["warnings"])


def test_an_upload_says_when_the_same_file_is_already_here(auth_headers):
    p = build_project(auth_headers)
    again = image(auth_headers, "logo copy", png((10, 80, 160)))  # the project's logo, uploaded again at the top level
    assert again["duplicates"] == [{"id": p["logo"]["id"], "name": "logo", "where": "Receipts"}]
    fresh = image(auth_headers, "something new", png((1, 1, 1)))
    assert fresh["duplicates"] == []
    css = stylesheet(auth_headers, "theme again")  # same CSS as the shared theme
    assert [d["name"] for d in css["duplicates"]] == ["theme"]


# --- who may, and what goes wrong ---------------------------------------------------------------------------------


def test_transfer_needs_its_permission(auth_headers, make_local_user):
    p = build_project(auth_headers)
    _, viewer = make_local_user("viewer", "pw-viewer-1", ("ROLE_REPORT_VIEWER",))
    assert client.post(f"/api/v1/projects/{p['root']['id']}/export", json={}, headers=viewer).status_code == 403
    _, author = make_local_user("author", "pw-author-1", ("ROLE_REPORT_ADMIN",))
    assert client.post(f"/api/v1/projects/{p['root']['id']}/export", json={}, headers=author).status_code == 200
    # A preview is its uploader's alone.
    plan = preview(auth_headers, export(auth_headers, p["root"]["id"]), org=ROOT_ORG_ID, expect=400)  # the name is taken here
    assert "already a folder" in plan["detail"]
    token = preview(auth_headers, export(auth_headers, p["root"]["id"]))["token"]
    assert client.post(f"/api/v1/projects/import/{token}", json={}, headers=author).status_code in (403, 404)


def test_not_a_package(auth_headers):
    assert "not a .zip" in preview(auth_headers, b"hello", expect=400)["detail"]
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("manifest.json", json.dumps({"kind": "something-else"}))
    assert "isn't an Aksor project package" in preview(auth_headers, out.getvalue(), expect=400)["detail"]


def test_a_failure_half_way_leaves_nothing_behind(auth_headers, monkeypatch):
    p = build_project(auth_headers)
    plan = preview(auth_headers, export(auth_headers, p["root"]["id"]))
    real = report_store.create_report
    calls = []

    def flaky(*args, **kwargs):
        calls.append(1)
        if len(calls) == 2:
            raise RuntimeError("disk full")
        return real(*args, **kwargs)

    monkeypatch.setattr(report_store, "create_report", flaky)
    resp = client.post(f"/api/v1/projects/import/{plan['token']}", json={"decisions": {}}, headers=auth_headers)
    assert resp.status_code == 500 and len(calls) == 2
    with db.SessionLocal() as session:
        assert session.query(db.Folder).filter(db.Folder.org_id == TARGET).count() == 0
        assert session.query(db.ImageResource).filter(db.ImageResource.org_id == TARGET).count() == 0
        assert session.query(db.ReportRow).filter(db.ReportRow.org_id == TARGET).count() == 0


def test_export_and_import_are_audited(auth_headers):
    p = build_project(auth_headers)
    done = apply(auth_headers, preview(auth_headers, export(auth_headers, p["root"]["id"]))["token"])
    with db.SessionLocal() as session:
        actions = {e.action: e for e in session.query(db.AuditEvent).filter(db.AuditEvent.action.like("project.%"))}
    assert actions["project.export"].summary.startswith('Exported project "Receipts" (2 templates, 4 resources')
    assert actions["project.import"].entity_id == done["project"]["id"]
