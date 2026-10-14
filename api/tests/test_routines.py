"""Stored procedures and functions (app/routines.py): a connection manager approves them on a SQL Server connection, a
report's query card calls one -- never SQL typed by the author -- and the call is built from the approved name with
every value bound. Checked here with a fake JDBC worker that answers by the exact statement it's sent, so the tests
also pin what is sent: the CALL text, the bound values, how many result sets are asked for. (The worker itself was
checked against a real SQL Server; see the CHANGELOG entry.)"""
from io import BytesIO

import httpx
import pytest
from docx import Document

from app import db, jdbc, jdbc_drivers, routines
from app.report_data import DataConfigError
from tests.test_jdbc import _configure, _connection, _pg, _register, _secret, _upload, http

_KHMER_BRANCH = "ភ្នំពេញ"
_KHMER_TITLE = "របាយការណ៍លក់"

APPROVED = [
    {"name": "dbo.GetBranchSales", "kind": "procedure",
     "parameters": [{"name": "@branch", "type": "nvarchar(10)"}, {"name": "@title", "type": "nvarchar(100)"}]},
    {"name": "dbo.fn_SalesFrom", "kind": "table_function", "parameters": [{"name": "@from", "type": "date"}]},
    {"name": "dbo.fn_BranchTotal", "kind": "scalar_function", "parameters": [{"name": "@branch", "type": "nvarchar(10)"}]},
]
_BRANCH = {"name": "branch", "label": "Branch", "type": "text", "required": True}
_FROM = {"name": "from_date", "label": "From", "type": "date", "required": True}


@pytest.fixture(autouse=True)
def driver_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(jdbc_drivers, "DRIVER_DIR", tmp_path / "jdbc_drivers")
    monkeypatch.setenv("JDBC_WORKER_URL", "http://jdbc-worker:9000")
    monkeypatch.setenv("JDBC_WORKER_TOKEN", "worker-token")


@pytest.fixture
def worker(monkeypatch):
    """A fake jdbc-worker (protocol 2): answers each statement it knows; records every request."""
    sent: list[dict] = []
    answers = {
        "{call [dbo].[GetBranchSales](?, ?)}": {
            "columns": ["code", "name", "title"], "rows": [["BR01", _KHMER_BRANCH, _KHMER_TITLE]],
            "more": [{"columns": ["sold_on", "amount"], "rows": [["2026-10-01", 40000.5], ["2026-10-02", 25000]]}],
        },
        "SELECT * FROM [dbo].[fn_SalesFrom](?)": {"columns": ["branch", "total"], "rows": [["BR01", 65000.5], ["BR02", 35000]]},
        "SELECT [dbo].[fn_BranchTotal](?) AS [value]": {"columns": ["value"], "rows": [[35000]]},
        routines.DISCOVER_SQL.strip(): {"columns": ["schema_name", "object_name", "object_type", "parameter_id", "parameter_name", "type_name", "max_length", "is_output"],
                                "rows": [["dbo", "CountSales", "P", 1, "@n", "int", 4, 1],
                                         ["dbo", "GetBranchSales", "P", 1, "@branch", "nvarchar", 20, 0],
                                         ["dbo", "GetBranchSales", "P", 2, "@title", "nvarchar", -1, 0],
                                         ["dbo", "fn_SalesFrom", "IF", 1, "@from", "date", 3, 0],
                                         ["dbo", "Weird Name", "P", None, None, None, None, None]]},
        routines.CAN_WRITE_SQL.strip(): {"columns": ["can_write"], "rows": [[1]]},
    }
    state = {"protocol": 2}

    def fake_post(url, json, headers, timeout):
        sent.append(json)
        statements = json.get("queries") or [{"sql": json["sql"], "args": json["args"]}]
        results = []
        for statement in statements:
            answer = dict(answers[statement["sql"].strip()])
            if not statement.get("call"):
                answer.pop("more", None)
            results.append(answer)
        if "queries" not in json:
            return httpx.Response(200, json=results[0])
        body = {"results": results, **({"protocol": state["protocol"]} if state["protocol"] else {})}
        return httpx.Response(200, json=body)

    monkeypatch.setattr(jdbc.httpx, "post", fake_post)
    return {"sent": sent, "state": state}


def _sqlserver(headers, routines_list=None):
    _secret(headers)
    driver = _upload(headers, name="mssql-jdbc", engine="sqlserver", klass="com.microsoft.sqlserver.jdbc.SQLServerDriver")
    conn = _connection(headers, name="erp-mssql", config=_pg(engine="sqlserver", port=1433, database="sales", driver_id=driver["id"]))
    if routines_list is not None:
        resp = http.put(f"/api/v1/connections/{conn['id']}/routines", json={"routines": routines_list}, headers=headers)
        assert resp.status_code == 200, resp.text
    return conn


def _report(headers, text: str) -> str:
    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return http.post("/api/v1/reports", files={"file": ("t.docx", buf.getvalue(), "application/octet-stream")},
                     data={"name": "Branch sales"}, headers=headers).json()["report_id"]


def _text(content: bytes) -> str:
    return "\n".join(p.text for p in Document(BytesIO(content)).paragraphs).replace("​", "")


_CARDS = [
    {"name": "info", "shape": "object", "routine": "dbo.GetBranchSales",
     "arguments": {"@branch": {"filter": "branch"}, "@title": {"value": _KHMER_TITLE}},
     "more_results": [{"name": "lines", "shape": "list"}]},
    {"name": "totals", "shape": "list", "routine": "dbo.fn_SalesFrom", "arguments": {"@from": {"filter": "from_date"}}},
    {"name": "other", "shape": "value", "routine": "dbo.fn_BranchTotal", "arguments": {"@branch": {"value": "BR02"}}},
]


# --- the approved list ----------------------------------------------------------------------------------------------


def test_the_approved_list_is_checked():
    assert routines.validate_routines(APPROVED, "sqlserver")[0]["parameters"][0] == {"name": "@branch", "type": "nvarchar(10)"}
    for raw, engine, fragment in [
        (APPROVED, "postgresql", "SQL Server connection"),
        ([{"name": "GetSales", "kind": "procedure"}], "sqlserver", "schema.name"),
        ([{"name": "dbo.Get Sales", "kind": "procedure"}], "sqlserver", "schema.name"),
        ([{"name": "dbo.x];DROP TABLE t;--", "kind": "procedure"}], "sqlserver", "schema.name"),
        ([{"name": "dbo.GetSales", "kind": "exec"}], "sqlserver", "stored procedure"),
        ([{"name": "dbo.GetSales", "kind": "procedure"}, {"name": "DBO.getsales", "kind": "procedure"}], "sqlserver", "twice"),
        ([{"name": "dbo.GetSales", "kind": "procedure", "parameters": [{"name": "branch"}]}], "sqlserver", "start with @"),
    ]:
        with pytest.raises(DataConfigError) as exc:
            routines.validate_routines(raw, engine)
        assert fragment in str(exc.value)


def test_approving_routines_is_audited_and_survives_an_edit(auth_headers, worker):
    conn = _sqlserver(auth_headers, APPROVED)
    # Editing the connection's settings without sending the list keeps it.
    edit = http.put(f"/api/v1/connections/{conn['id']}", json={
        "description": "ERP", "config": _pg(engine="sqlserver", port=1433, database="sales", driver_id=conn["config"]["driver_id"]),
    }, headers=auth_headers)
    assert edit.status_code == 200, edit.text
    assert [r["name"] for r in edit.json()["config"]["routines"]] == [r["name"] for r in APPROVED]
    # A report author sees what may be called (to pick one), with no credential.
    listed = next(c for c in http.get("/api/v1/connections", headers=auth_headers).json() if c["name"] == "erp-mssql")
    assert [r["name"] for r in listed["routines"]] == [r["name"] for r in APPROVED]
    with db.SessionLocal() as session:
        event = session.query(db.AuditEvent).filter(db.AuditEvent.action == "connection.routines").one()
    assert event.summary.startswith("Routines on erp-mssql: approved dbo.GetBranchSales, dbo.fn_BranchTotal, dbo.fn_SalesFrom")


def test_only_a_sql_server_connection_has_routines(auth_headers):
    _secret(auth_headers)
    pg = _connection(auth_headers, name="pg-db")
    resp = http.put(f"/api/v1/connections/{pg['id']}/routines", json={"routines": APPROVED}, headers=auth_headers)
    assert resp.status_code == 400 and "SQL Server" in resp.json()["detail"]


# --- calling them -----------------------------------------------------------------------------------------------------


def test_a_report_calls_a_procedure_and_functions(auth_headers, worker):
    _sqlserver(auth_headers, APPROVED)
    report_id = _report(auth_headers, "{{ info.title }} · {{ info.name }} · {{ lines | length }} lines · {{ totals[0].total }} · {{ other }}")
    saved = _configure(auth_headers, report_id, [_BRANCH, _FROM], {"type": "jdbc", "connection": "erp-mssql", "queries": _CARDS})
    assert saved["data_source"]["queries"][0]["routine"] == "dbo.GetBranchSales" and saved["data_source"]["queries"][0]["query"] is None

    run = http.post(f"/api/v1/reports/{report_id}/run", json={"parameters": {"branch": "BR01", "from_date": "2026-10-01"}, "format": "docx"},
                    headers=auth_headers)
    assert run.status_code == 200, run.text
    assert _text(run.content) == f"{_KHMER_TITLE} · {_KHMER_BRANCH} · 2 lines · 65000.5 · 35000"

    # One request to the worker (one JVM), and exactly this was sent: names bracket-quoted from the approved list,
    # every value bound -- the filter's and the fixed ones -- never pasted into the statement.
    [request] = worker["sent"]
    assert request["queries"] == [
        {"sql": "{call [dbo].[GetBranchSales](?, ?)}", "args": ["BR01", _KHMER_TITLE], "call": True, "result_sets": 2},
        {"sql": "SELECT * FROM [dbo].[fn_SalesFrom](?)", "args": ["2026-10-01"]},
        {"sql": "SELECT [dbo].[fn_BranchTotal](?) AS [value]", "args": ["BR02"]},
    ]


def test_a_filter_value_is_only_ever_a_bound_value(auth_headers, worker):
    _sqlserver(auth_headers, APPROVED)
    report_id = _report(auth_headers, "{{ info.code }}")
    _configure(auth_headers, report_id, [_BRANCH], {"type": "jdbc", "connection": "erp-mssql", "queries": [{**_CARDS[0], "more_results": None}]})
    nasty = "BR01'); DROP TABLE dbo.sales; --"
    http.post(f"/api/v1/reports/{report_id}/run", json={"parameters": {"branch": nasty}, "format": "docx"}, headers=auth_headers)
    [request] = worker["sent"]
    assert request["queries"][0]["sql"] == "{call [dbo].[GetBranchSales](?, ?)}" and request["queries"][0]["args"][0] == nasty


@pytest.mark.parametrize("card, fragment", [
    ({"routine": "dbo.NotApproved", "arguments": {}}, "isn't approved"),
    ({"routine": "dbo.fn_SalesFrom", "arguments": {}}, "give @from a filter or a fixed value"),
    ({"routine": "dbo.fn_SalesFrom", "arguments": {"@from": {"filter": "nope"}}}, "isn't defined"),
    ({"routine": "dbo.fn_SalesFrom", "arguments": {"@from": {"filter": "from_date", "value": "x"}}}, "a filter or a fixed value"),
    ({"routine": "dbo.fn_SalesFrom", "arguments": {"@from": {"filter": "from_date"}, "@extra": {"value": "1"}}}, "has no parameter @extra"),
    ({"routine": "dbo.fn_SalesFrom", "arguments": {"@from": {"filter": "from_date"}}, "more_results": [{"name": "x"}]}, "only a stored procedure"),
    ({"routine": "dbo.GetBranchSales", "arguments": {"@branch": {"filter": "from_date"}, "@title": {"value": "t"}},
      "more_results": [{"name": "rows"}]}, "Two results are named 'rows'"),
    ({"routine": "dbo.GetBranchSales; DROP TABLE x", "arguments": {}}, "choose the procedure or function"),
])
def test_a_routine_card_is_checked_when_saved(auth_headers, worker, card, fragment):
    _sqlserver(auth_headers, APPROVED)
    report_id = _register(auth_headers)
    queries = [{"name": "rows", "shape": "list", "routine": "dbo.fn_SalesFrom", "arguments": {"@from": {"filter": "from_date"}}},
               {"name": "two", "shape": "list", **card}]
    resp = _configure(auth_headers, report_id, [_FROM], {"type": "jdbc", "connection": "erp-mssql", "queries": queries}, expect=400)
    assert fragment in resp["detail"], resp


def test_a_routine_needs_a_sql_server_connection(auth_headers, worker):
    _secret(auth_headers)
    _connection(auth_headers, name="pg-db")
    report_id = _register(auth_headers)
    queries = [{"name": "rows", "shape": "list", "routine": "dbo.fn_SalesFrom", "arguments": {}}, {"name": "b", "shape": "value", "query": "SELECT 1"}]
    resp = _configure(auth_headers, report_id, [], {"type": "jdbc", "connection": "pg-db", "queries": queries}, expect=400)
    assert "SQL Server" in resp["detail"]


def test_withdrawing_approval_stops_it_on_the_next_run(auth_headers, worker):
    conn = _sqlserver(auth_headers, APPROVED)
    report_id = _report(auth_headers, "{{ other }}")
    _configure(auth_headers, report_id, [], {"type": "jdbc", "connection": "erp-mssql", "queries": [_CARDS[2]]})
    assert http.post(f"/api/v1/reports/{report_id}/run", json={"parameters": {}, "format": "docx"}, headers=auth_headers).status_code == 200
    http.put(f"/api/v1/connections/{conn['id']}/routines", json={"routines": APPROVED[:2]}, headers=auth_headers)
    run = http.post(f"/api/v1/reports/{report_id}/run", json={"parameters": {}, "format": "docx"}, headers=auth_headers)
    assert run.status_code == 502 and "no longer allows" in run.json()["detail"]
    assert len(worker["sent"]) == 1  # nothing was sent the second time


def test_an_older_worker_is_a_clear_error(auth_headers, worker):
    _sqlserver(auth_headers, APPROVED)
    report_id = _report(auth_headers, "{{ info.code }}")
    _configure(auth_headers, report_id, [_BRANCH], {"type": "jdbc", "connection": "erp-mssql", "queries": [_CARDS[0]]})
    worker["state"]["protocol"] = None  # a worker from before procedures: no protocol in its answer
    run = http.post(f"/api/v1/reports/{report_id}/run", json={"parameters": {"branch": "BR01"}, "format": "docx"}, headers=auth_headers)
    assert run.status_code == 502 and "too old to call stored procedures" in run.json()["detail"]


def test_more_names_than_result_sets_is_an_error_naming_the_query(auth_headers, worker):
    _sqlserver(auth_headers, APPROVED)
    report_id = _report(auth_headers, "{{ info.code }}")
    card = {**_CARDS[0], "more_results": [{"name": "lines", "shape": "list"}, {"name": "extra", "shape": "list"}]}
    _configure(auth_headers, report_id, [_BRANCH], {"type": "jdbc", "connection": "erp-mssql", "queries": [card]})
    run = http.post(f"/api/v1/reports/{report_id}/run", json={"parameters": {"branch": "BR01"}, "format": "docx"}, headers=auth_headers)
    assert run.status_code == 502 and "'info'" in run.json()["detail"] and "returned 2 result sets, but 3 are named" in run.json()["detail"]


def test_a_dataset_can_call_them_too(auth_headers, worker):
    _sqlserver(auth_headers, APPROVED)
    resp = http.post("/api/v1/datasets", headers=auth_headers, json={
        "name": "branch-sales", "title": "Branch sales", "parameters": [_BRANCH, _FROM],
        "data_source": {"type": "jdbc", "connection": "erp-mssql", "queries": _CARDS},
    })
    assert resp.status_code == 200, resp.text
    preview = http.post("/api/v1/datasets/preview", headers=auth_headers, json={
        "parameters": [_BRANCH, _FROM], "data_source": {"type": "jdbc", "connection": "erp-mssql", "queries": _CARDS},
        "values": {"branch": "BR01", "from_date": "2026-10-01"},
    })
    assert preview.status_code == 200, preview.text
    assert preview.json()["data"]["info"]["name"] == _KHMER_BRANCH and len(preview.json()["data"]["lines"]) == 2


# --- Discover -----------------------------------------------------------------------------------------------------------


def test_discover_lists_what_the_login_may_run(auth_headers, worker):
    conn = _sqlserver(auth_headers)
    resp = http.post(f"/api/v1/connections/{conn['id']}/routines/discover", headers=auth_headers)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["routines"] == [
        {"name": "dbo.GetBranchSales", "kind": "procedure",
         "parameters": [{"name": "@branch", "type": "nvarchar(10)"}, {"name": "@title", "type": "nvarchar(max)"}]},
        {"name": "dbo.fn_SalesFrom", "kind": "table_function", "parameters": [{"name": "@from", "type": "date"}]},
    ]  # CountSales (an OUTPUT parameter) and a name that would need quoting are left out
    assert body["skipped_output"] == ["dbo.CountSales"] and body["can_write"] is True
    # Discover approves nothing.
    assert http.get(f"/api/v1/connections/{conn['id']}", headers=auth_headers).json()["config"]["routines"] == []


def test_only_connection_managers_approve_or_discover(auth_headers, worker, make_local_user):
    conn = _sqlserver(auth_headers)
    _, author = make_local_user("author", "pw-author-1", ("ROLE_REPORT_ADMIN",))
    assert http.put(f"/api/v1/connections/{conn['id']}/routines", json={"routines": APPROVED}, headers=author).status_code == 403
    assert http.post(f"/api/v1/connections/{conn['id']}/routines/discover", headers=author).status_code == 403
