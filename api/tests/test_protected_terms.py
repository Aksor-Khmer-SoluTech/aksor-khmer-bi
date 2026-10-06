"""Protected term sets (reusable, org-scoped Khmer word-segmentation
overrides) and a report's own protected-terms-config layer -- see
specs/protected_terms_design.md.

Three things this file protects: (1) a set can only be created/edited/
deleted by someone holding `protected_terms:manage`, but merely *seeing*
what sets exist only needs that or `report:manage`; (2) a report can only
select sets from its own organization, re-checked server-side even though
a manager UI would only ever offer same-org sets; (3) the resolved terms
actually reach rendering and change segmentation output, and a stale
`set_ids` reference (a set deleted after being selected) degrades
gracefully instead of breaking the render.
"""
from datetime import datetime, timezone
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import db
from app.main import app
from app.rbac import ROOT_ORG_ID, create_organization

client = TestClient(app)

# Confirmed via aksor_khmer_ocr_segmenter.pipeline.process_text: ICU's
# default break iterator splits both of these into two syllable tokens
# (so they render with a break marker in the middle) unless they're
# injected as a protected term, in which case they stay one token.
GLOBAL_TERM = "សាកល្បងប្រាំបួន"
LOCAL_TERM = "សាកល្បងដប់មួយ"


# --- helpers ---------------------------------------------------------------


def _docx_bytes(text: str = "{{ note }}") -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _docx_text(content: bytes) -> str:
    return "\n".join(p.text for p in Document(BytesIO(content)).paragraphs)


def _register_report(headers, org_id=None) -> str:
    data = {"name": "Note"}
    if org_id is not None:
        data["org_id"] = org_id
    resp = client.post(
        "/api/v1/reports",
        files={"file": ("t.docx", _docx_bytes(), "application/octet-stream")},
        data=data,
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    return resp.json()["report_id"]


def _create_set(headers, org_id=ROOT_ORG_ID, name="Set", terms=None, exclude_terms=None, expect=200):
    resp = client.post(
        "/api/v1/protected-term-sets",
        json={"org_id": org_id, "name": name, "terms": terms or [], "exclude_terms": exclude_terms or []},
        headers=headers,
    )
    assert resp.status_code == expect, resp.text
    return resp.json() if expect < 300 else resp


@pytest.fixture
def org2():
    with db.SessionLocal() as session:
        create_organization(session, "org2", "Org Two")
        session.commit()
    return "org2"


@pytest.fixture
def report_manage_only_user(make_local_user):
    """A user with `report:manage` granted directly (not via a standard
    role -- every standard role that includes report:manage now also
    includes protected_terms:manage, so this is the only way to isolate
    "can see sets to pick one" from "can create/edit/delete sets")."""
    user_id, headers = make_local_user("reportmgr", "pw12345", role_names=())
    with db.SessionLocal() as session:
        session.add(
            db.UserPermissionGrant(
                user_id=user_id,
                permission_code="report:manage",
                granted_at=datetime.now(timezone.utc).isoformat(),
                is_active=True,
            )
        )
        session.commit()
    return user_id, headers


# --- CRUD + permission gates ------------------------------------------------


def test_create_protected_term_set_requires_permission(make_local_user):
    _, headers = make_local_user("noperm", "pw12345", role_names=())
    _create_set(headers, expect=403)


def test_create_and_get_protected_term_set(auth_headers):
    created = _create_set(auth_headers, name="Provinces", terms=[GLOBAL_TERM], exclude_terms=["foo"])
    assert created["name"] == "Provinces"
    assert created["terms"] == [GLOBAL_TERM]
    assert created["exclude_terms"] == ["foo"]
    assert created["org_id"] == ROOT_ORG_ID

    resp = client.get(f"/api/v1/protected-term-sets/{created['id']}", headers=auth_headers)
    assert resp.status_code == 200
    assert resp.json()["id"] == created["id"]


def test_create_protected_term_set_rejects_blank_name(auth_headers):
    resp = client.post(
        "/api/v1/protected-term-sets",
        json={"org_id": ROOT_ORG_ID, "name": "   ", "terms": [], "exclude_terms": []},
        headers=auth_headers,
    )
    assert resp.status_code == 400


def test_list_protected_term_sets_scoped_to_own_org(auth_headers, org2):
    _create_set(auth_headers, org_id=ROOT_ORG_ID, name="Root Set")
    _create_set(auth_headers, org_id=org2, name="Org2 Set")

    resp = client.get("/api/v1/protected-term-sets", params={"org_id": ROOT_ORG_ID}, headers=auth_headers)
    names = {row["name"] for row in resp.json()}
    assert "Root Set" in names
    assert "Org2 Set" not in names


def test_view_only_permission_can_list_and_get_but_not_write(report_manage_only_user, auth_headers):
    _, headers = report_manage_only_user
    created = _create_set(auth_headers, name="Viewable")

    resp = client.get("/api/v1/protected-term-sets", params={"org_id": ROOT_ORG_ID}, headers=headers)
    assert resp.status_code == 200
    assert any(row["id"] == created["id"] for row in resp.json())

    resp = client.get(f"/api/v1/protected-term-sets/{created['id']}", headers=headers)
    assert resp.status_code == 200

    resp = client.put(
        f"/api/v1/protected-term-sets/{created['id']}",
        json={"name": "Renamed", "terms": [], "exclude_terms": []},
        headers=headers,
    )
    assert resp.status_code == 403


def test_update_protected_term_set(auth_headers):
    created = _create_set(auth_headers, name="Original", terms=["a"])
    resp = client.put(
        f"/api/v1/protected-term-sets/{created['id']}",
        json={"name": "Renamed", "description": "d", "terms": ["b", "c"], "exclude_terms": []},
        headers=auth_headers,
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["name"] == "Renamed"
    assert body["terms"] == ["b", "c"]


def test_delete_protected_term_set(auth_headers):
    created = _create_set(auth_headers)
    resp = client.delete(f"/api/v1/protected-term-sets/{created['id']}", headers=auth_headers)
    assert resp.status_code == 204
    assert client.get(f"/api/v1/protected-term-sets/{created['id']}", headers=auth_headers).status_code == 404


def test_protected_term_set_org_isolation(auth_headers, org2, make_local_user):
    other_org_set = _create_set(auth_headers, org_id=org2, name="Org2 Only")
    _, headers = make_local_user("rootscoped", "pw12345", role_names=("ROLE_REPORT_ADMIN",), org_id=ROOT_ORG_ID)

    resp = client.get(f"/api/v1/protected-term-sets/{other_org_set['id']}", headers=headers)
    assert resp.status_code == 404


# --- a report's own protected-terms-config ----------------------------------


def test_get_protected_terms_config_defaults_to_empty(auth_headers):
    report_id = _register_report(auth_headers)
    resp = client.get(f"/api/v1/reports/{report_id}/protected-terms-config", headers=auth_headers)
    assert resp.status_code == 200
    assert resp.json() == {"set_ids": [], "terms": [], "exclude_terms": []}


def test_put_and_get_protected_terms_config(auth_headers):
    report_id = _register_report(auth_headers)
    term_set = _create_set(auth_headers, terms=[GLOBAL_TERM])
    resp = client.put(
        f"/api/v1/reports/{report_id}/protected-terms-config",
        json={"set_ids": [term_set["id"]], "terms": [LOCAL_TERM], "exclude_terms": []},
        headers=auth_headers,
    )
    assert resp.status_code == 200, resp.text
    assert resp.json() == {"set_ids": [term_set["id"]], "terms": [LOCAL_TERM], "exclude_terms": []}

    resp = client.get(f"/api/v1/reports/{report_id}/protected-terms-config", headers=auth_headers)
    assert resp.json()["set_ids"] == [term_set["id"]]


def test_protected_terms_config_rejects_nonexistent_set_id(auth_headers):
    report_id = _register_report(auth_headers)
    resp = client.put(
        f"/api/v1/reports/{report_id}/protected-terms-config",
        json={"set_ids": ["does-not-exist"], "terms": [], "exclude_terms": []},
        headers=auth_headers,
    )
    assert resp.status_code == 400


def test_protected_terms_config_rejects_cross_org_set_id(auth_headers, org2):
    report_id = _register_report(auth_headers)  # in ROOT_ORG_ID (default)
    other_org_set = _create_set(auth_headers, org_id=org2, name="Foreign")
    resp = client.put(
        f"/api/v1/reports/{report_id}/protected-terms-config",
        json={"set_ids": [other_org_set["id"]], "terms": [], "exclude_terms": []},
        headers=auth_headers,
    )
    assert resp.status_code == 400
    assert "different organization" in resp.json()["detail"]


def test_protected_terms_config_requires_report_manage_permission(make_local_user, auth_headers):
    """Gated by require_report_permission("manage") -- the same gate
    data-config uses -- not the separate protected_terms:manage
    permission (that one only governs the reusable-set library, not a
    single report's own selection of sets)."""
    _, headers = make_local_user("noreportaccess", "pw12345", role_names=())
    report_id = _register_report(auth_headers)
    resp = client.put(
        f"/api/v1/reports/{report_id}/protected-terms-config",
        json={"set_ids": [], "terms": [LOCAL_TERM], "exclude_terms": []},
        headers=headers,
    )
    assert resp.status_code == 403


# --- end-to-end render integration ------------------------------------------


def test_render_merges_global_set_and_report_own_terms(auth_headers):
    report_id = _register_report(auth_headers)
    term_set = _create_set(auth_headers, name="Global", terms=[GLOBAL_TERM])
    client.put(
        f"/api/v1/reports/{report_id}/protected-terms-config",
        json={"set_ids": [term_set["id"]], "terms": [LOCAL_TERM], "exclude_terms": []},
        headers=auth_headers,
    ).raise_for_status()

    note = f"{GLOBAL_TERM} {LOCAL_TERM}"
    resp = client.post(f"/api/v1/reports/{report_id}/render", params={"format": "docx"}, json={"note": note})
    assert resp.status_code == 200, resp.text
    rendered = _docx_text(resp.content)
    assert GLOBAL_TERM in rendered
    assert LOCAL_TERM in rendered


def test_render_without_protected_terms_config_splits_the_terms(auth_headers):
    """Sanity check for the assertions above: absent any config, these
    two made-up words are NOT expected to survive as one token -- proves
    the positive assertions above are actually exercising something."""
    report_id = _register_report(auth_headers)
    note = f"{GLOBAL_TERM} {LOCAL_TERM}"
    resp = client.post(f"/api/v1/reports/{report_id}/render", params={"format": "docx"}, json={"note": note})
    assert resp.status_code == 200, resp.text
    rendered = _docx_text(resp.content)
    assert GLOBAL_TERM not in rendered
    assert LOCAL_TERM not in rendered


def test_render_is_fail_soft_when_selected_set_is_deleted(auth_headers):
    report_id = _register_report(auth_headers)
    term_set = _create_set(auth_headers, name="Deletable", terms=[GLOBAL_TERM])
    client.put(
        f"/api/v1/reports/{report_id}/protected-terms-config",
        json={"set_ids": [term_set["id"]], "terms": [LOCAL_TERM], "exclude_terms": []},
        headers=auth_headers,
    ).raise_for_status()

    client.delete(f"/api/v1/protected-term-sets/{term_set['id']}", headers=auth_headers)

    note = f"{GLOBAL_TERM} {LOCAL_TERM}"
    resp = client.post(f"/api/v1/reports/{report_id}/render", params={"format": "docx"}, json={"note": note})
    assert resp.status_code == 200, resp.text  # never a 500 just because a referenced set vanished
    rendered = _docx_text(resp.content)
    assert LOCAL_TERM in rendered  # the report's own term is unaffected
    assert GLOBAL_TERM not in rendered  # the deleted set's term no longer applies


# --- GET /protected-term-sets/deployment-floor ------------------------------


def _sync_deployment_floor():
    from app.deployment_terms_sync import sync_deployment_terms

    with db.SessionLocal() as session:
        sync_deployment_terms(session)
        session.commit()


def test_deployment_floor_requires_authentication():
    resp = client.get("/api/v1/protected-term-sets/deployment-floor")
    assert resp.status_code == 401


def test_deployment_floor_requires_report_manage_or_protected_terms_manage(make_local_user):
    _sync_deployment_floor()
    _, headers = make_local_user("no_terms_perm", "pw12345", role_names=())
    resp = client.get("/api/v1/protected-term-sets/deployment-floor", headers=headers)
    assert resp.status_code == 403


def test_deployment_floor_visible_to_report_manage_and_protected_terms_manage(auth_headers):
    _sync_deployment_floor()
    resp = client.get("/api/v1/protected-term-sets/deployment-floor", headers=auth_headers)
    assert resp.status_code == 200
    body = resp.json()
    assert body["id"] == "default"
    assert body["version"] == 1
    assert "terms" in body and "exclude_terms" in body and "synced_at" in body


def test_deployment_floor_route_is_not_shadowed_by_set_id_route(auth_headers):
    """FastAPI matches routes in registration order -- if /deployment-floor
    were registered after /{set_id}, this request would 404 with "Protected
    term set not found" instead of returning the real, 200 deployment-floor
    shape."""
    _sync_deployment_floor()
    resp = client.get("/api/v1/protected-term-sets/deployment-floor", headers=auth_headers)
    assert resp.status_code == 200
    assert resp.json().get("detail") is None
