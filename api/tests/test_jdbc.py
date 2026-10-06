"""JDBC data sources: the SQL guard, JDBC URLs, host policy, uploaded drivers,
JDBC connections, and a report that runs a query (or serves static sample data)."""
import json
import sqlite3
import zipfile
from io import BytesIO

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select

from app import audit, db, jdbc, jdbc_drivers
from app.main import app
from app.report_data import DataConfigError, DataSourceError

http = TestClient(app)


# --- the SQL guard ---------------------------------------------------------------------


@pytest.mark.parametrize(
    "sql, names",
    [
        ("SELECT 1", []),
        ("select * from sales where d >= :from_date and d < :to_date;", ["from_date", "to_date"]),
        ("WITH t AS (SELECT 1 AS a) SELECT a FROM t WHERE a = :x OR :x = ''", ["x"]),
        ("SELECT d::date, ':notaparam', \"update\" FROM t -- :ignored\nWHERE b = :b", ["b"]),
        ("SELECT updated_at, created, deleted_flag FROM t", []),
        ("(SELECT 1) UNION (SELECT 2)", []),
    ],
)
def test_accepts_read_only_queries(sql, names):
    assert jdbc.validate_query(sql) == names


@pytest.mark.parametrize(
    "sql",
    [
        "",
        "   ",
        "DELETE FROM t",
        "UPDATE t SET a = 1",
        "SELECT 1; DROP TABLE t",
        "SELECT 1 INTO OUTFILE '/tmp/x'",
        "SELECT * FROM t FOR UPDATE",
        "WITH x AS (DELETE FROM t RETURNING *) SELECT * FROM x",
        "SELECT pg_read_file('/etc/passwd')",
        "SELECT sleep(100)",
        "SELECT /*! 1; DROP TABLE t */ 1",
        "EXEC sp_who",
        "SELECT '\\' ; DROP TABLE t --'",  # a backslash is not an escape: this reads as code
        "SELECT 'unclosed",
        "SELECT 1 /* never closed",
        "x" * (jdbc.MAX_QUERY_LENGTH + 1),
    ],
)
def test_rejects_anything_that_isnt_one_select(sql):
    with pytest.raises(DataConfigError):
        jdbc.validate_query(sql)


def test_keywords_inside_literals_and_comments_dont_count():
    sql = "SELECT 'drop table x; delete from y', /* update */ a FROM t -- insert"
    assert jdbc.validate_query(sql) == []


def test_compile_query_styles():
    values = {"a": "1", "b": "it's 100%"}
    sql = "SELECT '50%' AS pct FROM t WHERE x = :a AND y = :b AND z = :a"
    assert jdbc.compile_query(sql, "pyformat", values) == (
        "SELECT '50%%' AS pct FROM t WHERE x = %(a)s AND y = %(b)s AND z = %(a)s",
        {"a": "1", "b": "it's 100%"},
    )
    assert jdbc.compile_query(sql, "qmark", values) == (
        "SELECT '50%' AS pct FROM t WHERE x = ? AND y = ? AND z = ?",
        ["1", "it's 100%", "1"],
    )
    compiled, bound = jdbc.compile_query("SELECT :a", "named", values)
    assert compiled == "SELECT :a" and bound == {"a": "1"}


def test_compile_query_needs_every_placeholder_defined():
    with pytest.raises(DataSourceError):
        jdbc.compile_query("SELECT :missing", "qmark", {})


def test_json_safe_values():
    from datetime import date, datetime
    from decimal import Decimal

    assert jdbc.json_safe(Decimal("12")) == 12 and isinstance(jdbc.json_safe(Decimal("12")), int)
    assert jdbc.json_safe(Decimal("12.50")) == 12.5
    assert jdbc.json_safe(date(2026, 1, 2)) == "2026-01-02"
    assert jdbc.json_safe(datetime(2026, 1, 2, 3, 4)) == "2026-01-02T03:04:00"
    assert jdbc.json_safe(b"\x00\x01") == "AAE="


# --- JDBC URLs ----------------------------------------------------------------------------


@pytest.mark.parametrize(
    "config, url",
    [
        ({"engine": "postgresql", "host": "db.example.com", "port": 5432, "database": "sales", "ssl_mode": "verify-full"},
         "jdbc:postgresql://db.example.com:5432/sales?sslmode=verify-full"),
        ({"engine": "mysql", "host": "db", "port": 3306, "database": "sales", "ssl_mode": "require"},
         "jdbc:mysql://db:3306/sales?sslMode=REQUIRED"),
        ({"engine": "mariadb", "host": "db", "port": 3306, "database": "sales", "ssl_mode": "require"},
         "jdbc:mariadb://db:3306/sales?sslMode=trust"),
        ({"engine": "sqlserver", "host": "sql", "port": 1433, "database": "Sales", "ssl_mode": "verify-ca"},
         "jdbc:sqlserver://sql:1433;databaseName=Sales;encrypt=true;trustServerCertificate=false"),
        ({"engine": "db2", "host": "ibm", "port": 50000, "database": "SAMPLE", "ssl_mode": "require"},
         "jdbc:db2://ibm:50000/SAMPLE:sslConnection=true;"),
        ({"engine": "oracle", "host": "ora", "port": 1521, "database": "ORCLPDB1", "service_type": "service_name", "ssl_mode": "disable"},
         "jdbc:oracle:thin:@//ora:1521/ORCLPDB1"),
        ({"engine": "oracle", "host": "ora", "port": 1521, "database": "ORCL", "service_type": "sid", "ssl_mode": "disable"},
         "jdbc:oracle:thin:@ora:1521:ORCL"),
    ],
)
def test_build_then_parse_a_jdbc_url(config, url):
    assert jdbc.build_jdbc_url(config) == url
    parsed = jdbc.parse_jdbc_url(url)
    assert (parsed["engine"], parsed["host"], parsed["port"], parsed["database"]) == (
        config["engine"], config["host"], config["port"], config["database"]
    )


def test_oracle_over_tcps_round_trips():
    config = {"engine": "oracle", "host": "ora", "port": 2484, "database": "SVC", "service_type": "service_name", "ssl_mode": "require"}
    url = jdbc.build_jdbc_url(config)
    assert "PROTOCOL=tcps" in url
    parsed = jdbc.parse_jdbc_url(url)
    assert (parsed["host"], parsed["port"], parsed["database"], parsed["service_type"]) == ("ora", 2484, "SVC", "service_name")


@pytest.mark.parametrize(
    "url",
    ["postgresql://x/y", "jdbc:sqlite:/tmp/x.db", "jdbc:postgresql://user:pw@host/db", "jdbc:postgresql:host/db", "jdbc:oracle:oci:@x"],
)
def test_bad_jdbc_urls_are_refused(url):
    with pytest.raises(DataConfigError):
        jdbc.parse_jdbc_url(url)


def test_default_port_when_the_url_has_none():
    assert jdbc.parse_jdbc_url("jdbc:mysql://db/sales")["port"] == 3306


# --- where a connection may point -------------------------------------------------------------


@pytest.mark.parametrize("host", ["169.254.169.254", "metadata.google.internal", "0.0.0.0", "fe80::1", "", "bad host", "a" * 300])
def test_blocked_hosts(host):
    with pytest.raises(DataConfigError):
        jdbc.check_host(host)


@pytest.mark.parametrize("host", ["db.example.com", "10.0.0.5", "localhost", "127.0.0.1", "sql-1.internal.corp"])
def test_ordinary_hosts_pass(host):
    assert jdbc.check_host(host) == host


def test_allow_list_restricts_hosts(monkeypatch):
    monkeypatch.setenv("JDBC_ALLOWED_HOSTS", "*.corp.example.com, 10.1.0.0/16, exact.host")
    for ok in ("db.corp.example.com", "10.1.2.3", "exact.host"):
        assert jdbc.check_host(ok) == ok
    for bad in ("db.example.com", "10.2.0.1", "other.host"):
        with pytest.raises(DataConfigError):
            jdbc.check_host(bad)


def test_a_name_resolving_to_a_blocked_address_is_refused_at_run_time(monkeypatch):
    monkeypatch.setattr(jdbc.socket, "getaddrinfo", lambda *a, **k: [(2, 1, 6, "", ("169.254.169.254", 0))])
    with pytest.raises(DataSourceError):
        jdbc.assert_resolves_safely("looks-fine.example.com")


# --- helpers for the API tests ------------------------------------------------------------------


def _secret(headers, name="sales-db-password", value="s3cret-pw"):
    resp = http.post("/api/v1/secrets", json={"name": name, "value": value}, headers=headers)
    assert resp.status_code == 200, resp.text
    return resp.json()


def _jar(klass="org.example.Driver", entries=None) -> bytes:
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w") as jar:
        jar.writestr("META-INF/MANIFEST.MF", "Manifest-Version: 1.0\n")
        jar.writestr(klass.replace(".", "/") + ".class", b"\xca\xfe\xba\xbe")
        for entry in entries or []:
            jar.writestr(entry, b"x")
    return buf.getvalue()


@pytest.fixture(autouse=True)
def driver_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(jdbc_drivers, "DRIVER_DIR", tmp_path / "jdbc_drivers")
    return tmp_path / "jdbc_drivers"


def _upload(headers, *, name="Oracle 23", engine="oracle", klass="oracle.jdbc.OracleDriver", content=None, expect=200):
    resp = http.post(
        "/api/v1/jdbc/drivers",
        data={"name": name, "engine": engine, "driver_class": klass},
        files={"file": ("ojdbc11.jar", content if content is not None else _jar(klass), "application/java-archive")},
        headers=headers,
    )
    assert resp.status_code == expect, resp.text
    return resp.json()


def _pg(**overrides):
    return {
        "engine": "postgresql", "host": "127.0.0.1", "port": 5432, "database": "sales", "ssl_mode": "disable",
        "auth": {"type": "basic", "username": "report_ro", "password_secret": "sales-db-password"},
        **overrides,
    }


def _connection(headers, name="sales-db", config=None, expect=200):
    resp = http.post("/api/v1/connections", json={"name": name, "kind": "jdbc", "config": config or _pg()}, headers=headers)
    assert resp.status_code == expect, resp.text
    return resp.json()


@pytest.fixture
def sqlite_as_postgres(tmp_path, monkeypatch):
    """Runs the 'built-in PostgreSQL driver' against a SQLite file, so the whole
    path (validate -> resolve credentials -> compile -> execute -> JSON) is tested
    without a database server."""
    path = tmp_path / "sales.sqlite"
    conn = sqlite3.connect(path)
    conn.execute("CREATE TABLE sales (branch TEXT, name TEXT, total NUMERIC)")
    conn.executemany("INSERT INTO sales VALUES (?, ?, ?)", [("pp", "Phnom Penh", 42), ("sr", "Siem Reap", 7.5)])
    conn.commit()
    conn.close()
    seen = {}

    def connect(config, password):
        seen.update(password=password, config=config)
        return sqlite3.connect(path)

    monkeypatch.setitem(jdbc._NATIVE, "postgresql", (connect, "named"))
    return seen


# --- drivers --------------------------------------------------------------------------------------


def test_engines_list(auth_headers):
    engines = http.get("/api/v1/jdbc/engines", headers=auth_headers).json()
    assert [e["id"] for e in engines] == ["oracle", "postgresql", "mysql", "sqlserver", "mariadb", "db2"]
    built_in = {e["id"] for e in engines if e["built_in"]}
    assert built_in == {"postgresql", "mysql", "mariadb"}


def test_upload_list_download_delete_a_driver(auth_headers, driver_dir):
    content = _jar("oracle.jdbc.OracleDriver")
    created = _upload(auth_headers, content=content)
    assert created["engine"] == "oracle" and created["size_bytes"] == len(content) and len(created["sha256"]) == 64
    assert (driver_dir / f"{created['id']}.jar").read_bytes() == content

    assert [d["name"] for d in http.get("/api/v1/jdbc/drivers", headers=auth_headers).json()] == ["Oracle 23"]
    downloaded = http.get(f"/api/v1/jdbc/drivers/{created['id']}/download", headers=auth_headers)
    assert downloaded.status_code == 200 and downloaded.content == content

    assert http.delete(f"/api/v1/jdbc/drivers/{created['id']}", headers=auth_headers).status_code == 204
    assert not (driver_dir / f"{created['id']}.jar").exists()
    assert http.get("/api/v1/jdbc/drivers", headers=auth_headers).json() == []


def test_a_bad_upload_leaves_nothing_behind(auth_headers, driver_dir):
    _upload(auth_headers, content=b"not a zip", expect=400)
    _upload(auth_headers, content=_jar("some.other.Driver"), expect=400)  # class not in the jar
    _upload(auth_headers, content=b"", expect=400)
    _upload(auth_headers, engine="nosuch", expect=400)
    _upload(auth_headers, klass="not a class", expect=400)
    assert http.get("/api/v1/jdbc/drivers", headers=auth_headers).json() == []
    assert not list(driver_dir.glob("*")) if driver_dir.exists() else True


def test_upload_size_cap(auth_headers, monkeypatch):
    monkeypatch.setattr(jdbc_drivers, "MAX_BYTES", 100)
    _upload(auth_headers, content=_jar("oracle.jdbc.OracleDriver", entries=[f"pad{i}" for i in range(50)]), expect=400)


def test_driver_names_are_unique_per_organization(auth_headers):
    _upload(auth_headers)
    _upload(auth_headers, expect=409)


def test_uploading_needs_driver_manage(auth_headers, make_local_user):
    _, viewer = make_local_user("reporter", "pw-reporter-1", ("ROLE_REPORT_ADMIN",))
    _upload(viewer, expect=403)
    assert http.get("/api/v1/jdbc/drivers", headers=viewer).status_code == 403

    _, org_admin = make_local_user("orgadmin", "pw-orgadmin-1", ("ROLE_ORG_ADMIN",))
    created = _upload(org_admin)
    # connection:manage alone can list (to pick one) but not download or delete
    assert http.get(f"/api/v1/jdbc/drivers/{created['id']}/download", headers=viewer).status_code == 403
    assert http.delete(f"/api/v1/jdbc/drivers/{created['id']}", headers=viewer).status_code == 403


def test_upload_is_audited_with_its_hash(auth_headers):
    created = _upload(auth_headers)
    with db.SessionLocal() as session:
        event = session.scalars(select(db.AuditEvent).where(db.AuditEvent.action == "jdbc_driver.upload")).one()
    assert created["sha256"] in str(event.details)


# --- connections ------------------------------------------------------------------------------------


def test_create_a_postgres_connection_without_exposing_the_password(auth_headers):
    _secret(auth_headers)
    created = _connection(auth_headers)
    assert created["kind"] == "jdbc" and created["config"]["host"] == "127.0.0.1"
    assert created["config"]["auth"]["password_secret"] == "sales-db-password"
    assert "s3cret-pw" not in http.get(f"/api/v1/connections/{created['id']}", headers=auth_headers).text

    listed = http.get("/api/v1/connections", headers=auth_headers).json()
    assert listed[0]["engine"] == "postgresql"
    assert listed[0]["base_url"] == "jdbc:postgresql://127.0.0.1:5432/sales?sslmode=disable"
    assert listed[0]["credential"] == "secret" and listed[0]["secret_name"] == "sales-db-password"
    with db.SessionLocal() as session:
        assert "s3cret-pw" not in str(session.scalars(select(db.DataConnection)).one().config)


def test_connection_validation(auth_headers):
    _secret(auth_headers)
    assert "driver" in _connection(auth_headers, config=_pg(engine="oracle", database="X"), expect=400)["detail"].lower()
    assert _connection(auth_headers, config=_pg(port=70000), expect=422)
    for bad in (
        _pg(host="169.254.169.254"),
        _pg(database=""),
        _pg(auth={"type": "basic", "username": "", "password_secret": "sales-db-password"}),
        _pg(auth={"type": "basic", "username": "u", "password_secret": "no-such-secret"}),
        _pg(auth={"type": "basic", "username": "u"}),
        _pg(jdbc_url="jdbc:postgresql://other/db"),  # custom URLs need an uploaded driver
    ):
        assert _connection(auth_headers, config=bad, expect=400)


def test_a_connection_can_use_an_uploaded_driver_for_its_engine_only(auth_headers):
    _secret(auth_headers)
    oracle = _upload(auth_headers)
    config = _pg(engine="oracle", port=1521, database="ORCLPDB1", service_type="service_name", driver_id=oracle["id"])
    created = _connection(auth_headers, name="ora-db", config=config)
    assert created["config"]["driver_id"] == oracle["id"]

    mismatched = _connection(auth_headers, name="pg-with-oracle-jar", config=_pg(driver_id=oracle["id"]), expect=400)
    assert "Oracle" in mismatched["detail"]
    assert "doesn't exist" in _connection(auth_headers, name="x-db", config=_pg(driver_id="nope"), expect=400)["detail"]

    # the driver can't be deleted while a connection uses it
    assert http.delete(f"/api/v1/jdbc/drivers/{oracle['id']}", headers=auth_headers).status_code == 409


def test_custom_jdbc_url_decides_the_host(auth_headers):
    _secret(auth_headers)
    oracle = _upload(auth_headers)
    config = _pg(engine="oracle", port=1521, database="X", driver_id=oracle["id"], jdbc_url="jdbc:oracle:thin:@//169.254.169.254:1521/X")
    assert "link-local" in _connection(auth_headers, config=config, expect=400)["detail"]

    config["jdbc_url"] = "jdbc:oracle:thin:@//ora.example.com:1522/SVC"
    created = _connection(auth_headers, config=config)
    assert created["config"]["host"] == "ora.example.com" and created["config"]["port"] == 1522


def test_cannot_delete_a_secret_a_connection_uses(auth_headers):
    secret = _secret(auth_headers)
    _connection(auth_headers)
    assert http.delete(f"/api/v1/secrets/{secret['id']}", headers=auth_headers).status_code == 409
    used = http.get(f"/api/v1/secrets/{secret['id']}", headers=auth_headers).json()["used_by"]
    assert used[0]["kind"] == "connection"


def test_connection_test_success_and_failure(auth_headers, sqlite_as_postgres, monkeypatch):
    _secret(auth_headers)
    ok = http.post("/api/v1/connections/test", json={"config": _pg()}, headers=auth_headers).json()
    assert ok["ok"] is True and ok["message"] == "Connected"
    assert sqlite_as_postgres["password"] == "s3cret-pw"  # the Secret's value, resolved server-side

    def refuse(config, password):
        raise RuntimeError("FATAL: password authentication failed for user report_ro at 10.9.9.9")

    monkeypatch.setitem(jdbc._NATIVE, "postgresql", (refuse, "named"))
    failed = http.post("/api/v1/connections/test", json={"config": _pg()}, headers=auth_headers).json()
    assert failed["ok"] is False
    assert "10.9.9.9" not in failed["message"] and "report_ro" not in failed["message"]  # driver text stays in the log


def test_connection_test_needs_connection_manage(make_local_user):
    _, viewer = make_local_user("reporter", "pw-reporter-1", ("ROLE_REPORT_VIEWER",))
    assert http.post("/api/v1/connections/test", json={"config": _pg()}, headers=viewer).status_code == 403


# --- reports ----------------------------------------------------------------------------------------


def _docx() -> bytes:
    doc = Document()
    doc.add_paragraph("{{ rows[0].name }} = {{ rows[0].total }}")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _register(headers) -> str:
    resp = http.post("/api/v1/reports", files={"file": ("t.docx", _docx(), "application/octet-stream")}, data={"name": "Sales"}, headers=headers)
    assert resp.status_code == 200, resp.text
    return resp.json()["report_id"]


def _configure(headers, report_id, parameters, data_source, expect=200):
    resp = http.put(f"/api/v1/reports/{report_id}/data-config", json={"parameters": parameters, "data_source": data_source}, headers=headers)
    assert resp.status_code == expect, resp.text
    return resp.json()


def _text(content: bytes) -> str:
    return "\n".join(p.text for p in Document(BytesIO(content)).paragraphs).replace("​", "")


_BRANCH = {"name": "branch", "label": "Branch", "type": "text", "required": True}
_QUERY = "SELECT name, total FROM sales WHERE branch = :branch"


def test_a_report_runs_a_query_with_bound_parameters(auth_headers, sqlite_as_postgres):
    _secret(auth_headers)
    _connection(auth_headers)
    report_id = _register(auth_headers)
    saved = _configure(auth_headers, report_id, [_BRANCH], {"type": "jdbc", "connection": "sales-db", "query": _QUERY})
    assert {k: saved["data_source"][k] for k in ("type", "connection", "query", "root_key")} == {
        "type": "jdbc", "connection": "sales-db", "query": _QUERY, "root_key": None,
    }

    run = http.post(f"/api/v1/reports/{report_id}/run", json={"parameters": {"branch": "sr"}, "format": "docx"}, headers=auth_headers)
    assert run.status_code == 200, run.text
    assert _text(run.content) == "Siem Reap = 7.5"



def _resolved():
    return {"engine": "postgresql", "host": "127.0.0.1", "port": 5432, "database": "x", "ssl_mode": "disable",
            "username": "u", "password": "p", "driver": None, "jdbc_url": None}


def test_a_filter_value_that_looks_like_sql_is_just_a_value(sqlite_as_postgres):
    assert jdbc.run_query(_resolved(), _QUERY, {"branch": "x' OR '1'='1"}, None) == {"rows": []}
    assert jdbc.run_query(_resolved(), _QUERY, {"branch": "pp"}, None) == {"rows": [{"name": "Phnom Penh", "total": 42}]}


@pytest.mark.parametrize(
    "driver_error, expected",
    [
        ("connection timeout expired", "didn't answer in time"),
        ("connection to server at \"10.255.255.45\", port 6122 failed: Connection refused", "isn't accepting connections"),
        ("could not translate host name \"db.internal\" to address", "address can't be found"),
        ("FATAL: password authentication failed for user report_ro", "refused the saved login"),
        ("something odd", "Couldn't connect"),
    ],
)
def test_a_failed_connect_tells_the_user_why_without_naming_the_address(auth_headers, monkeypatch, caplog, driver_error, expected):
    def fail(config, password):
        raise RuntimeError(driver_error + " 10.255.255.45:6122")

    monkeypatch.setitem(jdbc._NATIVE, "postgresql", (fail, "named"))
    conn = {**_resolved(), "host": "10.255.255.45", "port": 6122}
    with caplog.at_level("INFO"), pytest.raises(jdbc.DataSourceError) as caught:
        jdbc.run_query(conn, _QUERY, {"branch": "x"}, None)
    message = str(caught.value)
    assert expected in message
    assert "10.255.255.45" not in message and "6122" not in message
    assert "10.255.255.45" not in caplog.text and "6122" not in caplog.text  # the log names it by alias
    assert jdbc.database_alias(conn) in caplog.text


def test_every_run_is_audited_with_how_many_parameters_were_chosen_and_how_it_ended(auth_headers, sqlite_as_postgres, monkeypatch):
    _secret(auth_headers)
    _connection(auth_headers)
    report_id = _register(auth_headers)
    _configure(auth_headers, report_id, [_BRANCH], {"type": "jdbc", "connection": "sales-db", "query": _QUERY})

    def events(action):
        with db.SessionLocal() as session:
            query = select(db.AuditEvent).where(db.AuditEvent.entity_id == report_id, db.AuditEvent.action == action)
            return [audit.event_to_dict(r) for r in session.execute(query).scalars()]

    ok = http.post(f"/api/v1/reports/{report_id}/run", json={"parameters": {"branch": "sr"}, "format": "docx"}, headers=auth_headers)
    assert ok.status_code == 200
    (row,) = events("report.run")
    assert row["actor_username"] and row["created_at"]
    assert row["details"]["parameters_selected"] == 1 and row["details"]["parameter_names"] == ["branch"]
    assert row["details"]["outcome"] == "success" and row["details"]["format"] == "docx"
    assert "sr" not in json.dumps(row["details"]["parameter_names"])  # names only, never the values

    def refuse(config, password):
        raise RuntimeError("connection timeout expired")

    monkeypatch.setitem(jdbc._NATIVE, "postgresql", (refuse, "named"))
    bad = http.post(f"/api/v1/reports/{report_id}/run", json={"parameters": {"branch": "sr"}, "format": "docx"}, headers=auth_headers)
    assert bad.status_code == 502 and "10." not in bad.text
    failed = events("report.run_failed")
    assert len(failed) == 1 and failed[0]["details"]["outcome"] == "data_source_failed"
    assert failed[0]["details"]["parameters_selected"] == 1


def test_jdbc_source_validation(auth_headers):
    _secret(auth_headers)
    _connection(auth_headers)
    report_id = _register(auth_headers)
    source = {"type": "jdbc", "connection": "sales-db", "query": _QUERY}
    for bad, fragment in (
        ({**source, "connection": None}, "needs a connection"),
        ({**source, "connection": "nope-db"}, "doesn't exist"),
        ({**source, "query": "DELETE FROM sales"}, "SELECT"),
        ({**source, "query": "SELECT 1 WHERE a = :undefined_filter"}, "undefined_filter"),
        ({**source, "root_key": "not valid"}, "result name"),
    ):
        detail = _configure(auth_headers, report_id, [_BRANCH], bad, expect=400)["detail"]
        assert fragment in detail, detail


def test_source_type_and_connection_kind_must_match(auth_headers):
    _secret(auth_headers)
    _connection(auth_headers)
    http.post("/api/v1/connections", json={"name": "erp-api", "config": {"base_url": "http://erp.test"}}, headers=auth_headers)
    report_id = _register(auth_headers)
    assert "isn't a database connection" in _configure(
        auth_headers, report_id, [_BRANCH], {"type": "jdbc", "connection": "erp-api", "query": _QUERY}, expect=400
    )["detail"]
    assert "database connection" in _configure(
        auth_headers, report_id, [], {"type": "rest", "connection": "sales-db", "url": "/x"}, expect=400
    )["detail"]


def test_root_key_names_what_the_template_loops_over(sqlite_as_postgres):
    assert jdbc.run_query(_resolved(), _QUERY, {"branch": "sr"}, "sales") == {"sales": [{"name": "Siem Reap", "total": 7.5}]}


def test_too_many_rows_is_an_error_not_a_silent_cut(sqlite_as_postgres, monkeypatch):
    monkeypatch.setattr(jdbc, "MAX_ROWS", 1)
    conn = {"engine": "postgresql", "host": "127.0.0.1", "port": 5432, "database": "x", "ssl_mode": "disable",
            "username": "u", "password": "p", "driver": None, "jdbc_url": None}
    with pytest.raises(DataSourceError, match="more than 1 rows"):
        jdbc.run_query(conn, "SELECT name FROM sales", {}, None)


def test_database_errors_stay_out_of_the_message(sqlite_as_postgres):
    conn = {"engine": "postgresql", "host": "127.0.0.1", "port": 5432, "database": "x", "ssl_mode": "disable",
            "username": "u", "password": "p", "driver": None, "jdbc_url": None}
    with pytest.raises(DataSourceError) as caught:
        jdbc.run_query(conn, "SELECT nonexistent_column FROM sales", {}, None)
    assert "nonexistent_column" not in str(caught.value)


def test_static_sample_data_source(auth_headers):
    report_id = _register(auth_headers)
    sample = {"rows": [{"name": "Demo branch", "total": 1234}]}
    saved = _configure(auth_headers, report_id, [], {"type": "static", "static_data": sample})
    assert (saved["data_source"]["type"], saved["data_source"]["static_data"]) == ("static", sample)
    run = http.post(f"/api/v1/reports/{report_id}/run", json={"parameters": {}, "format": "docx"}, headers=auth_headers)
    assert run.status_code == 200, run.text
    assert _text(run.content) == "Demo branch = 1234"

    assert "JSON object" in _configure(auth_headers, report_id, [], {"type": "static", "static_data": None}, expect=400)["detail"]
    too_big = {"rows": ["x" * 1000] * 2000}
    assert "too large" in _configure(auth_headers, report_id, [], {"type": "static", "static_data": too_big}, expect=400)["detail"]


# --- the driver service -------------------------------------------------------------------------------


def test_uploaded_driver_queries_go_to_the_worker(auth_headers, monkeypatch):
    monkeypatch.setenv("JDBC_WORKER_URL", "http://jdbc-worker:9000")
    monkeypatch.setenv("JDBC_WORKER_TOKEN", "worker-token")
    _secret(auth_headers)
    oracle = _upload(auth_headers)
    _connection(
        auth_headers, name="ora-db",
        config=_pg(engine="oracle", port=1521, database="ORCLPDB1", service_type="service_name", driver_id=oracle["id"]),
    )
    report_id = _register(auth_headers)
    _configure(
        auth_headers, report_id, [_BRANCH],
        {"type": "jdbc", "connection": "ora-db", "query": "SELECT name, total FROM sales WHERE branch = :branch AND :branch <> 'x'"},
    )

    sent = {}

    def fake_post(url, json, headers, timeout):
        sent.update(url=url, json=json, headers=headers)
        return httpx.Response(200, json={"columns": ["NAME", "TOTAL"], "rows": [["Phnom Penh", 42]]})

    monkeypatch.setattr(jdbc.httpx, "post", fake_post)
    run = http.post(f"/api/v1/reports/{report_id}/run", json={"parameters": {"branch": "pp"}, "format": "docx"}, headers=auth_headers)
    assert run.status_code == 200, run.text

    assert sent["url"] == "http://jdbc-worker:9000/query" and sent["headers"] == {"X-Worker-Token": "worker-token"}
    body = sent["json"]
    assert body["sql"] == "SELECT name, total FROM sales WHERE branch = ? AND ? <> 'x'" and body["args"] == ["pp", "pp"]
    assert body["jdbc_url"] == "jdbc:oracle:thin:@//127.0.0.1:1521/ORCLPDB1"
    assert body["driver_class"] == "oracle.jdbc.OracleDriver" and body["driver_file"] == f"{oracle['id']}.jar"
    assert body["driver_sha256"] == oracle["sha256"]
    assert (body["username"], body["password"]) == ("report_ro", "s3cret-pw")


def test_worker_not_configured_is_a_clear_error(monkeypatch):
    monkeypatch.delenv("JDBC_WORKER_URL", raising=False)
    conn = {"engine": "oracle", "host": "127.0.0.1", "port": 1521, "database": "X", "service_type": "sid", "ssl_mode": "disable",
            "username": "u", "password": "p", "jdbc_url": None,
            "driver": {"file": "a.jar", "sha256": "0" * 64, "class_name": "oracle.jdbc.OracleDriver"}}
    with pytest.raises(jdbc.WorkerNotConfigured, match="isn't set up"):
        jdbc.run_query(conn, "SELECT 1 FROM DUAL", {}, None)


def test_worker_failures_map_to_safe_messages(monkeypatch):
    monkeypatch.setenv("JDBC_WORKER_URL", "http://w")
    monkeypatch.setenv("JDBC_WORKER_TOKEN", "t")
    conn = {"engine": "oracle", "host": "127.0.0.1", "port": 1521, "database": "X", "service_type": "sid", "ssl_mode": "disable",
            "username": "u", "password": "p", "jdbc_url": None,
            "driver": {"file": "a.jar", "sha256": "0" * 64, "class_name": "oracle.jdbc.OracleDriver"}}
    for kind, text in (("connect", "Couldn't connect"), ("too_many_rows", "more than"), ("query", "rejected the query")):
        monkeypatch.setattr(
            jdbc.httpx, "post", lambda *a, _kind=kind, **k: httpx.Response(500, json={"kind": _kind, "detail": "ORA-12345 host=10.1.1.1"})
        )
        with pytest.raises(DataSourceError, match=text) as caught:
            jdbc.run_query(conn, "SELECT 1 FROM DUAL", {}, None)
        assert "ORA-12345" not in str(caught.value) and "10.1.1.1" not in str(caught.value)


# --- regressions from the bug hunt ---------------------------------------------------------------


@pytest.mark.parametrize(
    "sql",
    [
        "SELECT 1 # 1; DROP TABLE t",  # `#` is an operator in PostgreSQL, not a comment
        "SELECT 1--1; DROP TABLE t",  # in MySQL `--` needs a space after it to be a comment
        "SELECT 1 /*M!100100 ; DROP TABLE t */",  # MariaDB executable comment
    ],
)
def test_comment_syntax_that_engines_disagree_on_cant_hide_a_second_statement(sql):
    with pytest.raises(DataConfigError):
        jdbc.validate_query(sql)


def test_real_comments_and_operators_still_pass():
    assert jdbc.validate_query("SELECT a #> '{x}' FROM t -- trailing; drop\n WHERE b = :b") == ["b"]


def _oracle_with_url(auth_headers, url, expect):
    _secret(auth_headers)
    oracle = _upload(auth_headers)
    config = _pg(engine="oracle", port=1521, database="X", driver_id=oracle["id"], jdbc_url=url)
    return _connection(auth_headers, name=f"ora-{abs(hash(url)) % 9999}", config=config, expect=expect)


def test_every_address_in_a_custom_url_is_host_checked(auth_headers):
    descriptor = (
        "jdbc:oracle:thin:@(DESCRIPTION=(ADDRESS=(PROTOCOL=tcp)(HOST=169.254.169.254)(PORT=80))"
        "(ADDRESS=(PROTOCOL=tcp)(HOST=good.example.com)(PORT=1521))(CONNECT_DATA=(SERVICE_NAME=x)))"
    )
    assert "link-local" in _oracle_with_url(auth_headers, descriptor, 400)["detail"]


@pytest.mark.parametrize(
    "option", ["socketFactory=a.B", "autoDeserialize=true", "allowLoadLocalInfile=true", "password=hunter2", "loggerFile=/tmp/x"]
)
def test_dangerous_url_options_are_refused(auth_headers, option):
    assert "isn't allowed" in _oracle_with_url(auth_headers, f"jdbc:oracle:thin:@//ora.example.com:1521/x?{option}", 400)["detail"]


def test_a_choice_list_cant_use_a_database_connection(auth_headers):
    _secret(auth_headers)
    _connection(auth_headers)
    report_id = _register(auth_headers)
    parameter = {
        "name": "branch", "label": "Branch", "type": "text", "required": True,
        "options_source": {"connection": "sales-db", "url": "/x", "method": "GET", "value_field": "code"},
    }
    assert "database connection" in _configure(auth_headers, report_id, [parameter], None, expect=400)["detail"]


def test_sample_data_changes_are_audited_as_one_opaque_entry(auth_headers):
    report_id = _register(auth_headers)
    _configure(auth_headers, report_id, [], {"type": "static", "static_data": {f"k{i}": i for i in range(200)}})
    _configure(auth_headers, report_id, [], {"type": "static", "static_data": {f"k{i}": i + 1 for i in range(200)}})
    with db.SessionLocal() as session:
        events = session.scalars(select(db.AuditEvent).where(db.AuditEvent.action == "report.data_config_update").order_by(db.AuditEvent.created_at)).all()
    fields = [c["field"] for c in events[-1].changes]
    assert fields == ["data_source.static_data"] and events[-1].changes[0]["opaque"] is True


def test_org_admin_of_an_existing_organization_gets_new_permissions_on_startup():
    from app.rbac import create_organization, seed_defaults

    with db.SessionLocal() as session:
        org = create_organization(session, "acme", "Acme")
        role = session.scalar(select(db.Role).where(db.Role.org_id == "acme", db.Role.name == "ROLE_ORG_ADMIN"))
        session.query(db.RolePermission).filter_by(role_id=role.id, permission_code="driver:manage").delete()
        session.commit()
        seed_defaults(session)
        session.commit()
        granted = set(session.scalars(select(db.RolePermission.permission_code).where(db.RolePermission.role_id == role.id)))
    assert "driver:manage" in granted


# --- typed binding ---------------------------------------------------------------------------------


def test_typed_values_turns_each_filter_into_its_declared_type():
    import datetime as dt
    from decimal import Decimal

    out = jdbc.typed_values(
        {"d": "2026-09-30", "t": "08:30", "dt": "2026-09-30T08:15", "n": "12", "x": "12.50", "s": "12", "blank": ""},
        {"d": "date", "t": "time", "dt": "datetime", "n": "number", "x": "number", "s": "text", "blank": "text"},
    )
    assert out == {
        "d": dt.date(2026, 9, 30), "t": dt.time(8, 30), "dt": dt.datetime(2026, 9, 30, 8, 15),
        "n": 12, "x": Decimal("12.50"), "s": "12", "blank": "",
    }
    assert type(out["n"]) is int and out["x"].as_tuple().exponent == -2  # exact, not a float


@pytest.mark.parametrize("kind, value", [("date", "2026-13-45"), ("number", "nan"), ("number", "inf"), ("time", "25:00"), ("datetime", "soon")])
def test_typed_values_refuses_what_does_not_convert(kind, value):
    with pytest.raises(jdbc.DataSourceError, match="isn't a valid"):
        jdbc.typed_values({"f": value}, {"f": kind})


def _recording_driver(monkeypatch):
    seen = {}

    class Cursor:
        description = [("a",)]

        def execute(self, sql, bound):
            seen["bound"] = bound

        def fetchmany(self, n):
            return [(1,)]

    class Conn:
        def cursor(self):
            return Cursor()

        def close(self):
            pass

    monkeypatch.setitem(jdbc._NATIVE, "postgresql", (lambda config, password: Conn(), "pyformat"))
    return seen


def test_a_run_binds_real_types_only_when_the_source_asks_for_it(auth_headers, monkeypatch):
    import datetime as dt

    seen = _recording_driver(monkeypatch)
    _secret(auth_headers)
    _connection(auth_headers)
    report_id = _register(auth_headers)
    params = [{"name": "from", "label": "From", "type": "date", "required": True}]
    query = "SELECT a FROM t WHERE d >= :from"

    _configure(auth_headers, report_id, params, {"type": "jdbc", "connection": "sales-db", "query": query})
    http.post(f"/api/v1/reports/{report_id}/run", json={"parameters": {"from": "2026-09-30"}, "format": "docx"}, headers=auth_headers)
    assert seen["bound"] == {"from": "2026-09-30"}  # default: text, so an existing CAST / TO_DATE keeps working

    saved = _configure(auth_headers, report_id, params, {"type": "jdbc", "connection": "sales-db", "query": query, "typed_binding": True})
    assert saved["data_source"]["typed_binding"] is True
    http.post(f"/api/v1/reports/{report_id}/run", json={"parameters": {"from": "2026-09-30"}, "format": "docx"}, headers=auth_headers)
    assert seen["bound"] == {"from": dt.date(2026, 9, 30)}
