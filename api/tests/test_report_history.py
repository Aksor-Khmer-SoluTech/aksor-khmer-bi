"""Template file history: every upload is kept, the original (and any
earlier version) can be downloaded again, and the changelog says what
changed between versions. See report_store.py's versioning, GET
/reports/{id}/file and /changelog, and app/db/report_versions.py.
"""
import shutil
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select

from app import db, report_store
from app.main import app

client = TestClient(app)


def _docx(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _register(headers, text="Hello {{ name }} from {{ region }}", name="Quarterly Notice", filename="q3-draft.docx", note=None):
    data = {"name": name}
    if note:
        data["note"] = note
    resp = client.post(
        "/api/v1/reports", files={"file": (filename, _docx(text), "application/octet-stream")}, data=data, headers=headers
    )
    assert resp.status_code == 200, resp.text
    return resp.json()["report_id"], _docx(text)


def _replace(headers, report_id, text, note=None, filename="q3-final.docx"):
    resp = client.put(
        f"/api/v1/reports/{report_id}/file",
        files={"file": (filename, _docx(text), "application/octet-stream")},
        data={"note": note} if note else {},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


def _download(headers, report_id, version=None):
    params = {"version": version} if version is not None else {}
    return client.get(f"/api/v1/reports/{report_id}/file", params=params, headers=headers)


def _audit_actions(report_id):
    with db.SessionLocal() as session:
        rows = session.execute(
            select(db.AuditEvent).where(db.AuditEvent.entity_id == report_id).order_by(db.AuditEvent.created_at)
        ).scalars()
        return [r.action for r in rows]


# --- what gets kept -----------------------------------------------------------


def test_upload_records_version_one_with_fields_and_checksum(auth_headers):
    report_id, _ = _register(auth_headers, note="First cut for review")

    versions = report_store.list_versions(report_id)
    assert len(versions) == 1
    v1 = versions[0]
    assert v1["version"] == 1
    assert v1["created_by"] == "testadmin"
    assert v1["original_filename"] == "q3-draft.docx"
    assert v1["note"] == "First cut for review"
    assert v1["fields"] == ["name", "region"]
    assert len(v1["sha256"]) == 64
    assert v1["backfilled"] is False


def test_original_survives_a_replace_and_downloads_byte_for_byte(auth_headers):
    """The whole point: replacing the file used to destroy the original."""
    report_id, original_bytes = _register(auth_headers)
    _replace(auth_headers, report_id, "Hello {{ name }}, now with {{ branch }}")

    original = _download(auth_headers, report_id, version=1)
    assert original.status_code == 200
    # Compare parsed content, not raw bytes: python-docx stamps a created/modified
    # time into every save, so two saves of "the same" document differ. What
    # was stored is what was uploaded, which the sha256 check below pins exactly.
    assert Document(BytesIO(original.content)).paragraphs[0].text == "Hello {{ name }} from {{ region }}"
    stored = report_store.list_versions(report_id)[-1]
    import hashlib

    assert hashlib.sha256(original.content).hexdigest() == stored["sha256"]

    current = _download(auth_headers, report_id)
    assert Document(BytesIO(current.content)).paragraphs[0].text == "Hello {{ name }}, now with {{ branch }}"


def test_download_headers_name_the_version_and_type(auth_headers):
    report_id, _ = _register(auth_headers, name="Quarterly Notice")
    resp = _download(auth_headers, report_id)
    assert resp.headers["content-disposition"] == 'attachment; filename="quarterly-notice-v1.docx"'
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml")
    assert resp.headers["cache-control"] == "private, no-store"


def test_khmer_only_report_name_falls_back_to_the_report_id_in_the_filename(auth_headers):
    report_id, _ = _register(auth_headers, name="របាយការណ៍")
    resp = _download(auth_headers, report_id)
    assert resp.headers["content-disposition"] == f'attachment; filename="{report_id}-v1.docx"'


def test_unknown_version_and_unknown_report_are_404(auth_headers):
    report_id, _ = _register(auth_headers)
    assert _download(auth_headers, report_id, version=9).status_code == 404
    assert client.get("/api/v1/reports/nope/file", headers=auth_headers).status_code == 404


def test_a_template_that_fails_field_detection_still_uploads(auth_headers):
    """A broken Jinja tag must not make the file un-storable or un-downloadable."""
    report_id, _ = _register(auth_headers, text="{% if %} broken")
    v1 = report_store.list_versions(report_id)[0]
    assert v1["fields"] is None
    assert _download(auth_headers, report_id).status_code == 200


# --- the changelog --------------------------------------------------------------


def test_changelog_says_which_placeholders_a_version_added_and_removed(auth_headers):
    report_id, _ = _register(auth_headers)  # {{ name }}, {{ region }}
    _replace(auth_headers, report_id, "Hello {{ name }}, now with {{ branch }}", note="Swapped region for branch")

    log = client.get(f"/api/v1/reports/{report_id}/changelog", headers=auth_headers).json()
    assert log["current_version"] == 2
    assert log["original_available"] is True
    v2 = next(e["version"] for e in log["entries"] if e["kind"] == "version" and e["version"]["version"] == 2)
    assert v2["fields_added"] == ["branch"]
    assert v2["fields_removed"] == ["region"]
    assert v2["note"] == "Swapped region for branch"
    assert v2["identical_to_previous"] is False
    assert v2["size_delta"] is not None


def test_changelog_flags_a_reupload_of_an_identical_file(auth_headers):
    report_id, _ = _register(auth_headers)
    # Same bytes as v1: copy the stored snapshot back in through the API.
    same = _download(auth_headers, report_id, version=1).content
    resp = client.put(
        f"/api/v1/reports/{report_id}/file", files={"file": ("again.docx", same, "application/octet-stream")}, headers=auth_headers
    )
    assert resp.status_code == 200
    latest = report_store.list_versions(report_id)[0]
    assert latest["version"] == 2
    assert latest["identical_to_previous"] is True
    assert latest["fields_added"] == [] and latest["fields_removed"] == []


def test_changelog_merges_versions_with_other_changes_newest_first(auth_headers):
    report_id, _ = _register(auth_headers)
    client.patch(f"/api/v1/reports/{report_id}", json={"name": "Renamed Notice"}, headers=auth_headers)
    _replace(auth_headers, report_id, "Only {{ name }}")
    _download(auth_headers, report_id, version=1)

    log = client.get(f"/api/v1/reports/{report_id}/changelog", headers=auth_headers).json()
    kinds = [(e["kind"], (e["event"] or e["version"]).get("action") or f"v{e['version']['version']}") for e in log["entries"]]
    assert kinds == [
        ("event", "report.template_download"),
        ("version", "v2"),
        ("event", "report.update"),
        ("version", "v1"),
    ]
    # A version entry carries the audit event that says who/where -- and the
    # upload isn't *also* listed as a separate event.
    v2 = log["entries"][1]
    assert v2["upload_event"]["action"] == "report.file_replace"
    assert v2["upload_event"]["actor_username"] == "testadmin"
    assert not any(e["kind"] == "event" and e["event"]["action"] in ("report.create", "report.file_replace") for e in log["entries"])


def test_deleting_a_report_removes_its_files_but_keeps_the_audit_trail(auth_headers):
    report_id, _ = _register(auth_headers, name="Doomed")
    _replace(auth_headers, report_id, "v2 {{ x }}")
    versions_dir = report_store.STORE_DIR / report_id / "versions"
    assert versions_dir.exists()

    assert client.delete(f"/api/v1/reports/{report_id}", headers=auth_headers).status_code == 204

    assert not (report_store.STORE_DIR / report_id).exists()
    with db.SessionLocal() as session:
        assert session.execute(select(db.ReportVersion).where(db.ReportVersion.report_id == report_id)).first() is None
        deleted = session.execute(
            select(db.AuditEvent).where(db.AuditEvent.entity_id == report_id, db.AuditEvent.action == "report.delete")
        ).scalar_one()
    assert deleted.entity_label == "Doomed"
    assert deleted.details["version"] == 2
    assert len(deleted.details["sha256"]) == 64


# --- reports that predate version history ----------------------------------------


def _make_legacy(report_id):
    """Put a report back into the state the code left it in before this
    feature: version rows and the versions/ directory don't exist."""
    with db.SessionLocal() as session:
        session.query(db.ReportVersion).filter(db.ReportVersion.report_id == report_id).delete()
        session.commit()
    shutil.rmtree(report_store.STORE_DIR / report_id / "versions", ignore_errors=True)


def test_replacing_a_legacy_report_preserves_the_file_it_replaces(auth_headers):
    report_id, _ = _register(auth_headers, text="Legacy {{ a }}")
    _make_legacy(report_id)

    _replace(auth_headers, report_id, "New {{ b }}")

    v1 = _download(auth_headers, report_id, version=1)
    assert v1.status_code == 200
    assert Document(BytesIO(v1.content)).paragraphs[0].text == "Legacy {{ a }}"
    old = next(v for v in report_store.list_versions(report_id) if v["version"] == 1)
    assert old["backfilled"] is True
    assert old["created_by"] is None  # the uploader wasn't recorded back then
    with db.SessionLocal() as session:
        # v1's true upload time is known even for a legacy report.
        assert old["created_at"] == session.get(db.ReportRow, report_id).created_at


def test_legacy_report_replaced_before_history_says_the_original_is_gone(auth_headers):
    """Two replaces under the old code left version=3 and only the newest
    file. v1 and v2 were overwritten and can't be recovered -- say so, don't
    pretend."""
    report_id, _ = _register(auth_headers)
    with db.SessionLocal() as session:
        session.get(db.ReportRow, report_id).version = 3
        session.commit()
    _make_legacy(report_id)

    log = client.get(f"/api/v1/reports/{report_id}/changelog", headers=auth_headers).json()
    assert log["current_version"] == 3
    assert log["first_retained_version"] == 3
    assert log["original_available"] is False
    only = [e for e in log["entries"] if e["kind"] == "version"]
    assert len(only) == 1 and only[0]["version"]["backfilled"] is True and only[0]["upload_event"] is None

    gone = _download(auth_headers, report_id, version=1)
    assert gone.status_code == 404
    assert "replaced before version history was kept" in gone.json()["detail"]
    # ...while the current file is still downloadable.
    assert _download(auth_headers, report_id).status_code == 200


# --- who may, and that it's recorded ----------------------------------------------


def test_downloading_a_template_is_recorded_with_who_which_version_and_checksum(auth_headers):
    report_id, _ = _register(auth_headers)
    _replace(auth_headers, report_id, "v2 {{ x }}")

    _download(auth_headers, report_id, version=1)
    _download(auth_headers, report_id)

    with db.SessionLocal() as session:
        events = (
            session.execute(
                select(db.AuditEvent)
                .where(db.AuditEvent.entity_id == report_id, db.AuditEvent.action == "report.template_download")
                .order_by(db.AuditEvent.created_at)
            )
            .scalars()
            .all()
        )
    assert [e.details["version"] for e in events] == [1, 2]
    assert [e.details["is_current"] for e in events] == [False, True]
    assert "original upload" in events[0].summary
    assert events[0].details["sha256"] == report_store.list_versions(report_id)[-1]["sha256"]
    assert events[0].actor_username == "testadmin"


def test_only_a_manager_of_the_report_can_download_or_read_the_changelog(auth_headers, make_local_user):
    report_id, _ = _register(auth_headers)
    # ROLE_REPORT_VIEWER: may view and render, not manage.
    _, viewer = make_local_user("vera", "pw-vera-12345", ("ROLE_REPORT_VIEWER",))

    assert _download(viewer, report_id).status_code == 403
    assert _download(viewer, report_id, version=1).status_code == 403
    assert client.get(f"/api/v1/reports/{report_id}/changelog", headers=viewer).status_code == 403
    assert client.get(f"/api/v1/reports/{report_id}/file", headers={}).status_code == 401


def test_a_report_level_manage_grant_is_enough(auth_headers, make_local_user):
    report_id, _ = _register(auth_headers)
    user_id, headers = make_local_user("mo", "pw-mo-1234567", ("ROLE_USER",))
    assert _download(headers, report_id).status_code == 403

    grant = client.post(
        "/api/v1/grants/reports",
        json={"subject_type": "user", "subject_id": user_id, "report_id": report_id, "permission_level": "manage"},
        headers=auth_headers,
    )
    assert grant.status_code == 200, grant.text
    assert _download(headers, report_id).status_code == 200
    assert client.get(f"/api/v1/reports/{report_id}/changelog", headers=headers).status_code == 200


def test_reports_actions_are_all_recorded(auth_headers):
    report_id, _ = _register(auth_headers)
    client.patch(f"/api/v1/reports/{report_id}", json={"description": "d"}, headers=auth_headers)
    _replace(auth_headers, report_id, "x {{ y }}")
    client.delete(f"/api/v1/reports/{report_id}", headers=auth_headers)
    assert _audit_actions(report_id) == ["report.create", "report.update", "report.file_replace", "report.delete"]


def test_another_organizations_manager_cannot_read_or_download_the_template(auth_headers, make_local_user):
    """`report:manage` is a global permission, and the guard on these routes
    doesn't look at the report's organization -- the routes themselves must,
    since what they hand over is the template's contents and history."""
    report_id, _ = _register(auth_headers)  # lands in the root org
    client.post("/api/v1/organizations", json={"id": "acme", "name": "Acme"}, headers=auth_headers)
    _, acme_admin = make_local_user("acme-admin", "pw-acme-12345", ("ROLE_ORG_ADMIN",), org_id="acme")

    # Holds report:manage in its own org -- and is still told the report doesn't exist.
    assert _download(acme_admin, report_id).status_code == 404
    assert _download(acme_admin, report_id, version=1).status_code == 404
    assert client.get(f"/api/v1/reports/{report_id}/changelog", headers=acme_admin).status_code == 404
    # Nothing was recorded as downloaded.
    assert "report.template_download" not in _audit_actions(report_id)
    # The superuser (and the report's own org) still can.
    assert _download(auth_headers, report_id).status_code == 200


def _put_file(headers, report_id, text, **form):
    return client.put(
        f"/api/v1/reports/{report_id}/file",
        files={"file": ("new.docx", _docx(text), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data=form,
        headers=headers,
    )


def test_version_label_is_shown_but_the_sequence_still_counts(auth_headers):
    report_id, _ = _register(auth_headers)
    resp = _put_file(auth_headers, report_id, "a {{ x }}", version_label="v1.0.1")
    assert resp.status_code == 200
    assert resp.json()["version"] == 2 and resp.json()["version_label"] == "1.0.1"
    assert report_store.list_versions(report_id)[0]["version_label"] == "1.0.1"
    # Unlabelled uploads are still numbered, and the label belongs to its own version.
    assert _put_file(auth_headers, report_id, "b {{ x }}").json()["version_label"] is None


def test_version_label_must_be_valid_and_unique(auth_headers):
    report_id, _ = _register(auth_headers)
    assert _put_file(auth_headers, report_id, "a {{ x }}", version_label="1.0.1").status_code == 200
    assert _put_file(auth_headers, report_id, "b {{ x }}", version_label="1.0.1").status_code == 409
    assert _put_file(auth_headers, report_id, "c {{ x }}", version_label="bad label!").status_code == 400
    assert report_store.get_report(report_id)["version"] == 2  # rejected uploads changed nothing


def test_version_label_and_note_can_be_edited_after_upload(auth_headers):
    report_id, _ = _register(auth_headers)
    _put_file(auth_headers, report_id, "a {{ x }}", version_label="1.0.1")
    url = f"/api/v1/reports/{report_id}/versions/2"
    resp = client.patch(url, json={"version_label": "1.1.0", "note": "fixed footer"}, headers=auth_headers)
    assert resp.status_code == 200, resp.text
    assert resp.json()["version_label"] == "1.1.0" and resp.json()["note"] == "fixed footer"
    assert client.get(f"/api/v1/reports/{report_id}", headers=auth_headers).json()["version_label"] == "1.1.0"
    assert client.patch(f"/api/v1/reports/{report_id}/versions/1", json={"version_label": "1.1.0"}, headers=auth_headers).status_code == 409
    assert client.patch(url, json={}, headers=auth_headers).status_code == 400
    assert client.patch(f"/api/v1/reports/{report_id}/versions/9", json={"note": "x"}, headers=auth_headers).status_code == 404
    assert client.patch(url, json={"version_label": None}, headers=auth_headers).json()["version_label"] is None


def test_download_filename_uses_the_label(auth_headers):
    report_id, _ = _register(auth_headers)
    _put_file(auth_headers, report_id, "a {{ x }}", version_label="1.0.1")
    assert "-v1.0.1.docx" in _download(auth_headers, report_id).headers["content-disposition"]
