"""POST /reports/{id}/embed-run: an anonymous embed asking for a report by
parameters, authorized by an API client's id + secret (app/clients.py).

The property everything here protects: the report's data source is only ever
called for a request that carries valid client credentials for a client an
admin granted *this* report -- and never on the strength of anything the
embed's visitor typed. So most tests assert two things: the answer, and that
the data source was (or wasn't) called.

Nothing here sends an Authorization header: an embed has no account.
"""
import json
from dataclasses import dataclass, field
from io import BytesIO

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select

from app import clients, db, report_data
from app.routers.reports import _report_headers
from app.main import app

client = TestClient(app)

TOKEN_ENV = "ERP_TEST_TOKEN"
PARAMS = {"p_from": "2026-08-01", "p_to": "2026-08-31", "p_signed": "Sok Sophea"}

DEFINITIONS = [
    {"name": "p_from", "label": "From", "type": "date", "required": True},
    {"name": "p_to", "label": "To", "type": "date", "required": True},
    {"name": "p_signed", "label": "Signed by", "type": "text", "required": False},
]
DATA_SOURCE = {
    "url": "https://erp.example.test/api/sales?from={{ p_from }}&to={{ p_to }}",
    "method": "GET",
    "headers": {"X-Client-Id": "aksor-khmer-bi"},
    "auth": {"type": "bearer", "token_env": TOKEN_ENV},
}


@dataclass
class FakeErp:
    """The report's REST data source: records every request the server makes."""

    calls: list = field(default_factory=list)
    status: int = 200

    def handler(self, request: httpx.Request) -> httpx.Response:
        self.calls.append(request)
        if self.status != 200:
            return httpx.Response(self.status, json={"error": "boom"})
        return httpx.Response(200, json={"total": "1200", "office": "Phnom Penh"})


@pytest.fixture
def erp(monkeypatch) -> FakeErp:
    fake = FakeErp()
    monkeypatch.setattr(
        report_data, "_make_client", lambda: httpx.Client(transport=httpx.MockTransport(fake.handler), follow_redirects=False)
    )
    monkeypatch.setenv(TOKEN_ENV, "erp-secret-token")
    return fake


@pytest.fixture
def report(auth_headers) -> str:
    """A .docx report with code `sales`, three filters and a data source."""
    doc = Document()
    doc.add_paragraph("Sales {{ p_from }} to {{ p_to }}: {{ total }} at {{ office }}, signed {{ p_signed }}")
    buf = BytesIO()
    doc.save(buf)
    resp = client.post(
        "/api/v1/reports",
        files={"file": ("t.docx", buf.getvalue(), "application/octet-stream")},
        data={"name": "Sales", "code": "sales"},
        headers=auth_headers,
    )
    assert resp.status_code == 200, resp.text
    report_id = resp.json()["report_id"]
    cfg = client.put(
        f"/api/v1/reports/{report_id}/data-config",
        json={"parameters": DEFINITIONS, "data_source": DATA_SOURCE},
        headers=auth_headers,
    )
    assert cfg.status_code == 200, cfg.text
    return report_id


def _text(content: bytes) -> str:
    # The renderer inserts zero-width word-break marks for Khmer line breaking; drop them.
    return "\n".join(p.text for p in Document(BytesIO(content)).paragraphs).replace("\u200b", "")


def embed_run(ref, body, **kw):
    return client.post(f"/api/v1/reports/{ref}/embed-run", json=body, **kw)


def _denials() -> list:
    with db.SessionLocal() as session:
        return list(session.scalars(select(db.AccessDeniedEvent)))


@pytest.fixture
def granted(report, auth_headers) -> dict:
    """Credentials of a client the admin granted the `sales` report."""
    secret, _ = make_client(auth_headers, [report])
    return creds(secret)


# --- the happy path ---------------------------------------------------------


def test_a_granted_run_fetches_the_data_server_side_and_renders(report, erp, granted):
    resp = embed_run("sales", {"parameters": PARAMS, "format": "docx", **granted})
    assert resp.status_code == 200, resp.text
    assert _text(resp.content) == "Sales 2026-08-01 to 2026-08-31: 1200 at Phnom Penh, signed Sok Sophea"

    # The data came from the source, called with the report's own credentials...
    assert len(erp.calls) == 1
    call = erp.calls[0]
    assert str(call.url) == "https://erp.example.test/api/sales?from=2026-08-01&to=2026-08-31"
    assert call.headers["Authorization"] == "Bearer erp-secret-token"
    assert call.headers["X-Client-Id"] == "aksor-khmer-bi"


def test_defaults_to_pdf_and_honours_part(report, erp, granted):
    resp = embed_run("sales", {"parameters": PARAMS, "part": 1, **granted})
    assert resp.status_code == 200 and resp.headers["content-type"] == "application/pdf"
    assert embed_run("sales", {"parameters": PARAMS, "part": 9, **granted}).status_code == 400


# --- no credentials, no data -------------------------------------------------


def test_a_run_without_client_credentials_is_refused(report, erp):
    resp = embed_run("sales", {"parameters": PARAMS})
    assert resp.status_code == 401
    assert resp.json()["detail"] == "This report can only be run with API client credentials"
    assert erp.calls == []
    assert _denials() == []  # anyone can send these: logged, not written to the security feed


def test_rejections_of_anonymous_callers_are_logged_not_written_to_the_security_feed(report, erp):
    # Anyone can send bad credentials; a database row per attempt would let them grow the table.
    for _ in range(3):
        embed_run("sales", {"parameters": PARAMS, **creds("garbage", client_id="nobody-here")})
    assert _denials() == []
    assert erp.calls == []


def test_an_unknown_report_is_a_404_before_anything_else(erp, granted):
    assert embed_run("nope", {"parameters": PARAMS, **granted}).status_code == 404
    assert erp.calls == []


# --- the parameters are still validated ----------------------------------------


def test_the_parameters_are_still_validated_against_the_reports_definitions(report, erp, granted):
    resp = embed_run("sales", {"parameters": {**PARAMS, "p_from": "not-a-date"}, **granted})
    assert resp.status_code == 400
    assert erp.calls == []

    resp = embed_run("sales", {"parameters": {**PARAMS, "p_extra": "x"}, **granted})
    assert resp.status_code == 400 and "p_extra" in resp.json()["detail"]
    assert erp.calls == []


def test_a_format_the_template_cant_produce_is_refused(report, erp, granted):
    resp = embed_run("sales", {"parameters": PARAMS, "format": "xlsx", **granted})
    assert resp.status_code == 400
    assert erp.calls == []


def test_a_failing_data_source_is_a_readable_502_not_a_crash(report, erp, granted):
    erp.status = 500
    resp = embed_run("sales", {"parameters": PARAMS, **granted})
    assert resp.status_code == 502 and "data source" in resp.json()["detail"]


def test_the_data_source_credentials_are_never_in_the_response(report, erp, granted):
    resp = embed_run("sales", {"parameters": PARAMS, "format": "docx", **granted})
    assert b"erp-secret-token" not in resp.content
    assert "erp-secret-token" not in json.dumps(dict(resp.headers))


def test_run_and_render_are_unchanged_by_the_shared_helper(report, erp, auth_headers):
    run = client.post(f"/api/v1/reports/{report}/run", json={"parameters": PARAMS, "format": "docx"}, headers=auth_headers)
    assert run.status_code == 200 and "Phnom Penh" in _text(run.content)
    assert client.post(f"/api/v1/reports/{report}/render?format=docx", json={"total": "1", "office": "X", **PARAMS}).status_code == 200


# --- a run describes what it rendered ---------------------------------------


def test_a_run_says_which_report_and_template_type_it_rendered(report, erp, auth_headers):
    # So the embed page can start from parameters alone, with no lookup of the report first.
    secret, _ = make_client(auth_headers, [report])
    resp = embed_run("sales", {"parameters": PARAMS, "format": "docx", **creds(secret)})
    assert resp.status_code == 200
    assert resp.headers["X-Report-Name"] == "Sales" and resp.headers["X-Report-Ext"] == "docx"


def test_a_khmer_report_name_travels_percent_encoded():
    from urllib.parse import unquote

    headers = _report_headers({"name": "ប្រៀបធៀបចំណូល", "template_ext": "docx"})
    headers["X-Report-Name"].encode("ascii")  # a header value must be plain ASCII
    assert unquote(headers["X-Report-Name"]) == "ប្រៀបធៀបចំណូល"
    assert headers["X-Report-Ext"] == "docx"


# --- API clients: client id + secret (app/clients.py) -----------------------


@pytest.fixture(autouse=True)
def fresh_client_limiter(monkeypatch):
    monkeypatch.delenv(clients.ENV_LIMIT, raising=False)
    clients.limiter.reset()
    yield
    clients.limiter.reset()


def make_client(auth_headers, report_ids=(), client_id="partner-web", **extra) -> tuple[str, dict]:
    """A client through the real management API. Returns (secret, its ClientOut)."""
    resp = client.post(
        "/api/v1/clients",
        json={"client_id": client_id, "name": "Partner web", "report_ids": list(report_ids), **extra},
        headers=auth_headers,
    )
    assert resp.status_code == 200, resp.text
    return resp.json()["secret"], resp.json()["client"]


def creds(secret, client_id="partner-web") -> dict:
    return {"client_id": client_id, "client_secret": secret}


def test_a_granted_client_runs_the_report_from_parameters_alone(report, erp, auth_headers):
    secret, out = make_client(auth_headers, [report])
    resp = embed_run("sales", {"parameters": PARAMS, "format": "docx", **creds(secret)})
    assert resp.status_code == 200, resp.text
    assert _text(resp.content) == "Sales 2026-08-01 to 2026-08-31: 1200 at Phnom Penh, signed Sok Sophea"

    # The server called the source, with the report's own credentials -- never the client's.
    assert len(erp.calls) == 1
    assert erp.calls[0].headers["Authorization"] == "Bearer erp-secret-token"
    assert secret not in str(erp.calls[0].url) and secret not in str(erp.calls[0].headers)

    # ...and the use is visible to the admin who manages the client.
    assert client.get(f"/api/v1/clients/{out['id']}", headers=auth_headers).json()["last_used_at"]


def test_granting_a_client_is_all_it_takes_no_restart_no_setting(report, erp, auth_headers):
    assert embed_run("sales", {"parameters": PARAMS}).status_code == 401
    secret, _ = make_client(auth_headers, [report])
    assert embed_run("sales", {"parameters": PARAMS, **creds(secret)}).status_code == 200


def test_the_report_can_be_named_by_id_or_by_code(report, erp, auth_headers):
    secret, _ = make_client(auth_headers, [report])
    assert embed_run(report, {"parameters": PARAMS, **creds(secret)}).status_code == 200
    assert embed_run("sales", {"parameters": PARAMS, **creds(secret)}).status_code == 200


@pytest.mark.parametrize("who", ["wrong-secret", "unknown-client", "disabled-client"])
def test_bad_credentials_are_refused_with_one_message_and_nothing_is_fetched(report, erp, auth_headers, who):
    secret, out = make_client(auth_headers, [report])
    body = {
        "wrong-secret": creds(secret + "x"),
        "unknown-client": creds(secret, client_id="nobody-here"),
        "disabled-client": creds(secret),
    }[who]
    if who == "disabled-client":
        client.patch(f"/api/v1/clients/{out['id']}", json={"is_active": False}, headers=auth_headers)

    resp = embed_run("sales", {"parameters": PARAMS, **body})
    assert resp.status_code == 401 and resp.json()["detail"] == "Invalid client credentials"
    assert erp.calls == []


def test_a_client_can_only_run_the_reports_it_was_granted(report, erp, auth_headers):
    secret, _ = make_client(auth_headers, [])
    resp = embed_run("sales", {"parameters": PARAMS, **creds(secret)})
    assert resp.status_code == 403 and resp.json()["detail"] == "This client may not run this report"
    assert erp.calls == []
    # A genuine client asking for what it wasn't given is something the security feed shows.
    assert [(d.username, d.resource) for d in _denials()] == [("client:partner-web", report)]


def test_revoking_the_grant_takes_effect_on_the_next_run(report, erp, auth_headers):
    secret, out = make_client(auth_headers, [report])
    assert embed_run("sales", {"parameters": PARAMS, **creds(secret)}).status_code == 200
    client.put(f"/api/v1/clients/{out['id']}/reports", json={"report_ids": []}, headers=auth_headers)
    assert embed_run("sales", {"parameters": PARAMS, **creds(secret)}).status_code == 403


def test_a_rotated_secret_replaces_the_old_one_immediately(report, erp, auth_headers):
    secret, out = make_client(auth_headers, [report])
    new = client.post(f"/api/v1/clients/{out['id']}/rotate-secret", headers=auth_headers).json()["secret"]
    assert embed_run("sales", {"parameters": PARAMS, **creds(secret)}).status_code == 401
    assert embed_run("sales", {"parameters": PARAMS, **creds(new)}).status_code == 200


@pytest.mark.parametrize(
    "params",
    [
        {**PARAMS, "extra": "1"},  # undeclared parameter
        {"p_from": "2026-08-01"},  # a required one is missing
        {**PARAMS, "p_from": "not-a-date"},  # wrong shape for a date
        {**PARAMS, "p_signed": "x" * 500},  # over the length cap
    ],
)
def test_a_clients_run_still_validates_its_parameters(report, erp, auth_headers, params):
    secret, _ = make_client(auth_headers, [report])
    assert embed_run("sales", {"parameters": params, **creds(secret)}).status_code == 400
    assert erp.calls == []


def test_a_caller_cannot_steer_the_data_source_with_a_parameter_value(report, erp, auth_headers):
    secret, _ = make_client(auth_headers, [report])
    tricky = {**PARAMS, "p_signed": "a&from=1999-01-01#x"}
    assert embed_run("sales", {"parameters": tricky, **creds(secret)}).status_code == 200
    # p_signed isn't in the URL, and the values that are can only be dates -- the host is fixed by the config.
    assert erp.calls[0].url.host == "erp.example.test"
    assert erp.calls[0].url.query == b"from=2026-08-01&to=2026-08-31"


@pytest.mark.parametrize(
    "extra",
    [
        {"client_id": "partner-web"},  # id without a secret
        {"client_secret": "aksor_cs_x"},  # secret without an id
    ],
)
def test_the_client_id_and_secret_go_together(report, extra):
    assert embed_run("sales", {"parameters": PARAMS, **extra}).status_code == 422


def test_a_client_of_another_organization_cannot_run_the_report(report, erp, auth_headers):
    with db.SessionLocal() as session:
        org = db.Organization(id="acme", name="Acme", is_active=True, created_at="2026-01-01T00:00:00+00:00")
        session.add(org)
        session.commit()
    secret, out = make_client(auth_headers, [], client_id="acme-web", org_id="acme")
    # Even a grant row planted directly can't cross organizations.
    with db.SessionLocal() as session:
        session.add(db.ApiClientReport(client_pk=out["id"], report_id=report, granted_at="2026-01-01T00:00:00+00:00"))
        session.commit()
    assert embed_run("sales", {"parameters": PARAMS, **creds(secret, "acme-web")}).status_code == 403
    assert erp.calls == []


def test_client_runs_are_rate_limited_per_client_not_per_address(report, erp, auth_headers, monkeypatch):
    monkeypatch.setenv(clients.ENV_LIMIT, "2")
    secret, _ = make_client(auth_headers, [report])
    other_secret, _ = make_client(auth_headers, [report], client_id="second-client")

    body = {"parameters": PARAMS, **creds(secret)}
    assert [embed_run("sales", body).status_code for _ in range(3)] == [200, 200, 429]
    assert embed_run("sales", body).headers["Retry-After"]
    # Another client, from the same address, is unaffected.
    assert embed_run("sales", {"parameters": PARAMS, **creds(other_secret, "second-client")}).status_code == 200


def test_the_secret_is_not_echoed_in_any_response(report, erp, auth_headers):
    secret, _ = make_client(auth_headers, [report])
    ok = embed_run("sales", {"parameters": PARAMS, **creds(secret)})
    bad = embed_run("sales", {"parameters": PARAMS, **creds(secret + "x")})
    forbidden = embed_run("sales", {"parameters": {"p_from": "x"}, **creds(secret + "y")})
    for resp in (ok, bad, forbidden):
        assert secret.encode() not in resp.content and secret.encode() not in str(resp.headers).encode()


# --- parameter defaults -------------------------------------------------------


def test_a_parameter_the_embed_leaves_out_takes_its_default(report, erp, auth_headers):
    """A filter the embedder leaves out is filled from the report's own default."""
    definitions = [dict(d) for d in DEFINITIONS]
    definitions[2]["default_value"] = "H.E"
    cfg = client.put(
        f"/api/v1/reports/{report}/data-config", json={"parameters": definitions, "data_source": DATA_SOURCE}, headers=auth_headers
    )
    assert cfg.status_code == 200, cfg.text

    secret, _ = make_client(auth_headers, [report])
    params = {"p_from": "2026-08-01", "p_to": "2026-08-31"}
    resp = embed_run("sales", {"parameters": params, "format": "docx", **creds(secret)})
    assert resp.status_code == 200, resp.text
    assert _text(resp.content) == "Sales 2026-08-01 to 2026-08-31: 1200 at Phnom Penh, signed H.E"

    # Sent as an empty string it stays empty -- the embedder said "none".
    blank = {**params, "p_signed": ""}
    resp = embed_run("sales", {"parameters": blank, "format": "docx", **creds(secret)})
    assert _text(resp.content) == "Sales 2026-08-01 to 2026-08-31: 1200 at Phnom Penh, signed "
