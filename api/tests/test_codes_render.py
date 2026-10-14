"""QR codes and barcodes drawn at render time (doc_engine/codes.py), through the API: an HTML template may use
`src="{{ qrcode(...) }}"` (registration accepts it; a direct URL is still refused), the schema doesn't list the code
functions as fields (a real field named `barcode` still is), each template kind renders codes that scan, and a value
that can't be drawn fails the render with the reason -- for whoever manages the report."""
import io
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from PIL import Image

from app.main import app

zxingcpp = pytest.importorskip("zxingcpp")
client = TestClient(app)

INVOICE = "EACL2007091100012163"


def read(png: bytes) -> set[tuple[str, str]]:
    return {(r.format.name, r.text) for r in zxingcpp.read_barcodes(Image.open(io.BytesIO(png)))}


def _register(headers, filename: str, content: bytes, name="Codes"):
    resp = client.post("/api/v1/reports", files={"file": (filename, content, "application/octet-stream")}, data={"name": name}, headers=headers)
    return resp


def _docx(*paragraphs) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_an_html_template_may_draw_codes_but_not_fetch(auth_headers):
    html = '<p>{{ customer }}</p><img src="{{ qrcode(no) }}" style="width:40mm"><img src="{{ barcode(no, \'code128\') }}" style="width:80mm">'
    created = _register(auth_headers, "t.html", html.encode())
    assert created.status_code == 200, created.text
    report_id = created.json()["report_id"]
    assert client.get(f"/api/v1/reports/{report_id}/schema").json()["fields"] == ["customer", "no"]
    page = client.post(f"/api/v1/reports/{report_id}/render?format=png", json={"customer": "សុខ សុភា", "no": INVOICE})
    assert page.status_code == 200, page.text
    assert read(page.content) == {("QRCode", INVOICE), ("Code128", INVOICE)}
    # Text around a code call -- a URL slipped in -- is still a direct reference, and refused.
    bad = _register(auth_headers, "bad.html", b'<img src="{{ qrcode(a) }}https://example.com/x.png">')
    assert bad.status_code == 400 and "resource('name')" in bad.text


def test_word_schema_and_render(auth_headers):
    created = _register(auth_headers, "t.docx", _docx("{{ qrcode(no, width_mm=30) }}", "{{ barcode(sku, 'ean13') }}", "Shelf code: {{ barcode }}"))
    report_id = created.json()["report_id"]
    # `barcode` is also used as a value, so it's a field; qrcode is only ever called.
    assert client.get(f"/api/v1/reports/{report_id}/schema").json()["fields"] == ["barcode", "no", "sku"]
    resp = client.post(f"/api/v1/reports/{report_id}/render?format=docx", json={"no": INVOICE, "sku": "590123412345", "barcode": "A-17"})
    assert resp.status_code == 200, resp.text
    with zipfile.ZipFile(io.BytesIO(resp.content)) as z:
        found = set().union(*(read(z.read(n)) for n in z.namelist() if n.startswith("word/media/")))
    assert found == {("QRCode", INVOICE), ("EAN13", "5901234123457")}
    assert "Shelf code: A-17" in [p.text for p in Document(io.BytesIO(resp.content)).paragraphs]


def test_excel_render(auth_headers):
    wb = Workbook()
    wb.active["A1"], wb.active["B1"] = "Invoice", "{{ qrcode(no) }}"
    buf = io.BytesIO()
    wb.save(buf)
    report_id = _register(auth_headers, "t.xlsx", buf.getvalue()).json()["report_id"]
    assert client.get(f"/api/v1/reports/{report_id}/schema").json()["fields"] == ["no"]
    resp = client.post(f"/api/v1/reports/{report_id}/render?format=xlsx", json={"no": INVOICE})
    assert resp.status_code == 200, resp.text
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert [read(p._data()) for p in ws._images] == [{("QRCode", INVOICE)}]


def test_a_value_that_cant_be_drawn_says_why(auth_headers):
    report_id = _register(auth_headers, "t.docx", _docx("{{ barcode(sku, 'ean13') }}")).json()["report_id"]
    resp = client.post(f"/api/v1/reports/{report_id}/render?format=docx", json={"sku": "5901234123458"}, headers=auth_headers)
    assert resp.status_code == 422
    assert "check digit should be 7" in resp.json()["detail"]
    anonymous = client.post(f"/api/v1/reports/{report_id}/render?format=docx", json={"sku": "5901234123458"})
    assert anonymous.status_code == 422 and "check digit" not in anonymous.text  # the usual generic message
