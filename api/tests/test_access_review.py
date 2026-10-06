"""GET /api/v1/grants/access-review -- the first aggregate view of every
active report + folder grant in one call, backing the admin Access Review
page. See app/routers/grants.py's list_access_review_grants.
"""
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app import db
from app.main import app
from app.rbac import ROOT_ORG_ID, create_organization

client = TestClient(app)


def _register_report(auth_headers, name="Review Report", org_id: str | None = None) -> str:
    doc = Document()
    doc.add_paragraph("Hello {{ name }}")
    buf = BytesIO()
    doc.save(buf)
    data = {"name": name}
    resp = client.post(
        "/api/v1/reports",
        files={"file": ("t.docx", buf.getvalue(), "application/octet-stream")},
        data=data,
        headers=auth_headers,
    )
    assert resp.status_code == 200, resp.text
    report_id = resp.json()["report_id"]
    if org_id is not None:
        from app import report_store

        with db.SessionLocal() as session:
            row = session.get(db.ReportRow, report_id)
            row.org_id = org_id
            session.commit()
    return report_id


def _create_folder(auth_headers, name, org_id=ROOT_ORG_ID) -> str:
    resp = client.post("/api/v1/folders", json={"org_id": org_id, "name": name}, headers=auth_headers)
    assert resp.status_code == 200, resp.text
    return resp.json()["id"]


def _make_role_with_permission(auth_headers, name: str, permission: str, org_id=ROOT_ORG_ID) -> str:
    resp = client.post("/api/v1/roles", json={"org_id": org_id, "name": name}, headers=auth_headers)
    assert resp.status_code == 200, resp.text
    role_id = resp.json()["id"]
    resp = client.put(f"/api/v1/roles/{role_id}/permissions", json={"permissions": [permission]}, headers=auth_headers)
    assert resp.status_code == 200, resp.text
    return role_id


def _assign_role(auth_headers, user_id: str, role_id: str):
    resp = client.post("/api/v1/grants/roles", json={"user_id": user_id, "role_id": role_id}, headers=auth_headers)
    assert resp.status_code == 200, resp.text


def _grant_report(auth_headers, report_id: str, user_id: str, level="view", parameter_limits=None):
    body = {"subject_type": "user", "subject_id": user_id, "report_id": report_id, "permission_level": level}
    if parameter_limits is not None:
        body["parameter_limits"] = parameter_limits
    resp = client.post("/api/v1/grants/reports", json=body, headers=auth_headers)
    assert resp.status_code == 200, resp.text
    return resp.json()["id"]


def _grant_folder(auth_headers, folder_id: str, user_id: str, level="view"):
    resp = client.post(
        "/api/v1/grants/folders",
        json={"subject_type": "user", "subject_id": user_id, "folder_id": folder_id, "permission_level": level},
        headers=auth_headers,
    )
    assert resp.status_code == 200, resp.text
    return resp.json()["id"]


def _review(headers, org_id: str | None = None):
    params = {"org_id": org_id} if org_id is not None else {}
    return client.get("/api/v1/grants/access-review", params=params, headers=headers)


def test_requires_authentication():
    assert _review({}).status_code == 401


def test_caller_with_neither_permission_gets_403(make_local_user):
    _, headers = make_local_user("no_perms", "pw12345", role_names=())
    resp = _review(headers)
    assert resp.status_code == 403


def test_report_manage_only_sees_report_grants_not_folder_grants(auth_headers, make_local_user):
    report_id = _register_report(auth_headers, "R1")
    folder_id = _create_folder(auth_headers, "F1")
    target_user_id, _ = make_local_user("review_target", "pw12345", role_names=())
    _grant_report(auth_headers, report_id, target_user_id)
    _grant_folder(auth_headers, folder_id, target_user_id)

    role_id = _make_role_with_permission(auth_headers, "ReportManagerOnly", "report:manage")
    caller_id, headers = make_local_user("report_manager", "pw12345", role_names=())
    _assign_role(auth_headers, caller_id, role_id)

    resp = _review(headers)
    assert resp.status_code == 200
    grant_types = {g["grant_type"] for g in resp.json()}
    assert grant_types == {"report"}


def test_folder_manage_only_sees_folder_grants_not_report_grants(auth_headers, make_local_user):
    report_id = _register_report(auth_headers, "R2")
    folder_id = _create_folder(auth_headers, "F2")
    target_user_id, _ = make_local_user("review_target_2", "pw12345", role_names=())
    _grant_report(auth_headers, report_id, target_user_id)
    _grant_folder(auth_headers, folder_id, target_user_id)

    role_id = _make_role_with_permission(auth_headers, "FolderManagerOnly", "folder:manage")
    caller_id, headers = make_local_user("folder_manager", "pw12345", role_names=())
    _assign_role(auth_headers, caller_id, role_id)

    resp = _review(headers)
    assert resp.status_code == 200
    grant_types = {g["grant_type"] for g in resp.json()}
    assert grant_types == {"folder"}


def test_superuser_sees_every_org_grants_org_id_filters_correctly(auth_headers, make_local_user):
    with db.SessionLocal() as session:
        create_organization(session, "review-org", "Review Org")
        session.commit()

    root_report = _register_report(auth_headers, "RootReport")
    other_report = _register_report(auth_headers, "OtherReport", org_id="review-org")
    root_user_id, _ = make_local_user("root_target", "pw12345", role_names=(), org_id=ROOT_ORG_ID)
    other_user_id, _ = make_local_user("other_target", "pw12345", role_names=(), org_id="review-org")
    _grant_report(auth_headers, root_report, root_user_id)
    _grant_report(auth_headers, other_report, other_user_id)

    all_grants = _review(auth_headers).json()
    resource_ids = {g["resource_id"] for g in all_grants}
    assert {root_report, other_report} <= resource_ids

    scoped = _review(auth_headers, org_id="review-org").json()
    assert {g["resource_id"] for g in scoped} == {other_report}


def test_non_superuser_org_id_param_for_another_org_is_404(auth_headers, make_local_user):
    with db.SessionLocal() as session:
        create_organization(session, "review-org-2", "Review Org 2")
        session.commit()
    role_id = _make_role_with_permission(auth_headers, "ReportManagerScoped", "report:manage")
    caller_id, headers = make_local_user("scoped_manager", "pw12345", role_names=())
    _assign_role(auth_headers, caller_id, role_id)

    resp = _review(headers, org_id="review-org-2")
    assert resp.status_code == 404


def test_names_are_resolved_not_raw_ids(auth_headers, make_local_user):
    report_id = _register_report(auth_headers, "NamedReport")
    target_user_id, _ = make_local_user("named_target_user", "pw12345", role_names=())
    _grant_report(auth_headers, report_id, target_user_id)

    grants = _review(auth_headers).json()
    row = next(g for g in grants if g["resource_id"] == report_id)
    assert row["resource_name"] == "NamedReport"
    assert row["subject_name"] == "named_target_user"


def test_revoked_grant_is_excluded(auth_headers, make_local_user):
    report_id = _register_report(auth_headers, "RevokeMe")
    target_user_id, _ = make_local_user("revoke_target", "pw12345", role_names=())
    grant_id = _grant_report(auth_headers, report_id, target_user_id)

    assert report_id in {g["resource_id"] for g in _review(auth_headers).json()}
    resp = client.delete(f"/api/v1/grants/reports/{grant_id}", headers=auth_headers)
    assert resp.status_code == 204
    assert report_id not in {g["resource_id"] for g in _review(auth_headers).json()}


def test_report_grant_echoes_parameter_limits_folder_grant_has_none(auth_headers, make_local_user):
    doc = Document()
    doc.add_paragraph("{{ p_branch }}")
    buf = BytesIO()
    doc.save(buf)
    resp = client.post(
        "/api/v1/reports",
        files={"file": ("t.docx", buf.getvalue(), "application/octet-stream")},
        data={"name": "LimitedReport"},
        headers=auth_headers,
    )
    report_id = resp.json()["report_id"]
    client.put(
        f"/api/v1/reports/{report_id}/data-config",
        json={"parameters": [{"name": "p_branch", "options": [{"value": "BR01"}, {"value": "BR02"}]}]},
        headers=auth_headers,
    )
    folder_id = _create_folder(auth_headers, "LimitFolder")
    target_user_id, _ = make_local_user("limit_target", "pw12345", role_names=())
    _grant_report(auth_headers, report_id, target_user_id, level="render", parameter_limits={"p_branch": ["BR01"]})
    _grant_folder(auth_headers, folder_id, target_user_id)

    grants = _review(auth_headers).json()
    report_grant = next(g for g in grants if g["resource_id"] == report_id)
    folder_grant = next(g for g in grants if g["resource_id"] == folder_id)
    assert report_grant["parameter_limits"] == {"p_branch": ["BR01"]}
    assert folder_grant["parameter_limits"] is None
