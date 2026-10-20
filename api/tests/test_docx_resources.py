"""`{{ resource('name') }}` in a Word template: a picture from Resources > Images, mapped when the template is
registered -- so a letterhead many Word templates share is changed once, in the library, as html templates already
could. The body, headers and footers all count; a name must map to an image; replacing the image changes every
report that places it; a scheduled job places it too."""
import io
import json
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from app import image_store, job_executors, report_store, template_fields
from app.main import app
from app.rbac import ROOT_ORG_ID

client = TestClient(app)
EMU_PER_MM = 36000


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(image_store, "STORE_DIR", tmp_path / "images")


def png(color):
    out = io.BytesIO()
    Image.new("RGB", (40, 20), color).save(out, format="PNG")
    return out.getvalue()


def image(headers, name, content):
    resp = client.post("/api/v1/images", files={"file": (f"{name}.png", content, "image/png")},
                       data={"name": name, "org_id": ROOT_ORG_ID}, headers=headers)
    assert resp.status_code == 200, resp.text
    return resp.json()


def letterhead(body="{{ resource('seal', width_mm=30) }}", header="{{ resource('ministry', width_mm=50) }}"):
    """A Word template like the e-receipt: a picture in the page header, one in the body, and Khmer text."""
    d = Document()
    d.sections[0].header.paragraphs[0].text = header
    d.add_paragraph(body)
    d.add_paragraph("បង្កាន់ដៃអេឡិចត្រូនិកបញ្ជាក់ការបង់ប្រាក់ពន្ធ {{ RECEIPT_CODE }}")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def register(headers, content, bindings=None):
    data = {"name": "E-receipt"}
    if bindings is not None:
        data["resource_bindings"] = json.dumps(bindings)
    return client.post("/api/v1/reports", files={"file": ("r.docx", content, "application/octet-stream")}, data=data, headers=headers)


def media(docx: bytes) -> list[bytes]:
    with zipfile.ZipFile(io.BytesIO(docx)) as z:
        return [z.read(n) for n in z.namelist() if n.startswith("word/media/")]


def part(docx: bytes, prefix: str) -> str:
    with zipfile.ZipFile(io.BytesIO(docx)) as z:
        return " ".join(z.read(n).decode() for n in z.namelist() if n.startswith(prefix) and n.endswith(".xml"))


def render(headers, report_id):
    resp = client.post(f"/api/v1/reports/{report_id}/render?format=docx", json={"RECEIPT_CODE": "EACL2007091100012163"}, headers=headers)
    assert resp.status_code == 200, resp.text
    return resp.content


@pytest.fixture
def mapped(auth_headers):
    ministry = image(auth_headers, "ministry", png((10, 80, 160)))
    seal = image(auth_headers, "seal", png((20, 140, 60)))
    resp = register(auth_headers, letterhead(), {"ministry": {"kind": "image", "id": ministry["id"]}, "seal": {"kind": "image", "id": seal["id"]}})
    assert resp.status_code == 200, resp.text
    return resp.json(), ministry, seal


def test_names_are_found_in_headers_and_body_even_with_words_curly_quotes():
    content = letterhead(body="{{ resource(‘seal’) }} {{ resource(\"stamp\", height_mm=9) }}")
    assert template_fields.docx_resource_names(content) == ["ministry", "seal", "stamp"]


def test_each_name_must_be_mapped_to_an_image(auth_headers):
    ministry = image(auth_headers, "ministry", png((10, 80, 160)))
    unmapped = register(auth_headers, letterhead(), {"ministry": {"kind": "image", "id": ministry["id"]}})
    assert unmapped.status_code == 400 and "'seal'" in unmapped.json()["detail"]
    css = client.post("/api/v1/stylesheets", files={"file": ("t.css", b"h1{color:red}", "text/css")},
                      data={"name": "theme", "org_id": ROOT_ORG_ID}, headers=auth_headers).json()
    stylesheet = register(auth_headers, letterhead(), {"ministry": {"kind": "image", "id": ministry["id"]}, "seal": {"kind": "stylesheet", "id": css["id"]}})
    assert stylesheet.status_code == 400 and "must map to an image" in stylesheet.json()["detail"]


def test_the_mapping_step_sees_a_word_templates_names(auth_headers):
    resp = client.post("/api/v1/reports/parse-template", files={"file": ("r.docx", letterhead(), "application/octet-stream")}, headers=auth_headers)
    assert resp.status_code == 200, resp.text
    assert resp.json()["resources"] == ["ministry", "seal"]


def test_rendered_in_the_header_and_body_at_the_width_asked(auth_headers, mapped):
    report, ministry, seal = mapped
    assert report["resource_bindings"] == {"ministry": {"kind": "image", "id": ministry["id"]}, "seal": {"kind": "image", "id": seal["id"]}}
    out = render(auth_headers, report["report_id"])
    assert sorted(media(out)) == sorted([png((10, 80, 160)), png((20, 140, 60))])
    assert f'cx="{50 * EMU_PER_MM}"' in part(out, "word/header")
    assert f'cx="{30 * EMU_PER_MM}"' in part(out, "word/document")
    assert "EACL2007091100012163" in part(out, "word/document")
    fields = client.get(f"/api/v1/reports/{report['report_id']}/schema", headers=auth_headers).json()["fields"]
    assert "resource" not in fields and "RECEIPT_CODE" in fields


def test_replacing_the_library_image_changes_every_report_that_places_it(auth_headers, mapped):
    report, ministry, _ = mapped
    resp = client.put(f"/api/v1/images/{ministry['id']}/file", files={"file": ("new.png", png((200, 30, 30)), "image/png")}, headers=auth_headers)
    assert resp.status_code == 200, resp.text
    assert png((200, 30, 30)) in media(render(auth_headers, report["report_id"]))
    assert png((10, 80, 160)) not in media(render(auth_headers, report["report_id"]))


def test_a_new_version_can_map_a_name_it_adds(auth_headers, mapped):
    report, *_ = mapped
    kingdom = image(auth_headers, "kingdom", png((90, 90, 0)))
    content = letterhead(header="{{ resource('ministry') }} {{ resource('kingdom', width_mm=52) }}")
    url = f"/api/v1/reports/{report['report_id']}/file"
    unmapped = client.put(url, files={"file": ("r.docx", content, "application/octet-stream")}, headers=auth_headers)
    assert unmapped.status_code == 400 and "'kingdom'" in unmapped.json()["detail"]
    resp = client.put(url, files={"file": ("r.docx", content, "application/octet-stream")},
                      data={"resource_bindings": json.dumps({"kingdom": {"kind": "image", "id": kingdom["id"]}})}, headers=auth_headers)
    assert resp.status_code == 200, resp.text
    assert set(resp.json()["resource_bindings"]) == {"ministry", "seal", "kingdom"}
    assert png((90, 90, 0)) in media(render(auth_headers, report["report_id"]))


def test_a_scheduled_job_places_them_too(auth_headers, mapped):
    report, *_ = mapped
    report_store.update_data_config(report["report_id"], [], {"type": "static", "static_data": {"RECEIPT_CODE": "EACL1"}})
    out, _ = job_executors.render_with_data(report["report_id"], "docx", {})
    assert sorted(media(out)) == sorted([png((10, 80, 160)), png((20, 140, 60))])


def test_a_word_template_without_resources_is_unchanged(auth_headers):
    d = Document(); d.add_paragraph("{{ NAME }}"); buf = io.BytesIO(); d.save(buf)
    resp = register(auth_headers, buf.getvalue())
    assert resp.status_code == 200, resp.text
    assert resp.json()["resource_bindings"] is None
