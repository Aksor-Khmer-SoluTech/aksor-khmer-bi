"""Row-level access on datasets (Phase 2b): an open dataset leaves it to the reports built on it; a restricted one is
used only by the users and roles it grants, each only with the filter values those grants allow -- on top of the
report's own grants, so both must allow a value -- and only its "build" grantees may build a new report on it."""
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app import db
from app.main import app
from app.rbac import ROOT_ORG_ID

client = TestClient(app)

BRANCHES = [{"value": "BR01", "label": "ភ្នំពេញ"}, {"value": "BR02", "label": "សៀមរាប"}, {"value": "BR03", "label": "បាត់ដំបង"}]
SALES = {"title": "របាយការណ៍លក់", "rows": [{"branch": "BR01", "amount": 40000}, {"branch": "BR02", "amount": 35000}]}


def _dataset(headers, name="branch-sales") -> dict:
    resp = client.post("/api/v1/datasets", headers=headers, json={
        "name": name, "title": "Branch sales", "org_id": ROOT_ORG_ID,
        "parameters": [{"name": "p_branch", "label": "Branch", "options": BRANCHES}],
        "data_source": {"type": "static", "static_data": SALES},
        "measures": [{"key": "gross", "label": "Gross", "kind": "sum", "field": "amount", "format": "#,##0 ៛"}],
    })
    assert resp.status_code == 200, resp.text
    return resp.json()


def _report(headers, dataset="branch-sales") -> str:
    doc = Document()
    doc.add_paragraph("{{ title }} · {{ p_branch }} · {{ measures.gross }}")
    buf = BytesIO()
    doc.save(buf)
    resp = client.post("/api/v1/reports", files={"file": ("t.docx", buf.getvalue(), "application/octet-stream")},
                       data={"name": "Sales on a dataset"}, headers=headers)
    report_id = resp.json()["report_id"]
    put = client.put(f"/api/v1/reports/{report_id}/data-config", headers=headers,
                     json={"parameters": [], "data_source": {"type": "dataset", "dataset": dataset}})
    assert put.status_code == 200, put.text
    return report_id


def _run(headers, report_id, branch):
    return client.post(f"/api/v1/reports/{report_id}/run", json={"parameters": {"p_branch": branch}, "format": "docx"}, headers=headers)


def _text(content: bytes) -> str:
    return "\n".join(p.text for p in Document(BytesIO(content)).paragraphs).replace("​", "")


def _runner(admin, make_local_user, report_id, username="staff", limits=None, roles=()):
    """A user who may run the report (a render grant on it), optionally limited by the report's own grant."""
    user_id, headers = make_local_user(username, f"pw-{username}-1", roles)
    grant = client.post("/api/v1/grants/reports", headers=admin, json={
        "subject_type": "user", "subject_id": user_id, "report_id": report_id, "permission_level": "render",
        **({"parameter_limits": limits} if limits else {}),
    })
    assert grant.status_code == 200, grant.text
    return user_id, headers


def _restrict(admin, dataset_id, mode="restricted"):
    resp = client.put(f"/api/v1/datasets/{dataset_id}/access/mode", headers=admin, json={"mode": mode})
    assert resp.status_code == 200, resp.text
    return resp.json()


def _grant(admin, dataset_id, subject_id, subject_type="user", level="run", limits=None, expect=200, **extra):
    resp = client.post(f"/api/v1/datasets/{dataset_id}/access", headers=admin, json={
        "subject_type": subject_type, "subject_id": subject_id, "level": level,
        **({"parameter_limits": limits} if limits is not None else {}), **extra,
    })
    assert resp.status_code == expect, resp.text
    return resp.json()


def _run_form_options(headers, report_id):
    form = client.get(f"/api/v1/reports/{report_id}/run-form", headers=headers)
    assert form.status_code == 200, form.text
    return [o["value"] for o in form.json()["parameters"][0]["options"]]


def test_an_open_dataset_leaves_access_to_the_report(auth_headers, make_local_user):
    dataset = _dataset(auth_headers)
    report_id = _report(auth_headers)
    assert dataset["access_mode"] == "open"
    _, staff = _runner(auth_headers, make_local_user, report_id)
    assert _run(staff, report_id, "BR02").status_code == 200


def test_a_restricted_dataset_is_used_only_by_its_grants_and_within_their_limits(auth_headers, make_local_user):
    dataset = _dataset(auth_headers)
    report_id = _report(auth_headers)
    user_id, staff = _runner(auth_headers, make_local_user, report_id)

    assert _restrict(auth_headers, dataset["id"]) == {"mode": "restricted", "grants": []}
    denied = _run(staff, report_id, "BR01")
    assert denied.status_code == 403 and "Branch sales" in denied.json()["detail"]
    assert client.get(f"/api/v1/reports/{report_id}/run-form", headers=staff).status_code == 403

    grant = _grant(auth_headers, dataset["id"], user_id, limits={"p_branch": ["BR02"]})
    assert grant["subject_name"] == "staff" and grant["parameter_limits"] == {"p_branch": ["BR02"]}
    assert _run_form_options(staff, report_id) == ["BR02"]  # the run form only offers what both allow
    ok = _run(staff, report_id, "BR02")
    assert ok.status_code == 200 and _text(ok.content) == "របាយការណ៍លក់ · BR02 · 75,000 ៛"
    assert _run(staff, report_id, "BR01").status_code == 403

    # The dataset's owner and the superuser aren't limited by it.
    assert _run(auth_headers, report_id, "BR01").status_code == 200

    client.delete(f"/api/v1/datasets/{dataset['id']}/access/{grant['id']}", headers=auth_headers)
    assert _run(staff, report_id, "BR02").status_code == 403
    # Revoking the last grant doesn't open it up again: the mode is its own switch.
    assert client.get(f"/api/v1/datasets/{dataset['id']}/access", headers=auth_headers).json() == {"mode": "restricted", "grants": []}

    _restrict(auth_headers, dataset["id"], "open")
    assert _run(staff, report_id, "BR01").status_code == 200


def test_both_the_report_and_the_dataset_must_allow_a_value(auth_headers, make_local_user):
    dataset = _dataset(auth_headers)
    report_id = _report(auth_headers)
    user_id, staff = _runner(auth_headers, make_local_user, report_id, limits={"p_branch": ["BR01", "BR02"]})
    _restrict(auth_headers, dataset["id"])
    _grant(auth_headers, dataset["id"], user_id, limits={"p_branch": ["BR02", "BR03"]})
    assert _run_form_options(staff, report_id) == ["BR02"]
    assert _run(staff, report_id, "BR02").status_code == 200
    assert _run(staff, report_id, "BR01").status_code == 403  # the report allows it, the dataset doesn't
    assert _run(staff, report_id, "BR03").status_code == 403  # the dataset allows it, the report doesn't


def test_grants_add_up_across_a_person_and_their_roles(auth_headers, make_local_user):
    dataset = _dataset(auth_headers)
    report_id = _report(auth_headers)
    user_id, staff = _runner(auth_headers, make_local_user, report_id, roles=("ROLE_REPORT_VIEWER",))
    role_id = next(r["id"] for r in client.get("/api/v1/roles", params={"org_id": ROOT_ORG_ID}, headers=auth_headers).json()
                   if r["name"] == "ROLE_REPORT_VIEWER")
    _restrict(auth_headers, dataset["id"])
    _grant(auth_headers, dataset["id"], user_id, limits={"p_branch": ["BR01"]})
    _grant(auth_headers, dataset["id"], role_id, subject_type="role", limits={"p_branch": ["BR03"]})
    assert _run_form_options(staff, report_id) == ["BR01", "BR03"]
    _grant(auth_headers, dataset["id"], role_id, subject_type="role")  # one grant without a limit: unrestricted
    assert _run_form_options(staff, report_id) == ["BR01", "BR02", "BR03"]


def test_an_expired_grant_stops_counting(auth_headers, make_local_user):
    dataset = _dataset(auth_headers)
    report_id = _report(auth_headers)
    user_id, staff = _runner(auth_headers, make_local_user, report_id)
    _restrict(auth_headers, dataset["id"])
    grant = _grant(auth_headers, dataset["id"], user_id, expires_at="2999-01-01T00:00:00+00:00")
    assert _run(staff, report_id, "BR01").status_code == 200
    with db.SessionLocal() as session:
        session.get(db.DatasetAccessGrant, grant["id"]).expires_at = "2000-01-01T00:00:00+00:00"
        session.commit()
    assert _run(staff, report_id, "BR01").status_code == 403
    assert _grant(auth_headers, dataset["id"], user_id, expires_at="2000-01-01T00:00:00+00:00", expect=400)


def test_only_build_grantees_may_build_a_report_on_a_restricted_dataset(auth_headers, make_local_user):
    dataset = _dataset(auth_headers)
    user_id, author = make_local_user("author", "pw-author-1")
    assert client.post("/api/v1/grants/permissions", headers=auth_headers,
                       json={"user_id": user_id, "permission_code": "report:manage"}).status_code == 200
    report_id = _report(auth_headers, dataset="branch-sales")
    client.put(f"/api/v1/reports/{report_id}/data-config", headers=auth_headers, json={"parameters": [], "data_source": None})

    def choose():
        return client.put(f"/api/v1/reports/{report_id}/data-config", headers=author,
                          json={"parameters": [], "data_source": {"type": "dataset", "dataset": "branch-sales"}})

    _restrict(auth_headers, dataset["id"])
    assert [d["name"] for d in client.get("/api/v1/datasets", headers=author).json()] == []
    assert choose().status_code == 403
    _grant(auth_headers, dataset["id"], user_id, level="run")
    assert choose().status_code == 403  # running isn't building
    _grant(auth_headers, dataset["id"], user_id, level="build")
    listed = client.get("/api/v1/datasets", headers=author).json()
    assert [(d["name"], d["access_mode"], d["grant_count"]) for d in listed] == [("branch-sales", "restricted", 2)]
    assert choose().status_code == 200


def test_grants_are_checked_against_the_dataset(auth_headers, make_local_user):
    dataset = _dataset(auth_headers)
    user_id, _ = make_local_user("staff", "pw-staff-1")
    for limits, fragment in [
        ({"p_region": ["X"]}, "isn't a filter"),
        ({"p_branch": ["BR09"]}, "Not a choice"),
        ({"p_branch": []}, "empty"),
    ]:
        resp = _grant(auth_headers, dataset["id"], user_id, limits=limits, expect=400)
        assert fragment in resp["detail"]
    with db.SessionLocal() as session:
        session.add(db.Organization(id="other-org", name="Other", is_active=True, created_at="2026-01-01T00:00:00+00:00"))
        session.commit()
    other_id, _ = make_local_user("elsewhere", "pw-else-1", org_id="other-org")
    assert "another organization" in _grant(auth_headers, dataset["id"], other_id, expect=400)["detail"]
    assert _grant(auth_headers, dataset["id"], "no-such-user", expect=404)


def test_who_may_manage_a_datasets_access(auth_headers, make_local_user):
    dataset = _dataset(auth_headers)
    _, report_admin = make_local_user("author", "pw-author-1", ("ROLE_REPORT_ADMIN",))
    user_id, viewer = make_local_user("viewer", "pw-viewer-1", ("ROLE_REPORT_VIEWER",))
    assert client.get(f"/api/v1/datasets/{dataset['id']}/access", headers=report_admin).status_code == 200
    assert client.get(f"/api/v1/datasets/{dataset['id']}/access", headers=viewer).status_code == 403
    assert client.put(f"/api/v1/datasets/{dataset['id']}/access/mode", headers=viewer, json={"mode": "open"}).status_code == 403
    assert _grant(viewer, dataset["id"], user_id, expect=403)


def test_access_changes_are_audited_and_go_with_the_dataset(auth_headers, make_local_user):
    dataset = _dataset(auth_headers)
    user_id, _ = make_local_user("staff", "pw-staff-1")
    _restrict(auth_headers, dataset["id"])
    grant = _grant(auth_headers, dataset["id"], user_id, level="build", limits={"p_branch": ["BR02"]})
    client.delete(f"/api/v1/datasets/{dataset['id']}/access/{grant['id']}", headers=auth_headers)
    with db.SessionLocal() as session:
        actions = [e.action for e in session.query(db.AuditEvent).filter(db.AuditEvent.entity_id == dataset["id"]).order_by(db.AuditEvent.created_at)]
        summaries = [e.summary for e in session.query(db.AuditEvent).filter(db.AuditEvent.action == "dataset.access_grant")]
    assert actions[-3:] == ["dataset.access_mode", "dataset.access_grant", "dataset.access_revoke"]
    assert summaries == ["Granted build & run on dataset Branch sales to user staff (p_branch: BR02)"]
    assert client.delete(f"/api/v1/datasets/{dataset['id']}", headers=auth_headers).status_code == 204
    with db.SessionLocal() as session:
        assert session.query(db.DatasetAccessGrant).count() == 0


def test_dataset_grants_show_in_access_review(auth_headers, make_local_user):
    dataset = _dataset(auth_headers)
    other = _dataset(auth_headers, name="other-sales")
    user_id, _ = make_local_user("staff", "pw-staff-1")
    grant = _grant(auth_headers, dataset["id"], user_id, level="build", limits={"p_branch": ["BR02"]})
    rows = [g for g in client.get("/api/v1/grants/access-review", headers=auth_headers).json() if g["grant_type"] == "dataset"]
    assert [(g["subject_name"], g["resource_name"], g["permission_level"], g["parameter_limits"]) for g in rows] == [
        ("staff", "Branch sales", "build", {"p_branch": ["BR02"]})
    ]
    # Revoking is addressed through its own dataset: another dataset's id doesn't reach it.
    assert client.delete(f"/api/v1/datasets/{other['id']}/access/{grant['id']}", headers=auth_headers).status_code == 404
    assert client.delete(f"/api/v1/datasets/{dataset['id']}/access/{grant['id']}", headers=auth_headers).status_code == 204
    assert not [g for g in client.get("/api/v1/grants/access-review", headers=auth_headers).json() if g["grant_type"] == "dataset"]
