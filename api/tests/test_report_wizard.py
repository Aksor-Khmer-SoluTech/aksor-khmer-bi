"""The New report wizard (app/routers/report_wizard.py): run a data source unsaved, create a draft report from what
it returned -- with a generated starter template -- check an uploaded template against the data, publish. And a
draft stays out of sight of everyone who can't manage it until then."""
import io
import json
import zipfile

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import report_data, report_wizard
from app.main import app

client = TestClient(app)

INVOICES = {
    "invoices": [
        {"invoice_no": "INV-1", "customer_name": "Sokha Trading", "issued_on": "2026-09-02", "total_usd": 1240.0, "status": "paid"},
        # Unspaced Khmer: what has to come out the other end segmented, not mangled.
        {"invoice_no": "INV-2", "customer_name": "ហាងលក់ទំនិញស្រីពៅ", "issued_on": "2026-09-09", "total_usd": 312.75, "status": "unpaid"},
    ],
    "company": {"name": "Aksor Co", "phone": "012 345 678"},
}


@pytest.fixture
def upstream(monkeypatch):
    """A REST API answering with INVOICES and remembering the URLs it was asked for."""
    calls: list[str] = []

    def factory():
        def handler(request):
            calls.append(str(request.url))
            return httpx.Response(200, json=INVOICES)

        return httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=False)

    monkeypatch.setattr(report_data, "_make_client", factory)
    return calls


REST_SOURCE = {"type": "rest", "url": "https://erp.example.com/invoices?month={{ month }}", "method": "GET"}


def _preview(headers, source=REST_SOURCE, values=None, parameters=None):
    return client.post(
        "/api/v1/reports/data-preview",
        json={"data_source": source, "values": {"month": "2026-09"} if values is None else values, "parameters": parameters or []},
        headers=headers,
    )


def _draft(headers, fmt="docx", **extra):
    body = {"name": "Monthly invoices", "format": fmt, "data_source": REST_SOURCE, "sample": INVOICES, "values": {"month": "2026-09"}}
    body.update(extra)
    resp = client.post("/api/v1/reports/drafts", json=body, headers=headers)
    assert resp.status_code == 200, resp.text
    return resp.json()


# --- running a source unsaved ----------------------------------------------------------------------------------


def test_data_preview_runs_the_source_and_turns_its_filters_into_parameters(auth_headers, upstream):
    resp = _preview(auth_headers)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["data"] == INVOICES and body["truncated"] is False
    assert upstream == ["https://erp.example.com/invoices?month=2026-09"]
    # {{ month }} in the URL became a filter by itself -- nobody declared it first.
    assert [(p["name"], p["label"], p["type"]) for p in body["parameters"]] == [("month", "Month", "text")]
    paths = {f["path"]: f for f in body["fields"]}
    assert paths["invoices"]["kind"] == "list" and paths["invoices"]["count"] == 2
    assert paths["invoices[].total_usd"]["kind"] == "number"
    assert paths["invoices[].issued_on"]["kind"] == "date"
    assert paths["company.name"]["example"] == "Aksor Co"


def test_data_preview_asks_for_a_filter_value_before_running(auth_headers, upstream):
    first = _preview(auth_headers, values={}).json()
    assert first["missing"] == ["month"] and first["data"] == {} and upstream == []  # nothing fetched yet
    assert [p["name"] for p in first["parameters"]] == ["month"]
    # The page sends the filter back with a type and a test value.
    second = _preview(auth_headers, values={"month": "2026-09-01"}, parameters=[{"name": "month", "type": "date"}]).json()
    assert second["missing"] == [] and second["parameters"][0]["type"] == "date" and upstream


def test_data_preview_keeps_only_the_first_rows(auth_headers, monkeypatch):
    many = {"rows": [{"n": i} for i in range(500)]}
    source = {"type": "static", "static_data": many}
    body = _preview(auth_headers, source=source, values={}).json()
    assert len(body["data"]["rows"]) == report_wizard.SAMPLE_LIST_ITEMS and body["truncated"] is True


def test_data_preview_reports_a_bad_query_and_a_failing_source(auth_headers, monkeypatch):
    bad = _preview(auth_headers, source={"type": "jdbc", "connection": "nope", "query": "DELETE FROM invoices"})
    assert bad.status_code == 400

    def failing():
        return httpx.Client(transport=httpx.MockTransport(lambda request: httpx.Response(500)), follow_redirects=False)

    monkeypatch.setattr(report_data, "_make_client", failing)
    assert _preview(auth_headers).status_code == 502


def test_data_preview_needs_report_manage(make_local_user, upstream):
    _, headers = make_local_user("reader", "pw-reader-1", ("ROLE_REPORT_VIEWER",))
    assert _preview(headers).status_code == 403
    assert upstream == []  # refused before anything was fetched


# --- drafts and their starter template ---------------------------------------------------------------------------


def test_a_draft_starts_with_a_starter_template_its_source_and_its_sample(auth_headers, upstream):
    draft = _draft(auth_headers)
    assert draft["is_draft"] is True and draft["template_ext"] == "docx" and draft["version"] == 1
    assert draft["sample_context"]["month"] == "2026-09"  # the test value is kept with the data
    config = client.get(f"/api/v1/reports/{draft['report_id']}/data-config", headers=auth_headers).json()
    assert config["data_source"]["url"] == REST_SOURCE["url"]
    assert [p["name"] for p in config["parameters"]] == ["month"]

    # The starter places every field: the filter, the single values, a loop row over the list and a total.
    file = client.get(f"/api/v1/reports/{draft['report_id']}/file", headers=auth_headers)
    text = "\n".join(p.text for p in Document(io.BytesIO(file.content)).paragraphs)
    cells = [c.text for t in Document(io.BytesIO(file.content)).tables for r in t.rows for c in r.cells]
    assert "{{ month }}" in text and "{{ company.name }}" in text
    assert "{%tr for invoice in invoices %}" in cells and "{{ invoice.customer_name }}" in cells and "{%tr endfor %}" in cells
    assert '{{ invoices | sum(attribute="total_usd") }}' in cells

    # ...and it renders as it is, Khmer included.
    rendered = client.post(f"/api/v1/reports/{draft['report_id']}/render?format=docx", json=draft["sample_context"])
    assert rendered.status_code == 200, rendered.text
    xml = zipfile.ZipFile(io.BytesIO(rendered.content)).read("word/document.xml").decode()
    assert "ហាង" in xml and "{%" not in xml and "1552.75" in xml


@pytest.mark.parametrize("fmt", ["xlsx", "html"])
def test_starters_in_the_other_file_types_render(auth_headers, upstream, fmt):
    draft = _draft(auth_headers, fmt=fmt)
    out_format = "xlsx" if fmt == "xlsx" else "pdf"
    rendered = client.post(f"/api/v1/reports/{draft['report_id']}/render?format={out_format}", json=draft["sample_context"])
    assert rendered.status_code == 200, rendered.text
    if fmt == "xlsx":
        values = [c.value for row in load_workbook(io.BytesIO(rendered.content)).active.iter_rows() for c in row if c.value is not None]
        assert 1552.75 in values and 1240 in values  # numbers stay numbers, the total is computed
    else:
        assert rendered.content[:5] == b"%PDF-"


def test_the_starter_can_be_made_again_in_another_file_type(auth_headers, upstream):
    draft = _draft(auth_headers)
    again = client.post(f"/api/v1/reports/{draft['report_id']}/starter", json={"format": "xlsx"}, headers=auth_headers)
    assert again.status_code == 200, again.text
    assert again.json()["template_ext"] == "xlsx" and again.json()["version"] == 2
    log = client.get(f"/api/v1/reports/{draft['report_id']}/changelog", headers=auth_headers).json()
    assert log["current_version"] == 2 and log["original_available"] is True  # the docx starter is kept in the history


def test_a_draft_takes_a_code_and_refuses_a_taken_one(auth_headers, upstream):
    _draft(auth_headers, code="monthly-invoices")
    again = client.post(
        "/api/v1/reports/drafts",
        json={"name": "Again", "format": "docx", "data_source": REST_SOURCE, "sample": INVOICES, "code": "monthly-invoices"},
        headers=auth_headers,
    )
    assert again.status_code == 409


# --- checking a designed template --------------------------------------------------------------------------------


def _docx_with(paragraphs: list[str], rows: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if rows:
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for r, row in enumerate(rows):
            for c, text in enumerate(row):
                table.rows[r].cells[c].text = text
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _upload(headers, report_id, content: bytes, filename="design.docx"):
    resp = client.put(
        f"/api/v1/reports/{report_id}/file",
        files={"file": (filename, content, "application/octet-stream")},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_the_check_follows_loops_and_names_the_nearest_real_field(auth_headers, upstream):
    draft = _draft(auth_headers)
    design = _docx_with(
        ["Invoices for {{ month }}", "{{ compnay.name }}"],
        [
            ["{%tr for inv in invoices %}", ""],
            ["{{ inv.invoice_no }}", "{{ inv.customer }}"],
            ["{%tr endfor %}", ""],
        ],
    )
    _upload(auth_headers, draft["report_id"], design)
    check = client.get(f"/api/v1/reports/{draft['report_id']}/template-check", headers=auth_headers).json()
    assert check["checked"] is True
    assert "inv.invoice_no" in check["matched"] and "month" in check["matched"]
    assert {u["placeholder"]: u["suggestion"] for u in check["unknown"]} == {
        "compnay.name": "company.name",
        "inv.customer": "inv.customer_name",
    }
    assert "invoices[].status" in check["unused"] and "company.phone" in check["unused"]
    assert any(f["path"] == "invoices[].total_usd" for f in check["fields"])


def test_the_check_reads_html_and_xlsx_and_ignores_functions_and_methods(tmp_path):
    html = tmp_path / "t.html"
    html.write_text(
        "{% set kh = ''.maketrans('0123456789', '០១២៣៤៥៦៧៨៩') %}"
        "<img src=\"{{ resource('logo.png') }}\">"
        "{% for row in invoices %}{{ row.issued_on[:4] }} {{ row.customer_name.replace('Co', '') }} "
        "{{ '{:,.2f}'.format(row.total_usd) }} {{ loop.index }}{% endfor %}"
        "{{ invoices | sum(attribute='total_usd') }} {{ company.name | upper }}",
        encoding="utf-8",
    )
    result = report_wizard.check_template(html, "html", INVOICES)
    assert result["checked"] and result["unknown"] == []
    assert set(result["unused"]) == {"invoices[].invoice_no", "invoices[].status", "company.phone"}

    xlsx_bytes = report_wizard.starter_template("T", INVOICES, "xlsx")
    xlsx = tmp_path / "t.xlsx"
    xlsx.write_bytes(xlsx_bytes)
    assert report_wizard.check_template(xlsx, "xlsx", INVOICES)["unknown"] == []


def test_the_check_reports_a_syntax_error_instead_of_failing(tmp_path):
    broken = tmp_path / "t.html"
    broken.write_text("{% for x in invoices %}{{ x.invoice_no }}", encoding="utf-8")
    result = report_wizard.check_template(broken, "html", INVOICES)
    assert result["checked"] is False and "syntax error" in result["reason"]


def test_an_html_template_can_replace_a_docx_one(auth_headers, upstream):
    draft = _draft(auth_headers)
    page = "<html><body><h1>{{ month }}</h1>{% for i in invoices %}<p>{{ i.invoice_no }}</p>{% endfor %}</body></html>"
    replaced = _upload(auth_headers, draft["report_id"], page.encode("utf-8"), filename="design.html")
    assert replaced["template_ext"] == "html"
    broken = client.put(
        f"/api/v1/reports/{draft['report_id']}/file",
        files={"file": ("x.html", b"{{ resource('missing.png') }}", "text/html")},
        headers=auth_headers,
    )
    assert broken.status_code == 400  # a resource nobody mapped


# --- drafts are private until published --------------------------------------------------------------------------


def test_a_draft_is_hidden_from_readers_until_it_is_published(auth_headers, make_local_user, upstream):
    draft = _draft(auth_headers)
    rid = draft["report_id"]
    _, reader = make_local_user("reader", "pw-reader-1", ("ROLE_REPORT_VIEWER",))

    def listed(headers):
        return [r for r in client.get("/api/v1/reports/accessible", headers=headers).json() if r["report_id"] == rid]

    # The manager sees and runs it, marked as a draft; a reader doesn't know it exists.
    assert listed(auth_headers)[0]["is_draft"] is True
    assert listed(reader) == []
    assert client.get(f"/api/v1/reports/{rid}/run-form", headers=reader).status_code == 404
    assert client.post(f"/api/v1/reports/{rid}/run", json={"parameters": {"month": "2026-09"}, "format": "docx"}, headers=reader).status_code == 404
    manager_run = client.post(f"/api/v1/reports/{rid}/run", json={"parameters": {"month": "2026-09"}, "format": "docx"}, headers=auth_headers)
    assert manager_run.status_code == 200, manager_run.text

    published = client.post(f"/api/v1/reports/{rid}/publish", headers=auth_headers)
    assert published.status_code == 200 and published.json()["is_draft"] is False
    assert listed(reader)[0]["is_draft"] is False
    assert client.get(f"/api/v1/reports/{rid}/run-form", headers=reader).status_code == 200


def test_publishing_needs_manage_access(make_local_user, auth_headers, upstream):
    draft = _draft(auth_headers)
    _, reader = make_local_user("reader", "pw-reader-1", ("ROLE_REPORT_VIEWER",))
    assert client.post(f"/api/v1/reports/{draft['report_id']}/publish", headers=reader).status_code == 403


def test_registering_a_finished_template_still_publishes_it_at_once(auth_headers):
    resp = client.post(
        "/api/v1/reports",
        data={"name": "Direct"},
        files={"file": ("t.html", b"<p>{{ x }}</p>", "text/html")},
        headers=auth_headers,
    )
    assert resp.status_code == 200 and resp.json()["is_draft"] is False


def test_starter_text_is_never_read_as_template_code():
    # A name or a key with braces in it is shown, not run.
    blob = report_wizard.starter_template("Sales {{ evil }}", {"a{%b": 1, "rows": [{"x{{y": 2}]}, "html").decode()
    assert "{{ evil }}" not in blob and "Sales ((" in blob
    json.dumps(report_wizard.data_fields({"rows": [{"x": 1}, {"y": None}]}))  # merged keys, serialisable


def test_dates_and_ids_reach_the_template_unchanged_while_khmer_is_still_segmented(auth_headers, upstream):
    """The guide's date recipe slices the text (`d[8:10]`): that only works if nothing was inserted into it. Word
    breaks are for Khmer; a date, an ID or an English status has no use for them."""
    draft = _draft(auth_headers, fmt="xlsx")
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "{{ d[8:10] }}/{{ d[5:7] }}/{{ d[:4] }}"
    ws["A2"] = "{{ 'UNPAID' if s == 'not paid' else 'ok' }}"
    ws["A3"] = "{{ id }}"
    ws["A4"] = "{{ k }}"
    buf = io.BytesIO()
    wb.save(buf)
    _upload(auth_headers, draft["report_id"], buf.getvalue(), filename="t.xlsx")

    context = {"d": "2026-09-02", "s": "not paid", "id": "INV-2026-0131", "k": "ហាងលក់ទំនិញស្រីពៅ"}
    rendered = client.post(f"/api/v1/reports/{draft['report_id']}/render?format=xlsx", json=context)
    assert rendered.status_code == 200, rendered.text
    cells = [c.value for row in load_workbook(io.BytesIO(rendered.content)).active.iter_rows() for c in row]
    assert cells[:3] == ["02/09/2026", "UNPAID", "INV-2026-0131"]  # nothing invisible inside them
    assert "\u200b" in cells[3] and cells[3].replace("\u200b", "") == context["k"]  # Khmer still gets its break points
