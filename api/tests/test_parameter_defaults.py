"""Default values for free-text, number, date, datetime and time parameters:
what a manager may set, what the run form tells the portal, and what a run that
leaves the parameter out uses (app/report_data.py)."""
import re
from datetime import datetime, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import report_data
from app.main import app

http = TestClient(app)

TEMPLATE = "by={{ signedBy }} from={{ fromDate }} at={{ at }} stamp={{ stamp }}"
PARAMETERS = [
    {"name": "fromDate", "label": "From", "type": "date", "default_value": "now()"},
    {"name": "at", "type": "time", "required": False, "default_value": "now()"},
    {"name": "stamp", "type": "datetime", "required": False, "default_value": "now()"},
    {"name": "signedBy", "label": "Signed by", "type": "text", "required": False, "default_value": "H.E"},
]


def _validate(parameters):
    return report_data.validate_data_config(parameters, None)[0]


def _docx(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _text(content: bytes) -> str:
    return "\n".join(p.text for p in Document(BytesIO(content)).paragraphs).replace("​", "")


@pytest.fixture
def report_id(auth_headers) -> str:
    resp = http.post(
        "/api/v1/reports",
        files={"file": ("t.docx", _docx(TEMPLATE), "application/octet-stream")},
        data={"name": "Defaults"},
        headers=auth_headers,
    )
    assert resp.status_code == 200, resp.text
    rid = resp.json()["report_id"]
    saved = http.put(f"/api/v1/reports/{rid}/data-config", json={"parameters": PARAMETERS, "data_source": None}, headers=auth_headers)
    assert saved.status_code == 200, saved.text
    return rid


def _run(headers, rid, parameters):
    return http.post(f"/api/v1/reports/{rid}/run", json={"parameters": parameters, "format": "docx"}, headers=headers)


# --- what a manager may save ------------------------------------------------------------


def test_a_default_is_stored_as_given_and_absent_by_default():
    cleaned = _validate([
        {"name": "a", "default_value": "  H.E  "},
        {"name": "b"},
        {"name": "c", "default_value": "   "},
        {"name": "d", "default_value": None},
    ])
    assert [p["default_value"] for p in cleaned] == ["H.E", None, None, None]


@pytest.mark.parametrize(
    "type_, default",
    [
        ("text", "anything at all"),
        ("number", "12.5"),
        ("date", "2026-09-30"),
        ("time", "08:30"),
        ("time", "08:30:15"),
        ("datetime", "2026-09-30T08:30"),
        ("date", "now()"),
        ("time", "now()"),
        ("datetime", "now()"),
    ],
)
def test_defaults_that_fit_their_type_are_accepted(type_, default):
    assert _validate([{"name": "p", "type": type_, "default_value": default}])[0]["default_value"] == default


@pytest.mark.parametrize("spelling", ["now()", "NOW()", " now( ) ", "Now ()"])
def test_now_may_be_spelled_loosely_and_is_stored_canonically(spelling):
    assert _validate([{"name": "p", "type": "date", "default_value": spelling}])[0]["default_value"] == "now()"


@pytest.mark.parametrize(
    "type_, default, message",
    [
        ("number", "abc", "valid number"),
        ("number", "now()", "can't default to now()"),
        ("text", "now()", "can't default to now()"),  # a text box would print the words, not the time
        ("date", "30/09/2026", "valid date"),
        ("date", "2026-02-30", "valid date"),  # right shape, not a day
        ("date", "2026-13-01", "valid date"),
        ("time", "25:00", "valid time"),
        ("time", "8am", "valid time"),
        ("datetime", "2026-09-30", "valid datetime"),
        ("text", "x" * 201, "too long"),
    ],
)
def test_defaults_that_do_not_fit_their_type_are_refused_with_an_example(type_, default, message):
    with pytest.raises(report_data.DataConfigError, match=message) as info:
        _validate([{"name": "p", "type": type_, "default_value": default}])
    assert "'p'" in str(info.value)


def test_the_message_for_a_bad_date_offers_now():
    with pytest.raises(report_data.DataConfigError, match=r"2026-09-30, or use now\(\)"):
        _validate([{"name": "p", "type": "date", "default_value": "soon"}])


def test_a_choice_list_cannot_have_a_default():
    with pytest.raises(report_data.DataConfigError, match="choice list"):
        _validate([{"name": "p", "default_value": "A", "options": [{"value": "A"}]}])
    with pytest.raises(report_data.DataConfigError, match="choice list"):
        _validate([{"name": "p", "default_value": "A", "options_source": {"url": "http://x.test/", "value_field": "c"}}])


# --- what the run form says -----------------------------------------------------------------


def test_the_run_form_carries_defaults_as_configured_not_resolved(auth_headers, report_id):
    form = http.get(f"/api/v1/reports/{report_id}/run-form", headers=auth_headers).json()
    assert {p["name"]: p["default_value"] for p in form["parameters"]} == {
        "fromDate": "now()",
        "at": "now()",
        "stamp": "now()",
        "signedBy": "H.E",
    }


def test_the_data_config_round_trips_defaults(auth_headers, report_id):
    saved = http.get(f"/api/v1/reports/{report_id}/data-config", headers=auth_headers).json()
    assert [p["default_value"] for p in saved["parameters"]] == ["now()", "now()", "now()", "H.E"]


# --- what a run uses -----------------------------------------------------------------------


def test_a_run_that_leaves_a_parameter_out_uses_its_default(auth_headers, report_id):
    day = lambda: datetime.now(timezone.utc).strftime("%Y-%m-%d")  # noqa: E731
    before = day()
    resp = _run(auth_headers, report_id, {})
    after = day()
    assert resp.status_code == 200, resp.text
    text = _text(resp.content)
    assert re.fullmatch(
        rf"by=H\.E from=({before}|{after}) at=\d{{2}}:\d{{2}} stamp=({before}|{after})T\d{{2}}:\d{{2}}", text
    ), text


def test_a_value_that_is_sent_beats_the_default(auth_headers, report_id):
    text = _text(_run(auth_headers, report_id, {"signedBy": "Dr. Sok", "fromDate": "2026-01-02"}).content)
    assert "by=Dr. Sok" in text and "from=2026-01-02" in text


def test_an_empty_string_is_an_answer_not_a_gap(auth_headers, report_id):
    # The run form sends every filter, so clearing a pre-filled optional field means "none".
    text = _text(_run(auth_headers, report_id, {"signedBy": "", "at": "", "stamp": "", "fromDate": "2026-01-02"}).content)
    assert text == "by= from=2026-01-02 at= stamp="


def test_a_required_parameter_cleared_to_empty_is_still_refused(auth_headers, report_id):
    resp = _run(auth_headers, report_id, {"fromDate": ""})
    assert resp.status_code == 400 and "From" in resp.json()["detail"]


def test_a_required_parameter_with_a_default_no_longer_has_to_be_sent(auth_headers, report_id):
    assert _run(auth_headers, report_id, {}).status_code == 200


def test_a_default_is_still_checked_like_any_other_value(auth_headers):
    # Saved before the shape rules tightened, or edited straight into the database: the run still checks it.
    definitions = [{"name": "d", "type": "date", "required": True, "default_value": "yesterday"}]
    with pytest.raises(report_data.ParameterError, match="valid date"):
        report_data.resolve_run_parameters(definitions, {"d": None}, {})


# --- now() ---------------------------------------------------------------------------------


def test_now_resolves_by_type_in_the_shape_a_browser_control_submits():
    moment = datetime(2026, 9, 30, 8, 5, 42, tzinfo=timezone.utc)
    resolve = lambda type_: report_data.resolve_default({"type": type_, "default_value": "now()"}, now=moment)  # noqa: E731
    assert resolve("date") == "2026-09-30"
    assert resolve("time") == "08:05"
    assert resolve("datetime") == "2026-09-30T08:05"
    assert resolve("text") is None and resolve("number") is None  # now() means nothing there


def test_a_literal_default_is_returned_untouched_and_no_default_is_none():
    assert report_data.resolve_default({"type": "text", "default_value": "H.E"}) == "H.E"
    assert report_data.resolve_default({"type": "text"}) is None
    assert report_data.resolve_default({"type": "text", "default_value": ""}) is None


def test_now_is_read_in_the_configured_time_zone(monkeypatch):
    monkeypatch.setenv("REPORT_TIMEZONE", "Asia/Phnom_Penh")
    assert report_data._report_timezone() == ZoneInfo("Asia/Phnom_Penh")
    # 18:00 UTC on the 30th is already 01:00 on the 1st in Phnom Penh -- the reason the zone is configurable.
    moment = datetime(2026, 9, 30, 18, 0, tzinfo=timezone.utc).astimezone(report_data._report_timezone())
    assert report_data.resolve_default({"type": "date", "default_value": "now()"}, now=moment) == "2026-10-01"


def test_the_time_zone_defaults_to_utc_and_a_bad_name_falls_back_to_it(monkeypatch, caplog):
    monkeypatch.delenv("REPORT_TIMEZONE", raising=False)
    assert report_data._report_timezone() == timezone.utc
    monkeypatch.setenv("REPORT_TIMEZONE", "Mars/Olympus_Mons")
    assert report_data._report_timezone() == timezone.utc
    assert "Mars/Olympus_Mons" in caplog.text


# --- firstDayOfMonth() / lastDayOfMonth() -----------------------------------------------------


@pytest.mark.parametrize("spelling", ["firstDayOfMonth()", "firstDayOfMonth", " FIRSTDAYOFMONTH ( ) "])
def test_first_day_of_month_is_stored_in_one_spelling(spelling):
    assert _validate([{"name": "p", "type": "date", "default_value": spelling}])[0]["default_value"] == "firstDayOfMonth()"


@pytest.mark.parametrize("type_", ["time", "number", "text"])
@pytest.mark.parametrize("expression", ["firstDayOfMonth()", "lastDayOfMonth()"])
def test_month_days_only_make_sense_on_dates(type_, expression):
    with pytest.raises(report_data.DataConfigError, match="only works for date and datetime"):
        _validate([{"name": "p", "type": type_, "default_value": expression}])


def test_month_days_resolve_against_the_current_month():
    leap = datetime(2028, 2, 10, 15, 30, tzinfo=timezone.utc)
    resolve = lambda type_, expr, at: report_data.resolve_default({"type": type_, "default_value": expr}, now=at)  # noqa: E731
    assert resolve("date", "firstDayOfMonth()", leap) == "2028-02-01"
    assert resolve("date", "lastDayOfMonth()", leap) == "2028-02-29"
    assert resolve("datetime", "firstDayOfMonth()", leap) == "2028-02-01T00:00"
    assert resolve("datetime", "lastDayOfMonth()", leap) == "2028-02-29T23:59"
    assert resolve("date", "lastDayOfMonth()", datetime(2026, 12, 31, tzinfo=timezone.utc)) == "2026-12-31"


# --- a text value can't break out of a raw request body ---------------------------------------


@pytest.mark.parametrize(
    "content_type, value, expected",
    [
        ("application/json", 'x", "admin": true, "y": "', 'x\\", \\"admin\\": true, \\"y\\": \\"'),
        ("text/xml; charset=utf-8", "</q><admin/>", "&lt;/q&gt;&lt;admin/&gt;"),
        ("application/x-www-form-urlencoded", "a&role=admin", "a%26role%3Dadmin"),
        ("text/plain", "a&b", "a&b"),
    ],
)
def test_raw_body_values_are_escaped_for_the_content_type(content_type, value, expected):
    escape = report_data._body_escaper({"Content-Type": content_type})
    assert report_data._render_value("[{{ q }}]", report_data._env(), {"q": value}, escape) == f"[{expected}]"
