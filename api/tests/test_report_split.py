"""Splitting a report whose table is over the per-file row limit into several
files (app/report_split.py) -- the planning rules on their own, then the real
/render, /run and batch routes end to end (docx/xlsx: no LibreOffice needed)."""
import zipfile
from io import BytesIO

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import report_data, report_split
from app.main import app

client = TestClient(app)

TEMPLATE_LINE = (
    "Part {{ part_number }} of {{ part_count }} | rows {{ first_row }}-{{ last_row }} of {{ total_row_count }}"
    " | part {{ part_totals.amount }} | grand {{ grand_totals.amount }}"
)


def _rows(n: int) -> list[dict]:
    return [{"label": f"row{i}", "amount": 10} for i in range(1, n + 1)]


# --- planning: the rules, no HTTP ---------------------------------------------


def test_balanced_sizes_leave_no_stub_file():
    plan = report_split.plan_parts({"items": _rows(1001)}, {"items"}, limit=1000)
    assert [len(c["items"]) for c in plan] == [501, 500]  # not 1,000 + 1


def test_every_part_is_within_the_limit_and_nothing_is_lost_or_reordered():
    rows = _rows(47)
    plan = report_split.plan_parts({"items": rows}, {"items"}, limit=10)
    assert len(plan) == 5
    assert all(len(c["items"]) <= 10 for c in plan)
    assert [r for c in plan for r in c["items"]] == rows


def test_metadata_describes_each_part():
    plan = report_split.plan_parts({"items": _rows(5)}, {"items"}, limit=2)
    assert [(c["part_number"], c["part_count"]) for c in plan] == [(1, 3), (2, 3), (3, 3)]
    assert [(c["first_row"], c["last_row"], c["row_offset"]) for c in plan] == [(1, 2, 0), (3, 4, 2), (5, 5, 4)]
    assert all(c["total_row_count"] == 5 for c in plan)
    assert [c["part_row_count"] for c in plan] == [2, 2, 1]


def test_part_totals_and_grand_totals():
    rows = [{"label": "a", "amount": 10, "qty": 1}, {"label": "b", "amount": 20, "qty": 2}, {"label": "c", "amount": 30, "qty": 3}]
    plan = report_split.plan_parts({"items": rows}, {"items"}, limit=2)  # split 2 + 1: rows a+b, then c
    assert plan[0]["part_totals"] == {"amount": 30, "qty": 3}
    assert plan[1]["part_totals"] == {"amount": 30, "qty": 3}
    assert all(c["grand_totals"] == {"amount": 60, "qty": 6} for c in plan)
    assert sum(c["part_totals"]["amount"] for c in plan) == plan[0]["grand_totals"]["amount"]


def test_sums_ignore_booleans_strings_and_non_dict_rows_and_do_not_drift_on_floats():
    assert report_split._sums([{"a": True, "b": "5", "c": 2}, "not a dict", {"c": 3}]) == {"c": 5}
    assert report_split._sums([{"x": 0.1}] * 10) == {"x": 1.0}  # plain sum() gives 0.9999999999999999
    assert report_split._sums([{"n": 1}, {"n": 2.5}]) == {"n": 3.5}


def test_an_unsplit_report_still_gets_the_variables_so_templates_read_the_same():
    (only,) = report_split.plan_parts({"items": _rows(3)}, {"items"}, limit=10)
    assert (only["part_number"], only["part_count"], only["row_offset"]) == (1, 1, 0)
    assert only["part_totals"] == only["grand_totals"] == {"amount": 30}
    assert len(only["items"]) == 3


def test_no_repeating_table_means_the_data_is_untouched():
    data = {"customer": "x", "items": _rows(5000)}
    assert report_split.plan_parts(data, set(), limit=10) == [data]
    assert report_split.plan_parts({"customer": "x"}, {"items"}, limit=10) == [{"customer": "x"}]  # table absent from the data


def test_callers_own_keys_are_never_overridden():
    plan = report_split.plan_parts({"items": _rows(4), "grand_total": 999, "part_number": "mine"}, {"items"}, limit=2)
    assert all(c["grand_total"] == 999 and c["part_number"] == "mine" for c in plan)
    assert all(c["part_count"] == 2 for c in plan)  # the ones it didn't supply are still filled in


def test_two_oversized_tables_are_left_whole_with_a_warning(monkeypatch):
    warnings = []

    class _Log:
        def warning(self, *args, **kwargs):
            warnings.append(args)

    monkeypatch.setattr(report_split, "_log", _Log())
    data = {"a": _rows(5), "b": _rows(6)}
    assert report_split.plan_parts(data, {"a", "b"}, limit=2) == [data]
    assert warnings, "splitting one table would repeat the other in every part -- that should be logged"


def test_only_the_oversized_table_is_split_the_small_one_is_repeated():
    plan = report_split.plan_parts({"big": _rows(6), "small": _rows(2)}, {"big", "small"}, limit=3)
    assert len(plan) == 2
    assert all(len(c["big"]) == 3 and len(c["small"]) == 2 for c in plan)


def test_too_many_parts_is_refused_with_a_useful_message():
    with pytest.raises(report_split.TooManyPartsError) as info:
        report_split.plan_parts({"items": _rows(100)}, {"items"}, limit=10, parts_cap=5)
    assert (info.value.rows, info.value.parts, info.value.max_parts) == (100, 10, 5)
    assert "Narrow the filters" in str(info.value)


def test_settings_default_clamp_and_ignore_garbage(monkeypatch):
    monkeypatch.delenv("MAX_ROWS_PER_FILE", raising=False)
    assert report_split.max_rows_per_file() == 1000
    monkeypatch.setenv("MAX_ROWS_PER_FILE", "250")
    assert report_split.max_rows_per_file() == 250
    monkeypatch.setenv("MAX_ROWS_PER_FILE", "999999")
    assert report_split.max_rows_per_file() == report_split.ROWS_PER_FILE_CEILING  # can't be configured past the measured limit
    monkeypatch.setenv("MAX_ROWS_PER_FILE", "0")
    assert report_split.max_rows_per_file() == 1
    monkeypatch.setenv("MAX_ROWS_PER_FILE", "lots")
    assert report_split.max_rows_per_file() == 1000  # falls back rather than failing every render
    monkeypatch.setenv("MAX_REPORT_PARTS", "7")
    assert report_split.max_parts() == 7


def test_part_filenames_sort_correctly():
    assert report_split.part_filename("r1", 3, 12, "pdf") == "r1-part-03-of-12.pdf"
    assert report_split.part_filename("r1", 7, 120, "pdf") == "r1-part-007-of-120.pdf"


# --- finding the template's table --------------------------------------------


def _docx_with(*paragraph_runs: list[str]) -> bytes:
    doc = Document()
    for runs in paragraph_runs:
        p = doc.add_paragraph()
        for text in runs:
            p.add_run(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_loop_collections_from_a_docx_even_when_word_splits_the_tag_across_runs(tmp_path):
    path = tmp_path / "t.docx"
    path.write_bytes(_docx_with(["{%tr for it", "em in items %}"], ["{% for k, v in pairs %}"], ["{{ not_a_loop }}"]))
    assert report_split.loop_collections(path, "docx") == {"items", "pairs"}


def test_loop_collections_from_xlsx_and_html(tmp_path):
    wb = Workbook()
    wb.active["A1"] = "{% for line in lines %}"
    (tmp_path / "t.xlsx").write_bytes(b"")
    wb.save(tmp_path / "t.xlsx")
    assert report_split.loop_collections(tmp_path / "t.xlsx", "xlsx") == {"lines"}

    (tmp_path / "t.html").write_text("<ul>{% for x in rows %}<li>{{ x }}</li>{% endfor %}</ul>", encoding="utf-8")
    assert report_split.loop_collections(tmp_path / "t.html", "html") == {"rows"}


def test_an_unreadable_template_just_means_no_split(tmp_path):
    (tmp_path / "broken.docx").write_bytes(b"not a zip")
    assert report_split.loop_collections(tmp_path / "broken.docx", "docx") == set()
    assert report_split.loop_collections(tmp_path / "missing.docx", "docx") == set()


# --- the routes, end to end ----------------------------------------------------


def _table_docx() -> bytes:
    doc = Document()
    doc.add_paragraph(TEMPLATE_LINE)
    table = doc.add_table(rows=3, cols=1)
    table.cell(0, 0).text = "{%tr for item in items %}"
    table.cell(1, 0).text = "{{ row_offset + loop.index }}:{{ item.label }}"
    table.cell(2, 0).text = "{%tr endfor %}"
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _register(auth_headers, content=None, filename="t.docx", name="Ledger") -> str:
    resp = client.post(
        "/api/v1/reports",
        files={"file": (filename, content or _table_docx(), "application/octet-stream")},
        data={"name": name},
        headers=auth_headers,
    )
    assert resp.status_code == 200, resp.text
    return resp.json()["report_id"]


def _read(content: bytes) -> tuple[str, list[str]]:
    doc = Document(BytesIO(content))
    return doc.paragraphs[0].text, [row.cells[0].text for row in doc.tables[0].rows]


@pytest.fixture
def small_files(monkeypatch):
    monkeypatch.setenv("MAX_ROWS_PER_FILE", "2")


def test_render_splits_into_a_zip_of_parts(auth_headers, small_files):
    report_id = _register(auth_headers)
    resp = client.post(f"/api/v1/reports/{report_id}/render?format=docx", json={"items": _rows(5)})
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/zip"
    assert resp.headers["x-report-parts"] == "3"
    assert resp.headers["content-disposition"] == f'attachment; filename="{report_id}-parts.zip"'

    with zipfile.ZipFile(BytesIO(resp.content)) as zf:
        assert zf.namelist() == [f"{report_id}-part-0{n}-of-03.docx" for n in (1, 2, 3)]
        parts = [_read(zf.read(name)) for name in zf.namelist()]

    # each file carries its own rows, continuing the numbering, and the totals template variables
    assert parts[0] == ("Part 1 of 3 | rows 1-2 of 5 | part 20 | grand 50", ["1:row1", "2:row2"])
    assert parts[1] == ("Part 2 of 3 | rows 3-4 of 5 | part 20 | grand 50", ["3:row3", "4:row4"])
    assert parts[2] == ("Part 3 of 3 | rows 5-5 of 5 | part 10 | grand 50", ["5:row5"])


def test_part_param_returns_just_that_file(auth_headers, small_files):
    report_id = _register(auth_headers)
    resp = client.post(f"/api/v1/reports/{report_id}/render?format=docx&part=2", json={"items": _rows(5)})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert (resp.headers["x-report-parts"], resp.headers["x-report-part"]) == ("3", "2")
    assert resp.headers["content-disposition"] == f'attachment; filename="{report_id}-part-02-of-03.docx"'
    assert _read(resp.content) == ("Part 2 of 3 | rows 3-4 of 5 | part 20 | grand 50", ["3:row3", "4:row4"])


def test_a_part_that_does_not_exist_is_a_400(auth_headers, small_files):
    report_id = _register(auth_headers)
    resp = client.post(f"/api/v1/reports/{report_id}/render?format=docx&part=9", json={"items": _rows(5)})
    assert resp.status_code == 400
    assert "no part 9" in resp.json()["detail"]
    assert client.post(f"/api/v1/reports/{report_id}/render?format=docx&part=0", json={"items": _rows(5)}).status_code == 422


def test_a_report_under_the_limit_is_one_plain_file_and_still_gets_the_variables(auth_headers):
    report_id = _register(auth_headers)  # default limit 1,000
    resp = client.post(f"/api/v1/reports/{report_id}/render?format=docx", json={"items": _rows(3)})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert resp.headers["x-report-parts"] == "1"
    assert resp.headers["content-disposition"] == f'attachment; filename="{report_id}.docx"'
    assert _read(resp.content) == ("Part 1 of 1 | rows 1-3 of 3 | part 30 | grand 30", ["1:row1", "2:row2", "3:row3"])


def test_too_many_parts_is_a_400_before_anything_is_rendered(auth_headers, monkeypatch):
    monkeypatch.setenv("MAX_ROWS_PER_FILE", "2")
    monkeypatch.setenv("MAX_REPORT_PARTS", "2")
    report_id = _register(auth_headers)
    resp = client.post(f"/api/v1/reports/{report_id}/render?format=docx", json={"items": _rows(5)})
    assert resp.status_code == 400
    assert "Narrow the filters" in resp.json()["detail"]


def test_xlsx_reports_split_too(auth_headers, small_files):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "{% for item in items %}"
    ws.row_dimensions[1].hidden = True
    ws["A2"] = "{{ item.label }}"
    ws["A3"] = "{% endfor %}"
    ws.row_dimensions[3].hidden = True
    buf = BytesIO()
    wb.save(buf)
    report_id = _register(auth_headers, buf.getvalue(), filename="t.xlsx")

    resp = client.post(f"/api/v1/reports/{report_id}/render?format=xlsx", json={"items": _rows(3)})
    assert resp.status_code == 200 and resp.headers["x-report-parts"] == "2"
    with zipfile.ZipFile(BytesIO(resp.content)) as zf:
        assert zf.namelist() == [f"{report_id}-part-01-of-02.xlsx", f"{report_id}-part-02-of-02.xlsx"]
        labels = [
            [c.value for row in load_workbook(BytesIO(zf.read(n))).active.iter_rows() for c in row if c.value and str(c.value).startswith("row")]
            for n in zf.namelist()
        ]
    assert labels == [["row1", "row2"], ["row3"]]


def test_batch_refuses_a_record_that_would_need_splitting_before_rendering_any(auth_headers, small_files, monkeypatch):
    from app.routers import reports as reports_router

    rendered = []
    real = reports_router._render_one
    monkeypatch.setattr(reports_router, "_render_one", lambda *a, **k: rendered.append(1) or real(*a, **k))
    report_id = _register(auth_headers)
    resp = client.post(
        f"/api/v1/reports/{report_id}/render/batch?format=docx",
        json=[{"items": _rows(1)}, {"items": _rows(5)}],
    )
    assert resp.status_code == 400
    assert "Record 2" in resp.json()["detail"]
    assert rendered == [], "the oversized record must fail the request up front, not after rendering the first"


def test_batch_of_small_records_still_works(auth_headers, small_files):
    report_id = _register(auth_headers)
    resp = client.post(
        f"/api/v1/reports/{report_id}/render/batch?format=docx",
        json=[{"items": _rows(1)}, {"items": _rows(2)}],
    )
    assert resp.status_code == 200
    with zipfile.ZipFile(BytesIO(resp.content)) as zf:
        assert len(zf.namelist()) == 2
        assert _read(zf.read(zf.namelist()[1]))[1] == ["1:row1", "2:row2"]


# --- the authenticated /run route --------------------------------------------


@pytest.fixture
def ledger_source(monkeypatch):
    """The report's REST data source answering with 5 rows."""

    def factory():
        return httpx.Client(transport=httpx.MockTransport(lambda request: httpx.Response(200, json={"items": _rows(5)})), follow_redirects=False)

    monkeypatch.setattr(report_data, "_make_client", factory)


def _run(auth_headers, report_id, **body):
    return client.post(f"/api/v1/reports/{report_id}/run", json={"parameters": {}, "format": "docx", **body}, headers=auth_headers)


def test_run_splits_and_honours_part(auth_headers, small_files, ledger_source):
    report_id = _register(auth_headers)
    client.put(
        f"/api/v1/reports/{report_id}/data-config",
        json={"parameters": [], "data_source": {"url": "https://erp.example.test/ledger", "method": "GET"}},
        headers=auth_headers,
    )

    whole = _run(auth_headers, report_id)
    assert whole.status_code == 200, whole.text
    assert whole.headers["content-type"] == "application/zip" and whole.headers["x-report-parts"] == "3"

    one = _run(auth_headers, report_id, part=3)
    assert one.status_code == 200
    assert (one.headers["x-report-parts"], one.headers["x-report-part"]) == ("3", "3")
    assert _read(one.content) == ("Part 3 of 3 | rows 5-5 of 5 | part 10 | grand 50", ["5:row5"])

    assert _run(auth_headers, report_id, part=4).status_code == 400
    assert _run(auth_headers, report_id, part=0).status_code == 422
